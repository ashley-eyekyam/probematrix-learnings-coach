"""The 9M session engine (spec §7.4, §7.5; P5-13, OD5-1): deterministic suggestions for each 9M question
(nine-m-suggest), the session file <ws>/nine_m_session.json that records each question and its answer (nine-m-answer),
a scripted session that takes the recommended suggestion at every question (nine-m-replay --accept-recommended), the
nine-m-map.json side output of a nine-m-analysis pack and the blank actions template (nine-m-actions-template).

The model writes the nine-m-analysis pack from the session; the engine never writes the pack. Suggestions are
conditions, never character: discipline, accountability and consequences are never offered (§7.5), training comes
only with a stronger control and says it is weak on its own, and no suggestion carries an owner or a date (D7)."""
from __future__ import annotations

import csv
import io
import json
import re
import sys
from pathlib import Path

from .ninem import load_nine_m, NINE_M_CODES, NINE_M_NONE, NINE_M_SEED_FILE
from .util import _abs, _s, write_json
from .workspace import load_ws

NINE_M_SESSION_FILE = "nine_m_session.json"
NINE_M_MAP_FILE = "nine-m-map.json"
NINE_M_PACK = "nine-m-analysis"
# Per M, in this order (§7.4 step 3a-3f); 3g, the leading indicator, is optional and asked in free text.
NINE_M_STEPS = ("relevance", "subcategory", "condition", "driver", "barrier", "change")
NINE_M_GLOBAL_STEPS = ("scope", "cross-m", "gaps")      # §7.4 steps 2, 4 and 5
NINE_M_BUDGET_PER_M, NINE_M_BUDGET_GLOBAL = 6, 4      # at most 6 x 9 + 4 = 58 questions (§7.4)
NINE_M_TOP_MIN, NINE_M_TOP_MAX = 2, 4                 # 2-4 suggestions, plus None / unknown (§7.4)
NINE_M_SEED_MISSING = "{ws} has no nine_m_seed.json: run nine-m-seed first"
NINE_M_FREE_TEXT = ("Describe it in your own words (recorded as context supplied by the organisation, without a "
                    "finding ID)")
NINE_M_SKIP = "Skip the rest of this M (its remaining questions are recorded as not asked)"
NINE_M_WEAK = "(weak on its own)"
# Never offered (§7.5): a candidate whose text carries one of these words is dropped, whatever its source.
NINE_M_NEVER_RX = re.compile(r"\b(?:disciplin\w*|accountab\w*|consequences?|punish\w*|sanction\w*|blam\w*)", re.I)
# Suggestions that state no fact about the export, so they carry no finding ID (§7.4: only fact-bearing ones must).
NINE_M_PROCEDURAL = (NINE_M_NONE, NINE_M_FREE_TEXT, NINE_M_SKIP, "Not relevant: no finding in this export", "A list you choose",
                     "Keep every M as a separate learning", "No causal link joins", "Ask for evidence on each unknown",
                     "Record the unknowns as gaps", "No answer was unknown", "Ask for evidence anyway",
                     "Skip the subcategory", "Every M with seed weight", "The top four by top-rated findings")
# The four overlaps of §7.2: the skill asks which M owns a condition instead of counting it twice.
NINE_M_OVERLAPS = {"M7.02": "M9", "M7.10": "M1", "M6.04": "M1", "M1.02": "M6 and M7", "M1.03": "M7",
                   "M1.10": "M4", "M1.11": "M4", "M4.04": "M1", "M4.08": "M1"}
NINE_M_DATE_RX = re.compile(
    r"\b\d{4}-\d{2}-\d{2}\b|\b\d{1,2}/\d{1,2}/\d{2,4}\b|\b\d{1,2}(?:st|nd|rd|th)?\s+(?:jan|feb|mar|apr|may|jun|jul|aug|"
    r"sep|oct|nov|dec)[a-z]*\b|\b(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\s+\d{1,2}\b|\bby\s+(?:end of|"
    r"next|this)\s+\w+", re.I)
NINE_M_D7_NOTE = ("note: the answer names a date or a deadline; it is stored as given, but owners and due dates are "
                  "set by the organisation and stay out of the pack (D7)")
TEMPLATE_COLUMNS = ["Recommendation ID", "M", "Change", "HoC tier", "Corrective/Preventive", "Indicator", "Owner",
                    "Due date", "Status", "Notes"]
_NM_STOP = frozenset("""about after again also among and another any are because been before being between both but
by can could did does doing down during each either else every for from further had has have having here how into its
itself just like made make many more most much must not now off once only other our out over own same should since
some such than that the their them then there these they this those through too under until upon very was were what
when where whether which while who whom why will with within without would your yes partly were the job work used use
using""".split())
_NM_WORD_RX = re.compile(r"[a-z][a-z\-]{3,}")
_NM_SUB_RX = re.compile(r"\bM([1-9])\.(\d{2})\b")
_NM_STATE_CUES = (("Defeated", re.compile(r"bypass|overrid|defeat|disabl|inhibit|jumper", re.I)),
                  ("Absent", re.compile(r"\bno\b|\bnot (?:provided|available|installed|fitted|in place|done|"
                                        r"recorded)|absent|\black|without|missing|\bnone\b", re.I)),
                  ("Worked partly", re.compile(r"partly|partial|limited|delayed|\blate\b", re.I)),
                  ("Degraded", re.compile(r"inadequa|degrad|poor|insufficient|not maintained|faulty|fail", re.I)))


# ----------------------------------------------------------------------------- loading

def load_session_ws(ws) -> dict:
    """The workspace with its seed, causal graph and session, for the 9M commands. Exits 2 when there is no seed."""
    ws = _abs(ws)
    if not (ws / NINE_M_SEED_FILE).is_file():
        print(NINE_M_SEED_MISSING.format(ws=ws), file=sys.stderr)
        sys.exit(2)
    w = load_ws(ws)
    w.setdefault("nine_m_seed", json.loads((ws / NINE_M_SEED_FILE).read_text(encoding="utf-8")))
    cg = ws / "causal_graph.json"
    w["causal_graph"] = json.loads(cg.read_text(encoding="utf-8")) if cg.is_file() else {"nodes": [], "edges": []}
    return w


def read_session(ws) -> dict | None:
    p = _abs(ws) / NINE_M_SESSION_FILE
    if not p.is_file():
        return None
    try:
        s = json.loads(p.read_text(encoding="utf-8"))
    except ValueError:
        _nm_fail(f"{p} is not valid JSON; move it aside and start the session again")
    if not (isinstance(s, dict) and isinstance(s.get("records"), list)):
        _nm_fail(f"{p} is not a 9M session file (it needs 'records'); move it aside and start again")
    return s


def write_session(ws, s) -> None:
    write_json(_abs(ws) / NINE_M_SESSION_FILE, json.loads(json.dumps(s, sort_keys=True)))


def _nm_fail(msg, code=2):
    print(msg, file=sys.stderr)
    sys.exit(code)


def _nm_chapters() -> dict:
    return {c["code"]: c for c in load_nine_m()["chapters"]}


def subcategory_ids() -> dict:
    """{id: name} for the 97 subcategories of references/nine-m.json."""
    return {s["id"]: s["name"] for c in load_nine_m()["chapters"] for s in c["subcategories"]}


def is_procedural(text) -> bool:
    """True for a suggestion that states no fact about the export (None, free text, a scope or gap choice)."""
    return any(_s(text).startswith(p) for p in NINE_M_PROCEDURAL)


# ----------------------------------------------------------------------------- candidates

def _nm_tokens(*texts) -> set:
    return {t.strip("-") for x in texts for t in _NM_WORD_RX.findall(_s(x).lower())} - _NM_STOP


def _nm_trim(text, n=150) -> str:
    t = re.sub(r"\s+", " ", _s(text)).strip()
    if len(t) <= n:
        return t
    cut = t.rfind(" ", 0, n)
    return t[:cut if cut > 0 else n].rstrip(" ,;:") + "…"


def _nm_ranked(w, seed, m) -> list:
    """The findings with any score for m, in rating order: rating, then score for m, then level, then file order."""
    S = seed.get("findings", {})
    rows = [(i, f) for i, f in enumerate(w["findings"]) if (S.get(f["id"]) or {}).get("scores", {}).get(m)]
    rows.sort(key=lambda t: (-(t[1].get("rating") or 0), -S[t[1]["id"]]["scores"][m], -(t[1].get("level") or 0), t[0]))
    return [f for _, f in rows]


def _nm_sub_text(sub) -> set:
    return _nm_tokens(sub["name"], sub.get("definition"), *(sub.get("probes") or []))


def _nm_finding_text(f) -> set:
    return _nm_tokens(f.get("label"), f.get("short"), f.get("name"), f.get("subcategory"))


def _nm_focus(w, seed, m, subcategory) -> list:
    """The M's findings in rating order; with a subcategory, those that share words with it come first."""
    fs = _nm_ranked(w, seed, m)
    if not subcategory:
        return fs
    sub = next(s for s in _nm_chapters()[m]["subcategories"] if s["id"] == subcategory)
    st = _nm_sub_text(sub)
    return [f for f in fs if st & _nm_finding_text(f)] + [f for f in fs if not st & _nm_finding_text(f)]


def _nm_relevance(w, seed, m, fs, sub, top):
    title = _nm_chapters()[m]["title"]
    if not fs:
        return [{"text": f"Not relevant: no finding in this export points to {m} {title}", "cites": []},
                {"text": NINE_M_FREE_TEXT, "cites": []}]
    S = seed["findings"]
    prim = sum(1 for f in fs if S[f["id"]]["m_primary"] == m)
    lead = fs[0]
    out = [{"text": f"Relevant: {len(fs)} finding(s) point to {m} {title} ({prim} by their Nanocode, "
                    f"{len(fs) - prim} by their wording or cluster), led by {lead['code']} {lead.get('name') or ''}"
                    .rstrip(),
            "cites": [f["id"] for f in fs[:3]]}]
    codes = {}
    for f in fs:
        codes.setdefault(f["code"], []).append(f)
    if len(codes) > 1:
        code, mine = sorted(codes.items(), key=lambda kv: (-len(kv[1]), fs.index(kv[1][0])))[0]
        out.append({"text": f"Partly relevant: only the {code} {mine[0].get('name') or ''} findings".replace("  ", " "),
                    "cites": [f["id"] for f in mine[:3]]})
    out.append({"text": f"Not relevant: the {m} {title} signal in these findings is incidental (say why)",
                "cites": [fs[0]["id"]]})
    return out


def _nm_subcategory(w, seed, m, fs, sub, top):
    scored = []
    for s in _nm_chapters()[m]["subcategories"]:
        st = _nm_sub_text(s)
        hits = [(len(st & _nm_finding_text(f)), i, f) for i, f in enumerate(fs)]
        hits = [h for h in hits if h[0]]
        if hits:
            total = sum(h[0] for h in hits)
            hits.sort(key=lambda h: (-h[0], h[1]))
            scored.append((total, s, [h[2]["id"] for h in hits[:3]]))
    scored.sort(key=lambda t: (-t[0], t[1]["id"]))
    out = []
    for _, s, cites in scored:
        other = NINE_M_OVERLAPS.get(s["id"])
        text = f"{s['id']} {s['name']}" + (f" (overlaps {other}: say which M owns it)" if other else "")
        out.append({"text": text, "cites": cites})
    if len(out) < NINE_M_TOP_MIN:
        out += [{"text": "Skip the subcategory: record the condition without one", "cites": []},
                {"text": NINE_M_FREE_TEXT, "cites": []}]
    return out


def _nm_condition(w, seed, m, fs, sub, top):
    out, seen = [], set()
    for f in fs:
        if f["code"] in seen:
            continue
        seen.add(f["code"])
        same = [g["id"] for g in fs if g["code"] == f["code"]][:3]
        out.append({"text": f"Condition: {_nm_trim(f.get('short') or f.get('label'))} ({f['code']} "
                            f"{f.get('name') or ''})".replace(" )", ")"), "cites": same})
    return out


def _nm_level_name(edge_side_level) -> str:
    return f"level {edge_side_level}" if edge_side_level else "level ?"


def _nm_driver(w, seed, m, fs, sub, top):
    S = seed["findings"]
    mcodes = {f["code"] for f in fs if S[f["id"]]["m_primary"] == m or m in S[f["id"]]["m_secondary"]} or \
        {f["code"] for f in fs}
    graph = w.get("causal_graph") or {}
    names = {n.get("code"): n.get("name") or "" for n in graph.get("nodes") or [] if isinstance(n, dict)}
    ids = {f["id"] for f in w["findings"]}
    edges = [e for e in graph.get("edges") or [] if isinstance(e, dict) and e.get("from") != e.get("to")]
    into = [e for e in edges if e.get("to") in mcodes and e.get("from_level") in (3, 4)]
    outof = [e for e in edges if e.get("from") in mcodes and e.get("from_level") in (3, 4) and e not in into]
    out = []
    for e in sorted(into, key=lambda e: (-(e.get("mentions") or 0), e["from"], e["to"])) + \
            sorted(outof, key=lambda e: (-(e.get("mentions") or 0), e["from"], e["to"])):
        cites = [x for x in e.get("findings") or [] if x in ids][:3]
        if cites:
            out.append({"text": f"System driver: {e['from']} {names.get(e['from'], '')} ({_nm_level_name(e.get('from_level'))}) "
                                f"feeds {e['to']} {names.get(e['to'], '')}".replace("  ", " "), "cites": cites})
    if len(out) < NINE_M_TOP_MAX:
        for f in fs:
            if f.get("level") in (3, 4):
                out.append({"text": f"System driver: {f['code']} {f.get('name') or ''} (level {f['level']})"
                                    .replace("  ", " "), "cites": [f["id"]]})
    if not out:
        for f in fs[:2]:
            out.append({"text": f"Why it made sense: {_nm_trim(f.get('causal') or f.get('short'), 140)}",
                        "cites": [f["id"]]})
    return out


def _nm_state(text) -> str:
    return next((st for st, rx in _NM_STATE_CUES if rx.search(text)), "Degraded")


def _nm_quote(e, f) -> str:
    links = [lk for lk in e.get("links") or [] if lk.get("code") == f["code"]] or e.get("links") or [{}]
    q = _s(links[0].get("reference")) or _s(e.get("description"))
    return q.strip(" \"'\u201c\u201d\u2018\u2019")


def _nm_barrier(w, seed, m, fs, sub, top):
    """Evidence quotes, one per finding first (in rating order), then the findings' other evidence."""
    E = w.get("evidence") or {}
    out, seen = [], set()
    for rnd in (0, 1):
        for f in fs:
            for k, eid in enumerate(f.get("evidence") or []):
                e = E.get(eid)
                if not e or eid in seen or (rnd == 0 and k) or not _nm_quote(e, f):
                    continue
                if rnd == 0 and any(f["id"] in s["cites"] for s in out):
                    continue
                seen.add(eid)
                quote = _nm_quote(e, f)
                st = _nm_state(f"{quote} {f.get('short') or ''}")
                out.append({"text": f"{st}: {_nm_trim(f.get('short') or f.get('label'), 100)} ({eid}: "
                                    f"\"{_nm_trim(quote, 120)}\")", "cites": [f["id"], eid]})
            if len(out) >= NINE_M_TOP_MAX:
                break
        if len(out) >= NINE_M_TOP_MAX:
            break
    if len(out) < NINE_M_TOP_MIN:
        for f in fs:
            if f["id"] not in {c for s in out for c in s["cites"]}:
                out.append({"text": f"{_nm_state(f.get('short') or '')}: {_nm_trim(f.get('short') or f.get('label'), 120)}"
                                    " (no evidence quote)", "cites": [f["id"]]})
    return out


def _nm_change(w, seed, m, fs, sub, top):
    if not fs:
        return []
    f = fs[0]
    focus = f"{f['code']} condition ({_nm_trim(f.get('name') or f.get('short'), 70)})"
    cites = [g["id"] for g in fs[:3]]
    tiers = [("Elimination (preventive)", f"remove the hazard or the task, so that nobody has to catch the {focus}")]
    if m in ("M7", "M8"):
        tiers.append(("Substitution (preventive)", f"replace the equipment or material involved in the {focus} with "
                                                   f"a less hazardous one"))
    tiers += [("Engineering (corrective)", f"add a physical or automated control so that the {focus} cannot go "
                                           f"unnoticed"),
              ("Administrative (preventive)", f"change the procedure, permit or planning step that let the {focus} "
                                              f"arise"),
              (f"Engineering plus training {NINE_M_WEAK}, preventive",
               f"pair the engineered control with practice on the {focus}")]
    return [{"text": f"{t}: {body}", "cites": cites} for t, body in tiers]


def _nm_session_unknowns(session) -> list:
    return [r for r in (session or {}).get("records") or []
            if isinstance(r.get("answer"), dict) and r["answer"].get("text") == NINE_M_NONE]


def _nm_scope(w, seed, top):
    prof = seed.get("profile") or {}
    weighted = [m for m in NINE_M_CODES if _nm_ranked(w, seed, m)]
    lead = {m: _nm_ranked(w, seed, m)[:1] for m in weighted}
    four = sorted(weighted, key=lambda m: (-(prof.get(m) or {}).get("top_rated", 0),
                                          -(prof.get(m) or {}).get("findings", 0), m))[:4]
    four = [m for m in NINE_M_CODES if m in four]
    out = []
    if weighted:
        out.append({"text": f"Every M with seed weight, in order M1 to M9: {', '.join(weighted)}",
                    "cites": [lead[m][0]["id"] for m in weighted]})
    if four and four != weighted:
        out.append({"text": f"The top four by top-rated findings: {', '.join(four)}",
                    "cites": [lead[m][0]["id"] for m in four]})
    out.append({"text": "A list you choose (type the Ms, for example M3, M8)", "cites": []})
    if len(out) < NINE_M_TOP_MIN:
        out.append({"text": NINE_M_FREE_TEXT, "cites": []})
    return out


def _nm_code_m(seed, w) -> dict:
    """Each Nanocode's M: the primary M most of its findings carry (ties by M order)."""
    S = seed.get("findings", {})
    tally = {}
    for f in w["findings"]:
        m = (S.get(f["id"]) or {}).get("m_primary")
        if m:
            tally.setdefault(f["code"], {}).setdefault(m, 0)
            tally[f["code"]][m] += 1
    return {c: sorted(t.items(), key=lambda kv: (-kv[1], kv[0]))[0][0] for c, t in tally.items()}


def _nm_cross_m(w, seed, session, top):
    rel = [m for m in NINE_M_CODES if m_status(session or {}, m) == "relevant"]
    cm = _nm_code_m(seed, w)
    ids = {f["id"] for f in w["findings"]}
    chap = _nm_chapters()
    pairs = {}
    for e in (w.get("causal_graph") or {}).get("edges") or []:
        a, b = cm.get(e.get("from")), cm.get(e.get("to"))
        if a and b and a != b and a in rel and b in rel:
            k = (a, b)
            best = pairs.get(k)
            if best is None or (e.get("mentions") or 0) > (best.get("mentions") or 0):
                pairs[k] = e
    out = []
    for (a, b), e in sorted(pairs.items(), key=lambda kv: (-(kv[1].get("mentions") or 0), kv[0])):
        cites = [x for x in e.get("findings") or [] if x in ids][:3]
        if cites:
            out.append({"text": f"Keep {a} {chap[a]['title']} and {b} {chap[b]['title']} as one learning "
                                f"({e['from']} feeds {e['to']})", "cites": cites})
    out = out[:NINE_M_TOP_MAX - 1]
    if not out:
        out.append({"text": "No causal link joins the relevant Ms: keep them separate", "cites": []})
    out.append({"text": "Keep every M as a separate learning", "cites": []})
    return out


def _nm_gaps(session, top):
    unknown = _nm_session_unknowns(session)
    if unknown:
        what = ", ".join(f"{r['q']} ({r.get('m') or 'global'} {r['step']})" for r in unknown)
        return [{"text": f"Ask for evidence on each unknown: {what}", "cites": []},
                {"text": "Record the unknowns as gaps, with no evidence request now", "cites": []}]
    return [{"text": "No answer was unknown: no evidence request is needed", "cites": []},
            {"text": "Ask for evidence anyway, on the conditions that rest on a single source", "cites": []}]


_NM_STEP_FN = {"relevance": _nm_relevance, "subcategory": _nm_subcategory, "condition": _nm_condition,
               "driver": _nm_driver, "barrier": _nm_barrier, "change": _nm_change}


def suggest(w, seed, m, step, subcategory=None, top=4, session=None) -> list:
    """The suggestions for one question, as [{text, cites}], recommended first, always ending with None / unknown.

    Deterministic: candidates are taken in rating order and every fact-bearing suggestion cites at least one finding.
    m is M1..M9 for the per-M steps (NINE_M_STEPS) and None for the global ones (scope, cross-m, gaps), which read
    the session. subcategory (an id such as M8.02) steers the condition, driver, barrier and change candidates."""
    top = max(NINE_M_TOP_MIN, min(NINE_M_TOP_MAX, int(top or NINE_M_TOP_MAX)))
    if step in NINE_M_GLOBAL_STEPS:
        cands = (_nm_scope(w, seed, top) if step == "scope" else
                 _nm_cross_m(w, seed, session, top) if step == "cross-m" else _nm_gaps(session, top))
    elif step in NINE_M_STEPS:
        if m not in NINE_M_CODES:
            raise ValueError(f"--m must be one of M1-M9, not {m!r}")
        if subcategory and subcategory not in {s["id"] for s in _nm_chapters()[m]["subcategories"]}:
            raise ValueError(f"{subcategory} is not a subcategory of {m} (see references/nine-m.json)")
        fs = _nm_focus(w, seed, m, subcategory)
        cands = _NM_STEP_FN[step](w, seed, m, fs, subcategory, top)
        if step != "subcategory":                      # the framework's own names are kept verbatim (D9)
            cands = [c for c in cands if not NINE_M_NEVER_RX.search(c["text"])]
    else:
        raise ValueError(f"unknown step {step!r}; one of {', '.join(NINE_M_STEPS + NINE_M_GLOBAL_STEPS)}")
    out, seen = [], set()
    for c in cands:
        if c["text"] not in seen and c["text"] != NINE_M_NONE:
            seen.add(c["text"])
            out.append({"text": c["text"], "cites": list(c["cites"])})
    out = out[:top]
    for pad in (NINE_M_FREE_TEXT, NINE_M_SKIP if m else None):
        if pad and len(out) < NINE_M_TOP_MIN and pad not in seen:
            out.append({"text": pad, "cites": []})
    return out + [{"text": NINE_M_NONE, "cites": []}]


# ----------------------------------------------------------------------------- the session

def m_status(session, m) -> str:
    """relevant, not-relevant, unknown (answered None / unknown) or open (no answered relevance question)."""
    recs = [r for r in (session or {}).get("records") or []
            if r.get("m") == m and r.get("step") == "relevance" and isinstance(r.get("answer"), dict)]
    if not recs:
        return "open"
    text = _s(recs[-1]["answer"].get("text"))
    if text == NINE_M_NONE:
        return "unknown"
    return "not-relevant" if re.match(r"(?i)\s*(?:not relevant|no\b)", text) else "relevant"


def session_scope(session) -> list:
    """The Ms in scope, from the answered scope question (all nine when it is unknown or names none)."""
    rec = next((r for r in (session or {}).get("records") or [] if r.get("step") == "scope"
                and isinstance(r.get("answer"), dict)), None)
    if rec is None:
        return [m for m in NINE_M_CODES if any(r.get("m") == m for r in (session or {}).get("records") or [])]
    got = set(re.findall(r"\bM[1-9]\b", _s(rec["answer"].get("text")).upper()))
    return [m for m in NINE_M_CODES if m in got] or list(NINE_M_CODES)


def session_subcategories(session) -> dict:
    """{M: subcategory id} from each M's last answered subcategory question (an id of that M in the answer text)."""
    ids, out = subcategory_ids(), {}
    for r in (session or {}).get("records") or []:
        if r.get("step") == "subcategory" and isinstance(r.get("answer"), dict) and r.get("m"):
            hit = _NM_SUB_RX.search(_s(r["answer"].get("text")))
            sid = f"M{hit.group(1)}.{hit.group(2)}" if hit else None
            if sid in ids and sid.startswith(r["m"] + "."):
                out[r["m"]] = sid
    return out


def _nm_record(session, m, step, subcategory, suggestions) -> dict:
    n = len(session["records"])
    used = sum(1 for r in session["records"] if r.get("m") == m)
    budget = NINE_M_BUDGET_PER_M if m else NINE_M_BUDGET_GLOBAL
    if used >= budget:
        _nm_fail(f"The question budget for {m or 'the global questions'} is used up ({budget} questions; spec §7.4). "
                 f"Write what the session holds, or record the rest as gaps.")
    rec = {"q": f"Q{n + 1:02d}", "m": m, "step": step, "subcategory": subcategory, "suggestions": suggestions,
           "answer": None}
    session["records"].append(rec)
    return rec


def _nm_print(rec) -> None:
    head = f"{rec['q']} {rec['m'] or 'global'} {rec['step']}" + (f" ({rec['subcategory']})" if rec["subcategory"] else "")
    print(head)
    for i, s in enumerate(rec["suggestions"], 1):
        print(f"{i}. {s['text']}" + (f" [{', '.join(s['cites'])}]" if s["cites"] else ""))


def cmd_nine_m_suggest(a):
    if a.step in NINE_M_STEPS and not a.m:
        _nm_fail(f"--step {a.step} asks about one M: give --m M1..M9")
    w = load_session_ws(a.ws)
    session = read_session(a.ws)
    m = a.m if a.step in NINE_M_STEPS else None
    try:
        sug = suggest(w, w["nine_m_seed"], m, a.step, a.subcategory, a.top, session=session)
    except ValueError as e:
        _nm_fail(str(e))
    if getattr(a, "json", False):
        print(json.dumps(sug, indent=1, ensure_ascii=False))
    if not a.record:
        if not getattr(a, "json", False):
            _nm_print({"q": "(not recorded)", "m": m, "step": a.step, "subcategory": a.subcategory, "suggestions": sug})
        return
    if session is None:
        session = {"version": 1, "mode": "interview", "records": []}
    elif session.get("mode") == "replay":
        _nm_fail(f"{_abs(a.ws) / NINE_M_SESSION_FILE} is a replay (answers chosen for demonstration); move it aside "
                 f"to start an interview")
    rec = _nm_record(session, m, a.step, a.subcategory, sug)
    write_session(a.ws, session)
    if not getattr(a, "json", False):
        _nm_print(rec)
        print(f"Recorded as {rec['q']} in {_abs(a.ws) / NINE_M_SESSION_FILE}; record the answer with: "
              f"nine-m-answer <ws> --q {rec['q']} --choice N (or --text \"...\")")


def cmd_nine_m_answer(a):
    session = read_session(a.ws)
    if session is None:
        _nm_fail(f"{_abs(a.ws)} has no {NINE_M_SESSION_FILE} yet: ask a question with nine-m-suggest ... --record first")
    rec = next((r for r in session["records"] if r.get("q") == a.q), None)
    if rec is None:
        _nm_fail(f"{a.q} is not a question in {NINE_M_SESSION_FILE} "
                 f"(it holds {', '.join(r.get('q', '?') for r in session['records']) or 'none'})")
    if a.choice is not None:
        n = len(rec["suggestions"])
        if not 1 <= a.choice <= n:
            _nm_fail(f"--choice must be 1 to {n} for {a.q}")
        rec["answer"] = {"choice": a.choice, "text": rec["suggestions"][a.choice - 1]["text"], "source": "export"}
    else:
        rec["answer"] = {"choice": None, "text": a.text, "source": "user"}       # verbatim, owners and dates too
        if NINE_M_DATE_RX.search(a.text or ""):
            print(NINE_M_D7_NOTE, file=sys.stderr)
    write_session(a.ws, session)
    print(f"{a.q} answered: {rec['answer']['text']}")


def replay(w, ms=None) -> dict:
    """A whole session taking suggestion 1 at every question (OD5-1): scope, then per M in scope relevance and, for
    a relevant M, subcategory, condition, driver, barrier and change; then cross-m and gaps."""
    seed = w["nine_m_seed"]
    s = {"version": 1, "mode": "replay", "records": []}

    def ask(m, step, sub=None, choice=1, text=None):
        rec = _nm_record(s, m, step, sub, suggest(w, seed, m, step, sub, NINE_M_TOP_MAX, session=s))
        if text is None:
            rec["answer"] = {"choice": choice, "text": rec["suggestions"][choice - 1]["text"], "source": "export"}
        else:
            rec["answer"] = {"choice": None, "text": text, "source": "user"}
        return rec

    if ms:
        ask(None, "scope", text=", ".join(ms))
    else:
        ask(None, "scope")
    for m in session_scope(s):
        ask(m, "relevance")
        if m_status(s, m) != "relevant":
            continue
        ask(m, "subcategory")
        sub = session_subcategories(s).get(m)
        for step in ("condition", "driver", "barrier", "change"):
            ask(m, step, sub)
    if sum(1 for m in NINE_M_CODES if m_status(s, m) == "relevant") >= 2:
        ask(None, "cross-m")
    ask(None, "gaps")
    return s


def cmd_nine_m_replay(a):
    if not a.accept_recommended:
        _nm_fail("nine-m-replay writes a scripted session only with --accept-recommended (it takes suggestion 1 at "
                 "every question); for a real interview use nine-m-suggest --record and nine-m-answer")
    ms = None
    if a.ms:
        ms = [x.strip().upper() for x in a.ms.split(",") if x.strip()]
        bad = [x for x in ms if x not in NINE_M_CODES]
        if bad:
            _nm_fail(f"--ms takes M1-M9, not {', '.join(bad)}")
        ms = [m for m in NINE_M_CODES if m in ms]
    w = load_session_ws(a.ws)
    old = read_session(a.ws)
    if old is not None and old.get("mode") != "replay" and old.get("records"):
        _nm_fail(f"{_abs(a.ws) / NINE_M_SESSION_FILE} holds an interview ({len(old['records'])} question(s)); "
                 f"move it aside before a replay")
    s = replay(w, ms)
    write_session(a.ws, s)
    rel = [m for m in NINE_M_CODES if m_status(s, m) == "relevant"]
    nrel = [m for m in session_scope(s) if m not in rel]
    print(f"Replay (answers chosen for demonstration): {len(s['records'])} question(s); relevant: "
          f"{', '.join(rel) or 'none'}; not relevant: {', '.join(nrel) or 'none'}")
    print(f"Wrote {_abs(a.ws) / NINE_M_SESSION_FILE}")


# ----------------------------------------------------------------------------- side output and template

def _nm_section_m(sid) -> str | None:
    hit = re.fullmatch(r"m([1-9])", _s(sid))
    return f"M{hit.group(1)}" if hit else None


def nine_m_map(pack, session) -> dict:
    """outputs/nine-m-map.json from a nine-m-analysis pack whose items carry ids (with_item_ids): each group in a
    section m<n> is a condition of that M (an ungrouped learning is a condition of its own); a finding takes the M,
    subcategory and condition of the first condition that cites it, and the Ms of later ones as secondary."""
    subs = session_subcategories(session)
    conditions, findings = {}, {}
    for s in pack.get("sections") or []:
        m = _nm_section_m(s.get("id")) if isinstance(s, dict) else None
        if not m:
            continue
        for it in s.get("items") or []:
            if not isinstance(it, dict):
                continue
            cid = _s(it.get("group")) or (_s(it.get("id")) if it.get("kind") == "learning" else "")
            if not cid:
                continue
            c = conditions.setdefault(cid, {"m": m, "subcategory": subs.get(m), "learning_id": None,
                                            "recommendation_ids": []})
            if it.get("kind") == "learning" and c["learning_id"] is None:
                c["learning_id"] = _s(it.get("learning")) or _s(it.get("id")) or None
            if it.get("kind") == "recommendation" and _s(it.get("id")):
                c["recommendation_ids"].append(_s(it["id"]))
            for fid in it.get("cites") or []:
                if not (isinstance(fid, str) and fid.startswith("FND-")):
                    continue
                f = findings.get(fid)
                if f is None:
                    findings[fid] = {"m_primary": c["m"], "m_secondary": [], "subcategory": c["subcategory"],
                                     "condition_id": cid}
                elif c["m"] != f["m_primary"] and c["m"] not in f["m_secondary"]:
                    f["m_secondary"] = [x for x in NINE_M_CODES if x in f["m_secondary"] + [c["m"]]]
    return {"findings": findings, "conditions": conditions}


def _nm_corrective(text) -> str:
    hit = re.search(r"\b(corrective|preventive)\b", _s(text), re.I)
    return hit.group(1).capitalize() if hit else ""


def template_rows(pack) -> list:
    """One row per recommendation in a section m<n>, in TEMPLATE_COLUMNS order; Owner, Due date, Status and Notes are
    left blank for the organisation (D7)."""
    rows = []
    for s in pack.get("sections") or []:
        m = _nm_section_m(s.get("id")) if isinstance(s, dict) else None
        for it in (s.get("items") or []) if m else []:
            if isinstance(it, dict) and it.get("kind") == "recommendation":
                rows.append([_s(it.get("id")), m, _s(it.get("change")), _s(it.get("hoc_tier")),
                             _nm_corrective(it.get("change")), _s(it.get("indicator")), "", "", "", ""])
    return rows


def cmd_nine_m_actions_template(a):
    from .actions import _csv_cell
    from .handling import (effective_handling, is_draft, name_guard, refuse_internal_overwrite, share_ctx,
                           share_gate, shared_copy, SHARED_LINE, shared_name)
    from .pack.profiles import NO_OWNER_DUE
    from .pack.schema import _checked, record_check, with_item_ids
    from .render.exports import refuse_existing
    from .util import _need_openpyxl
    from .workspace import outputs_dir
    from . import BRAND, DRAFT_LABEL, STATUS_PHRASE
    w, pack, res = _checked(a, record_ok=False)
    owned = [_s(it.get("id")) or "a recommendation" for s in pack["sections"] for it in s["items"]
             if isinstance(it, dict) and (_s(it.get("owner")) or _s(it.get("by_when")))]
    if owned:
        _nm_fail(f"{', '.join(owned)}: {NO_OWNER_DUE}; remove owner and by_when, then run again", 1)
    mode = share_gate(pack, w, a, embeds=False, src=a.pack, ws=a.ws)
    record_check(a, pack, res, w)
    eh = effective_handling(pack, w, src=a.pack, ws=a.ws)
    share = share_ctx(a, w, mode) if mode != "internal" else None
    rpack = with_item_ids(shared_copy(pack, share, w) if share else pack)
    rows = template_rows(rpack)
    if not rows:
        _nm_fail(f"{a.pack} has no recommendation in a section m1 to m9; nothing to put in the template", 1)
    name_guard([c for r in rows for c in r], a)
    internal = outputs_dir(_abs(a.ws)) / f"{Path(a.pack).stem}.actions-template.{a.format}"
    out = Path(a.out) if a.out else shared_name(internal, mode)
    refuse_internal_overwrite(out, internal, mode)
    refuse_existing(out)                                   # the organisation fills it in: never overwritten
    out.parent.mkdir(parents=True, exist_ok=True)
    if a.format == "csv":
        buf = io.StringIO(newline="")
        cw = csv.writer(buf, lineterminator="\r\n")
        cw.writerow(TEMPLATE_COLUMNS)
        for r in rows:
            cw.writerow([_csv_cell(c) for c in r])
        out.write_bytes(buf.getvalue().encode("utf-8"))
    else:
        openpyxl = _need_openpyxl()
        from openpyxl.packaging.core import DocumentProperties
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from .register.xlsx import reset_docprops, SHARED_EPOCH, text_cells
        wb = openpyxl.Workbook()
        sh = wb.active
        sh.title = "9M actions"
        sh.append(TEMPLATE_COLUMNS)
        for r in rows:
            sh.append([c or None for c in r])
        fill, font = PatternFill("solid", fgColor="0B3954"), Font(color="FFFFFF", bold=True)
        for i, c in enumerate(sh[1], 1):
            c.fill, c.font = fill, font
            sh.column_dimensions[get_column_letter(i)].width = 40 if TEMPLATE_COLUMNS[i - 1] in ("Change", "Indicator") else 18
        for row in sh.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        sh.freeze_panes = "A2"
        sh.oddHeader.center.text = (STATUS_PHRASE + (f" · {DRAFT_LABEL}" if is_draft(eh) else "")
                                    + (f" · {SHARED_LINE.format(mode=mode)}" if share else ""))
        sh.oddFooter.center.text = BRAND
        if share:
            wb.properties = DocumentProperties(creator="", lastModifiedBy=None, created=SHARED_EPOCH,
                                               modified=SHARED_EPOCH)
        text_cells(wb)
        wb.save(out)
        if share:
            reset_docprops(out)
    if share:
        print(SHARED_LINE.format(mode=mode))
    print(f"9M actions template with {len(rows)} recommendation(s) written to {out}; Owner, Due date, Status and "
          f"Notes are blank for your organisation to fill in")
