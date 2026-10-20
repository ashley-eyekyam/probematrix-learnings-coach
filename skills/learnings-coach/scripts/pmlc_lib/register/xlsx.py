"""The register-xlsx command."""
from __future__ import annotations

import collections
import datetime
import json
import re
import sys
import zipfile
from pathlib import Path

from .. import BRAND, DRAFT_LABEL, STATUS_PHRASE, VERSION_STAMP
from ..classify import TIER_LABELS
from ..handling import (
    effective_handling, is_draft, REASON_NO_WORKSPACE, review_text, share_ctx, share_gate, shared_copy, SHARED_LINE,
    refuse_internal_overwrite, shared_name, SHARED_SOURCE
)
from ..register.check import HOC_TIERS, norm_entity, norm_hoc, norm_level
from ..util import _flat, _need_openpyxl
from ..workspace import is_standard_ws, load_ws
from .. import privacy


# ----------------------------------------------------------------------------- register-xlsx

REG_COLS = [
    ("lr_id", "LR-ID"), ("ml_id", "ML-ID"), ("title", "Title"), ("owner_entity", "Learning cluster (entity)"),
    ("tiers", "Tiers who must learn"), ("lesson", "Lesson statement"), ("context", "Context narrative"),
    ("finding_ids", "Finding IDs"), ("evidence_ids", "Evidence IDs"), ("evidence_strength", "Evidence strength"),
    ("confidence", "Confidence"), ("hfacs_level", "HFACS level / sub-category"), ("causal_role", "Causal role"),
    ("barrier", "Barrier (prevention / mitigation)"), ("hazard", "Hazard / energy"), ("activity", "Activity / phase"),
    ("performance_mode", "Human-performance mode"), ("org_system", "Organisational system"),
    ("pattern", "Cross-incident / recurrence pattern"), ("precursors", "Precursors and weak signals"),
    ("comparison", "Counterpart in comparator report"), ("reason_class", "Reason class (if none)"),
    ("worst_case", "Credible worst case"), ("transferability", "Transferability"), ("priority", "Priority"),
    ("audiences", "Audiences"), ("learning_level", "Learning level / loop"), ("formats", "Formats"),
]
# Phase 3 fields: a column appears only when some entry carries the field, in this order, after the last column above.
OPTIONAL_REG_COLS = [("urgency", "Urgency"), ("recommendation_class", "Recommendation class"),
                     ("regulatory_touchpoint", "Regulatory touchpoint")]
PRIORITIES = ("1", "2", "3")
COMPARISONS = ("yes", "partial", "none", "not compared")
SHARED_EPOCH = datetime.datetime(2000, 1, 1)
REC_COLS = [("lr_id", "LR-ID"), ("what", "What should change"), ("level_barrier", "HFACS level and barrier"),
            ("evidence_basis", "Evidence basis"), ("hoc_tier", "Hierarchy-of-controls tier"),
            ("leading_indicator", "Leading indicator")]


def refuse_pack_as_register(doc, path) -> None:
    """M9: a pack JSON given where a register goes exits 1 with one line, not a traceback."""
    if isinstance(doc, dict) and "pack" in doc and "entries" not in doc:
        sys.exit(f"{path} is a pack, not a register")


FORMULA_LEAD = ("=", "+", "-", "@")


def text_cells(wb) -> None:
    """M6: every cell string that starts with = + - or @ is stored as text, so a spreadsheet never runs it."""
    for sh in wb.worksheets:
        for row in sh.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith(FORMULA_LEAD):
                    c.data_type = "s"


def summary_groups(entries, recs) -> tuple:
    """The Summary counts shared by the xlsx Summary sheet and register-md: ([(label, Counter)], recommendations above
    the administrative tier)."""
    groups = [
        ("Learnings by cluster", collections.Counter(norm_entity(e.get("owner_entity")) for e in entries)),
        ("Learnings by HFACS level", collections.Counter(norm_level(e.get("hfacs_level")) for e in entries)),
        ("Learnings by priority", collections.Counter(_flat(e.get("priority")) or "unset" for e in entries)),
        ("Learnings by counterpart", collections.Counter(_flat(e.get("comparison")) or "not compared" for e in entries)),
        ("Learnings by tier that must act", collections.Counter(t for e in entries for t in (e.get("tier_messages") or {}))),
        ("Recommendations by control tier", collections.Counter(norm_hoc(r.get("hoc_tier")) for r in recs)),
    ]
    strong = sum(1 for r in recs if norm_hoc(r.get("hoc_tier")) in ("elimination", "substitution", "engineering"))
    return groups, strong


def _register_ws(a) -> tuple:
    """(workspace, its folder) for a register's forced-draft and sign-off checks: --ws, else the run's workspace when
    the register sits in <run>/outputs/ of the standard layout, else ({}, None): nothing can be checked, so the
    register stays draft with the reason 'no workspace to check'."""
    if getattr(a, "ws", None):
        return load_ws(a.ws), Path(a.ws)
    reg_dir = Path(a.register).resolve().parent
    guess = reg_dir.parent / "workspace"
    if reg_dir.name == "outputs" and is_standard_ws(guess) and (guess / "findings.json").exists():
        return load_ws(guess), guess
    return {}, None


def reset_docprops(path) -> None:
    """A shared workbook carries no author or timestamps. openpyxl stamps `modified` on save, so the saved file's
    core properties are rewritten: creator and last-modified-by blank, both dates the fixed epoch."""
    stamp = SHARED_EPOCH.strftime("%Y-%m-%dT%H:%M:%SZ")
    tmp = Path(str(path) + ".tmp")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "docProps/core.xml":
                t = data.decode("utf-8")
                t = re.sub(r"(<dcterms:(?:created|modified)\b[^>]*>)[^<]*", lambda m: m.group(1) + stamp, t)
                t = re.sub(r"<dc:creator>[^<]*</dc:creator>", "<dc:creator></dc:creator>", t)
                t = re.sub(r"<cp:lastModifiedBy>[^<]*</cp:lastModifiedBy>", "", t)
                data = t.encode("utf-8")
            zout.writestr(item, data)
    tmp.replace(path)


def cmd_register_xlsx(a):
    openpyxl = _need_openpyxl()
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.packaging.core import DocumentProperties
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.hyperlink import Hyperlink
    from openpyxl.worksheet.properties import PageSetupProperties
    from openpyxl.worksheet.table import Table, TableStyleInfo
    reg = json.loads(Path(a.register).read_text(encoding="utf-8"))
    refuse_pack_as_register(reg, a.register)
    rw, rws = _register_ws(a)
    hdoc = reg if isinstance(reg, dict) else {}
    mode = share_gate(hdoc, rw, a, embeds=False, extra=[] if rw else ["no workspace (pass --ws)"],
                      src=a.register, ws=rws)
    eh = effective_handling(hdoc, rw, src=a.register, ws=rws)
    if rws is None and REASON_NO_WORKSPACE not in eh["draft_reasons"]:
        eh["draft_reasons"].append(REASON_NO_WORKSPACE)
    share = share_ctx(a, rw, mode) if mode != "internal" else None
    if share:
        reg = shared_copy(reg, share, rw)
        if isinstance(reg, dict) and reg.get("source"):
            reg["source"] = SHARED_SOURCE
    entries = reg["entries"] if isinstance(reg, dict) else reg
    wb = openpyxl.Workbook()
    head_fill = PatternFill("solid", fgColor="0B3954")
    head_font = Font(color="FFFFFF", bold=True)

    def sheet(ws, cols, rows, table=None, lists=()):
        """One data sheet: header, rows, then (data sheets) an Excel Table, list validations {field: values}."""
        ws.append([c[1] for c in cols])
        for r in rows:
            ws.append([_flat(r.get(c[0])) for c in cols])
        for i, c in enumerate(cols, 1):
            ws.cell(1, i).fill = head_fill
            ws.cell(1, i).font = head_font
            ws.column_dimensions[get_column_letter(i)].width = 18 if len(c[1]) < 14 else 32
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "C2"
        if table:
            if ws.max_row > 1:      # a Table needs a data row; unique string headers are the columns' names
                t = Table(displayName=table, ref=ws.dimensions)
                t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
                ws.add_table(t)
            ws.print_title_rows = "1:1"
        last = max(ws.max_row, 2)
        for field, values in lists:
            i = next((n for n, c in enumerate(cols, 1) if c[0] == field), None)
            if i:
                dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
                dv.add(f"{get_column_letter(i)}2:{get_column_letter(i)}{last}")
                ws.add_data_validation(dv)

    ws = wb.active
    ws.title = "Learnings Register"
    reg_cols = REG_COLS + [c for c in OPTIONAL_REG_COLS if any(_flat(e.get(c[0])) for e in entries)]
    sheet(ws, reg_cols, entries, "LearningsRegister", [("priority", PRIORITIES), ("comparison", COMPARISONS)])
    lr_row = {e.get("lr_id"): n for n, e in enumerate(entries, 2)}
    recs = [dict(r, lr_id=e.get("lr_id")) for e in entries for r in (e.get("recommendations") or [])]
    rws = wb.create_sheet("Recommendations")
    sheet(rws, REC_COLS, recs, "Recommendations", [("hoc_tier", HOC_TIERS)])
    for n, r in enumerate(recs, 2):
        if r.get("lr_id") in lr_row:
            cell = rws.cell(n, 1)
            cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'Learnings Register'!A{lr_row[r['lr_id']]}")
            cell.font = Font(color="0563C1", underline="single")
    plan_cols = [("lr_id", "LR-ID"), ("title", "Title"), ("priority", "Priority")] + \
                [(t, f"{t} {TIER_LABELS[t]}") for t in TIER_LABELS]
    plan_rows = [dict({"lr_id": e.get("lr_id"), "title": e.get("title"), "priority": e.get("priority")},
                      **(e.get("tier_messages") or {})) for e in entries]
    sheet(wb.create_sheet("Learning Plan"), plan_cols, plan_rows, "LearningPlan")
    s = wb.create_sheet("Summary")
    s.append(["Count by", "Value", "Count"])
    for c in s[1]:
        c.fill, c.font = head_fill, head_font
    groups, strong = summary_groups(entries, recs)
    for lab, cnt in groups:
        for v, c in sorted(cnt.items(), key=lambda kv: (str(kv[0]))):
            s.append([lab, v, c])
    s.append(["Recommendations above administrative tier", f"{strong} of {len(recs)}",
              f"{100*strong/max(len(recs),1):.0f}%"])
    s.column_dimensions["A"].width, s.column_dimensions["B"].width, s.column_dimensions["C"].width = 38, 46, 10
    ab = wb.create_sheet("About")
    for line in [reg.get("title", "Learnings Register") if isinstance(reg, dict) else "Learnings Register",
                 reg.get("source", "") if isinstance(reg, dict) else "",
                 "Lessons and recommendations only: no owners, due dates or action tracking.",
                 STATUS_PHRASE, BRAND]:
        ab.append([line])
    review = review_text(eh) + (f". Draft because: {'; '.join(eh['draft_reasons'])}" if eh["draft_reasons"] else "")
    rows = [("Review status", review), ("Classification", eh["classification"]),
            ("Legal review", eh["legal_review"]), ("Distribution", ", ".join(eh["distribution"]) or "not set"),
            ("Version", str(eh["version"])), ("Retention review", eh.get("retention_review") or "not set")]
    if share:
        rows = [(k, privacy.share_text(v, share)) for k, v in rows] + [("Sharing", SHARED_LINE.format(mode=mode))]
    for k, v in rows:
        ab.append([k, v])
    ab.column_dimensions["A"].width = min(200, max(120, max(len(str(c.value)) for c in ab["A"] if c.value) + 2))
    ab.column_dimensions["B"].width = 60
    for sh in wb.worksheets:
        sh.oddHeader.center.text = STATUS_PHRASE + (f" · {DRAFT_LABEL}" if is_draft(eh) else "")
        sh.oddFooter.center.text = VERSION_STAMP
        sh.page_setup.orientation = "landscape"
        sh.page_setup.fitToWidth, sh.page_setup.fitToHeight = 1, 0
        sh.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    if share:
        wb.properties = DocumentProperties(creator="", lastModifiedBy=None, created=SHARED_EPOCH,
                                           modified=SHARED_EPOCH)
    internal = Path(a.register).resolve().parent / "learnings-register.xlsx"
    out = Path(a.out) if a.out else shared_name(internal, mode)
    refuse_internal_overwrite(out, internal, mode)
    out.parent.mkdir(parents=True, exist_ok=True)
    text_cells(wb)
    wb.save(out)
    if share:
        reset_docprops(out)
    print(f"Register with {len(entries)} entries and {len(recs)} recommendations written to {out}")
