"""Could it happen here? (spec §6.2, P6-5, P6-9): the site return template that `templates <ws> cihh` writes from a
pack's self_check items, the reader for a filled return, and `consolidate`, which turns the returns into the
corporate view (coverage and conditions, never a ranking of sites) and site-assessment events for the library.

D7: the engine writes no site name, answer, response or assessor; the sites fill them in. The template's Pack cell
holds the pack's content hash, so a return made from another pack's template is refused."""
from __future__ import annotations

import hashlib
import json
import re
import sys
import zipfile
from pathlib import Path

from .templates import OrgFileError

CIHH_SHEET = "Return"
CIHH_COLUMNS = ["lr_id", "condition", "evidence_to_check", "answer", "site_response", "assessed_by_role"]
CIHH_HEADER_ROW = 4                     # A1 Site, A2 Pack, A3 the banner; the columns from row 4, data from row 5
CIHH_ANSWERS = ("present", "partly", "absent", "n/a")
CIHH_NA_FORMS = ("na", "n.a.", "n.a", "n/a")
CIHH_NO_ITEMS = "the pack has no self_check items"
CIHH_CONSOLIDATED = "cihh-consolidated.json"
CIHH_EVENTS = "cihh-events.jsonl"
CIHH_NOT_RETURN = ("{file}: not a could-it-happen-here return ({why}); start from the template that "
                   "`templates <ws> cihh` wrote")
CIHH_VERSION = 1


# ----------------------------------------------------------------------------- the pack side

def _cihh_items(pack) -> list:
    """The pack's self_check items, in order, each with an id (with_item_ids)."""
    from .pack.schema import with_item_ids
    p = with_item_ids(pack)
    return [it for s in p.get("sections") or [] for it in s.get("items") or []
            if isinstance(it, dict) and it.get("kind") == "self_check"]


def _cihh_key(it) -> str:
    """A template row's key: the item's lr_id, else its item id."""
    return str(it.get("lr_id") or it.get("id") or "")


def cihh_rows(pack) -> list:
    """One row per self_check item, in CIHH_COLUMNS order; the last three cells blank (the site fills them in)."""
    return [[_cihh_key(it), str(it.get("condition") or ""), str(it.get("evidence_to_check") or ""), "", "", ""]
            for it in _cihh_items(pack)]


def _cihh_status(eh, share_mode) -> str:
    from . import DRAFT_LABEL, STATUS_PHRASE
    from .handling import is_draft, SHARED_LINE
    return (STATUS_PHRASE + (f" · {DRAFT_LABEL}" if eh and is_draft(eh) else "")
            + (f" · {SHARED_LINE.format(mode=share_mode)}" if share_mode not in (None, "internal") else ""))


def write_template(pack, w, out, share, eh=None) -> Path:
    """Write the return template for `pack` to `out` and return the path. share: a privacy.ShareCtx for a masked
    copy, or None. The Pack cell holds the content hash of `pack` itself (the internal pack), so returns made from a
    shared copy still consolidate against it. Never overwrites `out`."""
    from openpyxl.packaging.core import DocumentProperties
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    from . import BRAND
    from .handling import content_hash, effective_handling, shared_copy
    from .pack.profiles import BANNERS
    from .register.xlsx import reset_docprops, SHARED_EPOCH, text_cells
    from .render.exports import refuse_existing
    from .util import _need_openpyxl
    out = Path(out)
    refuse_existing(out)                                   # the sites fill it in: never overwritten
    openpyxl = _need_openpyxl()
    rows = cihh_rows(shared_copy(pack, share, w) if share else pack)
    wb = openpyxl.Workbook()
    sh = wb.active
    sh.title = CIHH_SHEET
    sh["A1"], sh["A2"], sh["B2"] = "Site", "Pack", content_hash(pack)
    sh["A3"] = BANNERS["cihh"]
    sh["A3"].font = Font(italic=True)
    for c, name in enumerate(CIHH_COLUMNS, 1):
        sh.cell(row=CIHH_HEADER_ROW, column=c, value=name)
    for r in rows:
        sh.append([x or None for x in r])
    fill, font = PatternFill("solid", fgColor="0B3954"), Font(color="FFFFFF", bold=True)
    for c in sh[CIHH_HEADER_ROW]:
        c.fill, c.font = fill, font
    for a in ("A1", "A2"):
        sh[a].font = Font(bold=True)
    for i, name in enumerate(CIHH_COLUMNS, 1):
        sh.column_dimensions[get_column_letter(i)].width = 45 if name in ("condition", "evidence_to_check",
                                                                          "site_response") else 18
    for row in sh.iter_rows(min_row=CIHH_HEADER_ROW + 1):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    if rows:
        dv = DataValidation(type="list", formula1='"' + ",".join(CIHH_ANSWERS) + '"', allow_blank=True)
        dv.add(f"D{CIHH_HEADER_ROW + 1}:D{CIHH_HEADER_ROW + len(rows)}")
        sh.add_data_validation(dv)
    sh.freeze_panes = f"A{CIHH_HEADER_ROW + 1}"
    eh = eh if eh is not None else effective_handling(pack, w)
    sh.oddHeader.center.text = _cihh_status(eh, share.mode if share else None)
    sh.oddFooter.center.text = BRAND
    if share:
        wb.properties = DocumentProperties(creator="", lastModifiedBy=None, created=SHARED_EPOCH,
                                           modified=SHARED_EPOCH)
    text_cells(wb)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    if share:
        reset_docprops(out)
    return out


def cmd_cihh_template(a):
    """templates <ws> cihh --pack P [--share external] [--out F]: the return template, one row per self_check."""
    from .handling import (effective_handling, name_guard, refuse_internal_overwrite, share_ctx, share_gate,
                           shared_copy, SHARED_LINE, shared_name)
    from .pack.schema import _checked, record_check
    from .util import _abs
    from .workspace import outputs_dir
    w, pack, res = _checked(a, record_ok=False)
    mode = share_gate(pack, w, a, embeds=False, src=a.pack, ws=a.ws)
    record_check(a, pack, res, w)
    share = share_ctx(a, w, mode) if mode != "internal" else None
    rows = cihh_rows(shared_copy(pack, share, w) if share else pack)
    name_guard([c for r in rows for c in r], a)
    internal = outputs_dir(_abs(a.ws)) / f"{Path(a.pack).stem}.cihh-template.xlsx"
    out = Path(a.out) if a.out else shared_name(internal, mode)
    refuse_internal_overwrite(out, internal, mode)
    write_template(pack, w, out, share, eh=effective_handling(pack, w, src=a.pack, ws=a.ws))
    if share:
        print(SHARED_LINE.format(mode=mode))
    if not rows:
        print(f"note: {CIHH_NO_ITEMS}; the template has its headings only")
    print(f"could-it-happen-here return template with {len(rows)} self-check(s) written to {out}; each site fills "
          f"in Site (B1), answer ({', '.join(CIHH_ANSWERS)}), site_response and assessed_by_role")


# ----------------------------------------------------------------------------- reading a return

def _cihh_cell(v) -> str:
    if v is None:
        return ""
    return str(v).strip()


def _cihh_answer(v) -> str:
    """A normalised answer: strip().lower(); N/A, na and n.a. all become n/a. Unknown values come back as given
    (lower-cased) for the caller to reject."""
    low = _cihh_cell(v).lower()
    return "n/a" if low in CIHH_NA_FORMS else low


def _cihh_site(v) -> str:
    return re.sub(r"\s+", " ", _cihh_cell(v))


def read_return(path) -> dict:
    """{"site", "pack_hash", "rows": [{lr_id, condition, evidence_to_check, answer, site_response, assessed_by_role,
    _row}]} from a filled return. Raises templates.OrgFileError with one line per problem (naming the file and row);
    its `ret` attribute holds the return with the good rows, or None when nothing in the file can be used."""
    import openpyxl
    from openpyxl.utils.exceptions import InvalidFileException
    path = Path(path)
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError, ValueError, TypeError):
        _cihh_raise([CIHH_NOT_RETURN.format(file=path, why="it is not an Excel workbook")], None)
    if CIHH_SHEET not in wb.sheetnames:
        _cihh_raise([CIHH_NOT_RETURN.format(file=path, why=f"it has no {CIHH_SHEET} sheet")], None)
    sh = wb[CIHH_SHEET]
    head = [_cihh_cell(sh.cell(row=CIHH_HEADER_ROW, column=c).value).lower() for c in range(1, len(CIHH_COLUMNS) + 1)]
    if (_cihh_cell(sh["A1"].value) != "Site" or _cihh_cell(sh["A2"].value) != "Pack" or head != CIHH_COLUMNS):
        _cihh_raise([CIHH_NOT_RETURN.format(file=path, why="the Site and Pack cells or the column headings are "
                                                           "missing")], None)
    site, problems, rows = _cihh_site(sh["B1"].value), [], []
    if not site:
        problems.append(f"{path} row 1: Site is blank; the site writes its name in B1")
    for n, cells in enumerate(sh.iter_rows(min_row=CIHH_HEADER_ROW + 1, max_col=len(CIHH_COLUMNS),
                                           values_only=True), CIHH_HEADER_ROW + 1):
        vals = [_cihh_cell(v) for v in cells] + [""] * (len(CIHH_COLUMNS) - len(cells))
        if not any(vals):
            continue
        row = dict(zip(CIHH_COLUMNS, vals))
        row["answer"], row["_row"] = _cihh_answer(row["answer"]), n
        if row["answer"] and row["answer"] not in CIHH_ANSWERS:
            problems.append(f"{path} row {n}: answer '{vals[3]}' must be one of {', '.join(CIHH_ANSWERS)} (or blank)")
            continue
        rows.append(row)
    ret = {"site": site, "pack_hash": _cihh_cell(sh["B2"].value), "rows": rows}
    if problems:
        _cihh_raise(problems, ret if site else None)
    return ret


def _cihh_raise(problems, ret):
    err = OrgFileError(problems, (ret or {}).get("rows"))
    err.ret = ret
    raise err


# ----------------------------------------------------------------------------- consolidating

def _cihh_match(ret, keys) -> tuple:
    """([(pack row index, return row)], problems) for one return: the n-th row with a key matches the pack's n-th
    template row with that key."""
    slots = {}
    for i, k in enumerate(keys):
        slots.setdefault(k, []).append(i)
    used, out, problems = {}, [], []
    for r in ret["rows"]:
        k = r["lr_id"]
        n = used.get(k, 0)
        if k not in slots:
            problems.append(f"row {r['_row']}: lr_id '{k}' is not a self-check of the pack")
            continue
        if n >= len(slots[k]):
            problems.append(f"row {r['_row']}: lr_id '{k}' appears more often than in the pack")
            continue
        used[k] = n + 1
        out.append((slots[k][n], r))
    return out, problems


def consolidate(returns, pack) -> dict:
    """The corporate view of the returns (P6-9): {version, returns, coverage, conditions, sites}. Coverage counts,
    per learning, the returns that answered its row (assessed) or left it blank or out (not_assessed); conditions
    count present and partly per self-check; sites lists the learnings each site assessed. Nothing ranks or scores
    a site. Rows the pack does not have are ignored here (cmd_consolidate reports them)."""
    items = _cihh_items(pack)
    keys = [_cihh_key(it) for it in items]
    coverage = {k: {"assessed": 0, "not_assessed": 0} for k in keys}
    conditions = [{"lr_id": k, "condition": str(it.get("condition") or ""), "present": 0, "partly": 0}
                  for k, it in zip(keys, items)]
    sites = {}
    for ret in returns:
        answered = {}
        for i, r in _cihh_match(ret, keys)[0]:
            if r["answer"]:
                answered[i] = r["answer"]
        for i, k in enumerate(keys):
            coverage[k]["assessed" if i in answered else "not_assessed"] += 1
            if answered.get(i) in ("present", "partly"):
                conditions[i][answered[i]] += 1
        done = sites.setdefault(ret["site"], [])
        done += [k for i, k in enumerate(keys) if i in answered and k not in done]
    return {"version": CIHH_VERSION, "returns": len(returns), "coverage": coverage, "conditions": conditions,
            "sites": sites}


def event_hash(event: dict) -> str:
    """sha256 over the canonical event (sorted keys, no spaces), without its event_hash."""
    body = {k: v for k, v in event.items() if k != "event_hash"}
    raw = json.dumps(body, sort_keys=True, separators=(",", ":"), ensure_ascii=False)
    return "sha256:" + hashlib.sha256(raw.encode("utf-8")).hexdigest()


def site_events(returns, pack) -> list:
    """One site_assessment event per answered row, ordered by site, then by the pack's row order."""
    keys = [_cihh_key(it) for it in _cihh_items(pack)]
    out = []
    for ret in sorted(returns, key=lambda r: r["site"].casefold()):
        for _, r in sorted(_cihh_match(ret, keys)[0], key=lambda x: x[0]):
            if not r["answer"]:
                continue
            e = {"type": "site_assessment", "lr_id": r["lr_id"], "site": ret["site"], "answer": r["answer"],
                 "site_response": r["site_response"], "source": "organisation"}
            e["event_hash"] = event_hash(e)
            out.append(e)
    return out


def _cihh_vet(paths, pack_hash, keys) -> tuple:
    """(good returns, problems): every file read, checked against the pack and the other returns."""
    good, problems, seen = [], [], {}
    for p in paths:
        try:
            ret = read_return(p)
        except OrgFileError as e:
            problems += e.problems
            ret = getattr(e, "ret", None)
            if ret is None:
                continue
        if ret["pack_hash"] != pack_hash:
            problems.append(f"{p} row 2: made from another pack's template (its Pack cell does not match the pack); "
                            "send the site the current template")
            continue
        fold = ret["site"].casefold()
        if fold in seen:
            problems.append(f"{p} row 1: site '{ret['site']}' already has a return ({seen[fold]})")
            continue
        seen[fold] = p
        matched, bad = _cihh_match(ret, keys)
        problems += [f"{p} {x}" for x in bad]
        ret = dict(ret, rows=[r for _, r in matched])
        good.append(ret)
    return good, problems


def _cihh_code_sites(view, events) -> tuple:
    """A shared copy codes each site as S01, S02... (sorted by name), in the view and the events."""
    names = sorted(view["sites"], key=str.casefold)
    code = {n: f"S{i:02d}" for i, n in enumerate(names, 1)}
    view = dict(view, sites={code[n]: v for n, v in view["sites"].items()})
    return view, [dict(e, site=code[e["site"]]) for e in events]


def _cihh_texts(v):
    if isinstance(v, str):
        yield v
    elif isinstance(v, dict):
        for k, x in v.items():
            yield str(k)
            yield from _cihh_texts(x)
    elif isinstance(v, list):
        for x in v:
            yield from _cihh_texts(x)


def cmd_consolidate(a):
    from .handling import content_hash, name_guard, share_ctx, share_gate, shared_copy, SHARED_LINE, shared_name
    from .pack.profiles import BANNERS
    from .pack.schema import load_pack
    from .util import _abs
    from .workspace import load_ws, outputs_dir
    w = load_ws(a.ws)
    pack, err = load_pack(a.pack)
    if err:
        sys.exit(f"{a.pack}: {err}")
    if not isinstance(pack, dict) or not isinstance(pack.get("sections"), list):
        sys.exit(f"{a.pack}: not a learning pack (no sections)")
    keys = [_cihh_key(it) for it in _cihh_items(pack)]
    if not keys:
        sys.exit(f"{a.pack}: {CIHH_NO_ITEMS}; nothing to consolidate")
    mode = share_gate(pack, w, a, embeds=False, src=a.pack, ws=a.ws)
    good, problems = _cihh_vet(a.returns, content_hash(pack), keys)
    for x in problems:
        print(x, file=sys.stderr)
    if problems and not a.keep_going:
        sys.exit(f"{len(problems)} problem(s) in the returns; nothing written (fix them, or pass --keep-going to "
                 "use the good rows)")
    if not good:
        sys.exit("no usable return; nothing written")
    view, events = consolidate(good, pack), site_events(good, pack)
    if mode != "internal":
        ctx = share_ctx(a, w, mode)
        view, events = _cihh_code_sites(view, events)
        view = shared_copy(view, ctx, w)
        events = [dict(shared_copy(e, ctx, w), site=e["site"]) for e in events]   # the code is no equipment tag
        for e in events:
            e["event_hash"] = event_hash(e)
    name_guard(list(_cihh_texts(view)) + list(_cihh_texts(events)), a)
    out = outputs_dir(_abs(a.ws))
    out.mkdir(parents=True, exist_ok=True)
    vpath, epath = shared_name(out / CIHH_CONSOLIDATED, mode), shared_name(out / CIHH_EVENTS, mode)
    vpath.write_text(json.dumps(view, sort_keys=True, indent=1, ensure_ascii=False) + "\n", encoding="utf-8")
    with open(epath, "w", encoding="utf-8", newline="\n") as fh:
        for e in events:
            fh.write(json.dumps(e, sort_keys=True, ensure_ascii=False) + "\n")
    if mode != "internal":
        print(SHARED_LINE.format(mode=mode))
    if problems:
        print(f"{len(problems)} problem(s) skipped (--keep-going): their rows or returns are left out")
    print(BANNERS["cihh"])
    print(f"consolidated {len(good)} return(s) over {len(keys)} self-check(s): coverage and conditions in {vpath}; "
          f"{len(events)} site assessment event(s) in {epath}")
