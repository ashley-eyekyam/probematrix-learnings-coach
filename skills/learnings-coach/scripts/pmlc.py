#!/usr/bin/env python3
"""pmlc — ProbeMatrix Learnings Coach engine.

Deterministic data work over a ProbeMatrix findings export (.xlsx), so that the
LLM skills spend their effort on judgement, not counting.

Subcommands
  ingest        xlsx -> JSON workspace in <export-dir>/pmlc-runs/<date>-<site>-<event>/workspace
  profile       Markdown profile of the workspace (counts, levels, roles, data-quality flags)
  classify      map 'Individuals Involved' roles to Learning Clusters (entity x tier x relation)
  select        compact Markdown/JSON bundle of findings for one audience or cluster
  causal        causal graph extracted from the AI's causal explanations (+ Causal Chain sheet)
  timeline      dated precursor candidates found in findings and evidence
  interfaces    entity x entity interface matrix (shared findings + cross-entity causal links)
  dashboard     self-contained HTML organisational learning dashboard
  register-xlsx write a Learnings Register JSON to a formatted .xlsx
  scan-names    deterministic zero-hit name scan over files (de-identification check)
  check-pack    check a learning pack JSON against the workspace
  sign-off      record a person's review of a pack or register (writes its handling block and content hash)
  check-register check a Learnings Register JSON against the workspace
  check-language just-culture and grounding lint (JC1-JC12) of a pack or register JSON
  pack          learning pack JSON -> <run>/outputs/<pack>.html and .md
  causal-item   the top causal links as a ready graph item for a pack (JSON on stdout)
  status        which derived files are current or stale, and the last check of each pack
  run-path      print a workspace's run folder, outputs folder and slug
  check-samples check samples/<slug>/ folders follow the contribution layout

Only dependency outside the standard library: openpyxl.
Run `python3 pmlc.py <subcommand> -h` for options.
"""
from __future__ import annotations

import argparse
import collections
import contextlib
import datetime as dt
import hashlib
import html
import io
import json
import os
import re
import sys
import types
from pathlib import Path

__version__ = "2.0.0a3"

HERE = Path(__file__).resolve().parent
if str(HERE) not in sys.path:
    sys.path.insert(0, str(HERE))
import pmlc_privacy as privacy  # noqa: E402  (the shared privacy core, next to this file)

REFS = HERE.parent / "references"
BRAND = (f"Generated with ProbeMatrix Learnings Coach v{__version__} (open source, MIT) from a ProbeMatrix "
         "four-level human-factors investigation. ProbeMatrix by Eyekyam: ashley@eyekyam.com")
# Fixed copy: this module is the only place these strings are written; renderers import them, and the
# dashboard JS reads them from the embedded `brand` payload (brand_payload()).
STATUS_PHRASE = "Learning purposes: not a statutory document"
AI_STATUS = "Learnings drafted by an AI agent from AI-assisted findings"
AI_JUSTIFICATION_LABEL = "AI justification (ProbeMatrix)"
AI_CHIP = "AI-drafted, not human-edited"
EXPORT_BANNER = ("This is the investigation's AI-drafted wording, shown as evidence; "
                 "it is not the coach's learning.")
DRAFT_LABEL = "DRAFT"
VERSION_STAMP = f"ProbeMatrix Learnings Coach v{__version__}"


def brand_payload() -> dict:
    """The fixed strings, for embedding in HTML data blocks as `data.brand`."""
    return {"status_phrase": STATUS_PHRASE, "ai_status": AI_STATUS, "ai_justification_label": AI_JUSTIFICATION_LABEL,
            "ai_chip": AI_CHIP, "export_banner": EXPORT_BANNER, "draft_label": DRAFT_LABEL, "brand": BRAND}


LEVEL_NAMES = {1: "Unsafe Acts", 2: "Preconditions", 3: "Supervision", 4: "Organisational Influences"}
CODE_RX = re.compile(r"\b([A-Z]{2,5})\s?(\d{1,2}\.\d{1,2})\b")


# ----------------------------------------------------------------------------- helpers

def _need_openpyxl():
    try:
        import openpyxl  # noqa: F401
        return openpyxl
    except ImportError:
        sys.exit("openpyxl is required: pip install openpyxl")


def _int(v, default=None):
    """An int from a cell: '436', 436 or the float cell '436.0'; anything else (4.5, nan, text) is `default`."""
    try:
        return int(str(v).strip())
    except (TypeError, ValueError):
        pass
    try:
        f = float(str(v).strip())
    except (TypeError, ValueError):
        return default
    return int(f) if f == f and abs(f) != float("inf") and f == int(f) else default


def _n(n, word) -> str:
    """'1 finding', '2 findings'."""
    return f"{n} {word}" if n == 1 else f"{n} {word}s"


def _s(v) -> str:
    return "" if v is None else str(v).strip()


def split_outside_parens(s: str) -> list[str]:
    """Split a comma-separated role list, ignoring commas inside parentheses."""
    out, buf, depth = [], "", 0
    for ch in s or "":
        if ch == "(":
            depth += 1
        elif ch == ")":
            depth = max(0, depth - 1)
        if ch == "," and depth == 0:
            if buf.strip():
                out.append(buf.strip())
            buf = ""
        else:
            buf += ch
    if buf.strip():
        out.append(buf.strip())
    return out


def norm_code(prefix: str, num: str) -> str:
    return f"{prefix} {num}"


def load_ws(ws_dir: Path) -> dict:
    ws_dir = Path(ws_dir)
    out = {}
    for name in ("findings", "evidence", "incident", "clusters"):
        p = ws_dir / f"{name}.json"
        if p.exists():
            out[name] = json.loads(p.read_text(encoding="utf-8"))
    sp = ws_dir / "singleton_roles.json"
    out["singletons"] = json.loads(sp.read_text(encoding="utf-8")).get("roles", []) if sp.exists() else []
    if "findings" not in out:
        sys.exit(f"No findings.json in {ws_dir}. Run: pmlc.py ingest <export.xlsx> --slug <site>-<event>")
    return out


def write_json(path: Path, obj) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_text(json.dumps(obj, indent=1, ensure_ascii=False), encoding="utf-8")
    os.replace(tmp, path)


# ----------------------------------------------------------------------------- manifest

MANIFEST = "manifest.json"
KNOWN_DERIVED = ("findings.json", "evidence.json", "incident.json", "clusters.json", "singleton_roles.json",
                 "causal_graph.json",
                 "causal_graph.mmd", "timeline_candidates.json", "timeline_event_day.json", "interfaces.json")
REBUILD = {"findings.json": "ingest", "evidence.json": "ingest", "incident.json": "ingest",
           "clusters.json": "classify", "singleton_roles.json": "classify", "causal_graph.json": "causal", "causal_graph.mmd": "causal",
           "timeline_candidates.json": "timeline", "timeline_event_day.json": "timeline",
           "interfaces.json": "interfaces"}


def file_sha256(path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


def _now() -> str:
    return dt.datetime.now().isoformat(timespec="seconds")


def load_manifest(ws) -> dict:
    try:
        m = json.loads((Path(ws) / MANIFEST).read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}
    if not isinstance(m, dict):
        return {}
    files = m.get("files")
    m["files"] = {k: {**v, **{x: {} for x in ("inputs", "sources") if x in v and not isinstance(v[x], dict)}}
                  for k, v in (files if isinstance(files, dict) else {}).items() if isinstance(v, dict)}
    packs = m.get("packs")
    m["packs"] = {k: v for k, v in (packs if isinstance(packs, dict) else {}).items() if isinstance(v, dict)}
    return m


def _save_manifest(ws, m: dict) -> None:
    m["engine_version"] = __version__
    m.setdefault("files", {})
    m.setdefault("packs", {})
    write_json(Path(ws) / MANIFEST, m)


def record_derived(ws, name: str, inputs: dict, sources: dict | None = None) -> None:
    """Record the sha256 of ws/name and of every input it was built from (read-modify-write)."""
    ws = Path(ws)
    m = load_manifest(ws)
    entry = {"sha256": file_sha256(ws / name), "inputs": inputs, "written": _now()}
    if sources:
        entry["sources"] = sources
    m.setdefault("files", {})[name] = entry
    _save_manifest(ws, m)


def record_export(ws, xlsx) -> None:
    ws = Path(ws)
    m = load_manifest(ws)
    m["export"] = {"file": os.path.basename(xlsx), "sha256": file_sha256(xlsx)}
    _save_manifest(ws, m)


def record_pack_check(ws, pack_name, pack_path, ok, n_problems) -> None:
    ws = Path(ws)
    m = load_manifest(ws)
    p = Path(pack_path)
    m.setdefault("packs", {})[pack_name] = {
        "sha256": file_sha256(p) if p.is_file() else "", "path": str(_abs(p)),
        "ok": bool(ok), "problems": n_problems, "checked": _now()}
    _save_manifest(ws, m)


def taxonomy_sources(path, overlays) -> list:
    return [Path(path) if path else REFS / "taxonomy.json", *[_overlay_path(x) for x in overlays]]


def taxonomy_hash(paths) -> str:
    return hashlib.sha256("\n".join(file_sha256(p) for p in paths).encode()).hexdigest()


def _current_input(ws: Path, name: str, entry: dict):
    """The input's hash now, 'none', or None when it cannot be checked (export) or is missing."""
    src = entry.get("sources", {})
    if name == "taxonomy":
        paths = [Path(p) for p in src.get("taxonomy", [])]
        return taxonomy_hash(paths) if paths and all(p.is_file() for p in paths) else None
    if name == "overrides":
        p = src.get("overrides")
        if not p and src.get("overrides_auto") and (ws / "role_overrides.json").is_file():
            return file_sha256(ws / "role_overrides.json")
        return "none" if not p else (file_sha256(p) if Path(p).is_file() else None)
    if name == "export":
        return "skip"
    p = ws / name
    return file_sha256(p) if p.is_file() else None


def stale_files(ws) -> list:
    """[(file, reason)] for recorded derived files whose recorded input hash differs from the current one."""
    ws = Path(ws)
    files = load_manifest(ws).get("files", {})
    out = {}
    for name, entry in files.items():
        for inp, rec in entry.get("inputs", {}).items():
            cur = _current_input(ws, inp, entry)
            if cur == "skip":
                continue
            if cur is None:
                out[name] = f"input {inp} is missing"
                break
            if cur != rec:
                out[name] = f"{inp} changed since this was built"
                break
    changed = True
    while changed:  # a file built from a stale derived file is stale too
        changed = False
        for name, entry in files.items():
            if name in out:
                continue
            for inp in entry.get("inputs", {}):
                if inp in out and inp in files:
                    out[name] = f"input {inp} is stale"
                    changed = True
                    break
    return sorted(out.items())


def warn_stale(ws) -> None:
    for name, reason in stale_files(ws):
        cmd = REBUILD.get(name, "ingest")
        print(f"warning: stale input {name} ({reason}); re-run pmlc {cmd} {ws}", file=sys.stderr)


def cmd_status(a):
    ws = Path(a.ws)
    if not ws.is_dir():
        sys.exit(f"{ws} is not a folder.")
    m = load_manifest(ws)
    recorded = m.get("files", {})
    stale = dict(stale_files(ws))
    print("| File | State | Reason |\n|---|---|---|")
    names = [n for n in KNOWN_DERIVED if (ws / n).exists() or n in recorded]
    for n in names:
        if not (ws / n).exists():
            state, reason = "stale", "file is missing"
        elif n not in recorded:
            state, reason = "unknown (built before manifest)", ""
        elif n in stale:
            state, reason = "stale", stale[n]
        else:
            state, reason = "current", ""
        print(f"| {n} | {state} | {reason} |")
    packs = dict(m.get("packs", {}))
    out = outputs_dir(ws)
    if out.is_dir():
        for p in sorted(out.glob("*.json")):
            try:
                d = json.loads(p.read_text(encoding="utf-8"))
            except (OSError, ValueError):
                continue
            if isinstance(d, dict) and isinstance(d.get("pack"), str):
                packs.setdefault(d["pack"], None)
    print("\n| Pack | Last check | Changed since |\n|---|---|---|")
    for name in sorted(packs):
        rec = packs[name]
        if not rec:
            print(f"| {name} | never | n/a |")
            continue
        last = ("ok" if rec.get("ok") else f"{rec.get('problems', 0)} problem(s)") + f" ({rec.get('checked', '')})"
        p = Path(rec.get("path", ""))
        changed = "missing" if not p.is_file() else ("yes" if file_sha256(p) != rec.get("sha256") else "no")
        print(f"| {name} | {last} | {changed} |")
    signed = m.get("sign_offs", {})
    if signed:
        print("\n| Signed off | Status | Role | Date | Changed since |\n|---|---|---|---|---|")
        for name in sorted(signed):
            r = signed[name]
            f = Path(r.get("path") or out / name)
            try:
                d = json.loads(f.read_text(encoding="utf-8"))
                changed = "yes" if _hash_mismatch(d) else "no"
            except (OSError, ValueError):
                changed = "missing"
            print(f"| {name} | {r.get('status')} | {r.get('reviewer_role')} | {r.get('reviewed_on')} | {changed} |")


# ----------------------------------------------------------------------------- handling and sign-off

HANDLING_DEFAULTS = {"status": "draft", "classification": "internal", "legal_review": "not-required",
                     "distribution": [], "version": "1.0"}
HANDLING_KEYS = ("status", "reviewer_role", "reviewed_on", "classification", "legal_review", "distribution",
                 "version", "retention_review", "content_hash")
HANDLING_ENUMS = {"status": ("draft", "reviewed", "approved"),
                  "classification": ("internal", "restricted", "external", "public"),
                  "legal_review": ("not-required", "requested", "cleared")}
REASON_IN_PROGRESS = "the investigation status is in progress"
REASON_UNEDITED = "no finding is marked as edited by a person"
REASON_EDITED_SINCE = "edited since sign-off"
REASON_NOT_SIGNED = "not signed off"
REASON_INVALID = "invalid handling block"
WARN_NO_HANDLING = "no handling block: rendered as draft (internal)"
WARN_EDITED_SINCE = "edited since sign-off: status is draft until signed off again"
_HASH_RX = re.compile(r"sha256:[0-9a-f]{64}")


def forced_draft_reasons(w) -> list:
    """Why nothing from this workspace may be reviewed or approved yet (empty when it may be)."""
    out = []
    status = _s(((w.get("incident") or {}).get("fields") or {}).get("Investigation Status"))
    if status.lower().startswith("in_progress"):
        out.append(REASON_IN_PROGRESS)
    F = w.get("findings") or []
    if F and all(_s(f.get("user_edited")).lower() in ("", "no") for f in F):
        out.append(REASON_UNEDITED)
    return out


def content_hash(doc: dict) -> str:
    """sha256 of the document without its handling block (key order and whitespace do not matter)."""
    body = {k: v for k, v in doc.items() if k != "handling"}
    raw = json.dumps(body, sort_keys=True, separators=(",", ":"), ensure_ascii=False)
    return "sha256:" + hashlib.sha256(raw.encode("utf-8")).hexdigest()


def _hash_mismatch(doc) -> bool:
    h = doc.get("handling") if isinstance(doc, dict) else None
    return isinstance(h, dict) and isinstance(h.get("content_hash"), str) and h["content_hash"] != content_hash(doc)


def check_handling(doc) -> tuple:
    """(problems, warnings) for a pack's or register's handling block."""
    if not isinstance(doc, dict):
        return [], []
    if "handling" not in doc:
        return [], [WARN_NO_HANDLING]
    h = doc["handling"]
    if not isinstance(h, dict):
        return ["handling: must be an object"], []
    P, W = [], []
    for k, v in h.items():
        if k not in HANDLING_KEYS:
            P.append(f"handling.{k}: unknown key; use {', '.join(HANDLING_KEYS)}")
        elif k in HANDLING_ENUMS and v not in HANDLING_ENUMS[k]:
            P.append(f"handling.{k}: must be one of {', '.join(HANDLING_ENUMS[k])}")
        elif k in ("reviewed_on", "retention_review") and not (isinstance(v, str) and iso_date(v) == v):
            P.append(f"handling.{k}: must be YYYY-MM-DD")
        elif k == "content_hash" and not (isinstance(v, str) and _HASH_RX.fullmatch(v)):
            P.append(f"handling.{k}: must be sha256:<64 hex digits>")
        elif k in ("reviewer_role", "version") and not (isinstance(v, str) and v.strip()):
            P.append(f"handling.{k}: must be non-empty text")
        elif k == "distribution" and not (isinstance(v, list) and all(isinstance(x, str) for x in v)):
            P.append(f"handling.{k}: must be a list of audience names")
    if h.get("status") in ("reviewed", "approved"):
        for k in ("reviewer_role", "reviewed_on", "content_hash"):
            if k not in h:
                P.append(f"handling.{k}: required when status is {h['status']}")
        if _hash_mismatch(doc):
            W.append(WARN_EDITED_SINCE)
    return P, W


def effective_handling(doc, w) -> dict:
    """The handling that applies: defaults merged with the block, forced down to draft when it must be."""
    h = doc.get("handling") if isinstance(doc, dict) else None
    out = {**HANDLING_DEFAULTS, **{k: v for k, v in (h if isinstance(h, dict) else {}).items() if k in HANDLING_KEYS}}
    out["distribution"] = list(out["distribution"]) if isinstance(out["distribution"], list) else []
    for k, ok in HANDLING_ENUMS.items():  # an invalid value never passes through as effective
        if out.get(k) not in ok:
            out[k] = HANDLING_DEFAULTS[k]
    reasons = forced_draft_reasons(w)
    problems = check_handling(doc)[0] if isinstance(h, dict) else []
    claims = isinstance(h, dict) and h.get("status") in ("reviewed", "approved")
    if problems:
        reasons.append(REASON_INVALID)
    if claims:
        if _hash_mismatch(doc):
            reasons.append(REASON_EDITED_SINCE)
        elif not (isinstance(h.get("content_hash"), str) and not problems):
            reasons.append(REASON_NOT_SIGNED)
    out["draft_reasons"] = list(dict.fromkeys(reasons))
    signed = claims and not reasons
    out["effective_status"] = h["status"] if signed else "draft"
    return out


RIBBON_HTML = f'<div class="ribbon" role="note">{DRAFT_LABEL}</div>'
# The per-page status in print: a page margin box, so every printed page carries it.
PAGE_FOOTER_CSS = f'@page{{@bottom-center{{content:"{STATUS_PHRASE} · ProbeMatrix Learnings Coach"}}}}'
# The DRAFT ribbon and handling line, shared by the pack page and the dashboard. In print the ribbon keeps a
# border and its text, so it never depends on colour alone.
HANDLING_CSS = PAGE_FOOTER_CSS + """
.ribbon{display:inline-block;border:2px solid var(--warn,#c2410c);color:var(--warn,#c2410c);font-weight:700;letter-spacing:.12em;padding:2px 12px;border-radius:4px;margin:0 0 8px}
p.handling{color:var(--mut);font-size:13px;margin:0 0 6px}
@media print{.ribbon{border:2px solid #000;color:#000;background:none}}"""


def is_draft(eh: dict) -> bool:
    """True when the DRAFT ribbon must show: anything short of an effective approval."""
    return eh["effective_status"] != "approved"


def review_text(eh: dict) -> str:
    """'approved by <role> on <date>', or the bare effective status (no reviewer for a status that does not hold)."""
    st = eh["effective_status"]
    who = (f" by {eh['reviewer_role']} on {eh['reviewed_on']}"
           if st != "draft" and eh.get("reviewer_role") and eh.get("reviewed_on") else "")
    return st + who


def handling_lines(eh: dict) -> list:
    """The handling line, then 'Draft because: ...' when there are reasons."""
    out = [f"Review: {review_text(eh)} · Classification: {eh['classification']} · Version {eh['version']}"]
    if eh["draft_reasons"]:
        out.append("Draft because: " + "; ".join(eh["draft_reasons"]))
    return out


def status_line(status="") -> str:
    """A document's status text, always carrying STATUS_PHRASE and AI_STATUS (each once)."""
    s = _s(status)
    bits = ([] if STATUS_PHRASE in s else [STATUS_PHRASE]) + ([s] if s else [])
    return " · ".join(bits + ([] if AI_STATUS in s else [AI_STATUS]))


def record_sign_off(ws, path, h: dict) -> None:
    m = load_manifest(Path(ws))
    m.setdefault("sign_offs", {})[Path(path).name] = {
        **{k: h[k] for k in ("status", "reviewer_role", "reviewed_on", "content_hash") if k in h},
        "path": str(Path(path).resolve())}
    _save_manifest(ws, m)


def cmd_sign_off(a):
    w = load_ws(a.ws)
    path = Path(a.file)
    try:
        doc = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError) as e:
        sys.exit(f"cannot read {path} as JSON ({e.__class__.__name__})")
    if isinstance(doc, list):
        sys.exit('wrap the entries as {"entries": [...]} before signing off')
    if not isinstance(doc, dict):
        sys.exit(f"{path} must be a JSON object")
    reasons = forced_draft_reasons(w)
    if reasons:
        sys.exit("cannot sign off: " + "; ".join(reasons))
    if "pack" in doc:
        problems = _pack_module().check_pack(doc, w).problems
    elif "entries" in doc:
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            try:
                cmd_check_register(argparse.Namespace(ws=a.ws, register=str(path)))
                problems = []
            except SystemExit as e:
                problems = [x for x in buf.getvalue().splitlines() if x.startswith(" - ")] if e.code else []
                if e.code and not problems:
                    problems = [f"check-register failed ({e.code})"]
    else:
        sys.exit(f"{path} is neither a pack (no 'pack' key) nor a register (no 'entries')")
    if problems:
        sys.exit("cannot sign off: the file has problems; run the check and fix them first\n" + "\n".join(problems[:20]))
    h = dict(doc.get("handling") or {}) if isinstance(doc.get("handling"), dict) else {}
    h.pop("content_hash", None)
    h.update(status=a.status, reviewer_role=a.role, reviewed_on=a.date or dt.date.today().isoformat())
    for k in ("classification", "legal_review", "version", "retention_review"):
        if getattr(a, k) is not None:
            h[k] = getattr(a, k)
    if a.distribution:
        h["distribution"] = list(a.distribution)
    doc["handling"] = h
    h["content_hash"] = content_hash(doc)
    bad, _ = check_handling(doc)
    if bad:
        sys.exit("cannot sign off: " + "; ".join(bad))
    write_json(path, doc)
    record_sign_off(a.ws, path, h)
    print(f"Signed off {path}: {h['status']} by {h['reviewer_role']} on {h['reviewed_on']}")


# ----------------------------------------------------------------------------- run folders

RUNS_DIR = "pmlc-runs"
SLUG_PART_RX = re.compile(r"^[a-z0-9]+(?:-[a-z0-9]+)*$")
RUN_SLUG_RX = re.compile(r"^\d{4}-\d{2}-\d{2}(?:-[a-z0-9]+)+$")


def iso_date(v) -> str | None:
    """'2020-09-24' (or a value starting with it, e.g. an Excel datetime) -> '2020-09-24'; else None."""
    m = re.match(r"\s*(\d{4}-\d{2}-\d{2})(?!\d)", _s(v))
    if not m:
        return None
    try:
        dt.date.fromisoformat(m.group(1))
    except ValueError:
        return None
    return m.group(1)


def run_slug(date, slug) -> str:
    """Incident date + '<site>-<event>' -> run folder name, e.g. 2020-09-24-hazira-gas-release-fire."""
    d = iso_date(date)
    if not d:
        raise ValueError(f"not an ISO date (YYYY-MM-DD): {date!r}")
    if not SLUG_PART_RX.match(slug or ""):
        raise ValueError("slug must be lowercase words joined by hyphens, "
                         f"e.g. hazira-gas-release-fire (got {slug!r})")
    return f"{d}-{slug}"


def valid_run_slug(s) -> bool:
    return bool(RUN_SLUG_RX.match(s or "")) and iso_date(s[:10]) is not None


def _abs(path) -> Path:
    """Absolute path without following symlinks (folder names are what the layout rule checks)."""
    return Path(os.path.abspath(path))


def is_standard_ws(ws) -> bool:
    """True for <anything>/pmlc-runs/<valid slug>/workspace."""
    p = _abs(ws)
    return p.name == "workspace" and p.parent.parent.name == RUNS_DIR and valid_run_slug(p.parent.name)


def outputs_dir(ws) -> Path:
    """Where a workspace's outputs go: the sibling outputs/ in the standard layout, ws/outputs otherwise."""
    p = _abs(ws)
    return p.parent / "outputs" if is_standard_ws(p) else p / "outputs"


def cmd_run_path(a):
    ws = _abs(a.ws)
    std = is_standard_ws(ws)
    print(f"run:       {ws.parent if std else ws}")
    print(f"workspace: {ws}")
    print(f"outputs:   {outputs_dir(ws)}")
    print(f"slug:      {ws.parent.name if std else '(legacy layout: not under pmlc-runs/<slug>/workspace)'}")
    if (ws / "findings.json").is_file():
        w = load_ws(ws)
        print(f"top_rating: {top_rating(w)}")
        print(f"incident_date: {incident_date(w) or 'unknown'}")


def sample_folder_problems(samples) -> list:
    """Contribution rule: samples/<YYYY-MM-DD>-<site>-<event>/ with README.md and <slug>_probematrix-findings.xlsx."""
    problems = []
    for d in sorted(p for p in Path(samples).iterdir() if p.is_dir()):
        if not valid_run_slug(d.name):
            problems.append(f"{d.name}: folder name must be YYYY-MM-DD-<site>-<event> "
                            "(lowercase words joined by hyphens)")
            continue
        if not (d / "README.md").is_file():
            problems.append(f"{d.name}: missing README.md")
        if not (d / f"{d.name}_probematrix-findings.xlsx").is_file():
            problems.append(f"{d.name}: missing {d.name}_probematrix-findings.xlsx")
    return problems


def cmd_check_samples(a):
    if not Path(a.dir).is_dir():
        sys.exit(f"No such folder: {a.dir}")
    probs = sample_folder_problems(a.dir)
    for p in probs:
        print(" -", p)
    print(f"{len(probs)} problem(s) in {a.dir}")
    sys.exit(1 if probs else 0)


def find_sheet(wb, *candidates):
    for c in candidates:
        for n in wb.sheetnames:
            if n.strip().lower() == c.lower():
                return wb[n]
    for c in candidates:
        for n in wb.sheetnames:
            if c.lower() in n.lower():
                return wb[n]
    return None


def rows_as_dicts(ws):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = [_s(h) for h in rows[0]]
    return [dict(zip(header, r)) for r in rows[1:] if any(c is not None for c in r)]


# ----------------------------------------------------------------------------- export schema

KNOWN_COLUMNS = {
    "Identified Findings": ["Finding ID", "Nanocode", "Name", "Level", "Level Name", "Category", "Subcategory",
                            "Classification", "Rating", "Instance ID", "Instance Label", "Short Description",
                            "AI Justification", "Causal Explanation", "Individuals Involved", "Linked Evidence",
                            "Evidence to Collect", "Key Questions", "Suggested Interviews", "Causal Hypothesis",
                            "User Edited", "Created At", "Updated At"],
    "Evidence Links": ["Finding Code", "Finding Name", "Evidence ID", "Evidence Type", "Relevance",
                       "Specific Reference", "Evidence Description"],
    "Causal Chain": ["Link ID", "From Code", "From Level", "To Code", "To Level", "Relationship", "Confidence",
                     "Basis", "Evidence Support", "User Confirmed", "Notes"],
}
KNOWN_SUMMARY_FIELDS = [
    "Incident Number", "Incident Title", "Incident Type", "Date", "Time", "Location", "Severity", "Outcome",
    "Investigation Status", "Investigating Officer", "Consequence Type", "Persons Affected", "Fatalities",
    "Injuries", "Risk Rating", "Potential Severity", "Likelihood of Recurrence", "Reportable to Regulator",
    "Regulatory Bodies", "Regulatory Reference", "Department", "Area/Zone", "Shift", "Weather Conditions",
    "Lighting Conditions", "Supervisor at Time", "Investigation Team", "Due Date", "Completed Date",
    "Verification Required",
]
REQUIRED_COLUMNS = ["Finding ID", "Level", "Rating", "Individuals Involved", "Linked Evidence"]
CAPA_PLANNER_SHEETS = ("Chapter Headers", "Corrective Actions", "Preventive Actions", "Residual Risks")


def _header(ws) -> list[str]:
    if ws is None:
        return []
    for row in ws.iter_rows(values_only=True):
        return [h for h in (_s(c) for c in row) if h]
    return []


def _known_sheet(wb, name):
    """The sheet for a KNOWN_COLUMNS key, accepting the alias ingest accepts for the findings sheet."""
    return find_sheet(wb, name, "Findings") if name == "Identified Findings" else find_sheet(wb, name)


def schema_report(wb) -> dict:
    """Fingerprint plus missing and unknown columns for the three columnar sheets."""
    headers = {n: _header(_known_sheet(wb, n)) for n in KNOWN_COLUMNS}
    canon = "\n".join(list(wb.sheetnames) + ["\t".join(headers[n]) for n in KNOWN_COLUMNS])
    missing, unknown = {}, {}
    for n, known in KNOWN_COLUMNS.items():
        if _known_sheet(wb, n) is None:
            continue
        m = [c for c in known if c not in headers[n]]
        u = [c for c in headers[n] if c not in known]
        if m:
            missing[n] = m
        if u:
            unknown[n] = u
    return {"fingerprint": hashlib.sha256(canon.encode("utf-8")).hexdigest()[:12],
            "missing": missing, "unknown": unknown}


def analysis_stats(ws):
    """Read the `Level | Level Name | Total Findings` block and its Total row, or None."""
    if ws is None:
        return None
    by_level, total, in_block = {}, None, False
    for row in ws.iter_rows(values_only=True):
        cells = list(row) + [None] * 3
        if not in_block:
            in_block = (_s(cells[0]) == "Level" and _s(cells[1]) == "Level Name"
                        and _s(cells[2]) == "Total Findings")
            continue
        first = cells[0]
        if _s(first).lower() == "total":
            total = _int(cells[2])
            break
        lvl, n = _int(first), _int(cells[2])
        if lvl is not None and n is not None:
            by_level[lvl] = n
    if total is None:
        return None
    return {"total": total, "by_level": by_level}


def _fmt_cols(d: dict) -> str:
    return "; ".join(f"{s}: {', '.join(c)}" for s, c in d.items()) or "none"


# ----------------------------------------------------------------------------- ingest

def ingest_target(a, fields: dict) -> Path:
    """Resolve where ingest writes the workspace; exits with a clear message if the folder rule is broken."""
    if a.legacy_out:
        if not is_standard_ws(a.legacy_out):
            print(f"Warning: {a.legacy_out} is outside the standard layout "
                  f"<export-dir>/{RUNS_DIR}/<YYYY-MM-DD>-<site>-<event>/workspace.", file=sys.stderr)
        return _abs(a.legacy_out)
    if a.out:
        if not is_standard_ws(a.out):
            sys.exit(f"--out must be <dir>/{RUNS_DIR}/<YYYY-MM-DD>-<site>-<event>/workspace (got {a.out}). "
                     "Use --slug instead, or --legacy-out for a non-standard folder.")
        return _abs(a.out)
    if not a.slug:
        sys.exit("Give --slug <site>-<event>, e.g. --slug hazira-gas-release-fire "
                 "(the date comes from --date or the summary sheet's Date).")
    date = a.date or fields.get("Date", "")
    if not iso_date(date):
        src = "--date" if a.date else "the summary sheet's Date"
        sys.exit(f"No valid incident date from {src}: {date!r}. Pass --date YYYY-MM-DD.")
    try:
        slug = run_slug(date, a.slug)
    except ValueError as err:
        sys.exit(str(err))
    root = _abs(a.root) if a.root else _abs(a.xlsx).parent
    return root / RUNS_DIR / slug / "workspace"


def cmd_ingest(a):
    openpyxl = _need_openpyxl()
    wb = openpyxl.load_workbook(a.xlsx, read_only=True, data_only=True)
    names = {n.strip().lower() for n in wb.sheetnames}
    if sum(1 for c in CAPA_PLANNER_SHEETS if c.lower() in names) >= 2:
        sys.exit('This is a CAPA Planner export, not a findings export. '
                 'Export "Identified Findings" from the investigation instead.')

    fs = find_sheet(wb, "Identified Findings", "Findings")
    if fs is None:
        sys.exit("Could not find an 'Identified Findings' sheet in the workbook.")
    schema = schema_report(wb)
    req_missing = [c for c in REQUIRED_COLUMNS if c in schema["missing"].get("Identified Findings", [])]
    if a.strict and req_missing:
        sys.exit("--strict: Identified Findings is missing required column(s): " + ", ".join(req_missing))
    findings = []
    for r in rows_as_dicts(fs):
        fid = _s(r.get("Finding ID"))
        if not fid:
            continue
        lvl = _int(r.get("Level"))
        findings.append({
            "id": fid,
            "code": _s(r.get("Nanocode")),
            "name": _s(r.get("Name")),
            "level": lvl,
            "level_name": _s(r.get("Level Name")) or LEVEL_NAMES.get(lvl, ""),
            "category": _s(r.get("Category")),
            "subcategory": _s(r.get("Subcategory")),
            "classification": _s(r.get("Classification")),
            "rating": _int(r.get("Rating")),
            "instance": _s(r.get("Instance ID")),
            "label": _s(r.get("Instance Label")),
            "short": _s(r.get("Short Description")),
            "justification": _s(r.get("AI Justification")),
            "causal": _s(r.get("Causal Explanation")),
            "roles": split_outside_parens(_s(r.get("Individuals Involved"))),
            "evidence": [e.strip() for e in _s(r.get("Linked Evidence")).split(",") if e.strip()],
            "evidence_to_collect": _s(r.get("Evidence to Collect")),
            "key_questions": _s(r.get("Key Questions")),
            "suggested_interviews": _s(r.get("Suggested Interviews")),
            "causal_hypothesis": _s(r.get("Causal Hypothesis")),
            "user_edited": _s(r.get("User Edited")),
        })

    evidence = {}
    es = find_sheet(wb, "Evidence Links", "Evidence")
    if es is not None:
        for r in rows_as_dicts(es):
            eid = _s(r.get("Evidence ID"))
            if not eid:
                continue
            e = evidence.setdefault(eid, {"id": eid, "type": _s(r.get("Evidence Type")),
                                          "description": _s(r.get("Evidence Description")), "links": []})
            e["links"].append({"code": _s(r.get("Finding Code")), "relevance": _s(r.get("Relevance")),
                               "reference": _s(r.get("Specific Reference"))})

    incident = {"fields": {}, "narrative": ""}
    isht = find_sheet(wb, "Incident Summary")
    if isht is not None:
        grab_next = False
        for row in isht.iter_rows(values_only=True):
            cells = [c for c in row if c is not None]
            if not cells:
                continue
            if grab_next and len(cells) == 1 and len(_s(cells[0])) > 200:
                incident["narrative"] = _s(cells[0])
                grab_next = False
                continue
            k = _s(row[0]) if row else ""
            v = _s(row[1]) if len(row) > 1 else ""
            if k.upper() == "INCIDENT NARRATIVE":
                grab_next = True
                if len(v) > 200:
                    incident["narrative"] = v
                    grab_next = False
                continue
            if k and v and k not in ("Field",):
                incident["fields"][k] = v

    causal_sheet = []
    cs = find_sheet(wb, "Causal Chain")
    if cs is not None:
        for r in rows_as_dicts(cs):
            if _s(r.get("From Code")) and _s(r.get("To Code")):
                causal_sheet.append({k: _s(v) for k, v in r.items()})
    incident["causal_chain_sheet"] = causal_sheet
    votes = []
    vs = find_sheet(wb, "Collaboration Data")
    if vs is not None:
        mode = None
        for row in vs.iter_rows(values_only=True):
            first = _s(row[0]) if row else ""
            if first.upper() in ("VOTES", "COMMENTS"):
                mode = first.upper()
                continue
            if mode == "VOTES" and first and first != "Finding ID" and len(row) > 2:
                votes.append({"finding": first, "vote": _s(row[1]), "reason": _s(row[2])})
    incident["review_votes"] = votes
    incident["date"] = iso_date(a.date) or iso_date(incident["fields"].get("Date"))
    incident["source_file"] = os.path.basename(a.xlsx)
    incident["sheets"] = wb.sheetnames
    incident["schema"] = schema
    stats = analysis_stats(find_sheet(wb, "Analysis Statistics"))
    incident["analysis_stats"] = stats

    out = ingest_target(a, incident["fields"])
    if is_standard_ws(out) and (out / "findings.json").exists() and not a.force:
        sys.exit(f"{out} already holds a workspace. Pass --force to overwrite it.")
    write_json(out / "findings.json", findings)
    write_json(out / "evidence.json", evidence)
    write_json(out / "incident.json", incident)
    record_export(out, a.xlsx)
    for n in ("findings.json", "evidence.json", "incident.json"):
        record_derived(out, n, {"export": file_sha256(a.xlsx)})
    print(f"Ingested {len(findings)} findings, {len(evidence)} evidence items, "
          f"{len(causal_sheet)} causal-chain rows -> {out}")
    print(f"schema: fingerprint {schema['fingerprint']}; missing: {_fmt_cols(schema['missing'])}; "
          f"unknown: {_fmt_cols(schema['unknown'])}")
    if stats is None:
        print("note: no Analysis Statistics block recognised; cross-check skipped")
    elif stats["total"] == len(findings):
        print(f"cross-check: Analysis Statistics agrees ({stats['total']})")
    else:
        print(f"warning: Analysis Statistics says {stats['total']} findings, ingested {len(findings)}")
    if is_standard_ws(out):
        outputs_dir(out).mkdir(parents=True, exist_ok=True)
        print(f"Run folder: {out.parent}\n  workspace: {out}\n  outputs:   {outputs_dir(out)}")


# ----------------------------------------------------------------------------- profile

def cmd_profile(a):
    w = load_ws(a.ws)
    F, E, inc = w["findings"], w.get("evidence", {}), w.get("incident", {})
    C = collections.Counter
    lvl = C(f["level"] for f in F)
    rat = C(f["rating"] for f in F)
    sub = C((f["level"], f["subcategory"]) for f in F)
    codes = C(f"{f['code']} {f['name']}" for f in F)
    roles = C(r for f in F for r in f["roles"])
    etypes = C(e["type"] for e in E.values())
    fields = inc.get("fields", {})
    L = []
    L.append(f"# Investigation profile: {fields.get('Incident Title', inc.get('source_file', ''))}\n")
    L.append("| Field | Value |\n|---|---|")
    for k in ("Incident Number", "Incident Title", "Incident Type", "Date", "Location", "Severity",
              "Investigation Status"):
        if k in fields:
            L.append(f"| {k} | {fields[k]} |")
    n = len(F)
    up = lvl.get(3, 0) + lvl.get(4, 0)
    L.append(f"\n**Findings:** {n} across {len(codes)} distinct factor codes. "
             f"**Evidence items:** {len(E)}. **Distinct role strings:** {len(roles)}.")
    if n:
        L.append(f"**Share at Supervision + Organisational levels:** {up} of {n} ({100*up/n:.0f}%).\n")
    L.append("## By level\n\n| Level | Name | Findings |\n|---|---|---:|")
    for k in sorted(x for x in lvl if x is not None):
        L.append(f"| {k} | {LEVEL_NAMES.get(k, '')} | {lvl[k]} |")
    L.append("\n## By rating (higher = stronger causal contribution)\n\n| Rating | Findings |\n|---|---:|")
    for k in sorted(x for x in rat if x is not None):
        L.append(f"| {k} | {rat[k]} |")
    L.append("\n## Sub-categories\n\n| Level | Sub-category | Findings |\n|---|---|---:|")
    for (lv, sc), c in sub.most_common():
        L.append(f"| {lv} | {sc} | {c} |")
    L.append(f"\n## Most frequent factors (top {a.top})\n\n| Factor | Instances |\n|---|---:|")
    for k, c in codes.most_common(a.top):
        L.append(f"| {k} | {c} |")
    L.append(f"\n## Most-named roles (top {a.top}, raw strings)\n\n| Role | Findings |\n|---|---:|")
    for k, c in roles.most_common(a.top):
        L.append(f"| {k} | {c} |")
    L.append("\n## Evidence types\n\n| Type | Items |\n|---|---:|")
    for k, c in etypes.most_common():
        L.append(f"| {k or '(blank)'} | {c} |")
    flags = []
    for col, lab in (("evidence_to_collect", "Evidence to Collect"), ("key_questions", "Key Questions"),
                     ("suggested_interviews", "Suggested Interviews"), ("causal_hypothesis", "Causal Hypothesis")):
        if n and not any(f[col] for f in F):
            flags.append(f"Column '{lab}' is empty on every finding (the investigator-coach skill can draft it).")
    if not inc.get("causal_chain_sheet"):
        flags.append("Causal Chain sheet is empty; use `pmlc.py causal` to extract links from the causal explanations.")
    if n and all(f["user_edited"].lower() in ("", "no") for f in F):
        flags.append("No finding is marked 'User Edited'. Confirm the human review status before publishing learnings.")
    narr = inc.get("narrative", "")[:4000]
    sdate = fields.get("Date", "")
    if sdate and narr:
        try:
            d = dt.date.fromisoformat(sdate[:10])
            nd = [x[0] for x in _parse_dates(narr) if x[1] in (0, 1)]
            if nd and d not in nd:
                flags.append(f"Summary-sheet Date {sdate} does not appear in the narrative, which mentions "
                             f"{', '.join(sorted({x.isoformat() for x in nd})[:3])}. Confirm the date and pass "
                             f"`--date` to `dashboard`.")
        except ValueError:
            pass
    itype = fields.get("Incident Type", "")
    if itype and narr and itype.lower() not in narr.lower():
        flags.append(f"Incident Type '{itype}' is not mentioned in the narrative; check it before quoting it.")
    downs = [v for v in inc.get("review_votes", []) if _s(v.get("vote")).lower() in ("downgrade", "reject", "remove")]
    if downs:
        flags.append("Reviewer downgrade votes on: " + "; ".join(f"{v.get('finding')} ({v.get('reason', '')})" for v in downs)
                     + ". Exclude or flag these findings in learnings.")
    if fields.get("Investigation Status", "").lower().startswith("in_progress"):
        flags.append("Investigation Status is 'in_progress'; label outputs as draft learnings.")
    L.append("\n## Data-quality flags\n")
    L += [f"- {x}" for x in flags] or ["- None."]
    L.append(f"\n---\n*{BRAND}*")
    text = "\n".join(L)
    if a.out:
        Path(a.out).write_text(text, encoding="utf-8")
        print(f"Profile written to {a.out}")
    else:
        print(text)


# ----------------------------------------------------------------------------- classify

def _overlay_path(x) -> Path:
    """A bare name resolves to references/overlays/<name>.json; anything with a slash or .json is a path."""
    if x.endswith(".json") or "/" in x or os.sep in x:
        return Path(x)
    return REFS / "overlays" / f"{x}.json"


def _load_taxonomy(path=None, overlays=()) -> dict:
    """The taxonomy rules with compiled regexes. Each overlay (name or path) ORs its extra regexes onto the base
    rules with the same IDs: match = (?:base)|(?:extra)."""
    p = Path(path) if path else REFS / "taxonomy.json"
    t = json.loads(p.read_text(encoding="utf-8"))
    t["_file"] = str(path) if path else "taxonomy.json"
    t["_overlays"] = []
    for x in overlays:
        op = _overlay_path(x)
        if not op.exists():
            sys.exit(f"Overlay {x} not found at {op}. Pass a name from references/overlays/ or a path to a .json file.")
        try:
            ov = json.loads(op.read_text(encoding="utf-8"))
        except ValueError:
            sys.exit(f"overlay {op.stem}: not valid JSON ({op}). Fix the file or pass another overlay.")
        if not isinstance(ov, dict):
            sys.exit(f"overlay {op.stem}: must be a JSON object with an `extend` block ({op}).")
        name = ov.get("name") or op.stem
        ext = ov.get("extend") or {}
        if not isinstance(ext, dict):
            sys.exit(f"overlay {name}: `extend` must be an object keyed by entities, tiers or relations.")
        for group in ("entities", "tiers", "relations"):
            rules = {r["id"]: r for r in t[group]}
            for rid, extra in (ext.get(group) or {}).items():
                if rid not in rules:
                    sys.exit(f"Overlay {name} extends {rid}, which is not in the {group} of taxonomy {t['_file']}.")
                rules[rid]["match"] = f"(?:{rules[rid]['match']})|(?:{extra})"
        t["_overlays"].append(name)
    for group in ("entities", "tiers", "relations"):
        for rule in t[group]:
            rule["_rx"] = re.compile(rule["match"], re.I)
            rule["_ex"] = re.compile(rule["exclude"], re.I) if rule.get("exclude") else None
            rule["_wk"] = re.compile(rule["weak"], re.I) if rule.get("weak") else None
    return t


def map_role(role: str, tax: dict) -> dict:
    ents = [e["id"] for e in tax["entities"]
            if e["_rx"].search(role) and not (e["_ex"] and e["_ex"].search(role))]
    tier = next((t["id"] for t in tax["tiers"] if t["_rx"].search(role)), None) \
        or next((t["id"] for t in tax["tiers"] if t["_wk"] and t["_wk"].search(role)), "T9")
    rel = next((r["id"] for r in tax["relations"] if r["_rx"].search(role)), tax["default_relation"]["id"])
    return {"F": ents, "T": tier, "R": rel, "source": "rule"}


def entities_for_tags(clusters: dict, tags: list) -> list:
    """Entity IDs (taxonomy order) carrying any of the tags. Exits when a tag is unknown or clusters.json has no tags."""
    tx = clusters.get("taxonomy", {})
    tagmap = tx.get("tags") or {}
    in_use = sorted({t for ts in tagmap.values() for t in ts})
    if not in_use:
        sys.exit(f"Taxonomy {tx.get('file', 'unknown')} defines no tags: add `tags` to its entities (see "
                 "references/taxonomy.json), or re-run classify if clusters.json predates v2.0.")
    missing = [t for t in tags if t not in in_use]
    if missing:
        sys.exit(f"No entity has tag(s) {missing} in taxonomy {tx.get('file', 'unknown')}. Tags in use: "
                 f"{', '.join(in_use)}.")
    return [e for e, ts in tagmap.items() if set(ts) & set(tags)]


def _taxonomy_label(file, overlays) -> str:
    return " + ".join([file, *overlays])


OVERRIDE_SHAPE = '{"<role>": {"F": ["F05"], "T": "T3", "R": "owner"}}'


def load_overrides(path: Path, tax: dict | None = None) -> dict:
    """{role: {F: [ids], T: id, R: id}} from an overrides file, with F normalised to a list. Exits with a sentence
    naming the file (and the role) when the file is missing, is not JSON or has the wrong shape, or, given the
    taxonomy, when an ID is not one of its entities (or F99), tiers (or T9) or relations (or the default)."""
    p = Path(path)
    if not p.is_file():
        sys.exit(f"Overrides file {path} not found. Pass an existing file or --no-overrides.")
    try:
        raw = json.loads(p.read_text(encoding="utf-8"))
    except json.JSONDecodeError as err:
        sys.exit(f"{path} is not valid JSON ({err}). Fix it or pass --no-overrides.")
    if not isinstance(raw, dict):
        sys.exit(f"{path} must be a JSON object of roles, shaped {OVERRIDE_SHAPE}.")
    allowed = None
    if tax is not None:
        allowed = {"F": {e["id"] for e in tax["entities"]} | {"F99"}, "T": {t["id"] for t in tax["tiers"]} | {"T9"},
                   "R": {r["id"] for r in tax["relations"]} | {tax["default_relation"]["id"]}}
    out = {}
    for role, o in raw.items():
        where = f"{path}: the override for role {role!r}"
        if not isinstance(o, dict):
            sys.exit(f"{where} must be an object shaped {OVERRIDE_SHAPE}, not {json.dumps(o)}.")
        extra = sorted(set(o) - {"F", "T", "R"})
        if extra:
            sys.exit(f"{where} has unknown key(s) {', '.join(extra)}; only F, T and R are read.")
        o = dict(o)
        if isinstance(o.get("F"), str):
            o["F"] = [o["F"]]
        if "F" in o and not (isinstance(o["F"], list) and all(isinstance(x, str) for x in o["F"])):
            sys.exit(f"{where} has F {json.dumps(o['F'])}; F must be a list of entity IDs such as [\"F05\"].")
        for k in ("T", "R"):
            if k in o and not isinstance(o[k], str):
                sys.exit(f"{where} has {k} {json.dumps(o[k])}; {k} must be one ID as a string.")
        if allowed:
            for k in ("F", "T", "R"):
                bad = [x for x in ([] if k not in o else o[k] if k == "F" else [o[k]]) if x not in allowed[k]]
                if bad:
                    sys.exit(f"{where} has unknown {k} ID(s) {', '.join(bad)}; taxonomy {tax['_file']} has "
                             f"{', '.join(sorted(allowed[k]))}.")
        out[role] = o
    return out


def _same_taxonomy_file(ws, prev, taxonomy) -> bool:
    """True when this classify's taxonomy file is the one the previous classify recorded (compared by real path)."""
    src = load_manifest(ws).get("files", {}).get("clusters.json", {}).get("sources", {}).get("taxonomy") or []
    now = Path(taxonomy) if taxonomy else REFS / "taxonomy.json"
    return bool(src) and Path(src[0]).resolve() == now.resolve()


def _previous_taxonomy(ws: Path, prev: dict):
    """(taxonomy path or None for the base one, [overlay names or paths]) that the last classify used. Prefers the
    manifest's absolute sources; a built-in overlay comes back as its bare name."""
    src = load_manifest(ws).get("files", {}).get("clusters.json", {}).get("sources", {}).get("taxonomy") or []
    names = list(prev.get("overlays", []))
    file = prev.get("file")
    tax = None
    custom = bool(file) and file != "taxonomy.json"
    if src and file == "taxonomy.json" and Path(src[0]).resolve() != (REFS / "taxonomy.json").resolve():
        custom = True  # a custom file that happens to be called taxonomy.json, recorded by its absolute path
    if custom:
        tax = src[0] if src and Path(src[0]).is_file() else file
    paths = src[1:] if len(src) == len(names) + 1 else [None] * len(names)
    builtin = (REFS / "overlays").resolve()
    overlays = []
    for name, op in zip(names, paths):
        if op and Path(op).is_file():
            overlays.append(Path(op).stem if Path(op).resolve().parent == builtin else op)
        else:
            overlays.append(name)
    return tax, overlays


def cmd_classify(a):
    ws = Path(a.ws)
    w = load_ws(ws)
    if a.no_overlay and a.overlay:
        sys.exit("--no-overlay and --overlay cannot be combined; pass one or the other.")
    prev = w.get("clusters", {}).get("taxonomy", {})
    taxonomy, overlays = a.taxonomy, list(a.overlay or ())
    reused, reused_tax = False, False
    if (not a.taxonomy or not a.overlay) and "file" in prev:
        rec_tax, rec_overlays = _previous_taxonomy(ws, prev)
        if not a.taxonomy:  # a flag for only one half keeps the other half from the previous classify
            taxonomy, reused_tax = rec_tax, bool(rec_tax)
        if not a.overlay and not a.no_overlay:
            overlays = rec_overlays
        if taxonomy and not Path(taxonomy).is_file():
            sys.exit(f"The previous classify used taxonomy {taxonomy}, which is no longer there. Pass --taxonomy "
                     f"<file> to choose one (the base taxonomy is {REFS / 'taxonomy.json'}).")
        reused = reused_tax or bool(overlays and not a.overlay)
        if reused:
            print(f"Reusing taxonomy {taxonomy or 'taxonomy.json'} and overlays {', '.join(overlays) or 'none'} "
                  "from the previous classify (pass --taxonomy/--overlay to change, --no-overlay for none).")
    tax = _load_taxonomy(taxonomy, overlays)
    if reused_tax:
        tax["_file"] = prev["file"]  # the same file under the name it was recorded with
    if "file" in prev:
        was, now = _taxonomy_label(prev["file"], prev.get("overlays", [])), _taxonomy_label(tax["_file"], tax["_overlays"])
        if was != now and _same_taxonomy_file(ws, prev, taxonomy) and prev.get("overlays", []) == tax["_overlays"]:
            was = now  # the same file reached by a relative path one time and an absolute one the next
        if was != now:
            print(f"warning: previous classify used {was}; this run uses {now}")
    ov_path, remembered = None, False
    if a.no_overrides and a.overrides:
        print("warning: --no-overrides wins: ignoring --overrides", file=sys.stderr)
    if not a.no_overrides:
        if a.overrides:
            ov_path = a.overrides
        elif (ws / "role_overrides.json").exists():
            ov_path = str(ws / "role_overrides.json")
        else:  # an explicit --overrides file from the previous classify is remembered like the taxonomy is
            rec = load_manifest(ws).get("files", {}).get("clusters.json", {}).get("sources", {})
            if rec.get("overrides") and not rec.get("overrides_auto") and Path(rec["overrides"]).is_file():
                ov_path = rec["overrides"]
                remembered = True
    overrides = load_overrides(Path(ov_path), tax) if ov_path else {}
    role_map = {}
    for f in w["findings"]:
        for r in f["roles"]:
            if r in role_map:
                continue
            m = map_role(r, tax)
            if r in overrides:
                o = overrides[r]
                m = {"F": o.get("F", m["F"]), "T": o.get("T", m["T"]), "R": o.get("R", m["R"]), "source": "override"}
            role_map[r] = m
    applied = sum(1 for m in role_map.values() if m["source"] == "override")
    stale = sorted(k for k in overrides if k not in role_map)
    names = {e["id"]: e["name"] for e in tax["entities"]}
    names["F99"] = "Unmapped"
    tier_names = {t["id"]: t["name"] for t in tax["tiers"]}
    tier_names["T9"] = "Unresolved"
    per_f, ent_idx = {}, collections.defaultdict(list)
    for f in w["findings"]:
        Fs, Ts, Rs = set(), set(), set()
        for r in f["roles"]:
            m = role_map[r]
            Fs.update(m["F"] or ["F99"])
            Ts.add(m["T"])
            Rs.add(m["R"])
        if not f["roles"]:
            Fs.add("F99")
        per_f[f["id"]] = {"F": sorted(Fs), "T": sorted(Ts), "R": sorted(Rs)}
        for e in Fs:
            ent_idx[e].append(f["id"])
    byid = {f["id"]: f for f in w["findings"]}
    top_rating = max((f["rating"] or 0) for f in w["findings"]) if w["findings"] else 0
    entities = {}
    for e, ids in sorted(ent_idx.items()):
        lv = collections.Counter(byid[i]["level"] for i in ids)
        entities[e] = {"name": names.get(e, e), "findings": ids, "n": len(ids),
                       "by_level": {str(k): lv.get(k, 0) for k in (1, 2, 3, 4)},
                       "top_rated": sum(1 for i in ids if byid[i]["rating"] == top_rating)}
    unresolved = sorted(r for r, m in role_map.items() if not m["F"] or m["T"] == "T9")
    out = {"taxonomy": {"entities": names, "tiers": tier_names,
                        "relations": {r["id"]: r["name"] for r in tax["relations"]} | {
                            tax["default_relation"]["id"]: tax["default_relation"]["name"]},
                        "tags": {e["id"]: list(e.get("tags", [])) for e in tax["entities"]},
                        "file": tax["_file"], "overlays": tax["_overlays"]},
           "entities": entities, "findings": per_f, "role_map": role_map, "unresolved_roles": unresolved,
           "overrides": {"file": ov_path, "applied": applied, "stale": stale}}
    write_json(ws / "clusters.json", out)
    write_json(ws / "singleton_roles.json", privacy.build_singletons(role_map, w["findings"], tax, TIER_LABELS))
    tpaths = taxonomy_sources(taxonomy, overlays)
    for derived in ("clusters.json", "singleton_roles.json"):
        record_derived(ws, derived,
                       {"findings.json": file_sha256(ws / "findings.json"), "taxonomy": taxonomy_hash(tpaths),
                        "overrides": file_sha256(ov_path) if ov_path else "none"},
                       {"taxonomy": [str(_abs(p)) for p in tpaths],
                        "overrides": str(_abs(ov_path)) if ov_path else None,
                        "overrides_auto": not a.no_overrides and not a.overrides and not remembered})
    if ov_path:
        print(f"Applied {applied} override(s) from {ov_path}")
    if stale:
        print(f"warning: {len(stale)} override(s) match no role: {', '.join(stale)}")
    tc = collections.Counter(t for v in per_f.values() for t in v["T"])
    print(f"Classified {len(role_map)} distinct roles; {len(unresolved)} need resolution "
          f"(no entity or no tier). Wrote {ws/'clusters.json'}\n")
    print("| Cluster | Name | Findings | L1/L2/L3/L4 |\n|---|---|---:|---|")
    for e, d in entities.items():
        bl = d["by_level"]
        print(f"| {e} | {d['name']} | {d['n']} | {bl['1']}/{bl['2']}/{bl['3']}/{bl['4']} |")
    print("\n| Tier | Findings touching |\n|---|---:|")
    for t in sorted(tc):
        print(f"| {t} {tier_names.get(t, '')} | {tc[t]} |")
    if unresolved:
        p = ws / "unresolved_roles.json"
        write_json(p, unresolved)
        print(f"\nUnresolved roles listed in {p}. Resolve them into an overrides JSON "
              f"({{role: {{F: [...], T: ..., R: ...}}}}) in {ws/'role_overrides.json'} and re-run classify; "
              f"that file is picked up automatically.")
    else:
        (ws / "unresolved_roles.json").unlink(missing_ok=True)


# ----------------------------------------------------------------------------- select

def _evidence_for(f, E, max_quote=400):
    out = []
    for eid in f["evidence"]:
        e = E.get(eid)
        if not e:
            out.append({"id": eid})
            continue
        link = next((lk for lk in e["links"] if lk["code"] == f["code"]), e["links"][0] if e["links"] else {})
        out.append({"id": eid, "type": e["type"], "relevance": link.get("relevance", ""),
                    "reference": (link.get("reference", "") or "")[:max_quote]})
    return out


SELECT_DEFAULTS = dict(entity=None, tag=None, tier=None, relation=None, level=None, code=None, text=None,
                       subcategory=None, prefix=None, min_rating=None, top_rated=False, entity_mode="any",
                       exclude_downgraded=False, max=None)


def select_args(**kw) -> types.SimpleNamespace:
    """The namespace _filter reads, with every filter at its default; callers pass only what they set."""
    unknown = set(kw) - set(SELECT_DEFAULTS)
    if unknown:
        raise TypeError(f"unknown select filter(s): {', '.join(sorted(unknown))}")
    return types.SimpleNamespace(**{**SELECT_DEFAULTS, **kw})


def top_rating(w) -> int:
    """The highest rating in the workspace (0 when there are no findings)."""
    return max(((f["rating"] or 0) for f in w["findings"]), default=0)


def incident_date(w) -> str | None:
    """The incident date (YYYY-MM-DD): recorded at ingest, else the summary sheet's Date field, else None."""
    inc = w.get("incident", {})
    return inc.get("date") or iso_date(inc.get("fields", {}).get("Date"))


def _filter(w, a):
    F = w["findings"]
    clusters = w.get("clusters", {})
    cl = clusters.get("findings", {})
    entity = set(a.entity or [])
    tag_sets = [set(entities_for_tags(clusters, [t])) for t in (a.tag or [])]
    all_mode = a.entity_mode == "all"
    if not all_mode:
        entity = entity.union(*tag_sets)
    prefixes = {p.upper() for p in a.prefix or []}
    subs = {x.lower() for x in a.subcategory or []}
    top = top_rating(w) if a.top_rated else None
    dropped = downgraded_findings(w) if a.exclude_downgraded else set()
    out = []
    for f in F:
        c = cl.get(f["id"], {"F": [], "T": [], "R": []})
        have = set(c["F"])
        if all_mode:
            if (entity or tag_sets) and not (entity <= have and all(ts & have for ts in tag_sets)):
                continue
        elif entity and not entity & have:
            continue
        if a.tier and not set(a.tier) & set(c["T"]):
            continue
        if a.relation and not set(a.relation) & set(c["R"]):
            continue
        if a.level and f["level"] not in a.level:
            continue
        if a.min_rating and (f["rating"] or 0) < a.min_rating:
            continue
        if top is not None and (f["rating"] or 0) != top:
            continue
        if a.code and f["code"] not in a.code:
            continue
        if prefixes and f["code"].split(" ")[0].upper() not in prefixes:
            continue
        if subs and f["subcategory"].lower() not in subs:
            continue
        if f["id"] in dropped:
            continue
        if a.text and not re.search(a.text, " ".join([f["short"], f["justification"], f["causal"], f["label"]]), re.I):
            continue
        out.append((f, c))
    out.sort(key=lambda fc: (-(fc[0]["rating"] or 0), -(fc[0]["level"] or 0), fc[0]["code"]))
    return out[: a.max] if a.max else out


SELECT_FIELDS = ["label", "short", "justification", "causal", "roles", "clusters", "evidence", "code", "name",
                 "level", "subcategory", "rating"]
COUNT_BY = ["entity", "tier", "relation", "level", "subcategory", "code"]


def _group_keys(f, c, how):
    """The group(s) a finding falls in; an entity/tier/relation key repeats the finding once per value."""
    if how == "level":
        return [f"L{f['level']}"]
    if how == "subcategory":
        return [f["subcategory"] or "(none)"]
    if how == "code":
        return [f["code"]]
    return list(c[{"entity": "F", "tier": "T", "relation": "R"}[how]]) or ["(none)"]


def _md_block(f, c, E, fields, brief, head):
    def on(k):
        return fields is None or k in fields

    parts = [f["id"]]
    cn = " ".join(str(f[k]) for k in ("code", "name") if on(k))
    lv = " ".join(x for x in ([f"L{f['level']}"] if on("level") else []) + ([f["subcategory"]] if on("subcategory") else []))
    for p in (cn, lv, f"rating {f['rating']}" if on("rating") else ""):
        if p:
            parts.append(p)
    L = [f"{head} " + " · ".join(parts)]
    if on("label"):
        L.append(f"**Instance:** {f['label']}\n")
    if on("short"):
        L.append(f"**Short:** {f['short']}\n")
    if not brief:
        if on("justification"):
            L.append(f"**Justification:** {f['justification']}\n")
        if on("causal"):
            L.append(f"**Causal explanation:** {f['causal']}\n")
    if on("roles"):
        L.append(f"**Roles:** {'; '.join(f['roles'])}  ")
    if on("clusters"):
        L.append(f"**Clusters:** F={','.join(c['F'])} T={','.join(c['T'])} R={','.join(c['R'])}\n")
    if not brief and on("evidence"):
        for e in _evidence_for(f, E):
            L.append(f"- `{e['id']}` ({e.get('type', '?')}): {e.get('relevance', '')} "
                     f"Quote: \"{e.get('reference', '')}\"")
    L.append("")
    return "\n".join(L)


def _json_item(f, c, E, fields):
    d = {k: f[k] for k in ("id", "code", "name", "level", "subcategory", "rating", "label", "short",
                           "justification", "causal", "roles")}
    d["clusters"] = c
    if fields is None or "evidence" in fields:
        d["evidence"] = _evidence_for(f, E)
    return d if fields is None else {k: v for k, v in d.items() if k == "id" or k in fields}


def _fit_count(n, fits):
    """The largest k in 0..n with fits(k) (0 when even nothing fits); fits is near-monotone in k."""
    lo, hi = 0, n
    while lo < hi:
        mid = (lo + hi + 1) // 2
        lo, hi = (mid, hi) if fits(mid) else (lo, mid - 1)
    while lo and not fits(lo):
        lo -= 1
    return lo


SELECT_FILTER_FLAGS = ("--entity", "--tier", "--relation", "--level", "--min-rating", "--top-rated", "--code",
                       "--text", "--tag", "--subcategory", "--prefix", "--entity-mode", "--exclude-downgraded", "--max")
PRESET_FILE = REFS / "select-presets.json"
SELECT_LINE = re.compile(r"PMLC select <ws>((?: [^`\n]*)?)")


def skill_select_lines(text: str) -> list[list[str]]:
    """The argument lists of the `PMLC select <ws> ...` gather calls in a SKILL.md, in order.

    Each call runs to the end of its line or its closing backtick. A trailing `# comment` and any `--out ...`
    are dropped. A call with a `<...>` placeholder left (such as <F-id>) is a template, not a preset."""
    import shlex
    calls = []
    for m in SELECT_LINE.finditer(text):
        rest = re.sub(r"\s+#.*$", "", m.group(1))
        rest = re.sub(r"\s--out\s+\S+", "", rest)
        if re.search(r"<[^>]*>", rest):
            continue
        calls.append(shlex.split(rest))
    return calls


def load_presets() -> dict:
    return json.loads(PRESET_FILE.read_text(encoding="utf-8"))


def _select_presets(a):
    """select --preset skill[:name]: run the skill's recorded gather calls with the real select parser."""
    given = [t for t in getattr(a, "_argv", []) if t.split("=")[0] in SELECT_FILTER_FLAGS]
    changed = [k for k, v in SELECT_DEFAULTS.items() if getattr(a, k) != v]
    if given or changed:
        sys.exit("--preset cannot be combined with filters; edit the preset or drop --preset.")
    skill, _, name = a.preset.partition(":")
    presets = load_presets()
    if skill not in presets:
        sys.exit(f"Unknown preset skill {skill}. Skills with presets: {', '.join(sorted(presets))}")
    calls = [p for p in presets[skill] if not name or p["name"] == name]
    if not calls:
        sys.exit(f"Unknown preset {a.preset}. Names for {skill}: {', '.join(p['name'] for p in presets[skill])}")
    out_flags = []
    if any(t.split("=")[0] == "--format" for t in getattr(a, "_argv", [])):
        out_flags += ["--format", a.format]
    for flag, val in (("--fields", a.fields), ("--group-by", a.group_by), ("--count-by", a.count_by)):
        if val is not None:
            out_flags += [flag, str(val)]
    if a.brief:
        out_flags.append("--brief")
    import contextlib
    import io
    parts = []
    for p in calls:
        head = f"# Preset {skill}: {p['name']}\n\n"
        budget = []
        if a.budget_chars is not None:  # one budget for the whole bundle: each call gets what the earlier ones left
            used = sum(len(x) for x in parts) + len(parts) + len(head)
            budget = ["--budget-chars", str(max(0, a.budget_chars - used))]
        sub = build_parser().parse_args(["select", a.ws, *p["args"], *out_flags, *budget])
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            cmd_select(sub)
        parts.append(head + buf.getvalue())
    text = "\n".join(parts)
    if a.out:
        Path(a.out).write_text(text, encoding="utf-8")
        print(f"Wrote {len(calls)} preset bundle(s) to {a.out}")
    else:
        sys.stdout.write(text)


def cmd_select(a):
    if getattr(a, "preset", None):
        return _select_presets(a)
    w = load_ws(a.ws)
    grouping = a.group_by or a.count_by
    if (a.entity or a.tag or a.tier or a.relation or grouping in ("entity", "tier", "relation")) \
            and "clusters" not in w:
        sys.exit("Filtering or grouping by entity/tag/tier/relation needs clusters.json. "
                 "Run: pmlc.py classify <ws>")
    fields = None
    if a.fields:
        fields = {x.strip() for x in a.fields.split(",") if x.strip()}
        bad = fields - set(SELECT_FIELDS) - {"id"}
        if bad:
            sys.exit(f"Unknown --fields value(s): {', '.join(sorted(bad))}. Allowed: {', '.join(SELECT_FIELDS)}")
    E = w.get("evidence", {})
    sel = _filter(w, a)
    skipped = []
    if a.exclude_downgraded:
        would = {f["id"] for f, _ in _filter(w, types.SimpleNamespace(**{**vars(a), "exclude_downgraded": False,
                                                                         "max": None}))}
        skipped = sorted(downgraded_findings(w) & would)
    note = f"Excluded (reviewer downgrade): {', '.join(skipped) or 'none'}"
    nsel = len(sel)
    if a.count_by:
        ignored = [n for n, v in (("--budget-chars", a.budget_chars), ("--fields", a.fields),
                                  ("--group-by", a.group_by)) if v]
        if ignored:
            print(f"note: --count-by ignores {', '.join(ignored)}", file=sys.stderr)
        counts = collections.Counter(k for f, c in sel for k in _group_keys(f, c, a.count_by))
        rows = sorted(counts.items(), key=lambda kv: (-kv[1], kv[0]))
        if a.format == "json":
            text = json.dumps(dict(rows), indent=1, ensure_ascii=False) + "\n"
        else:
            text = "\n".join([f"# Findings by {a.count_by} ({nsel} selected)\n", f"| {a.count_by} | findings |",
                              "|---|---:|"] + [f"| {k} | {n} |" for k, n in rows] + [""])
        kept = nsel
    else:
        md = a.format != "json"
        head = "###" if a.group_by else "##"
        blocks = [_md_block(f, c, E, fields, a.brief, head) if md else _json_item(f, c, E, fields) for f, c in sel]

        def render(k):
            if md:
                els = [f"# Selected findings ({k})\n",
                       f"Filters: entity={a.entity} tag={a.tag} tier={a.tier} relation={a.relation} level={a.level} "
                       f"min_rating={a.min_rating} top_rated={a.top_rated} entity_mode={a.entity_mode} "
                       f"code={a.code} prefix={a.prefix} subcategory={a.subcategory} text={a.text}\n"]
                if a.group_by:
                    groups = {}
                    for (f, c), b in zip(sel[:k], blocks[:k]):
                        for g in _group_keys(f, c, a.group_by):
                            groups.setdefault(g, []).append(b)
                    for g in sorted(groups):
                        els.append(f"## {g}\n")
                        els.extend(groups[g])
                else:
                    els.extend(blocks[:k])
                els.extend(footers_for(k))
                return "\n".join(els) + "\n"
            if a.group_by:
                groups = {}
                for (f, c), b in zip(sel[:k], blocks[:k]):
                    for g in _group_keys(f, c, a.group_by):
                        groups.setdefault(g, []).append(b)
                data = {g: groups[g] for g in sorted(groups)}
            else:
                data = blocks[:k]
            return json.dumps(data, indent=1, ensure_ascii=False) + "\n"

        def dropped_line(k):
            ids = [f["id"] for f, _ in sel[k:]]
            low = sel[k - 1][0]["rating"] if k else "none"
            return (f"Dropped for budget: {_n(len(ids), 'finding')} (lowest kept rating {low}): "
                    f"{', '.join(ids[:20])}{'…' if len(ids) > 20 else ''}")

        def footers_for(k):
            out = [dropped_line(k)] if k < nsel else []
            return out + ([note] if a.exclude_downgraded else [])

        kept = nsel
        if a.budget_chars is not None:
            kept = _fit_count(nsel, lambda k: len(render(k)) <= a.budget_chars)
        text = render(kept)
        if kept < nsel:
            print(dropped_line(kept), file=sys.stderr)
        if kept == 0 and nsel and a.budget_chars is not None and len(text) > a.budget_chars:
            print(f"select: budget {a.budget_chars} is below the minimum output ({len(text)} chars); "
                  "no findings kept", file=sys.stderr)
        if not md and a.exclude_downgraded:
            print(note, file=sys.stderr)
    if a.count_by and a.exclude_downgraded:
        if a.format == "json":
            print(note, file=sys.stderr)
        else:
            text += note + "\n"
    if a.out:
        Path(a.out).write_text(text, encoding="utf-8")
        print(f"Wrote {_n(kept, 'finding')} to {a.out}")
    else:
        sys.stdout.write(text)
    print(f"select: {_n(kept, 'finding')}, {len(text)} chars", file=sys.stderr)


# ----------------------------------------------------------------------------- causal

def cmd_causal(a):
    ws = Path(a.ws)
    w = load_ws(ws)
    F = w["findings"]
    present = {f["code"]: f["level"] for f in F}
    names = {f["code"]: f["name"] for f in F}
    edges = collections.Counter()
    evid = collections.defaultdict(list)
    for f in F:
        text = f["causal"] + " " + (f["justification"] if a.include_justification else "")
        refs = {norm_code(p, n) for p, n in CODE_RX.findall(text)} - {f["code"]}
        for r in refs:
            if r not in present:
                continue
            a_, b_ = f["code"], r
            la, lb = present[a_] or 0, present[b_] or 0
            # Convention: higher HFACS level (organisational) -> lower level (acts) = "enables".
            src, dst = (a_, b_) if la >= lb else (b_, a_)
            edges[(src, dst)] += 1
            evid[(src, dst)].append(f["id"])
    for row in w.get("incident", {}).get("causal_chain_sheet", []):
        k = (row.get("From Code"), row.get("To Code"))
        edges[k] += 1
        evid[k].append("causal-chain-sheet")
    items = [{"from": s, "to": t, "from_level": present.get(s), "to_level": present.get(t),
              "mentions": c, "findings": evid[(s, t)][:12]} for (s, t), c in edges.most_common()]
    write_json(ws / "causal_graph.json", {"nodes": [{"code": c, "name": names[c], "level": present[c]}
                                                    for c in sorted(present)], "edges": items})
    top = items[: a.top]
    M = ["flowchart TB"]
    used = {e["from"] for e in top} | {e["to"] for e in top}
    for lv in (4, 3, 2, 1):
        nodes = sorted(c for c in used if present.get(c) == lv)
        if nodes:
            M.append(f'  subgraph L{lv}["L{lv} {LEVEL_NAMES[lv]}"]')
            for c in nodes:
                M.append(f'    {c.replace(" ", "_").replace(".", "_")}["{c}<br/>{html.escape(names[c][:40])}"]')
            M.append("  end")
    for e in top:
        s = e["from"].replace(" ", "_").replace(".", "_")
        t = e["to"].replace(" ", "_").replace(".", "_")
        M.append(f"  {s} -->|{e['mentions']}| {t}")
    (ws / "causal_graph.mmd").write_text("\n".join(M), encoding="utf-8")
    ins = {n: file_sha256(ws / n) for n in ("findings.json", "incident.json") if (ws / n).exists()}
    for n in ("causal_graph.json", "causal_graph.mmd"):
        record_derived(ws, n, ins)
    linked = sum(1 for f in F if ({norm_code(p, n) for p, n in CODE_RX.findall(f['causal'])} - {f['code']}) & set(present))
    print(f"{linked} of {len(F)} findings name other factors in their causal explanation. "
          f"{len(items)} distinct links. Wrote causal_graph.json and causal_graph.mmd (top {a.top}).")
    print("\n| From | To | Mentions |\n|---|---|---:|")
    for e in top[:25]:
        print(f"| {e['from']} (L{e['from_level']}) | {e['to']} (L{e['to_level']}) | {e['mentions']} |")


# ----------------------------------------------------------------------------- timeline

MONTHS = "jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec"
DATE_RXS = [
    re.compile(r"\b(\d{1,2})[./-](\d{1,2})[./-](\d{4})\b"),
    re.compile(rf"\b(\d{{1,2}})[ -]({MONTHS})[a-z]*[ -,]*(\d{{4}})\b", re.I),
    re.compile(rf"\b({MONTHS})[a-z]*\.?\s+(\d{{4}})\b", re.I),
    re.compile(r"\b(19[5-9]\d|20[0-4]\d)\b"),
]
MIDX = {m: i + 1 for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"])}


def _parse_dates(text):
    hits, spans = [], []
    for i, rx in enumerate(DATE_RXS):
        for m in rx.finditer(text):
            if any(s <= m.start() < e for s, e in spans):
                continue
            try:
                if i == 0:
                    d = dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
                elif i == 1:
                    d = dt.date(int(m.group(3)), MIDX[m.group(2).lower()[:3]], int(m.group(1)))
                elif i == 2:
                    d = dt.date(int(m.group(2)), MIDX[m.group(1).lower()[:3]], 1)
                else:
                    d = dt.date(int(m.group(1)), 1, 1)
            except (ValueError, KeyError):
                continue
            spans.append((m.start(), m.end()))
            hits.append((d, i, m.start(), m.end()))
    return hits


CLOCK_RX = re.compile(r"\b([01]?\d|2[0-3])[:.]([0-5]\d)(?::[0-5]\d)?\s*(?:hours?\b|hrs?\b|h\b)?")
PRECISION_RANK = {"day": 0, "month": 1, "year": 2}


def months_between(a: dt.date, b: dt.date) -> int:
    """Whole months from a to b (negative when a is after b)."""
    return (b.year - a.year) * 12 + b.month - a.month - (b.day < a.day)


def _clock_times(text, start=0, end=None):
    """Clock times in text[start:end] as (match start, match end, 'HH:MM'). A '.' separator needs an hrs/h suffix, so
    decimals and dd.mm.yyyy dates are not read as times."""
    out = []
    for m in CLOCK_RX.finditer(text, start, len(text) if end is None else end):
        if text[m.end(1)] == "." and not re.search(r"h(ours?|rs?)?$", m.group(0).rstrip()):
            continue
        out.append((m.start(), m.end(), f"{int(m.group(1)):02d}:{m.group(2)}"))
    return out


def _event_day_sequence(w, day, sources):
    """One entry per minute: {time, mentions, sources, snippets}, sorted by time.

    A time counts when its sentence holds no day-precision date other than `day`. A sentence with no date at all
    counts only when `day` is the incident date, since undated times in an export describe the incident day."""
    found = {}
    undated_ok = day == incident_date(w)
    for sid, text in sources:
        cuts = [0] + [m.end() for m in re.finditer(r"\. |\n", text)] + [len(text)]
        for s, e, hhmm in _clock_times(text):
            i = max(j for j in range(len(cuts) - 1) if cuts[j] <= s)
            sentence = text[cuts[i]: cuts[i + 1]]
            days = [d.isoformat() for d, prec, _, _ in _parse_dates(sentence) if prec <= 1]
            if any(d != day for d in days) or (not days and not undated_ok):
                continue
            ev = found.setdefault(hhmm, {"time": hhmm, "mentions": 0, "sources": [], "snippets": []})
            ev["mentions"] += 1
            if sid not in ev["sources"]:
                ev["sources"].append(sid)
            snip = re.sub(r"\s+", " ", text[max(0, s - 100): e + 100]).strip()
            if len(ev["snippets"]) < 3 and snip not in ev["snippets"]:
                ev["snippets"].append(snip)
    return sorted(found.values(), key=lambda x: x["time"])


def cmd_timeline(a):
    ws = Path(a.ws)
    w = load_ws(ws)
    inc_date = incident_date(w)
    before = a.before
    if before and before != "incident" and not iso_date(before):
        print('--before: use YYYY-MM-DD or "incident"', file=sys.stderr)
        sys.exit(2)
    if before == "incident":
        before = inc_date
        if not before:
            sys.exit("No incident date in this workspace; pass --date to ingest.")
    day = a.on
    if day == "incident":
        day = inc_date
        if not day:
            sys.exit("No incident date in this workspace; pass --date to ingest.")
    elif day:
        try:
            dt.date.fromisoformat(day)
        except ValueError:
            sys.exit("--on takes a YYYY-MM-DD date or the word incident.")
    findings = w["findings"]
    evidence = w.get("evidence", {})
    if a.entity:
        cl = w.get("clusters", {}).get("findings")
        if cl is None:
            sys.exit("Needs clusters.json. Run: pmlc.py classify <ws>")
        keep = {f["id"] for f in findings if set(a.entity) & set(cl.get(f["id"], {}).get("F", []))}
        keep_ev = {e for f in findings if f["id"] in keep for e in f["evidence"]}
        findings = [f for f in findings if f["id"] in keep]
        evidence = {k: v for k, v in evidence.items() if v["id"] in keep_ev}
    events = collections.OrderedDict()
    sources = [(f["id"], f["code"], " ".join([f["label"], f["short"], f["justification"], f["causal"]]))
               for f in findings]
    sources += [(e["id"], e["type"], e["description"]) for e in evidence.values()]
    min_rank = PRECISION_RANK[a.min_precision]
    for sid, kind, text in sources:
        for d, prec, s, e in _parse_dates(text):
            if before and d.isoformat() > before:
                continue
            precision = ["day", "day", "month", "year"][prec]
            if PRECISION_RANK[precision] > min_rank:
                continue
            snippet = text[max(0, s - 140): e + 140].replace("\n", " ").strip()
            key = (d.isoformat(), prec, re.sub(r"\W+", " ", text[max(0, s - 40): e + 40]).lower()[:60])
            ev = events.setdefault(key, {"date": d.isoformat(), "precision": precision,
                                         "snippet": snippet, "sources": []})
            if precision == "day" and "time" not in ev:
                times = _clock_times(text, e, e + 40)
                if times:
                    ev["time"] = times[0][2]
            if sid not in ev["sources"]:
                ev["sources"].append(sid)
    items = sorted(events.values(), key=lambda x: (x["date"], x["precision"]))
    # collapse near-duplicate snippets per date
    merged = collections.OrderedDict()
    for it in items:
        k = (it["date"], it["precision"])
        if k in merged and len(merged[k]) >= a.per_date:
            continue
        merged.setdefault(k, []).append(it)
    flat = [x for v in merged.values() for x in v]
    inc = dt.date.fromisoformat(inc_date) if inc_date else None
    for it in flat:
        d = dt.date.fromisoformat(it["date"])
        if it["precision"] == "year":
            it["months_before"] = None
            it["years_before"] = inc.year - d.year if inc else None
        else:
            it["months_before"] = months_between(d, inc) if inc else None
    write_json(ws / "timeline_candidates.json", flat)
    record_derived(ws, "timeline_candidates.json",
                   {n: file_sha256(ws / n) for n in ("findings.json", "evidence.json") if (ws / n).exists()})
    print(f"{len(flat)} dated precursor candidates written to {ws/'timeline_candidates.json'} "
          f"(the precursor-timeline skill curates these).\n")
    print("| Date | Precision | Sources | Snippet |\n|---|---|---|---|")
    for it in flat[: a.show]:
        when = it["date"] + (f" {it['time']}" if it.get("time") else "")
        print(f"| {when} | {it['precision']} | {', '.join(it['sources'][:4])} | {it['snippet'][:160]} |")
    if day:
        texts = [(x[0], x[2]) for x in sources]
        texts.append(("narrative", w.get("incident", {}).get("narrative", "")))
        seq = _event_day_sequence(w, day, texts)
        write_json(ws / "timeline_event_day.json", seq)
        record_derived(ws, "timeline_event_day.json",
                       {n: file_sha256(ws / n) for n in ("findings.json", "evidence.json", "incident.json")
                        if (ws / n).exists()})
        print(f"\n{len(seq)} clock times on {day} written to {ws/'timeline_event_day.json'}.\n")
        print("| Time | Mentions | Sources | Snippet |\n|---|---:|---|---|")
        for it in seq[: a.show]:
            print(f"| {it['time']} | {it['mentions']} | {', '.join(it['sources'][:4])} | {it['snippets'][0][:160]} |")


# ----------------------------------------------------------------------------- interfaces

def cmd_interfaces(a):
    ws = Path(a.ws)
    w = load_ws(ws)
    if "clusters" not in w:
        sys.exit("Needs clusters.json. Run: pmlc.py classify <ws>")
    cl = w["clusters"]
    names = cl["taxonomy"]["entities"]
    pair = collections.Counter()
    examples = collections.defaultdict(list)
    for fid, c in cl["findings"].items():
        ents = [e for e in c["F"] if e != "F99"]
        for i, x in enumerate(ents):
            for y in ents[i + 1:]:
                k = tuple(sorted((x, y)))
                pair[k] += 1
                examples[k].append(fid)
    code_ents = collections.defaultdict(set)
    for f in w["findings"]:
        code_ents[f["code"]].update(e for e in cl["findings"][f["id"]]["F"] if e != "F99")
    # Cross-entity causal mention: finding f (owned by entity x) names factor r in its causal
    # explanation, and r is owned (in this investigation) by entity y != x. Counted once per
    # (finding, referenced factor, entity pair).
    cross = collections.Counter()
    present = {f["code"] for f in w["findings"]}
    for f in w["findings"]:
        own = {e for e in cl["findings"][f["id"]]["F"] if e != "F99"}
        refs = {norm_code(p, n) for p, n in CODE_RX.findall(f["causal"])} & present - {f["code"]}
        seen = set()
        for r in refs:
            for x in own:
                for y in code_ents[r] - own:
                    k = tuple(sorted((x, y)))
                    if (r, k) not in seen:
                        seen.add((r, k))
                        cross[k] += 1
    ents = sorted({e for k in pair for e in k} | {e for k in cross for e in k})
    out = {"pairs": [{"a": k[0], "b": k[1], "shared_findings": pair[k], "cross_causal_mentions": cross.get(k, 0),
                      "examples": examples[k][:10]} for k in sorted(set(pair) | set(cross),
                                                                    key=lambda k: -(pair.get(k, 0) * 3 + cross.get(k, 0)))]}
    write_json(ws / "interfaces.json", out)
    record_derived(ws, "interfaces.json", {n: file_sha256(ws / n) for n in ("findings.json", "clusters.json")})
    print("| Interface | Shared findings | Cross-entity causal mentions |\n|---|---:|---:|")
    for p in out["pairs"][: a.top]:
        print(f"| {p['a']} {names.get(p['a'], '')} ↔ {p['b']} {names.get(p['b'], '')} | "
              f"{p['shared_findings']} | {p['cross_causal_mentions']} |")
    print(f"\nWrote {ws/'interfaces.json'} ({len(out['pairs'])} pairs across {len(ents)} entities).")


# ----------------------------------------------------------------------------- dashboard

DASHBOARD_ASSETS = HERE / "dashboard_assets"


EVD_RX = re.compile(r"\bEVD-\d{3,}\b")


def linked_evidence(f: dict) -> set:
    """Evidence a finding is linked to: its Linked Evidence column plus the EVD IDs its AI Justification names."""
    return set(f["evidence"]) | set(EVD_RX.findall(f["justification"]))


def downgraded_findings(w: dict) -> set:
    """Finding IDs with a reviewer vote to downgrade, reject or remove them (incident.json review_votes)."""
    return {v.get("finding") for v in w.get("incident", {}).get("review_votes", [])
            if _s(v.get("vote")).lower() in ("downgrade", "reject", "remove")}


def citation_records(w: dict, finding_ids, evidence_ids=()) -> tuple:
    """Full workspace records for the given findings (each with its clusters and reviewer votes), and for the
    given evidence plus every evidence item those findings link. Unknown IDs are skipped; first-seen order kept.
    Shared by the Detailed register tab and learning packs so both embed identical records."""
    F = {f["id"]: f for f in w["findings"]}
    E = w.get("evidence", {})
    clf = w.get("clusters", {}).get("findings", {})
    votes = collections.defaultdict(list)
    for v in w.get("incident", {}).get("review_votes", []):
        votes[v.get("finding")].append(v)
    findings = {}
    for fid in finding_ids:
        if fid in F and fid not in findings:
            findings[fid] = dict(F[fid], clusters=clf.get(fid, {}), votes=votes.get(fid, []))
    eids = []
    for x in list(evidence_ids) + [x for f in findings.values() for x in (f.get("evidence") or [])]:
        if x not in eids:
            eids.append(x)
    return findings, {x: E[x] for x in eids if x in E}


def build_detail_data(w: dict, reg) -> dict:
    """Everything the Detailed register tab shows: the register as given, plus the full workspace
    records of every finding and evidence item it cites (each once) and the findings no learning covers."""
    entries = reg["entries"] if isinstance(reg, dict) else reg
    meta = {k: v for k, v in reg.items() if k != "entries"} if isinstance(reg, dict) else {}
    F = {f["id"]: f for f in w["findings"]}
    downs = downgraded_findings(w)

    derived, cited = {}, []
    for e in entries:
        hocs = [norm_hoc(r.get("hoc_tier")) for r in (e.get("recommendations") or [])]
        ranked = [t for t in HOC_TIERS if t in hocs]
        derived[_s(e.get("lr_id"))] = {"level": norm_level(e.get("hfacs_level")),
                                       "entity": norm_entity(e.get("owner_entity")),
                                       "hoc": hocs, "strongest": ranked[0] if ranked else "unclassified"}
        cited += [fid for fid in (e.get("finding_ids") or []) if fid in F and fid not in cited]
    uncovered = [{"id": fid, "label": f.get("label", ""), "rating": f.get("rating"), "level": f.get("level"),
                  "reason": "reviewer downgrade vote" if fid in downs else "not covered by any learning"}
                 for fid, f in F.items() if fid not in cited]

    findings, evidence = citation_records(w, cited + [u["id"] for u in uncovered],
                                          [x for e in entries for x in (e.get("evidence_ids") or [])])
    tax = w.get("clusters", {}).get("taxonomy", {})
    return {"register": meta, "entries": entries, "derived": derived, "findings": findings, "evidence": evidence,
            "taxonomy": {"entities": tax.get("entities", {}), "tiers": tax.get("tiers", {})},
            "brand": brand_payload(), "tier_labels": TIER_LABELS, "labels": dict(REG_COLS), "hoc_order": HOC_TIERS,
            "totals": {"findings": len(F), "evidence": len(w.get("evidence", {}))}, "uncovered": uncovered}


def embed_json(obj) -> str:
    """JSON safe to place inside <script type="application/json">: no '<', '>' or '&' characters."""
    s = json.dumps(obj, ensure_ascii=False, separators=(",", ":"))
    return s.replace("<", "\\u003c").replace(">", "\\u003e").replace("&", "\\u0026")


PAGE_TOKENS_CSS = """:root{--bg:#f7f8fa;--fg:#111827;--mut:#5b6472;--card:#fff;--line:#e3e6eb;--acc:#0077b6}
@media (prefers-color-scheme:dark){:root:not([data-theme=light]){--bg:#0d1117;--fg:#e6edf3;--mut:#9aa4b2;--card:#161b22;--line:#2a313c;--acc:#00bfff}}
:root[data-theme=dark]{--bg:#0d1117;--fg:#e6edf3;--mut:#9aa4b2;--card:#161b22;--line:#2a313c;--acc:#00bfff}
body{margin:0;background:var(--bg);color:var(--fg);font:15px/1.5 system-ui,-apple-system,Segoe UI,Roboto,sans-serif}"""
# Wide screens: containers use the space, prose keeps a readable measure (owner request 2026-10-15).
WIDE_CSS = "@media (min-width:1400px){main{max-width:min(calc(100vw - 32px),1600px)} main p,main li{max-width:75ch}}"


def cmd_dashboard(a):
    ws = Path(a.ws)
    w = load_ws(ws)
    if "clusters" not in w:
        sys.exit("Needs clusters.json. Run: pmlc.py classify <ws>")
    warn_stale(ws)
    F = {f["id"]: f for f in w["findings"]}
    cl = w["clusters"]
    names, tiers = cl["taxonomy"]["entities"], cl["taxonomy"]["tiers"]
    fields = dict(w.get("incident", {}).get("fields", {}))
    for k, v in (("Incident Title", a.title), ("Date", a.date), ("Location", a.location)):
        if v:
            fields[k] = v
    ent_ids = [e for e in cl["entities"] if e != "F99"] + (["F99"] if "F99" in cl["entities"] else [])
    tier_ids = ["T1", "T2", "T3", "T4", "T5", "T9"]
    et = collections.Counter()
    el = collections.Counter()
    for fid, c in cl["findings"].items():
        for e in c["F"]:
            el[(e, F[fid]["level"])] += 1
            for t in c["T"]:
                et[(e, t)] += 1
    mx_et = max(et.values() or [1])
    mx_el = max(el.values() or [1])
    lvl = collections.Counter(f["level"] for f in F.values())
    rat = collections.Counter(f["rating"] for f in F.values())
    top_codes = collections.Counter(f"{f['code']} · {f['name']}" for f in F.values()).most_common(12)

    def cell(v, mx):
        if not v:
            return '<td class="z">·</td>'
        op = 0.15 + 0.85 * v / mx
        return f'<td style="--o:{op:.2f}"><span>{v}</span></td>'

    def table(rows, cols, colnames, counter, mx):
        h = ["<table><thead><tr><th></th>" + "".join(f"<th>{html.escape(colnames[c])}</th>" for c in cols) + "</tr></thead><tbody>"]
        for r in rows:
            h.append(f"<tr><th>{r} · {html.escape(names.get(r, r))}</th>" + "".join(cell(counter.get((r, c), 0), mx) for c in cols) + "</tr>")
        h.append("</tbody></table>")
        return "".join(h)

    tier_lab = {t: f"{t} {tiers.get(t, '')}" for t in tier_ids}
    lvl_lab = {k: f"L{k} {LEVEL_NAMES[k]}" for k in (1, 2, 3, 4)}
    n = len(F)
    up = lvl.get(3, 0) + lvl.get(4, 0)
    title = html.escape(fields.get("Incident Title", "ProbeMatrix investigation"))
    top = max((r for r in rat if r is not None), default=None)
    kpis = [("Findings", n), ("Evidence items", len(w.get("evidence", {}))),
            ("Learning clusters", len([e for e in ent_ids if e != "F99"])),
            ("At supervision + organisational level", f"{100*up/max(n,1):.0f}%"),
            ("Highest-rated findings", rat[top] if top is not None else 0)]
    bars = "".join(f'<div class="bar"><span>{html.escape(k)}</span><i style="width:{100*c/top_codes[0][1]:.0f}%"></i><b>{c}</b></div>'
                   for k, c in top_codes) if top_codes else ""
    learn_html = ""
    detail = None
    hdoc = {}
    if a.register:
        reg = json.loads(Path(a.register).read_text(encoding="utf-8"))
        hdoc = reg if isinstance(reg, dict) else {}
        entries = reg["entries"] if isinstance(reg, dict) else reg
        entries = sorted(entries, key=lambda e: (_int(e.get("priority"), 9), _s(e.get("lr_id"))))
        detail = build_detail_data(w, reg)
        recs = [r for e in entries for r in (e.get("recommendations") or [])]
        strong = sum(1 for r in recs if norm_hoc(r.get("hoc_tier")) in ("elimination", "substitution", "engineering"))
        kpis = [("Learnings", len(entries)),
                ("Priority-1 learnings", sum(1 for e in entries if _int(e.get("priority")) == 1)),
                ("Recommendations", len(recs)),
                ("Fixes above administrative tier", f"{100*strong/max(len(recs),1):.0f}%")] + kpis
        tiers_used = [t for t in TIER_LABELS if any(t in (e.get("tier_messages") or {}) for e in entries)]
        rows = []
        for e in entries:
            tm = e.get("tier_messages") or {}
            cells = "".join(
                f'<td style="--o:.55" title="{html.escape(tm[t])}"><span>●</span></td>' if t in tm else '<td class="z">·</td>'
                for t in tiers_used)
            lid = html.escape(_s(e.get("lr_id")))
            rows.append(f'<tr><th><a class="lr-link" href="#detail/{lid}">{lid} · {html.escape(_s(e.get("title"))[:90])}</a></th>'
                        f"<td>{html.escape(_s(e.get('priority')))}</td><td>{html.escape(norm_entity(e.get('owner_entity')))}</td>{cells}</tr>")
        matrix = ("<table><thead><tr><th>Learning</th><th>Priority</th><th>Owner</th>"
                  + "".join(f"<th>{t} {TIER_LABELS[t]}</th>" for t in tiers_used) + "</tr></thead><tbody>"
                  + "".join(rows) + "</tbody></table>")
        per_tier = []
        for t in tiers_used:
            items = "".join(f"<li><b>{html.escape(_s(e.get('lr_id')))}</b> {html.escape(e['tier_messages'][t])}</li>"
                            for e in entries if t in (e.get("tier_messages") or {}))
            n = sum(1 for e in entries if t in (e.get("tier_messages") or {}))
            per_tier.append(f"<details><summary>{t} {TIER_LABELS[t]} · {n} learnings</summary><ul>{items}</ul></details>")
        hoc = collections.Counter(norm_hoc(r.get("hoc_tier")) for r in recs)
        hoc_bars = "".join(f'<div class="bar"><span>{t}</span><i style="width:{100*hoc[t]/max(hoc.values()):.0f}%"></i><b>{hoc[t]}</b></div>'
                           for t in HOC_TIERS + ["unclassified"] if hoc.get(t))
        learn_html = f"""<h2>Learnings and who must act on each</h2>
<p class="m">One row per learning from the Learnings Register; ● means that tier has a message (hover to read it).</p>
<div class="wrap">{matrix}</div>
<h2>Learning plan by level</h2>
<div class="wrap tiers">{''.join(per_tier)}</div>
<h2>Strength of recommended fixes</h2>
<div class="wrap" style="padding:12px">{hoc_bars}</div>"""
    overview = f"""<div class="k">{''.join(f'<div><b>{v}</b><span>{html.escape(k)}</span></div>' for k, v in kpis)}</div>
{learn_html}
<h2>Who learns, at which level of the organisation</h2>
<p class="m">Findings touching each functional entity, by organisational tier. A finding can involve several entities and tiers.</p>
<div class="wrap">{table(ent_ids, tier_ids, tier_lab, et, mx_et)}</div>
<h2>What kind of learning each entity owns</h2>
<p class="m">Findings by four-level human-factors category. Rows weighted to Supervision and Organisational Influences need system changes, not retraining.</p>
<div class="wrap">{table(ent_ids, [1, 2, 3, 4], lvl_lab, el, mx_el)}</div>
<h2>Most frequent factors</h2>
<div class="wrap" style="padding:12px">{bars}</div>"""
    if detail is not None:
        body = f"""<nav class="tabbar" role="tablist" aria-label="Views">
<button type="button" role="tab" id="t-overview" aria-controls="tab-overview" aria-selected="true">Overview</button>
<button type="button" role="tab" id="t-detail" aria-controls="tab-detail" aria-selected="false" tabindex="-1">Detailed register</button>
</nav>
<section id="tab-overview" role="tabpanel" aria-labelledby="t-overview">
{overview}
</section>
<section id="tab-detail" role="tabpanel" aria-labelledby="t-detail" hidden>
<noscript><p class="m">The detailed register needs JavaScript. The same data is in learnings-register.xlsx and learnings-register.json.</p></noscript>
</section>"""
        extra_css = (DASHBOARD_ASSETS / "detail.css").read_text(encoding="utf-8")
        scripts = (f'<script type="application/json" id="pmlc-data">{embed_json(detail)}</script>\n'
                   f'<script>\n{(DASHBOARD_ASSETS / "detail.js").read_text(encoding="utf-8")}\n</script>')
    else:
        body = (overview + '\n<p class="m">Run <code>pmlc.py dashboard --register &lt;learnings-register.json&gt;</code> '
                "to add the detailed register tab.</p>")
        extra_css = scripts = ""
    eh = effective_handling(hdoc, w)
    ribbon = RIBBON_HTML + "\n" if is_draft(eh) else ""
    handling = "".join(f'<p class="handling">{html.escape(x)}</p>' for x in handling_lines(eh))
    doc = f"""<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Organisational Learning Map: {title}</title>
<style>
{PAGE_TOKENS_CSS}
main{{max-width:1180px;margin:0 auto;padding:24px 16px 48px}}
{WIDE_CSS}
h1{{font-size:24px;margin:0 0 4px}} h2{{font-size:17px;margin:32px 0 8px}} p.m{{color:var(--mut);margin:0 0 16px}}
.k{{display:grid;grid-template-columns:repeat(auto-fit,minmax(170px,1fr));gap:12px}}
.k div{{background:var(--card);border:1px solid var(--line);border-radius:10px;padding:12px 14px}}
.k b{{display:block;font-size:24px}} .k span{{color:var(--mut);font-size:13px}}
.wrap{{overflow-x:auto;background:var(--card);border:1px solid var(--line);border-radius:10px}}
table{{border-collapse:collapse;width:100%;font-size:13px}} th,td{{padding:6px 8px;border-bottom:1px solid var(--line);text-align:center}}
tbody th{{text-align:left;font-weight:500;white-space:nowrap}} thead th{{font-weight:600;font-size:12px;color:var(--mut)}}
td{{background:color-mix(in srgb,var(--acc) calc(var(--o,0)*100%),transparent)}} td span{{font-weight:600}} td.z{{color:var(--mut);background:none}}
.bar{{display:grid;grid-template-columns:minmax(160px,340px) 1fr 40px;gap:8px;align-items:center;margin:4px 0;font-size:13px}}
.bar i{{display:block;height:12px;background:var(--acc);border-radius:3px}} .bar b{{text-align:right}}
.tiers details{{padding:8px 12px;border-bottom:1px solid var(--line)}} .tiers summary{{cursor:pointer;font-weight:600}} .tiers li{{margin:4px 0}}
footer{{margin-top:40px;color:var(--mut);font-size:12px;border-top:1px solid var(--line);padding-top:12px}}
{HANDLING_CSS}
{extra_css}
</style></head><body><main>
{ribbon}<h1>Organisational Learning Map</h1>
<p class="m"><b>{html.escape(status_line())}</b></p>
{handling}
<p class="m">{title} · {html.escape(fields.get('Location', ''))} · {html.escape(fields.get('Date', ''))}. Learnings from one ProbeMatrix investigation, distributed across the organisation by who owns them.</p>
{body}
<footer>{html.escape(BRAND)}</footer>
</main>
{scripts}
</body></html>"""
    out = Path(a.out) if a.out else outputs_dir(ws) / "organisation-learning-map.html"
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(doc, encoding="utf-8")
    print(f"Dashboard written to {out}")


# ----------------------------------------------------------------------------- register-xlsx

REG_COLS = [
    ("lr_id", "LR-ID"), ("ml_id", "ML-ID"), ("title", "Title"), ("owner_entity", "Learning cluster (entity)"),
    ("tiers", "Tiers who must learn"), ("lesson", "Lesson statement"), ("context", "Context narrative"),
    ("finding_ids", "Finding IDs"), ("evidence_ids", "Evidence IDs"), ("evidence_strength", "Evidence strength"),
    ("confidence", "Confidence"), ("hfacs_level", "HFACS level / sub-category"), ("causal_role", "Causal role"),
    ("barrier", "Barrier (prevention / mitigation)"), ("hazard", "Hazard / energy"), ("activity", "Activity / phase"),
    ("performance_mode", "Human-performance mode"), ("org_system", "Organisational system"),
    ("pattern", "Cross-incident / recurrence pattern"), ("precursors", "Precursors and weak signals"),
    ("comparison", "Counterpart in comparator report"), ("reason_class", "Reason class (if none)"),
    ("worst_case", "Credible worst case"), ("transferability", "Transferability"), ("priority", "Priority"),
    ("audiences", "Audiences"), ("learning_level", "Learning level / loop"), ("formats", "Formats"),
]
REC_COLS = [("lr_id", "LR-ID"), ("what", "What should change"), ("level_barrier", "HFACS level and barrier"),
            ("evidence_basis", "Evidence basis"), ("hoc_tier", "Hierarchy-of-controls tier"),
            ("leading_indicator", "Leading indicator")]


TIER_LABELS = {"T1": "Frontline", "T2": "Supervisors", "T3": "Function managers", "T4": "Site leadership",
               "T5": "Corporate / board"}
HOC_TIERS = ["elimination", "substitution", "engineering", "administrative", "training", "ppe"]


def norm_level(v):
    m = re.search(r"\bL\s?([1-4])\b|level\s?([1-4])", _flat(v), re.I)
    return f"L{m.group(1) or m.group(2)} {LEVEL_NAMES[int(m.group(1) or m.group(2))]}" if m else "unclassified"


def norm_hoc(v):
    s = _flat(v).lower()
    for t in HOC_TIERS:
        if t in s or (t == "ppe" and "protective equipment" in s) or (t == "administrative" and "procedur" in s):
            return t
    return "unclassified"


def norm_entity(v):
    s = _flat(v)
    m = re.match(r"\s*(F\d{2}[a-z]?)\b", s)
    return m.group(1) if m else (s or "unset")


def _flat(v):
    if isinstance(v, list):
        return "; ".join(_flat(x) for x in v)
    if isinstance(v, dict):
        return "; ".join(f"{k}: {_flat(x)}" for k, x in v.items())
    return "" if v is None else str(v)


def _register_ws(a) -> dict:
    """The workspace for a register's forced-draft check: --ws, else the run's workspace when the register sits in
    <run>/outputs/ of the standard layout, else none (no forced-draft reasons can be read)."""
    if getattr(a, "ws", None):
        return load_ws(a.ws)
    guess = Path(a.register).resolve().parent.parent / "workspace"
    return load_ws(guess) if is_standard_ws(guess) and (guess / "findings.json").exists() else {}


def cmd_register_xlsx(a):
    openpyxl = _need_openpyxl()
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    reg = json.loads(Path(a.register).read_text(encoding="utf-8"))
    entries = reg["entries"] if isinstance(reg, dict) else reg
    eh = effective_handling(reg if isinstance(reg, dict) else {}, _register_ws(a))
    wb = openpyxl.Workbook()
    head_fill = PatternFill("solid", fgColor="0B3954")
    head_font = Font(color="FFFFFF", bold=True)

    def sheet(ws, cols, rows):
        ws.append([c[1] for c in cols])
        for r in rows:
            ws.append([_flat(r.get(c[0])) for c in cols])
        for i, c in enumerate(cols, 1):
            ws.cell(1, i).fill = head_fill
            ws.cell(1, i).font = head_font
            ws.column_dimensions[get_column_letter(i)].width = 18 if len(c[1]) < 14 else 32
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "C2"

    ws = wb.active
    ws.title = "Learnings Register"
    sheet(ws, REG_COLS, entries)
    recs = [dict(r, lr_id=e.get("lr_id")) for e in entries for r in (e.get("recommendations") or [])]
    sheet(wb.create_sheet("Recommendations"), REC_COLS, recs)
    plan_cols = [("lr_id", "LR-ID"), ("title", "Title"), ("priority", "Priority")] + \
                [(t, f"{t} {TIER_LABELS[t]}") for t in TIER_LABELS]
    plan_rows = [dict({"lr_id": e.get("lr_id"), "title": e.get("title"), "priority": e.get("priority")},
                      **(e.get("tier_messages") or {})) for e in entries]
    sheet(wb.create_sheet("Learning Plan"), plan_cols, plan_rows)
    s = wb.create_sheet("Summary")
    s.append(["Count by", "Value", "Count"])
    for c in s[1]:
        c.fill, c.font = head_fill, head_font
    groups = [
        ("Learnings by cluster", collections.Counter(norm_entity(e.get("owner_entity")) for e in entries)),
        ("Learnings by HFACS level", collections.Counter(norm_level(e.get("hfacs_level")) for e in entries)),
        ("Learnings by priority", collections.Counter(_flat(e.get("priority")) or "unset" for e in entries)),
        ("Learnings by counterpart", collections.Counter(_flat(e.get("comparison")) or "not compared" for e in entries)),
        ("Learnings by tier that must act", collections.Counter(t for e in entries for t in (e.get("tier_messages") or {}))),
        ("Recommendations by control tier", collections.Counter(norm_hoc(r.get("hoc_tier")) for r in recs)),
    ]
    for lab, cnt in groups:
        for v, c in sorted(cnt.items(), key=lambda kv: (str(kv[0]))):
            s.append([lab, v, c])
    strong = sum(1 for r in recs if norm_hoc(r.get("hoc_tier")) in ("elimination", "substitution", "engineering"))
    s.append(["Recommendations above administrative tier", f"{strong} of {len(recs)}",
              f"{100*strong/max(len(recs),1):.0f}%"])
    s.column_dimensions["A"].width, s.column_dimensions["B"].width, s.column_dimensions["C"].width = 38, 46, 10
    ab = wb.create_sheet("About")
    for line in [reg.get("title", "Learnings Register") if isinstance(reg, dict) else "Learnings Register",
                 reg.get("source", "") if isinstance(reg, dict) else "",
                 "Lessons and recommendations only: no owners, due dates or action tracking.",
                 STATUS_PHRASE, BRAND]:
        ab.append([line])
    review = review_text(eh) + (f". Draft because: {'; '.join(eh['draft_reasons'])}" if eh["draft_reasons"] else "")
    for k, v in [("Review status", review), ("Classification", eh["classification"]),
                 ("Legal review", eh["legal_review"]), ("Distribution", ", ".join(eh["distribution"]) or "not set"),
                 ("Version", str(eh["version"])), ("Retention review", eh.get("retention_review") or "not set")]:
        ab.append([k, v])
    ab.column_dimensions["A"].width = 120
    ab.column_dimensions["B"].width = 60
    for sh in wb.worksheets:
        sh.oddHeader.center.text = STATUS_PHRASE + (f" · {DRAFT_LABEL}" if is_draft(eh) else "")
        sh.oddFooter.center.text = VERSION_STAMP
    out = Path(a.out) if a.out else Path(a.register).resolve().parent / "learnings-register.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"Register with {len(entries)} entries and {len(recs)} recommendations written to {out}")


# ----------------------------------------------------------------------------- check-register

REGISTER_VOCAB = {
    "evidence_strength": ("direct", "corroborated", "inferred"),
    "confidence": ("high", "medium", "low"),
    "causal_role": ("initiating", "contributory", "escalating", "failed mitigation"),
    "transferability": ("same site", "other sites", "industry-wide"),
    "comparison": ("yes", "partial", "none", "not compared"),
}
PERFORMANCE_MODES = ("skill-based", "rule-based", "knowledge-based", "n/a")


def _vals(v) -> list:
    """Non-empty string values of a field; a list is taken element by element."""
    items = v if isinstance(v, (list, tuple)) else [v]
    return [x for x in (_s(i) for i in items if i is not None) if x]


def _fold_spelling(s) -> str:
    s = _s(s)
    return re.sub(r"\s+", " ", s.lower().replace("isation", "ization")).strip()


def register_warnings(entries: list, w: dict) -> list:
    """Controlled-vocabulary drift as 'warning: ...' lines. Absent or empty fields are not checked."""
    E = w.get("evidence", {})
    subs = {_fold_spelling(f["subcategory"]) for f in w.get("findings", []) if _s(f.get("subcategory"))}
    out = []
    for e in entries:
        lid = e.get("lr_id", "?")
        for field, allowed in REGISTER_VOCAB.items():
            for v in _vals(e.get(field)):
                if v.lower() not in allowed:
                    out.append(f"warning: {lid}: {field} {v!r} is not one of {', '.join(allowed)}")
        for v in _vals(e.get("performance_mode")):
            if not v.lower().startswith(PERFORMANCE_MODES):
                out.append(f"warning: {lid}: performance_mode {v!r} does not start with one of "
                           f"{', '.join(PERFORMANCE_MODES)}")
        tm = e.get("tier_messages")
        for k in (tm if isinstance(tm, dict) else {}):
            if k not in TIER_LABELS:
                out.append(f"warning: {lid}: tier_messages key {k!r} is not one of {', '.join(TIER_LABELS)}")
        bases = [e.get("evidence_basis")] + [r.get("evidence_basis") for r in (e.get("recommendations") or [])
                                              if isinstance(r, dict)]
        for eid in sorted({x for b in bases for v in _vals(b) for x in EVD_RX.findall(v)}):
            if eid not in E:
                out.append(f"warning: {lid}: evidence_basis names {eid}, which is not in the workspace")
        for lv in (_vals(e.get("hfacs_level")) if subs else []):
            m = re.match(r"\s*L\s?[1-4]\b[\s:.\-]*(.*)$", lv, re.I | re.S)
            rest = _fold_spelling(m.group(1)) if m else ""
            if rest and not any(rest == x or re.match(re.escape(x) + r"[\s(:]", rest) for x in subs):
                out.append(f"warning: {lid}: hfacs_level sub-category {lv!r} matches no sub-category in the export")
    return out

def cmd_check_register(a):
    w = load_ws(a.ws)
    reg = json.loads(Path(a.register).read_text(encoding="utf-8"))
    entries = reg["entries"] if isinstance(reg, dict) else reg
    F = {f["id"]: f for f in w["findings"]}
    E = w.get("evidence", {})
    cl = w.get("clusters", {}).get("findings", {})
    downs = downgraded_findings(w)
    problems, covered = [], set()
    for e in entries:
        lid = e.get("lr_id", "?")
        fids = e.get("finding_ids") or []
        if not fids:
            problems.append(f"{lid}: no finding_ids")
        for fid in fids:
            if fid not in F:
                problems.append(f"{lid}: unknown finding {fid}")
            else:
                covered.add(fid)
                if fid in downs:
                    problems.append(f"{lid}: cites {fid}, which has a reviewer downgrade vote")
        eids = e.get("evidence_ids") or []
        if not eids:
            problems.append(f"{lid}: no evidence_ids")
        linked = {x for fid in fids if fid in F for x in linked_evidence(F[fid])}
        for eid in eids:
            if eid not in E:
                problems.append(f"{lid}: unknown evidence {eid}")
            elif linked and eid not in linked:
                problems.append(f"{lid}: {eid} is not linked to any of its findings")
        if norm_level(e.get("hfacs_level")) == "unclassified":
            problems.append(f"{lid}: hfacs_level should start with L1-L4")
        if _s(e.get("priority")) not in ("1", "2", "3"):
            problems.append(f"{lid}: priority should be 1, 2 or 3")
        if not e.get("tier_messages"):
            problems.append(f"{lid}: no tier_messages")
        recs = e.get("recommendations") or []
        if not recs:
            problems.append(f"{lid}: no recommendations")
        for i, r in enumerate(recs, 1):
            if norm_hoc(r.get("hoc_tier")) == "unclassified":
                problems.append(f"{lid} rec {i}: hoc_tier not one of {HOC_TIERS}")
            if not _s(r.get("leading_indicator")):
                problems.append(f"{lid} rec {i}: no leading_indicator")
        # owner suggestion: entity with most findings in this learning
        cnt = collections.Counter(x for fid in fids for x in cl.get(fid, {}).get("F", []) if x != "F99")
        if cnt and norm_entity(e.get("owner_entity")) not in [k for k, _ in cnt.most_common(3)]:
            problems.append(f"{lid}: owner {norm_entity(e.get('owner_entity'))} is not among the top entities of its "
                            f"findings ({', '.join(k for k, _ in cnt.most_common(3))})")
    lint_hits = _lint_module().lint(reg, w, kind="register", **_lint_module().cli_options(a))
    problems += [_lint_module().as_problem(h) for h in lint_hits if h.severity == "error"]
    uncovered = [fid for fid in F if fid not in covered and fid not in downs]
    hp, hw = check_handling(reg)
    problems += hp
    print(f"{_n(len(entries), 'learning')}; {len(covered)} of {_n(len(F), 'finding')} covered "
          f"({100*len(covered)/max(len(F),1):.0f}%); {len(problems)} problem(s).")
    for p in problems[:80]:
        print(" -", p)
    if uncovered:
        print(f"Uncovered findings ({len(uncovered)}), highest rated first:")
        for fid in sorted(uncovered, key=lambda x: -(F[x]['rating'] or 0))[:20]:
            print(f"   {fid} {F[fid]['code']} r{F[fid]['rating']} {F[fid]['short'][:90]}")
    for line in register_warnings(entries, w):
        print(line)
    for line in hw:
        print(f"warning: {line}")
    for h in lint_hits:
        if h.severity != "error":
            print("warning: " + _lint_module().as_problem(h))
    sys.exit(1 if problems else 0)


# ----------------------------------------------------------------------------- scan-names

TEXT_TYPES = (".md", ".txt", ".json", ".csv", ".html", ".mmd", ".py", ".yaml", ".yml")


def _scan_files(path: Path):
    if path.is_dir():
        return [p for p in sorted(path.rglob("*")) if p.is_file() and not any(x.startswith(".git") for x in p.parts)]
    return [path]


_OFFICE_ZIP = (".docx", ".pptx", ".xlsx", ".xlsm")
_MAX_MEMBER = 50 * 1024 * 1024
_XML_BREAK = re.compile(r"<(?:w:(?:tab|br|cr)|a:br)\b[^>]*>")
_XML_PARA_END = re.compile(r"</(?:w|a):p>")
_XML_RUN_TAG = re.compile(r"</?(?:w|a):[^>]*>")
_XML_TAG = re.compile(r"<[^>]+>")
# Parts that can carry a name beyond the cell/body text: properties, comments, people, headers/footers, notes.
_META_PART = re.compile(r"(?i)^docProps/.*\.xml$|(?:^|/)(?:threaded|modern)?comments?[^/]*\.xml$"
                        r"|(?:^|/)modernComment_[^/]*\.xml$|^word/people\.xml$|^ppt/authors\.xml$|^xl/persons/person[^/]*\.xml$"
                        r"|/(?:header|footer)\d*\.xml$|/(?:footnotes|endnotes|notesSlide\d*)\.xml$")
_BODY_PART = re.compile(r"(?i)^(?:word/document|ppt/slides/slide\d+)\.xml$")


_NAME_ATTR = re.compile(r"""(?<![\w.:-])(?:[\w.-]+:)?(?:author|name|displayName|initials|creator|lastModifiedBy)\s*=\s*(?:"([^"]*)"|'([^']*)')""",
                        re.I)


def _xml_attr_names(raw: str) -> str:
    """The values of author, name, initials, creator and lastModifiedBy attributes (any namespace prefix), one per
    line: comment and people parts often hold a name only in an attribute."""
    return "\n".join(html.unescape(m.group(1) if m.group(1) is not None else m.group(2))
                      for m in _NAME_ATTR.finditer(raw))


def _xml_text(raw: str) -> str:
    """Plain text of an Office XML part. Run text inside a paragraph is joined with no separator, so a name split
    across runs ('Ra' + 'vi Sharma') is whole again; a paragraph end is a newline; other tags become a space."""
    raw = _XML_BREAK.sub(" ", raw)
    raw = _XML_PARA_END.sub("\n", raw)
    raw = _XML_RUN_TAG.sub("", raw)
    return html.unescape(_XML_TAG.sub(" ", raw))


def _office_xml_texts(path: Path, unread: list):
    """(location, text) for the metadata, comment, header/footer parts of an Office zip, plus the body of a .docx or
    .pptx (an .xlsx body is read cell by cell). A file that cannot be read, or has an oversize member, goes on
    the unread list so it is never reported as clean."""
    import zipfile
    body = path.suffix.lower() in (".docx", ".pptx")
    try:
        with zipfile.ZipFile(path) as z:
            for zi in z.infolist():
                n = zi.filename
                if not (_META_PART.search(n) or (body and _BODY_PART.search(n))):
                    continue
                if zi.file_size > _MAX_MEMBER:
                    print(f"warning: {path}:{n} is over 50 MB and was skipped")
                    unread.append(str(path))
                    continue
                raw = z.read(n).decode("utf-8", errors="ignore")
                text = _xml_text(raw)
                if _META_PART.search(n):
                    text += "\n" + _xml_attr_names(raw)
                yield f"{path}:{n}", text
    except Exception as e:
        print(f"warning: {path} could not be read ({type(e).__name__})")
        unread.append(str(path))


def _iter_texts(path: Path, unread: list):
    if path.is_dir():
        for p in _scan_files(path):
            yield from _iter_texts(p, unread)
        return
    ext = path.suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        try:
            openpyxl = _need_openpyxl()
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            cells = [(f"{path}:{ws.title}!{c.coordinate}", c.value) for ws in wb.worksheets
                     for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
        except SystemExit:
            raise
        except Exception as e:
            print(f"warning: {path} could not be read ({type(e).__name__})")
            unread.append(str(path))
            cells = []
        yield from cells
    if ext in _OFFICE_ZIP:
        yield from _office_xml_texts(path, unread)
    elif ext in TEXT_TYPES:
        for i, line in enumerate(path.read_text(encoding="utf-8", errors="ignore").splitlines(), 1):
            yield f"{path}:{i}", line


def _pdf_meta(path: Path):
    """(location, text) for the Author and Creator lines of a PDF, from pdfinfo. None when pdfinfo is absent,
    False when it failed or timed out."""
    import shutil
    import subprocess
    exe = shutil.which("pdfinfo")
    if not exe:
        return None
    try:
        r = subprocess.run([exe, str(path.resolve())], capture_output=True, text=True, errors="ignore", timeout=30)
    except (subprocess.TimeoutExpired, OSError):
        return False
    if r.returncode != 0:
        return False
    out = []
    for line in r.stdout.splitlines():
        k, _, v = line.partition(":")
        if k.strip() in ("Author", "Creator"):
            out.append((f"{path}:pdf:{k.strip()}", v.strip()))
    return out


def cmd_scan_names(a):
    names = privacy.load_names(a.names)
    if not names:
        sys.exit("Names file is empty.")
    scans = [("", privacy.name_pattern(names, variants=a.variants))]
    if a.patterns:
        scans += [(f"pattern:{k}", re.compile(v)) for k, v in privacy.PATTERNS.items()]
    hits = 0

    def check(loc, text):
        nonlocal hits
        for label, rx in scans:
            for m in rx.finditer(text):
                hits += 1
                if hits <= a.show:
                    tag = f" [{label}]" if label else ""
                    print(f"HIT{tag} {loc}: '{m.group(0)}' … {text[max(0, m.start()-50):m.end()+50]!r}")

    pdf_missing = False
    bad = []
    for target in a.paths:
        for loc, text in _iter_texts(Path(target), bad):
            check(loc, text)
        for f in _scan_files(Path(target)):
            if f.suffix.lower() == ".pdf":
                meta = _pdf_meta(f)
                if meta is None:
                    pdf_missing = True
                elif meta is False:
                    print(f"note: pdfinfo failed or timed out on {f}: PDF metadata not checked")
                else:
                    for loc, text in meta:
                        check(loc, text)
    if pdf_missing:
        print("note: pdfinfo not found: PDF metadata not checked")
    print(f"\n{hits} hit(s).")
    readable = TEXT_TYPES + _OFFICE_ZIP
    unread = [str(f) for t in a.paths for f in _scan_files(Path(t)) if f.suffix.lower() not in readable]
    unread += [f for f in dict.fromkeys(bad) if f not in unread]
    if unread:
        print(f"{len(unread)} file(s) not scanned (this type cannot be read), so zero hits does not cover them:")
        for f in unread:
            print(" -", f)
        print("Scan their text instead, e.g. `pdftotext -layout report.pdf report.txt`, then scan report.txt.")
    sys.exit(1 if hits else 2 if unread else 0)


# ----------------------------------------------------------------------------- learning packs

def _pack_module():
    """pmlc_pack lives next to this file; import it lazily so the core engine never depends on it."""
    if str(HERE) not in sys.path:
        sys.path.insert(0, str(HERE))
    import pmlc_pack
    return pmlc_pack


def _lint_module():
    _pack_module()
    import pmlc_lint
    return pmlc_lint


def cmd_check_pack(a):
    _pack_module().cmd_check_pack(a)


def cmd_check_language(a):
    _lint_module().cmd_check_language(a)


def cmd_pack(a):
    _pack_module().cmd_pack(a)


def cmd_causal_item(a):
    _pack_module().cmd_causal_item(a)


# ----------------------------------------------------------------------------- main

def _non_negative_int(v):
    n = int(v)
    if n < 0:
        raise argparse.ArgumentTypeError("--budget-chars must be 0 or more")
    return n


def _lint_flags(s):
    s.add_argument("--share", choices=["internal", "external", "public"],
                   help="lint for this audience (default: the handling classification, else internal)")
    s.add_argument("--names", help="names file, one per line (default: $PMLC_NAMES)")
    s.add_argument("--orgs", help="organisation names file, one per line (default: $PMLC_ORGS)")


def build_parser():
    p = argparse.ArgumentParser(prog="pmlc", description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--version", action="version", version=f"pmlc {__version__}")
    sp = p.add_subparsers(dest="cmd", required=True)

    s = sp.add_parser("ingest", help="xlsx export -> JSON workspace in <export-dir>/pmlc-runs/<slug>/workspace")
    s.add_argument("xlsx")
    where = s.add_mutually_exclusive_group()
    where.add_argument("--slug", help="<site>-<event>: lowercase words joined by hyphens, e.g. hazira-gas-release-fire")
    s.add_argument("--date", help="incident date YYYY-MM-DD (default: the summary sheet's Date)")
    s.add_argument("--root", help="folder that holds pmlc-runs/ (default: the export's folder)")
    where.add_argument("--out", help="explicit workspace path; must be <dir>/pmlc-runs/<YYYY-MM-DD>-<site>-<event>/workspace")
    where.add_argument("--legacy-out", help="write the workspace to any folder (pre-0.2 layout; not for new runs)")
    s.add_argument("--force", action="store_true", help="overwrite an existing workspace in the run folder")
    s.add_argument("--strict", action="store_true", help="exit 1 when a required column is missing")
    s.set_defaults(fn=cmd_ingest)

    s = sp.add_parser("profile", help="Markdown profile")
    s.add_argument("ws")
    s.add_argument("--top", type=int, default=15)
    s.add_argument("--out")
    s.set_defaults(fn=cmd_profile)

    s = sp.add_parser("classify", help="roles -> Learning Clusters")
    s.add_argument("ws")
    s.add_argument("--taxonomy", help="alternative taxonomy.json")
    s.add_argument("--overlay", action="append", metavar="NAME_OR_PATH",
                   help="extra site role words OR-ed onto the taxonomy (repeatable); a bare name means "
                        "references/overlays/<name>.json, e.g. india-process-site")
    s.add_argument("--no-overlay", action="store_true",
                   help="use no overlay. Without --taxonomy/--overlay/--no-overlay, classify reuses the taxonomy "
                        "and overlays the previous classify used")
    s.add_argument("--overrides", help="JSON {role: {F:[..], T:.., R:..}} from the LLM resolution pass "
                                          "(default: <ws>/role_overrides.json when it exists)")
    s.add_argument("--no-overrides", action="store_true", help="ignore any overrides file")
    s.set_defaults(fn=cmd_classify)

    s = sp.add_parser("select", help="filtered findings bundle")
    s.add_argument("ws")
    s.add_argument("--entity", nargs="*", help="e.g. F02 F09b")
    s.add_argument("--tag", nargs="+", help="entity tags, e.g. emergency or supply contractor; unioned with --entity")
    s.add_argument("--tier", nargs="*", help="e.g. T2 T3")
    s.add_argument("--relation", nargs="*", help="actor approver owner absent norm_setter learner challenger")
    s.add_argument("--level", nargs="*", type=int, help="1 2 3 4")
    s.add_argument("--min-rating", type=int)
    s.add_argument("--top-rated", action="store_true", help="only findings at the workspace's top rating (run-path prints it)")
    s.add_argument("--entity-mode", choices=["any", "all"], default="any",
                   help="any: a finding matches if it has any requested entity (default). all: it must have every "
                        "--entity and, for each --tag, at least one of that tag's entities (the shared findings)")
    s.add_argument("--exclude-downgraded", action="store_true",
                   help="drop findings a reviewer voted to downgrade, reject or remove, and say which")
    s.add_argument("--subcategory", nargs="+", help="HFACS subcategory, case-insensitive exact match")
    s.add_argument("--prefix", nargs="+", help="Nanocode prefix before the space, e.g. SV OR")
    s.add_argument("--code", nargs="*")
    s.add_argument("--text", help="regex over short/justification/causal/label")
    s.add_argument("--max", type=int)
    s.add_argument("--brief", action="store_true", help="omit justification, causal text and evidence quotes")
    s.add_argument("--budget-chars", type=_non_negative_int, help="keep findings in rating order while the output stays within "
                                                    "N characters, and say how many were dropped; the header and "
                                                    "footer lines are always printed")
    s.add_argument("--fields", help="comma-separated: " + ", ".join(SELECT_FIELDS) + " (id is always kept)")
    s.add_argument("--group-by", choices=["level", "subcategory", "code", "entity"],
                   help="group the findings under one heading (md) or key (json) per value")
    s.add_argument("--count-by", choices=COUNT_BY, help="print only a count table (md) or dict (json)")
    s.add_argument("--format", choices=["md", "json"], default="md")
    s.add_argument("--out")
    s.add_argument("--preset", metavar="SKILL[:NAME]",
                   help="run a skill's recorded gather calls (references/select-presets.json); takes only output "
                        "flags. Each call keeps its own --format unless you pass --format, which overrides all. "
                        "With several calls the output is one bundle per call under '# Preset <skill>: <name>' "
                        "headings, so json output is concatenated bundles, not one JSON value. --budget-chars "
                        "caps the whole bundle, headings included: each call gets what the earlier calls left, and a "
                        "call with nothing left prints only its minimum output (with the below-minimum warning)")
    s.set_defaults(fn=cmd_select)

    s = sp.add_parser("causal", help="extract causal graph")
    s.add_argument("ws")
    s.add_argument("--top", type=int, default=60)
    s.add_argument("--include-justification", action="store_true")
    s.set_defaults(fn=cmd_causal)

    s = sp.add_parser("timeline", help="dated precursor candidates")
    s.add_argument("ws")
    s.add_argument("--before", help="ISO date, or the word incident for the workspace's incident date; keep only events up to it")
    s.add_argument("--per-date", type=int, default=4)
    s.add_argument("--show", type=int, default=60)
    s.add_argument("--on", metavar="DAY|incident",
                   help="also write timeline_event_day.json: every clock time for that day (YYYY-MM-DD, or incident). "
                        "Times in undated sentences count only for the incident date; for any other day a time "
                        "counts only when its sentence gives that day's date")
    s.add_argument("--entity", nargs="+", metavar="F", help="keep only findings in these entities, and their linked evidence")
    s.add_argument("--min-precision", choices=["day", "month", "year"], default="year",
                   help="coarsest date precision to keep (year keeps everything)")
    s.set_defaults(fn=cmd_timeline)

    s = sp.add_parser("interfaces", help="entity x entity interfaces")
    s.add_argument("ws")
    s.add_argument("--top", type=int, default=25)
    s.set_defaults(fn=cmd_interfaces)

    s = sp.add_parser("dashboard", help="HTML organisational learning map")
    s.add_argument("ws")
    s.add_argument("--out", help="default: <run>/outputs/organisation-learning-map.html")
    s.add_argument("--register", help="Learnings Register JSON: adds learnings, tier plan and control-strength sections")
    s.add_argument("--title", help="override the incident title shown in the header")
    s.add_argument("--date", help="override the incident date (use when the summary sheet is wrong)")
    s.add_argument("--location", help="override the location shown in the header")
    s.set_defaults(fn=cmd_dashboard)

    s = sp.add_parser("check-register", help="validate a Learnings Register JSON against the workspace")
    s.add_argument("ws")
    s.add_argument("register")
    _lint_flags(s)
    s.set_defaults(fn=cmd_check_register)

    s = sp.add_parser("check-language", help="just-culture and grounding lint (JC1-JC12) of a pack or register JSON")
    s.add_argument("ws")
    s.add_argument("json", help="a learning pack or Learnings Register JSON")
    _lint_flags(s)
    s.set_defaults(fn=cmd_check_language)

    s = sp.add_parser("register-xlsx", help="Learnings Register JSON -> xlsx")
    s.add_argument("register")
    s.add_argument("--out", help="default: learnings-register.xlsx next to the register JSON")
    s.add_argument("--ws", help="workspace for the review status (default: the run's workspace in the standard layout)")
    s.set_defaults(fn=cmd_register_xlsx)

    s = sp.add_parser("status", help="show which workspace files are current or stale, and pack check history")
    s.add_argument("ws")
    s.set_defaults(fn=cmd_status)

    s = sp.add_parser("run-path", help="print the run folder, workspace, outputs folder and slug of a workspace")
    s.add_argument("ws")
    s.set_defaults(fn=cmd_run_path)

    s = sp.add_parser("check-samples", help="check samples/<slug>/ folders follow the contribution layout")
    s.add_argument("dir", nargs="?", default=str(HERE.parents[2] / "samples"))
    s.set_defaults(fn=cmd_check_samples)

    s = sp.add_parser("check-pack", help="check a learning pack JSON against the workspace")
    s.add_argument("ws")
    s.add_argument("pack")
    _lint_flags(s)
    s.set_defaults(fn=cmd_check_pack)

    s = sp.add_parser("pack", help="learning pack JSON -> <run>/outputs/<pack>.html (and .md)")
    s.add_argument("ws")
    s.add_argument("pack")
    s.add_argument("--out", help="output folder (default: <run>/outputs)")
    s.add_argument("--no-md", action="store_true", help="write only the HTML")
    s.set_defaults(fn=cmd_pack)

    s = sp.add_parser("causal-item", help="top causal links as a pack graph item (JSON on stdout)")
    s.add_argument("ws")
    s.add_argument("--top", type=int, default=40, help="links to keep (default 40)")
    s.set_defaults(fn=cmd_causal_item)

    s = sp.add_parser("sign-off", help="record a person's review of a pack or register (writes its handling block)")
    s.add_argument("ws")
    s.add_argument("file")
    s.add_argument("--role", required=True, help="the reviewer's role (never a name)")
    s.add_argument("--status", choices=["reviewed", "approved"], default="reviewed")
    s.add_argument("--classification", choices=list(HANDLING_ENUMS["classification"]))
    s.add_argument("--legal-review", dest="legal_review", choices=list(HANDLING_ENUMS["legal_review"]))
    s.add_argument("--distribution", nargs="+", metavar="AUDIENCE")
    s.add_argument("--retention-review", dest="retention_review", metavar="YYYY-MM-DD")
    s.add_argument("--version")
    s.add_argument("--date", metavar="YYYY-MM-DD", help="review date (default: today)")
    s.set_defaults(fn=cmd_sign_off)

    s = sp.add_parser("scan-names", help="zero-hit name scan (exit 1 on hits, 2 if a file could not be read)")
    s.add_argument("names", help="text file, one name per line")
    s.add_argument("paths", nargs="+")
    s.add_argument("--show", type=int, default=50)
    s.add_argument("--variants", action="store_true",
                   help="also match surname, initial + surname and honorific + surname for each name")
    s.add_argument("--patterns", action="store_true",
                   help="also flag email, phone and staff-number patterns (an aid, not proof of de-identification)")
    s.set_defaults(fn=cmd_scan_names)

    return p


def main(argv=None):
    a = build_parser().parse_args(argv)
    a._argv = list(sys.argv[1:] if argv is None else argv)
    a.fn(a)


if __name__ == "__main__":
    main()
