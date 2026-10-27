#!/usr/bin/env python3
"""Write the UAT input files into docs/uat/inputs/ (fictional event, no real data).

Run from the repository root:  python3 tools/make-uat-inputs.py [--out DIR]   (default: docs/uat/inputs)

It writes four synthetic ProbeMatrix findings exports, each testing one ingest, check or sign-off path:
  uat-votes.xlsx           Collaboration Data VOTES with a reviewer downgrade on FND-0004
  uat-renamed-header.xlsx  the Identified Findings column "Rating" renamed to "Score"
  uat-capa-planner.xlsx    Corrective Actions and Preventive Actions sheets (a CAPA Planner export)
  uat-approvable.xlsx      Phase 2: completed, FND-0001 edited by a person (so sign-off is allowed), with fictional
                           names, an organisation, a tag, a clock time, a full date, a numbered role, an email and a
                           staff number for the privacy and sharing tests
two small learning packs to check against the uat-votes workspace:
  uat-votes-pack.json      cites the downgraded FND-0004 without a data_quality note (check-pack fails)
  uat-unlinked-pack.json   cites EVD-0003 next to FND-0002, which does not link it (check-pack warns)
and, for Phase 2, against the uat-approvable workspace:
  uat-bad-pack.json        the L3-01 test pack (tests/fixtures/packs/bad-l3-01.json): fails the language check
  uat-shareable-pack.json  a clean pack that can be signed off and shared
  uat-register.json        a clean two-learning register that can be signed off and shared
  uat-names.txt, uat-orgs.txt, uat-redact.json  names, organisations and a redaction map (all fictional)
and, for Phase 3 (docs/uat/v2.0-consolidated-uat.md):
  uat-hazira-register.json     the Hazira sample register with every why_not_stronger removed, so that
                               check-register --strict names LR-08 and LR-12 whatever the shipped sample holds
  uat-hazira-cards.json        the two-card Hazira test pack (tests/fixtures/hazira/print-2-card.json) whose cards
                               name the learnings LR-02 and LR-10
  uat-unknown-lr-pack.json     a Hazira pack whose second item names LR-99, which no register holds
  uat-adoption-filled.csv      adoption.csv for the Hazira register as Excel saves it: a byte-order mark, CRLF line
                               ends, reordered and extra columns, mixed-case responses and trailing blank rows
  uat-distribution-filled.csv  distribution.csv for uat-hazira-cards: two dated deliveries and one planned
  uat-org-chart.csv            an organisation's departments for the uat-approvable export (invented names)
  uat-comparator.json, uat-comparison-map.json  a made-up company report on the pump fire and its mapping to the
                               two learnings of uat-register.json
and, for Phase 4:
  uat-brand.json, uat-brand-logo.png  a valid brand: an invented organisation, an accent of 5.8:1 against white and a
                               1x1 PNG logo
  uat-brand-bad.json           a brand whose accent has 1.7:1 against white, so pack --brand refuses it
  uat-print-cards.json         the two Hazira print cards with the card fields, two toolbox talks and three checks
                               (one a hold point), for the print layouts
  uat-quiz.json                five Hazira questions: single, multi, truefalse, short and one discussion question
and, for Phase 5:
  uat-weak-fix-register.json   the Hazira sample register with LR-08's administrative fix reworded as a reminder memo,
                               so rec-review flags LR-08 as a weak fix for a deep cause
  uat-capa-planner-9m.xlsx     a CAPA Planner workbook as the planner writes it (four sheets, no Identified
                               Findings), which nine-m-seed and setup refuse
  uat-process-export.xlsx      the synthetic process-industry export of tests/conftest.py (30 invented findings), for
                               setup --industry process
and, for Phase 6:
  uat-no-urgency-register.json the Hazira sample register without urgency, so alert-grade suggests a grade from
                               priority and the safety-alert skill asks the user (OD6-1)
  uat-library-register.json    the pump-fire register (uat-register.json) whose LR-01 context names an invented site,
                               a full date and a clock time, for library publish on a synthetic run: the public lines
                               must code the site (S01) and keep the month only

The export shape follows tests/test_pmlc.py (FIND_COLS, ROWS, make_export), duplicated here so this tool
needs only the standard library and openpyxl. The workbooks are rebuilt from scratch on every run.
"""
import argparse
import csv
import io
import json
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required: pip install openpyxl")

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "uat" / "inputs"
sys.path.insert(0, str(ROOT / "skills" / "learnings-coach" / "scripts"))
from pmlc import STATUS_PHRASE as STATUS  # noqa: E402

FIND_COLS = ["Finding ID", "Nanocode", "Name", "Level", "Level Name", "Category", "Subcategory", "Classification",
             "Rating", "Instance ID", "Instance Label", "Short Description", "AI Justification", "Causal Explanation",
             "Individuals Involved", "Linked Evidence", "Evidence to Collect", "Key Questions", "Suggested Interviews",
             "Causal Hypothesis", "User Edited", "Created At", "Updated At"]

ROWS = [
    ["FND-0001", "AB 1.1", "Step Omitted", 1, "Unsafe Acts", "Unsafe Acts", "Errors", "substantiated", 3, 1,
     "Pump restarted without leak test", "Pump restarted without leak test after seal change on 12/03/2024.",
     "Work order shows no test [EVD-0001].", "Enabled by SV 1.4 and OR 2.2.",
     "Pump Technician (seal change crew), Maintenance Supervisor (return-to-service approval)", "EVD-0001",
     None, None, None, None, "No", None, None],
    ["FND-0002", "SV 1.4", "No Guidance Provided", 3, "Supervision", "Supervision", "Supervision Inadequate",
     "substantiated", 4, 1, "No return-to-service checklist", "No checklist for return to service.",
     "Procedure index has no entry [EVD-0002].", "Led to AB 1.1.",
     "Maintenance Supervisor, Maintenance Manager (procedure owner)", "EVD-0002", None, None, None, None, "No",
     None, None],
    ["FND-0003", "OR 2.2", "Spares Programme Failure", 4, "Organisational Influences", "Organisational Influences",
     "Resource Management", "substantiated", 4, 1, "Seals bought locally without specification",
     "Local seals without specification since March 2022.", "Purchase records [EVD-0003].",
     "This condition enabled SV 1.4 and AB 1.1.",
     "Stores/Procurement Officer, Plant Manager (budget owner), Reliability Engineer (absent)", "EVD-0003, EVD-0001",
     None, None, None, None, "No", None, None],
    ["FND-0004", "PE 2.1", "Alarm Not Configured", 2, "Preconditions", "Preconditions", "Technological Environment",
     "substantiated", 2, 1, "Gas alarm display-only", "Gas detector alarm had no shutdown.", "DCS list [EVD-0002].",
     "Delayed response; see OR 2.2.", "Control Room Operators, Emergency Response planners", "EVD-0002",
     None, None, None, None, "No", None, None],
]


def base_workbook(summary=(), narrative=None):
    wb = openpyxl.Workbook()
    s = wb.active
    s.title = "Incident Summary"
    for r in [("INCIDENT SUMMARY", None), (None, None), ("Field", "Value"), ("Incident Title", "Synthetic pump fire"),
              ("Date", "2024-03-20"), ("Investigation Status", "completed"), *summary]:
        s.append(r)
    if narrative:
        s.append(("INCIDENT NARRATIVE", None))
        s.append((narrative, None))
    f = wb.create_sheet("Identified Findings")
    f.append(FIND_COLS)
    for r in ROWS:
        f.append(r)
    e = wb.create_sheet("Evidence Links")
    e.append(["Finding Code", "Finding Name", "Evidence ID", "Evidence Type", "Relevance", "Specific Reference",
              "Evidence Description"])
    e.append(["AB 1.1", "Step Omitted", "EVD-0001", "document", "Shows no test", "no leak test recorded", "Work order"])
    e.append(["SV 1.4", "No Guidance Provided", "EVD-0002", "document", "No checklist", "index has no entry", "Index"])
    e.append(["OR 2.2", "Spares Programme Failure", "EVD-0003", "log_record", "Local buys", "bought locally", "PO list"])
    c = wb.create_sheet("Causal Chain")
    c.append(["Link ID", "From Code", "From Level", "To Code", "To Level", "Relationship", "Confidence", "Basis",
              "Evidence Support", "User Confirmed", "Notes"])
    c.append(["No causal links found"] + [None] * 10)
    return wb


def votes(wb):
    c = wb.create_sheet("Collaboration Data")
    for r in [("VOTES",), ("Finding ID", "Vote", "Reason"), ("FND-0004", "downgrade", "weak evidence")]:
        c.append(r)


def renamed_header(wb):
    for c in wb["Identified Findings"][1]:
        if c.value == "Rating":
            c.value = "Score"


def capa_planner(wb):
    wb.create_sheet("Corrective Actions")
    wb.create_sheet("Preventive Actions")


# Fictional people (no real person, and no one connected with the Hazira incident). The redaction map gives each a role.
NAMES = ["Fictional Officer", "Ms Jane Placeholder"]
REDACT = {"Fictional Officer": "Investigating Officer", "Ms Jane Placeholder": "Document Controller"}
ORGS = ["Acme Contractors Pvt Ltd", "Acme Contractors"]
APPROVABLE_SUMMARY = [("Investigating Officer", "Fictional Officer")]
APPROVABLE_NARRATIVE = (  # over 200 characters, so ingest reads it as the narrative
    "A seal on a transfer pump was changed on the night shift. The pump was restarted without a leak test, the new "
    "seal leaked, and the vapour found an ignition source near the pump. The fire was put out by the site fire team. "
    "The regulator was told the same morning.")


def approvable(wb):
    """Completed, FND-0001 edited by a person, plus the fictional identifying details the Phase 2 tests mask."""
    f = wb["Identified Findings"]
    col = {c.value: c.column for c in f[1]}
    rows = {f.cell(r, 1).value: r for r in range(2, f.max_row + 1)}
    f.cell(rows["FND-0001"], col["User Edited"]).value = "Yes"
    f.cell(rows["FND-0001"], col["AI Justification"]).value = (
        "Work order from Acme Contractors Pvt Ltd shows pump P-101 restarted at 03:12 on 2024-03-20 with no leak "
        "test recorded [EVD-0001].")
    f.cell(rows["FND-0002"], col["AI Justification"]).value = (
        "Procedure index has no entry [EVD-0002]; Ms Jane Placeholder confirmed the gap.")
    f.cell(rows["FND-0004"], col["Individuals Involved"]).value = (
        "Control Room Operators, Emergency Response planners, Security Guard - Gate 3")
    e = wb["Evidence Links"]
    e.cell(3, 7).value = "Index; queries to docs.control@example.com, staff no. 4471"


WORKBOOKS = {"uat-votes.xlsx": votes, "uat-renamed-header.xlsx": renamed_header, "uat-capa-planner.xlsx": capa_planner,
             "uat-approvable.xlsx": approvable}
SUMMARY = {"uat-approvable.xlsx": (APPROVABLE_SUMMARY, APPROVABLE_NARRATIVE)}


def pack(name, title, items, **top):
    p = {"pack": name, "schema": 1, "title": title, "audience": "UAT testers", "status": STATUS,
         "bottom_line": [{"text": "No check stopped an untested pump returning to service.", "cites": ["FND-0002"]}],
         "sections": [{"id": "learnings", "heading": "Learnings", "items": items}]}
    p.update(top)
    return p


PACKS = {
    "uat-votes-pack.json": pack("uat-votes-pack", "UAT votes pack: synthetic pump fire", [
        {"kind": "learning", "lesson": "A gas alarm that only displays does not stop the release.",
         "cites": ["FND-0004"]}]),
    "uat-unlinked-pack.json": pack("uat-unlinked-pack", "UAT unlinked evidence pack: synthetic pump fire", [
        {"kind": "learning", "lesson": "Return to service needs a checklist with a hold point.",
         "cites": ["FND-0002", "EVD-0003"]}]),
}

# The L3-01 test pack, as tests/fixtures/packs/bad-l3-01.json: a name, blame words, an outside fact, a percentage
# and a retraining fix.
BAD_PACK = {
    "pack": "uat-bad-pack", "schema": 1, "title": "UAT bad pack: the L3-01 test pack", "audience": "Site leadership",
    "status": STATUS,
    "bottom_line": [{"text": "John Smith was careless and should be disciplined", "cites": ["FND-0001"]}],
    "sections": [{"id": "what-to-change", "heading": "What to change", "items": [
        {"kind": "para", "text": "As at the Northfield refinery fire, 80% of such fires start at flanges",
         "cites": ["FND-0001"]},
        {"kind": "recommendation", "change": "Retrain the fitter and hold him accountable",
         "indicator": "Share of fitters retrained", "hoc_tier": "training", "cites": ["FND-0002"]}]}]}

SHAREABLE_PACK = {
    "pack": "uat-shareable-pack", "schema": 1, "title": "UAT shareable pack: synthetic pump fire",
    "audience": "Site leadership", "status": STATUS,
    "bottom_line": [
        {"text": "A pump went back into service after a seal change with no leak test.", "cites": ["FND-0001"]},
        {"text": "No return-to-service checklist set a hold point for the test.", "cites": ["FND-0002"]}],
    "sections": [{"id": "learnings", "heading": "Learnings", "items": [
        {"kind": "learning", "lesson": "When a joint is disturbed, a leak test before restart is the barrier.",
         "cites": ["FND-0001", "EVD-0001"]},
        {"kind": "learning", "lesson": "Seals bought without a specification weaken every joint they go into.",
         "cites": ["FND-0003", "EVD-0003"]},
        {"kind": "learning", "lesson": "A gas alarm that only displays leaves the response to whoever sees it.",
         "cites": ["FND-0004"]},
        {"kind": "recommendation", "change": "Block restart until a signed leak-test record is attached.",
         "indicator": "Share of restarts with a leak-test record", "hoc_tier": "engineering",
         "cites": ["FND-0001", "FND-0002"]}]}]}

REGISTER = {"title": "UAT register: synthetic pump fire", "source": "uat-approvable.xlsx", "status": STATUS,
            "entries": [
    {"lr_id": "LR-01", "title": "Return to service without a leak test", "finding_ids": ["FND-0001", "FND-0002"],
     "evidence_ids": ["EVD-0001", "EVD-0002"], "priority": 1, "owner_entity": "F02 Mechanical Maintenance",
     "hfacs_level": "L3 Supervision Inadequate", "tiers": ["T1", "T2"],
     "lesson": "When a joint is disturbed, a leak test before restart is the barrier, and a checklist holds it.",
     "context": "A pump went back into service after a seal change with no leak test recorded.",
     "tier_messages": {"T1": "Do not restart without the leak-test record.", "T2": "Hold point before acceptance."},
     "recommendations": [{"what": "Block restart until a signed leak-test record is attached",
                          "hoc_tier": "engineering", "leading_indicator": "Share of restarts with a leak-test record"}]},
    {"lr_id": "LR-02", "title": "Seals without a specification", "finding_ids": ["FND-0003"],
     "evidence_ids": ["EVD-0003"], "priority": 2, "owner_entity": "F06 Materials Management",
     "hfacs_level": "L4 Resource Management", "tiers": ["T3"],
     "lesson": "Spares bought without a specification weaken every joint they go into.",
     "context": "Seals were bought locally without a specification.",
     "tier_messages": {"T3": "Stock only seals that match the specification."},
     "recommendations": [{"what": "Stock only specification-approved seals, so a local substitute cannot be issued",
                          "hoc_tier": "substitution", "leading_indicator": "Share of seal issues that match the specification"}]}]}

PHASE2_JSON = {"uat-bad-pack.json": BAD_PACK, "uat-shareable-pack.json": SHAREABLE_PACK, "uat-register.json": REGISTER,
               "uat-redact.json": REDACT}
TEXT_FILES = {
    "uat-names.txt": "# Fictional names for the UAT (no real person). One name per line; # starts a comment.\n"
                     + "\n".join(NAMES) + "\n",
    "uat-orgs.txt": "# Fictional organisation names for the UAT. One per line.\n" + "\n".join(ORGS) + "\n",
}


# ----------------------------------------------------------------------------- Phase 3 inputs

SAMPLE_REGISTER = ROOT / "samples" / "2020-09-24-hazira-gas-release-fire" / "outputs" / "learnings-register.json"
PRINT_2_CARD = ROOT / "tests" / "fixtures" / "hazira" / "print-2-card.json"


def hazira_register():
    """The sample register without any why_not_stronger: LR-08 and LR-12 are priority 1 with only administrative and
    training fixes, so check-register --strict names them (Ruling P3-3)."""
    r = json.loads(SAMPLE_REGISTER.read_text(encoding="utf-8"))
    for e in r["entries"]:
        e.pop("why_not_stronger", None)
    return r


def hazira_cards():
    p = json.loads(PRINT_2_CARD.read_text(encoding="utf-8"))
    p["pack"] = "uat-hazira-cards"
    p["title"] = "UAT learning cards: Hazira gas release and fire, 2020"
    cards = p["sections"][0]["items"]
    cards[0]["learning"] = "LR-02"    # FND-3AB996FC: return to service without a test to design pressure
    cards[1]["learning"] = "LR-10"    # FND-01C3DE98: gas detection that only displayed an alarm
    return p


UNKNOWN_LR_PACK = {
    "pack": "uat-unknown-lr-pack", "schema": 1, "title": "UAT unknown learning pack: Hazira", "audience": "UAT testers",
    "status": STATUS,
    "bottom_line": [{"text": "A pressure-boundary repair went back into service with no test to design pressure.",
                     "cites": ["FND-3AB996FC"], "learning": "LR-02"}],
    "sections": [{"id": "learnings", "heading": "Learnings", "items": [
        {"kind": "learning", "lesson": "A repair that disturbs a pressure boundary needs a test before restart.",
         "cites": ["FND-3AB996FC"], "learning": "LR-02"},
        {"kind": "learning", "lesson": "A seal with no specification on file invites a substitute part.",
         "cites": ["FND-01C3DE98"], "learning": "LR-99"}]}]}

# The organisation's answers (invented; roles only). Responses use Excel's capitals; the loader folds case.
ADOPTION_FILLED = {
    "LR-01": ("Accepted", "Head of Maintenance", "Q1 2027", "", "agreed at the site leadership meeting"),
    "LR-02": ("Accepted", "Head of Maintenance", "Q1 2027", "", ""),
    "LR-03": ("Modified", "Maintenance Manager", "Q2 2027",
              "Independent check added to the existing permit rather than a new form", ""),
    "LR-09": ("Accepted", "Head of Projects", "2027 shutdown", "", ""),
    "LR-10": ("Accepted", "Instrumentation Manager", "2027 shutdown", "", ""),
    "LR-12": ("Modified", "Head of Operations", "Q2 2027", "Simulator time shared with a sister site", ""),
    "LR-24": ("Rejected", "HSE Manager", "", "Covered by the corporate drill standard now in revision", ""),
}


def adoption_filled(register) -> bytes:
    """adoption.csv as Excel saves it: BOM, CRLF, columns reordered, an extra Notes column, two blank rows."""
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\r\n")
    w.writerow(["lr_id", "title", "response", "owning_role", "target_window", "reason", "Notes"])
    for e in register["entries"]:
        resp, role, window, reason, note = ADOPTION_FILLED.get(e["lr_id"], ("", "", "", "", ""))
        w.writerow([e["lr_id"], e["title"], resp, role, window, reason, note])
    w.writerow([""] * 7)
    w.writerow([""] * 7)
    return b"\xef\xbb\xbf" + buf.getvalue().encode("utf-8")


DISTRIBUTION_FILLED = [
    ["pack", "audience", "channel", "date", "delivered_by_role", "attendees", "contractors_included",
     "questions_raised", "concerns_escalated"],
    ["uat-hazira-cards", "Frontline crews (T1)", "Toolbox talk", "2026-11-05", "Shift In-charge", "14", "no",
     "Where is the test record kept?", ""],
    ["uat-hazira-cards", "Frontline crews (T1)", "Toolbox talk", "2026-11-06", "Maintenance Supervisor", "9", "yes",
     "", "Spare seals in the store carry no part numbers"],
    ["uat-hazira-cards", "Frontline crews (T1)", "Toolbox talk", "", "Shift In-charge", "", "", "", ""],
]

# Invented departments for the uat-approvable pump fire: one extends F06, two become new entities F20 and F21.
ORG_CHART = [
    ["department", "entity", "match", "tier"],
    ["Stores and Purchase", "F06", "Stores/Procurement", "T3"],
    ["Site Security", "", "Security Guard", "T1"],
    ["Reliability Cell", "", "Reliability Engineer", "T3"],
]

COMPARATOR = {
    "source": {"title": "Synthetic pump fire: site investigation report", "issuer_role": "Site investigation team",
               "statutory_status": "non-statutory company report", "kind": "company-report"},
    "items": [
        {"id": "OBS-01", "type": "observation", "section": "3.1", "group": "Maintenance",
         "text": "The pump was restarted after a seal change with no leak test recorded.", "hoc_tier": ""},
        {"id": "RC-01", "type": "root_cause", "section": "4.1", "group": "Maintenance",
         "text": "No return-to-service checklist required a leak test.", "hoc_tier": ""},
        {"id": "RC-02", "type": "root_cause", "section": "4.2", "group": "Stores",
         "text": "Seals were bought locally without a specification.", "hoc_tier": ""},
        {"id": "REC-01", "type": "recommendation", "section": "5.1", "group": "Maintenance",
         "text": "Introduce a return-to-service checklist with a leak-test hold point.", "hoc_tier": "administrative"},
        {"id": "REC-02", "type": "recommendation", "section": "5.2", "group": "Maintenance",
         "text": "Brief the seal-change crews on the leak-test step.", "hoc_tier": "training"},
        {"id": "REC-03", "type": "recommendation", "section": "5.3", "group": "Control systems",
         "text": "Link the pump-area gas detector to an automatic pump trip.", "hoc_tier": "engineering"},
        {"id": "ACT-01", "type": "action", "section": "6.1", "group": "Stores",
         "text": "Seal stock checked against the specification.", "hoc_tier": "administrative"}]}

COMPARISON_MAP = [
    {"lr_id": "LR-01", "ref": "OBS-01", "match": "yes"},
    {"lr_id": "LR-01", "ref": "RC-01", "match": "yes"},
    {"lr_id": "LR-01", "ref": "REC-01", "match": "partial"},
    {"lr_id": "LR-02", "ref": "RC-02", "match": "yes"},
    {"lr_id": "LR-02", "ref": "ACT-01", "match": "partial"},
]


def _csv_text(rows) -> str:
    buf = io.StringIO()
    csv.writer(buf, lineterminator="\n").writerows(rows)
    return buf.getvalue()


def phase3_files():
    """{file name: bytes} for the Phase 3 inputs."""
    reg = hazira_register()
    dump = lambda o: (json.dumps(o, indent=1, ensure_ascii=False) + "\n").encode("utf-8")   # noqa: E731
    return {"uat-hazira-register.json": dump(reg), "uat-hazira-cards.json": dump(hazira_cards()),
            "uat-unknown-lr-pack.json": dump(UNKNOWN_LR_PACK), "uat-adoption-filled.csv": adoption_filled(reg),
            "uat-distribution-filled.csv": _csv_text(DISTRIBUTION_FILLED).encode("utf-8"),
            "uat-org-chart.csv": _csv_text(ORG_CHART).encode("utf-8"),
            "uat-comparator.json": dump(COMPARATOR), "uat-comparison-map.json": dump(COMPARISON_MAP)}


# ----------------------------------------------------------------------------- Phase 4 inputs

BRAND_ACCENT = "#1A7070"                     # 5.8:1 against white: passes WCAG AA
BRAND = {"organisation": "Example Process Terminal (UAT)", "accent": BRAND_ACCENT, "logo": "uat-brand-logo.png"}
BRAND_BAD = {"organisation": "Example Process Terminal (UAT)", "accent": "#7FD3CF"}   # 1.7:1: refused


def one_pixel_png(hex_colour: str) -> bytes:
    """A 1x1 RGB PNG of one colour, written with the standard library only."""
    import struct
    import zlib

    def chunk(tag: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + tag + data + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)
    rgb = bytes(int(hex_colour.lstrip("#")[i:i + 2], 16) for i in (0, 2, 4))
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(b"\x00" + rgb)) + chunk(b"IEND", b""))


def print_cards():
    """The two Hazira print cards with the Phase 4 card fields, two toolbox talks and three stop-and-ask checks (one a
    hold point): the input for the cards, toolbox and checklist print layouts."""
    p = json.loads(PRINT_2_CARD.read_text(encoding="utf-8"))
    p["pack"] = "uat-print-cards"
    p["title"] = "UAT print cards: Hazira gas release and fire, 2020"
    c1, c2 = p["sections"][0]["items"]
    c1["did_you_know"] = ["Operating pressure is lower than design pressure, so a joint that holds at operating "
                          "pressure has not been proved [EVD-2894]."]
    c1["what_you_can_do"] = ["Ask to see the test record at design pressure before you open the line."]
    c2["did_you_know"] = ["The detectors raised an alarm on a panel; nothing tripped by itself [EVD-8657]."]
    c2["what_you_can_do"] = ["Learn which alarms on your unit trip equipment and which need you to act."]
    p["sections"].append({"id": "toolbox-talks", "heading": "Toolbox talks", "items": [
        {"kind": "toolbox",
         "opening_question": "Who here has put a meter back in line after a gasket change?",
         "story": "A metering line went back into service after a gasket change. It was pressurised only to operating "
                  "pressure, and the meter went back online the next day [EVD-2894].",
         "warning_signs": ["A seal or gasket that is not the original part",
                           "No test record at design pressure"],
         "one_change": "No meter goes back in line until the test at design pressure is on the record.",
         "crew_question": "Where on our unit is the test record kept, and who can see it?",
         "cites": ["FND-3AB996FC", "EVD-2894"]},
        {"kind": "toolbox",
         "opening_question": "What happens on our unit when a gas detector alarms?",
         "story": "The gas detectors showed an alarm on an annunciator panel for display only. There was no automatic "
                  "shutdown and no fire suppression [EVD-8657].",
         "warning_signs": ["A gas alarm that only shows on a panel", "No automatic action after a gas alarm"],
         "one_change": "Each crew member can name which alarms trip equipment and which need a person to act.",
         "crew_question": "Which of our alarms would need one of us to act at once?",
         "cites": ["FND-01C3DE98", "EVD-8657"]}]})
    p["sections"].append({"id": "stop-and-ask", "heading": "Stop and ask", "items": [
        {"kind": "check", "hold_point": True, "when": "Before a high-pressure meter goes back in line",
         "text": "The test at design pressure is recorded and witnessed.", "cites": ["FND-3AB996FC", "EVD-2894"]},
        {"kind": "check", "when": "A gasket in your hand is not the original part",
         "text": "Ask the supervisor for the specification before you fit it.", "cites": ["FND-3AB996FC"]},
        {"kind": "check", "when": "A gas alarm shows on the panel",
         "text": "Tell the control room what you see, and act on the alarm without waiting for a trip.",
         "cites": ["FND-01C3DE98", "EVD-8657"]}]})
    return p


def _opts(*pairs):
    return [{"text": t, "correct": c} for t, c in pairs]


QUIZ = {
    "pack": "uat-quiz", "schema": 1, "title": "UAT quiz: Hazira gas release and fire, 2020",
    "audience": "Operators and fitters (T1)", "status": STATUS,
    "bottom_line": [{"text": "A repaired joint is proved only by a recorded test at design pressure.",
                     "cites": ["FND-3AB996FC", "EVD-2894"]}],
    "sections": [{"id": "questions", "heading": "Check your understanding", "items": [
        {"kind": "question", "type": "single",
         "question": "A gasket on a high-pressure meter has been replaced. What proves the joint before return to "
                     "service?",
         "options": _opts(("A recorded test at design pressure, with a witness", True),
                          ("Bringing the line up to operating pressure", False),
                          ("A visual check that the bolts are tight", False)),
         "why": "Knowing which test proves a repaired joint.",
         "answer": "A recorded test at design pressure, with a witness; operating pressure does not prove it.",
         "good_evidence": "The learner names the test pressure and where the record is kept.",
         "cites": ["FND-3AB996FC", "EVD-2894"]},
        {"kind": "question", "type": "multi",
         "question": "Which of these are warning signs before a return to service?",
         "options": _opts(("A seal or gasket that is not the original part", True),
                          ("No test record at design pressure", True),
                          ("A permit signed by the issuer", False)),
         "why": "Recognising the conditions that should stop a return to service.",
         "answer": "A part that is not the original, and no test record at design pressure.",
         "good_evidence": "Both warning signs named without prompting.",
         "cites": ["FND-904FCA63", "EVD-2894"]},
        {"kind": "question", "type": "truefalse",
         "question": "The gas detectors at the terminal shut the plant down automatically.",
         "options": _opts(("True", False), ("False", True)),
         "why": "Knowing what the gas alarm did and did not do.",
         "answer": "False: the detectors showed an alarm on a panel only, with no automatic shutdown.",
         "good_evidence": "The learner says who must act when an alarm only displays.",
         "cites": ["FND-01C3DE98", "EVD-8657"]},
        {"kind": "question", "type": "short",
         "question": "In two words: at what pressure must a repaired high-pressure joint be tested?",
         "options": _opts(("design pressure", True), ("the design pressure", True)),
         "why": "Remembering the test pressure.",
         "answer": "Design pressure.",
         "good_evidence": "The learner gives the words without a prompt.",
         "cites": ["FND-212DC9B7", "EVD-2894"]},
        {"kind": "question",
         "question": "On your unit, who can stop a return to service when the test record is missing?",
         "why": "Knowing who holds the stop in practice.",
         "answer": "Anyone on the job can ask for the record, and the supervisor holds the return to service until "
                   "it is there.",
         "good_evidence": "The learner names a role and the place the record is kept.",
         "cites": ["FND-15279F97"]}]}]}


def phase4_files():
    """{file name: bytes} for the Phase 4 inputs."""
    dump = lambda o: (json.dumps(o, indent=1, ensure_ascii=False) + "\n").encode("utf-8")   # noqa: E731
    return {"uat-brand.json": dump(BRAND), "uat-brand-logo.png": one_pixel_png(BRAND_ACCENT),
            "uat-brand-bad.json": dump(BRAND_BAD), "uat-print-cards.json": dump(print_cards()),
            "uat-quiz.json": dump(QUIZ)}


# ----------------------------------------------------------------------------- Phase 5 inputs

PROCESS_ROWS = ROOT / "tests" / "fixtures" / "process-export" / "findings.json"
WEAK_FIX_LR, WEAK_FIX_TEXT = "LR-08", ("Issue a reminder memo to control-room shift leaders on the isolation procedure "
                                       "for major releases.")


def weak_fix_register():
    """The sample register with LR-08's administrative fix reworded as a reminder memo: a weaker change (P5-14), so
    LR-08, a level 3 cause, has nothing stronger than training and rec-review flags 'weak fix for a deep cause'."""
    r = json.loads(SAMPLE_REGISTER.read_text(encoding="utf-8"))
    e = next(x for x in r["entries"] if x["lr_id"] == WEAK_FIX_LR)
    rec = next(x for x in e["recommendations"] if x["hoc_tier"] == "administrative")
    rec["what"] = WEAK_FIX_TEXT
    return r


def capa_planner_9m() -> bytes:
    """A ProbeMatrix CAPA Planner workbook as the planner writes it: its four sheets and no Identified Findings, so
    nine-m-seed (and setup) refuse it. Invented rows only."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, head, row in [("Corrective Actions", ["ID", "Chapter", "Action"], ["CA-1", "M3", "Invented corrective row"]),
                             ("Preventive Actions", ["ID", "Chapter", "Action"], ["PA-1", "M9", "Invented preventive row"]),
                             ("Chapter Headers", ["Chapter", "Title"], ["M1", "Invented header"]),
                             ("Residual Risks", ["ID", "Risk"], ["RR-1", "Invented residual risk"])]:
        s = wb.create_sheet(title)
        s.append(head)
        s.append(row)
    return _xlsx_bytes(wb)


def process_export() -> bytes:
    """The synthetic process-industry export of tests/conftest.py process_export (P5-16): 30 invented findings whose
    roles use process-safety words. Without --industry process, 10 of them are F99; with it, none."""
    rows = json.loads(PROCESS_ROWS.read_text(encoding="utf-8"))
    wb = openpyxl.Workbook()
    s = wb.active
    s.title = "Incident Summary"
    for r in [("INCIDENT SUMMARY", None), (None, None), ("Field", "Value"),
              ("Incident Title", "Synthetic process unit release"), ("Date", "2024-05-02"),
              ("Investigation Status", "completed")]:
        s.append(r)
    f = wb.create_sheet("Identified Findings")
    f.append(FIND_COLS[:16] + ["User Edited"])
    for i, r in enumerate(rows):
        f.append([r["id"], r["code"], r["name"], r["level"], "Level", "Level", r["sub"], "substantiated", 3 + i % 2,
                  f"I-{i + 1}", r["name"], "Described briefly.", "Supported by EVD-0001.", "Described briefly.",
                  r["role"], "EVD-0001", "No"])
    e = wb.create_sheet("Evidence Links")
    e.append(["Finding Code", "Finding Name", "Evidence ID", "Evidence Type", "Relevance", "Specific Reference",
              "Evidence Description"])
    e.append(["AV 1.1", "Violation - Routine", "EVD-0001", "document", "Shows the step", "step skipped", "Work order"])
    c = wb.create_sheet("Causal Chain")
    c.append(["Link ID", "From Code", "From Level", "To Code"])
    c.append(["No causal links found", None, None, None])
    return _xlsx_bytes(wb)


def _xlsx_bytes(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def phase5_files():
    """{file name: bytes} for the Phase 5 inputs."""
    dump = lambda o: (json.dumps(o, indent=1, ensure_ascii=False) + "\n").encode("utf-8")   # noqa: E731
    return {"uat-weak-fix-register.json": dump(weak_fix_register()), "uat-capa-planner-9m.xlsx": capa_planner_9m(),
            "uat-process-export.xlsx": process_export()}


LIBRARY_CONTEXT = ("At Kestrel Works the pump went back into service on 2024-03-20 at 03:12 after a seal change, "
                   "with no leak test recorded.")


def no_urgency_register():
    """The sample register with every urgency removed: alert-grade returns grade_source "suggested" (P6-3)."""
    r = json.loads(SAMPLE_REGISTER.read_text(encoding="utf-8"))
    for e in r["entries"]:
        e.pop("urgency", None)
    return r


def library_register():
    """The pump-fire register with an invented site name, a full date and a clock time in LR-01's context, so a
    library publish of the synthetic run shows the site coded and the date kept to the month (P6-12)."""
    r = json.loads(json.dumps(REGISTER))
    r["title"] = "UAT library register: synthetic pump fire"
    r["entries"][0]["context"] = LIBRARY_CONTEXT
    return r


def phase6_files():
    """{file name: bytes} for the Phase 6 inputs."""
    dump = lambda o: (json.dumps(o, indent=1, ensure_ascii=False) + "\n").encode("utf-8")   # noqa: E731
    return {"uat-no-urgency-register.json": dump(no_urgency_register()),
            "uat-library-register.json": dump(library_register())}


def main():
    ap = argparse.ArgumentParser(description="Write the UAT input files.")
    ap.add_argument("--out", type=Path, default=OUT, help="output folder (default: docs/uat/inputs)")
    out = ap.parse_args().out
    out.mkdir(parents=True, exist_ok=True)
    for name, edit in WORKBOOKS.items():
        wb = base_workbook(*SUMMARY.get(name, ()))
        edit(wb)
        wb.save(out / name)
        print(f"wrote {out / name}")
    for name, p in {**PACKS, **PHASE2_JSON}.items():
        (out / name).write_text(json.dumps(p, indent=1) + "\n", encoding="utf-8")
        print(f"wrote {out / name}")
    for name, text in TEXT_FILES.items():
        (out / name).write_text(text, encoding="utf-8")
        print(f"wrote {out / name}")
    for name, data in {**phase3_files(), **phase4_files(), **phase5_files(), **phase6_files()}.items():
        (out / name).write_bytes(data)
        print(f"wrote {out / name}")


if __name__ == "__main__":
    main()
