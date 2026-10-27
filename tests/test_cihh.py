"""Phase 6, Task 3: the `self_check` kind (P6-5), `templates <ws> cihh` and `consolidate` (P6-9). Synthetic
fixtures only: the returns are built by write_return() from the template the engine writes. D7: the site, answer,
response and assessor cells stay blank in a template; the consolidated view never ranks a site."""
import json
import re
import shutil

import openpyxl
import pytest

from test_pack import item, pack_with, render_fixture
from test_pmlc import run
from test_privacy import _cli, share_ws, sign
import pmlc  # noqa: E402
import pmlc_pack as pp  # noqa: E402

from pmlc_lib import cihh  # noqa: E402
from pmlc_lib.handling import content_hash  # noqa: E402
from pmlc_lib.pack import schema as PS  # noqa: E402
from pmlc_lib.pack.profiles import BANNERS  # noqa: E402
from pmlc_lib.render.pptx import pptx_slides_text  # noqa: E402
from pmlc_lib.templates import OrgFileError  # noqa: E402

SC = {"kind": "self_check", "id": "sc-test", "condition": "A joint is returned to service without a leak test",
      "evidence_to_check": "Work orders after a seal change carry a signed leak-test record", "lr_id": "LR-01",
      "look_in": "Permit file and work-order system", "cites": ["FND-0001"]}
COLUMNS = ["lr_id", "condition", "evidence_to_check", "answer", "site_response", "assessed_by_role"]


def sc(n, lr, condition, evidence="The record", **extra):
    return dict(SC, id=f"sc-{n}", lr_id=lr, condition=condition, evidence_to_check=evidence, **extra)


THREE = [sc(1, "LR-01", "Joints returned to service untested"),
         sc(2, "LR-02", "Spares bought without a specification"),
         sc(3, "LR-03", "Alarms shelved without review")]


def cihh_pack(items=THREE):
    p = pack_with(*items)
    p["pack"] = "cihh-test"           # not could-it-happen-here: that pack's profile (Task 8) shapes its sections
    return p


@pytest.fixture(scope="module")
def base(tmp_path_factory):
    root = tmp_path_factory.mktemp("cihh-base")
    ws, orgs = share_ws(root)
    return root, ws, orgs


@pytest.fixture()
def env(base, tmp_path):
    """A fresh copy of the synthetic run for each test: (ws, outputs folder, orgs file, pack path)."""
    root, ws, orgs = base
    shutil.copytree(root / "pmlc-runs", tmp_path / "pmlc-runs")
    ws_ = tmp_path / ws.relative_to(root)
    out = ws_.parent / "outputs"
    out.mkdir(exist_ok=True)
    pack = out / "cihh-test.json"
    pack.write_text(json.dumps(cihh_pack()), encoding="utf-8")
    return ws_, out, orgs, pack


def template(env_, *extra):
    ws_, out, _, pack = env_
    r = _cli("templates", str(ws_), "cihh", "--pack", str(pack), *extra)
    assert r.returncode == 0, r.stderr + r.stdout
    return out / "cihh-test.cihh-template.xlsx", r


def write_return(tpl, dest, site, answers, responses=None, role="Site HSE lead"):
    """A filled return: a copy of the template with B1 set and one answer per data row (None leaves it blank)."""
    wb = openpyxl.load_workbook(tpl)
    sh = wb["Return"]
    sh["B1"] = site
    for i, ans in enumerate(answers):
        r = 5 + i
        sh.cell(row=r, column=4).value = ans
        sh.cell(row=r, column=5).value = (responses or {}).get(i, "Checked on the last two shutdowns" if ans else None)
        sh.cell(row=r, column=6).value = role if ans else None
    wb.save(dest)
    return dest


def rows_of(path):
    sh = openpyxl.load_workbook(path)["Return"]
    return [[c.value for c in r] for r in sh.iter_rows(min_row=1)]


# ----------------------------------------------------------------------------- the kind

def test_self_check_kind(ws_synth):
    spec = PS.KINDS["self_check"]
    assert spec["req"] == ["condition", "evidence_to_check"] and spec["opt"] == ["lr_id", "look_in"]
    assert spec["cite"] is True
    roles = dict(spec["fields"])
    assert roles == {"condition": "title", "lr_id": "badge", "evidence_to_check": "body", "look_in": "body"}
    assert pp.check_pack(pack_with(SC), ws_synth).problems == []
    bad = {k: v for k, v in SC.items() if k != "evidence_to_check"}
    assert "sections[0].items[0].evidence_to_check: required for self_check" in pp.check_pack(pack_with(bad),
                                                                                             ws_synth).problems
    nocite = {k: v for k, v in SC.items() if k != "cites"}
    assert any("a self_check needs at least one" in p for p in pp.check_pack(pack_with(nocite), ws_synth).problems)
    badlr = dict(SC, lr_id="L1")
    assert any(p.startswith("sections[0].items[0].lr_id") and "LR-" in p
               for p in pp.check_pack(pack_with(badlr), ws_synth).problems)
    pack = pack_with(SC)
    html = render_fixture("x", "html", pack=pack)[0].data
    md = render_fixture("x", "md", pack=pack)[0].data
    for text in (html, md):
        assert SC["condition"] in text and SC["evidence_to_check"] in text and "LR-01" in text
        assert SC["look_in"] in text
    docx = next(o for o in render_fixture("x", "docx", pack=pack) if o.name.endswith(".docx"))
    import io
    import zipfile
    body = zipfile.ZipFile(io.BytesIO(docx.data)).read("word/document.xml").decode("utf-8")
    assert SC["condition"] in body and SC["evidence_to_check"] in body
    pptx = next(o for o in render_fixture("x", "pptx", pack=pack) if o.name.endswith(".pptx"))
    assert SC["condition"] in json.dumps(pptx_slides_text(pptx.data), ensure_ascii=False)


@pytest.fixture(scope="module")
def ws_synth(base):
    return pmlc.load_ws(base[1])


# ----------------------------------------------------------------------------- templates cihh

def test_cihh_template_rows(env):
    tpl, r = template(env)
    rows = rows_of(tpl)
    assert rows[0][0] == "Site" and rows[1][0] == "Pack"
    assert rows[1][1] == content_hash(cihh_pack())
    assert rows[3] == COLUMNS
    data = rows[4:]
    assert [d[0] for d in data] == ["LR-01", "LR-02", "LR-03"]
    assert [d[1] for d in data] == [x["condition"] for x in THREE]
    assert "3 self-check" in r.stdout


def test_cihh_template_org_columns_blank(env):
    tpl, _ = template(env)
    rows = rows_of(tpl)
    assert rows[0][1] in (None, "")
    for d in rows[4:]:
        assert d[3:6] == [None, None, None], d


def test_cihh_template_never_overwrites(env):
    tpl, _ = template(env)
    before = tpl.read_bytes()
    ws_, _, _, pack = env
    r = _cli("templates", str(ws_), "cihh", "--pack", str(pack))
    assert r.returncode == 1
    assert "already exists and may hold your organisation's entries" in r.stderr
    assert tpl.read_bytes() == before


def test_cihh_template_needs_pack(env):
    ws_ = env[0]
    r = _cli("templates", str(ws_), "cihh")
    assert r.returncode != 0 and "--pack" in r.stderr and "Traceback" not in r.stderr


def test_cihh_template_shared_copy(env):
    ws_, out, orgs, pack = env
    doc = cihh_pack([sc(1, "LR-01", "Pump P-101 restarted at 14:05 without a leak test"), THREE[1]])
    pack.write_text(json.dumps(doc), encoding="utf-8")
    refused = _cli("templates", str(ws_), "cihh", "--pack", str(pack), "--share", "external", "--orgs", str(orgs))
    assert refused.returncode == 1 and "cannot share externally" in refused.stderr
    sign(ws_, pack, "--classification", "external", "--legal-review", "cleared")
    r = _cli("templates", str(ws_), "cihh", "--pack", str(pack), "--share", "external", "--orgs", str(orgs))
    assert r.returncode == 0, r.stderr
    assert sorted(x.name for x in out.glob("*.xlsx")) == ["cihh-test.cihh-template.external.xlsx"]
    rows = rows_of(out / "cihh-test.cihh-template.external.xlsx")
    text = json.dumps(rows)
    assert "P-101" not in text and "14:05" not in text
    signed = json.loads(pack.read_text(encoding="utf-8"))
    assert rows[1][1] == content_hash(signed)                      # the internal pack's hash: returns consolidate
    assert "Shared external" in r.stdout


def test_cihh_template_name_guard(env, tmp_path):
    ws_, _, _, pack = env
    names = tmp_path / "names.txt"
    names.write_text("Tarvo Quillane\n", encoding="utf-8")
    doc = cihh_pack([sc(1, "LR-01", "Tarvo Quillane restarts the pump without a test")])
    pack.write_text(json.dumps(doc), encoding="utf-8")
    r = _cli("templates", str(ws_), "cihh", "--pack", str(pack), "--names", str(names))
    assert r.returncode == 1 and "Tarvo Quillane" in r.stderr           # the pack check (JC1) refuses it first
    assert not list(ws_.parent.glob("outputs/*.xlsx"))


def test_templates_cihh_thin(env):
    ws_, out, _, pack = env
    pack.write_text(json.dumps(cihh_pack([item("para", text="No transferable learning yet.",
                                                cites=["FND-0001"])])), encoding="utf-8")
    tpl, r = template(env)
    assert "the pack has no self_check items" in r.stdout
    rows = rows_of(tpl)
    assert rows[3] == COLUMNS and len(rows) == 4


def test_template_function_direct(env, tmp_path):
    ws_, _, _, _ = env
    w = pmlc.load_ws(ws_)
    p = cihh.write_template(cihh_pack(), w, tmp_path / "t.xlsx", None)
    assert p == tmp_path / "t.xlsx" and rows_of(p)[4][0] == "LR-01"
    with pytest.raises(SystemExit):
        cihh.write_template(cihh_pack(), w, tmp_path / "t.xlsx", None)


# ----------------------------------------------------------------------------- read_return

def test_read_return(env, tmp_path):
    tpl, _ = template(env)
    f = write_return(tpl, tmp_path / "a.xlsx", "  North Works ", [" Present", "N/A", None])
    got = cihh.read_return(f)
    assert got["site"] == "North Works" and got["pack_hash"] == content_hash(cihh_pack())
    assert [r["answer"] for r in got["rows"]] == ["present", "n/a", ""]
    assert got["rows"][0]["_row"] == 5


@pytest.mark.parametrize("raw", ["N/A", "na", "n.a.", "NA ", " n/a"])
def test_read_return_na_forms(env, tmp_path, raw):
    tpl, _ = template(env)
    f = write_return(tpl, tmp_path / "a.xlsx", "North Works", [raw, None, None])
    assert cihh.read_return(f)["rows"][0]["answer"] == "n/a"


def test_read_return_errors_carry_rows(env, tmp_path):
    tpl, _ = template(env)
    f = write_return(tpl, tmp_path / "a.xlsx", "North Works", ["present", "maybe", "absent"])
    with pytest.raises(OrgFileError) as e:
        cihh.read_return(f)
    assert e.value.problems == [f"{f} row 6: answer 'maybe' must be one of present, partly, absent, n/a (or blank)"]
    assert [r["answer"] for r in e.value.ret["rows"]] == ["present", "absent"]


# ----------------------------------------------------------------------------- consolidate

def three_returns(env_, tmp_path):
    tpl, _ = template(env_)
    a = write_return(tpl, tmp_path / "north.xlsx", "North Works", ["present", "absent", "partly"])
    b = write_return(tpl, tmp_path / "south.xlsx", "South Terminal", ["partly", None, "present"])
    c = write_return(tpl, tmp_path / "east.xlsx", "East Depot", ["absent", "n/a", None])
    return a, b, c


def consolidated(out, mode=None):
    name = "cihh-consolidated.json" if mode is None else f"cihh-consolidated.{mode}.json"
    return json.loads((out / name).read_text(encoding="utf-8"))


def test_consolidate_coverage(env, tmp_path):
    ws_, out, _, pack = env
    files = three_returns(env, tmp_path)
    r = _cli("consolidate", str(ws_), *map(str, files), "--pack", str(pack))
    assert r.returncode == 0, r.stderr
    d = consolidated(out)
    assert d["version"] == 1 and d["returns"] == 3
    assert d["coverage"] == {"LR-01": {"assessed": 3, "not_assessed": 0}, "LR-02": {"assessed": 2, "not_assessed": 1},
                             "LR-03": {"assessed": 2, "not_assessed": 1}}
    assert d["conditions"] == [
        {"lr_id": "LR-01", "condition": THREE[0]["condition"], "present": 1, "partly": 1},
        {"lr_id": "LR-02", "condition": THREE[1]["condition"], "present": 0, "partly": 0},
        {"lr_id": "LR-03", "condition": THREE[2]["condition"], "present": 1, "partly": 1}]
    assert d["sites"] == {"East Depot": ["LR-01", "LR-02"], "North Works": ["LR-01", "LR-02", "LR-03"],
                          "South Terminal": ["LR-01", "LR-03"]}
    assert sorted(d) == ["conditions", "coverage", "returns", "sites", "version"]
    assert "3 return(s)" in r.stdout
    assert BANNERS["cihh"] in r.stdout


def test_consolidate_function_direct(env, tmp_path):
    files = three_returns(env, tmp_path)
    d = cihh.consolidate([cihh.read_return(f) for f in files], cihh_pack())
    assert d["returns"] == 3 and d["coverage"]["LR-01"] == {"assessed": 3, "not_assessed": 0}


def test_consolidate_no_league_table(env, tmp_path):
    ws_, out, _, pack = env
    files = three_returns(env, tmp_path)
    assert _cli("consolidate", str(ws_), *map(str, files), "--pack", str(pack)).returncode == 0
    d = consolidated(out)
    for site, v in d["sites"].items():
        assert isinstance(v, list) and all(isinstance(x, str) and re.fullmatch(r"LR-\d{2,}", x) for x in v), site

    def keys(o):
        if isinstance(o, dict):
            for k, v in o.items():
                yield k
                yield from keys(v)
        elif isinstance(o, list):
            for v in o:
                yield from keys(v)
    for k in keys(d):
        assert not re.search(r"rank|score|total", k, re.I), k


def _fails(r, *needles):
    assert r.returncode == 1, r.stdout + r.stderr
    assert "Traceback" not in r.stderr
    for n in needles:
        assert n in r.stderr, r.stderr


def test_consolidate_rejects_missing_site(env, tmp_path):
    ws_, out, _, pack = env
    a, b, c = three_returns(env, tmp_path)
    bad = write_return(template_path(env), tmp_path / "nosite.xlsx", None, ["present", "present", "present"])
    r = _cli("consolidate", str(ws_), str(a), str(bad), "--pack", str(pack))
    _fails(r, f"{bad} row 1: Site is blank")
    assert not (out / "cihh-consolidated.json").exists()
    r = _cli("consolidate", str(ws_), str(a), str(bad), "--pack", str(pack), "--keep-going")
    assert r.returncode == 0, r.stderr
    assert f"{bad} row 1: Site is blank" in r.stderr
    d = consolidated(out)
    assert d["returns"] == 1 and list(d["sites"]) == ["North Works"]


def template_path(env_):
    return env_[1] / "cihh-test.cihh-template.xlsx"


def test_consolidate_rejects_duplicate_site(env, tmp_path):
    ws_, out, _, pack = env
    a, b, c = three_returns(env, tmp_path)
    dup = write_return(template_path(env), tmp_path / "north2.xlsx", "north works ", ["absent", "absent", "absent"])
    r = _cli("consolidate", str(ws_), str(a), str(b), str(dup), "--pack", str(pack))
    _fails(r, f"{dup} row 1: site 'north works' already has a return ({a})")
    r = _cli("consolidate", str(ws_), str(a), str(b), str(dup), "--pack", str(pack), "--keep-going")
    assert r.returncode == 0
    d = consolidated(out)
    assert d["returns"] == 2 and d["conditions"][0]["present"] == 1


def test_consolidate_rejects_bad_answer(env, tmp_path):
    ws_, out, _, pack = env
    tpl = template(env)[0]
    bad = write_return(tpl, tmp_path / "bad.xlsx", "West Yard", ["yes", "present", "sometimes"])
    r = _cli("consolidate", str(ws_), str(bad), "--pack", str(pack))
    _fails(r, f"{bad} row 5: answer 'yes' must be one of present, partly, absent, n/a (or blank)",
           f"{bad} row 7: answer 'sometimes'")
    assert len([x for x in r.stderr.splitlines() if " row " in x]) == 2
    r = _cli("consolidate", str(ws_), str(bad), "--pack", str(pack), "--keep-going")
    assert r.returncode == 0
    d = consolidated(out)
    assert d["sites"] == {"West Yard": ["LR-02"]} and d["conditions"][1]["present"] == 1


def test_consolidate_rejects_other_pack(env, tmp_path):
    ws_, out, _, pack = env
    a, _, _ = three_returns(env, tmp_path)
    wb = openpyxl.load_workbook(a)
    wb["Return"]["B2"] = "sha256:" + "0" * 64
    other = tmp_path / "other.xlsx"
    wb.save(other)
    r = _cli("consolidate", str(ws_), str(other), "--pack", str(pack))
    _fails(r, f"{other} row 2: made from another pack's template")
    r = _cli("consolidate", str(ws_), str(a), str(other), "--pack", str(pack), "--keep-going")
    assert r.returncode == 0 and consolidated(out)["returns"] == 1


def test_consolidate_rejects_not_a_return(env, tmp_path):
    ws_, out, _, pack = env
    a, _, _ = three_returns(env, tmp_path)
    junk = tmp_path / "notes.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = "Minutes of the meeting"
    wb.save(junk)
    text = tmp_path / "return.xlsx"
    text.write_text("not a workbook", encoding="utf-8")
    r = _cli("consolidate", str(ws_), str(junk), str(text), "--pack", str(pack))
    _fails(r, f"{junk}: not a could-it-happen-here return", f"{text}: not a could-it-happen-here return")
    r = _cli("consolidate", str(ws_), str(a), str(junk), "--pack", str(pack), "--keep-going")
    assert r.returncode == 0 and consolidated(out)["returns"] == 1


def test_consolidate_unknown_lr_row(env, tmp_path):
    ws_, out, _, pack = env
    a, _, _ = three_returns(env, tmp_path)
    wb = openpyxl.load_workbook(a)
    wb["Return"]["A6"] = "LR-99"
    f = tmp_path / "edited.xlsx"
    wb.save(f)
    r = _cli("consolidate", str(ws_), str(f), "--pack", str(pack))
    _fails(r, f"{f} row 6: lr_id 'LR-99' is not a self-check of the pack")


def test_consolidate_nothing_usable(env, tmp_path):
    ws_, out, _, pack = env
    junk = tmp_path / "x.xlsx"
    junk.write_text("x", encoding="utf-8")
    r = _cli("consolidate", str(ws_), str(junk), "--pack", str(pack), "--keep-going")
    _fails(r, "no usable return")
    assert not (out / "cihh-consolidated.json").exists()


def test_consolidate_pack_without_self_check(env, tmp_path):
    ws_, out, _, pack = env
    a, _, _ = three_returns(env, tmp_path)
    pack.write_text(json.dumps(cihh_pack([item("para", text="t", cites=["FND-0001"])])), encoding="utf-8")
    r = _cli("consolidate", str(ws_), str(a), "--pack", str(pack))
    _fails(r, "the pack has no self_check items")


def test_cihh_events_shape(env, tmp_path):
    ws_, out, _, pack = env
    files = three_returns(env, tmp_path)
    assert _cli("consolidate", str(ws_), *map(str, files), "--pack", str(pack)).returncode == 0
    lines = (out / "cihh-events.jsonl").read_text(encoding="utf-8").splitlines()
    assert len(lines) == 7                                   # the answered rows: 3 + 2 + 2
    keys = {"type", "lr_id", "site", "answer", "site_response", "source", "event_hash"}
    for ln in lines:
        e = json.loads(ln)
        assert set(e) == keys and e["type"] == "site_assessment" and e["source"] == "organisation"
        assert ln == json.dumps(e, sort_keys=True, ensure_ascii=False)
        body = {k: v for k, v in e.items() if k != "event_hash"}
        assert e["event_hash"] == cihh.event_hash(body)
        assert re.fullmatch(r"sha256:[0-9a-f]{64}", e["event_hash"])
    assert {json.loads(x)["site"] for x in lines} == {"North Works", "South Terminal", "East Depot"}
    first = (out / "cihh-events.jsonl").read_bytes()
    assert _cli("consolidate", str(ws_), *map(str, files), "--pack", str(pack)).returncode == 0
    assert (out / "cihh-events.jsonl").read_bytes() == first          # stable across runs


def test_consolidate_shared_copy(env, tmp_path):
    ws_, out, orgs, pack = env
    files = three_returns(env, tmp_path)
    wb = openpyxl.load_workbook(files[0])
    wb["Return"]["E5"] = "Bharat Seals fitted the seal at 14:05"
    wb["Return"]["E6"] = "north  works checked this; South Terminal uses the same permit"
    wb.save(files[0])
    r = _cli("consolidate", str(ws_), *map(str, files), "--pack", str(pack), "--share", "external",
             "--orgs", str(orgs))
    assert r.returncode == 1 and "cannot share externally" in r.stderr
    sign(ws_, pack, "--classification", "external", "--legal-review", "cleared")
    r = _cli("consolidate", str(ws_), *map(str, files), "--pack", str(pack), "--share", "external",
             "--orgs", str(orgs))
    assert r.returncode == 0, r.stderr
    assert not (out / "cihh-consolidated.json").exists() and not (out / "cihh-events.jsonl").exists()
    d = consolidated(out, "external")
    assert sorted(d["sites"]) == ["S01", "S02", "S03"]
    ev = (out / "cihh-events.external.jsonl").read_text(encoding="utf-8")
    assert {json.loads(x)["site"] for x in ev.splitlines()} == {"S01", "S02", "S03"}
    assert all(json.loads(x)["event_hash"] == cihh.event_hash(json.loads(x)) for x in ev.splitlines())
    assert d["sites"]["S02"] == ["LR-01", "LR-02", "LR-03"]           # North Works, in name order
    for leak in ("North Works", "north  works", "South Terminal", "East Depot", "Bharat Seals", "14:05"):
        assert leak.casefold() not in ev.casefold() and leak.casefold() not in json.dumps(d).casefold(), leak
    assert "S02 checked this; S03 uses the same permit" in ev           # a response naming sites gets their codes
    assert "Shared external" in r.stdout


def test_consolidate_name_guard(env, tmp_path):
    ws_, out, _, pack = env
    files = three_returns(env, tmp_path)
    wb = openpyxl.load_workbook(files[0])
    wb["Return"]["E5"] = "Checked by Tarvo Quillane"
    wb.save(files[0])
    names = tmp_path / "names.txt"
    names.write_text("Tarvo Quillane\n", encoding="utf-8")
    r = _cli("consolidate", str(ws_), *map(str, files), "--pack", str(pack), "--names", str(names))
    assert r.returncode == 1 and "a listed name appears" in r.stderr
    assert not (out / "cihh-events.jsonl").exists()


def test_cihh_rows_without_lr_id(env, tmp_path):
    """A self_check with no lr_id is keyed by its item id; two checks of one learning stay two conditions."""
    ws_, out, _, pack = env
    items = [dict(sc(1, "LR-01", "First condition")), sc(2, "LR-01", "Second condition")]
    nolr = {k: v for k, v in sc(3, "LR-03", "Third condition").items() if k != "lr_id"}
    pack.write_text(json.dumps(cihh_pack(items + [nolr])), encoding="utf-8")
    tpl, _ = template(env)
    assert [r[0] for r in rows_of(tpl)[4:]] == ["LR-01", "LR-01", "sc-3"]
    f = write_return(tpl, tmp_path / "a.xlsx", "North Works", ["present", "partly", "absent"])
    assert _cli("consolidate", str(ws_), str(f), "--pack", str(pack)).returncode == 0
    d = consolidated(out)
    assert d["coverage"]["LR-01"] == {"assessed": 2, "not_assessed": 0}
    assert [(c["condition"], c["present"], c["partly"]) for c in d["conditions"]] == [
        ("First condition", 1, 0), ("Second condition", 0, 1), ("Third condition", 0, 0)]


def test_artefacts_registered():
    from pmlc_lib import contracts
    art = contracts.load_artefacts()
    assert art["cihh-template"] == {"file": "outputs/<pack>.cihh-template.xlsx", "producer": "engine:templates",
                                    "org_filled": True}
    assert art["cihh-consolidated"] == {"file": "outputs/cihh-consolidated.json", "producer": "engine:consolidate",
                                        "org_filled": False}


def test_consolidate_help():
    assert "--keep-going" in run("consolidate", "-h")


def test_templates_refuses_cihh_flags_elsewhere(env):
    """templates adoption|distribution|verification refuse the cihh-only flags instead of ignoring them."""
    ws_, _, orgs, pack = env
    for flags in (["--pack", str(pack)], ["--share", "external"], ["--names", str(orgs)], ["--orgs", str(orgs)]):
        r = _cli("templates", str(ws_), "adoption", *flags)
        assert r.returncode != 0 and "is for templates cihh only" in r.stderr, flags


def test_plan_engine_commands_cihh_template():
    from pmlc_lib.plan import _engine_commands
    assert _engine_commands("templates", ["adoption", "cihh-template"]) == [
        "PMLC templates <ws> adoption",
        "PMLC templates <ws> cihh --pack <run>/outputs/could-it-happen-here.json"]


def test_plan_engine_templates_step_leaves_out_cihh(env):
    """full-cascade's bare templates step lists the three organisation files only: the CIHH template needs --pack,
    so only the could-it-happen-here skill's own instructions write it."""
    ws_ = env[0]
    r = _cli("plan", str(ws_), "--workflow", "full-cascade")
    assert r.returncode == 0, r.stderr
    plan = json.loads((ws_ / "plan.json").read_text(encoding="utf-8"))
    step = next(s for s in plan["steps"] if s["skill"] == "engine:templates")
    assert "cihh" not in json.dumps(step) and step["produces"] == "adoption,distribution,verification"
