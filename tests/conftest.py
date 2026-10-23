"""Session-wide test helpers."""
import os
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "skills" / "learnings-coach" / "scripts"))

PROBE_SVG = ('<?xml version="1.0" encoding="UTF-8"?>\n<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 40 20" '
             'width="40" height="20"><rect width="40" height="20" fill="#fff"/></svg>\n')


def probe_png():
    """None when `--format png` really renders on this machine, else the reason it does not (final review C1).

    A browser on PATH is not enough: on a stock Ubuntu 24.04 runner Chrome is found, yet its sandbox may refuse to
    start. So the probe renders one tiny SVG through the engine's own code path (with its --no-sandbox retry), and
    the png tests skip with the reason it records, so a CI log says why."""
    from pmlc_lib import render
    from pmlc_lib.render import graphs
    browser = render.find_browser()
    if not browser:
        return "no Chromium-family browser"
    data, why = graphs._gx_png(browser, PROBE_SVG)
    return None if data else f"{browser}: {why}"


def png_gate(reason):
    """Skip when png does not render (reason is set); fail instead under PMLC_REQUIRE_PNG=1 (N1/X20), so a CI leg or
    the local gate that has a browser can never pass by skipping the png tests."""
    if not reason:
        return
    msg = f"png does not render here: {reason}"
    if os.environ.get("PMLC_REQUIRE_PNG") == "1":
        pytest.fail(msg)
    pytest.skip(msg)


@pytest.fixture(scope="session")
def png_probe():
    return probe_png()


@pytest.fixture
def needs_png(png_probe):
    """Skip (a plain pytest.skip, not a Node marker) unless png really renders here; fail under PMLC_REQUIRE_PNG=1."""
    png_gate(png_probe)


def process_export(d: Path, name: str = "process-export.xlsx") -> Path:
    """A synthetic process-industry export (P5-16): 30 findings whose roles use process-safety vocabulary (permits,
    isolation, alarms, relief systems, turnaround, MOC). Rows come from tests/fixtures/process-export/findings.json.
    Invented content only (D3)."""
    import json
    import openpyxl
    rows = json.loads((Path(__file__).parent / "fixtures" / "process-export" / "findings.json").read_text())
    wb = openpyxl.Workbook()
    s = wb.active
    s.title = "Incident Summary"
    for r in [("INCIDENT SUMMARY", None), (None, None), ("Field", "Value"), ("Incident Title", "Synthetic process unit release"),
              ("Date", "2024-05-02"), ("Investigation Status", "completed")]:
        s.append(r)
    f = wb.create_sheet("Identified Findings")
    f.append(["Finding ID", "Nanocode", "Name", "Level", "Level Name", "Category", "Subcategory", "Classification",
              "Rating", "Instance ID", "Instance Label", "Short Description", "AI Justification", "Causal Explanation",
              "Individuals Involved", "Linked Evidence", "User Edited"])
    for i, r in enumerate(rows):
        f.append([r["id"], r["code"], r["name"], r["level"], "Level", "Level", r["sub"], "substantiated", 3 + i % 2,
                  f"I-{i + 1}", r["name"], "Described briefly.", "Supported by EVD-0001.", "Described briefly.",
                  r["role"], "EVD-0001", "No"])
    e = wb.create_sheet("Evidence Links")
    e.append(["Finding Code", "Finding Name", "Evidence ID", "Evidence Type", "Relevance", "Specific Reference",
              "Evidence Description"])
    e.append(["AV 1.1", "Violation - Routine", "EVD-0001", "document", "Shows the step", "step skipped", "Work order"])
    c = wb.create_sheet("Causal Chain")
    c.append(["Link ID", "From Code", "From Level", "To Code"])
    c.append(["No causal links found", None, None, None])
    x = d / name
    wb.save(x)
    return x
