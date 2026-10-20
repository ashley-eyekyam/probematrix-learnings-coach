"""The register workbook (Phase 3, Task 6): tables, validation lists, print setup, LR links, share, run folder."""
import copy

import openpyxl

from test_pmlc import MINI_REG, run
from test_privacy import _cli, shws as shws, sign, write_doc  # noqa: F401,F811
from pmlc_lib.register.check import HOC_TIERS

STATUS = "Learning purposes: not a statutory document"


def two_entry_reg():
    reg = copy.deepcopy(MINI_REG)
    second = copy.deepcopy(reg["entries"][0])
    second["lr_id"], second["title"], second["priority"], second["comparison"] = "LR-02", "Second", 2, "partial"
    reg["entries"].append(second)
    return reg


def build(tmp_path, reg=None, name="reg"):
    rp = write_doc(tmp_path, reg or two_entry_reg(), f"{name}.json")
    out = tmp_path / f"{name}.xlsx"
    run("register-xlsx", str(rp), "--out", str(out))
    return openpyxl.load_workbook(out)


def col_of(ws, header):
    return next(c.column_letter for c in ws[1] if c.value == header)


def test_xlsx_tables_and_filters(tmp_path):
    wb = build(tmp_path)
    for title, name in [("Learnings Register", "LearningsRegister"), ("Recommendations", "Recommendations"),
                        ("Learning Plan", "LearningPlan")]:
        ws = wb[title]
        assert list(ws.tables) == [name], title
        assert ws.tables[name].ref == ws.dimensions
        assert ws.tables[name].tableStyleInfo.name == "TableStyleMedium2"
    assert not wb["Summary"].tables and not wb["About"].tables


def dv_for(ws, col):
    rng = f"{col}2:{col}{ws.max_row}"
    return [d for d in ws.data_validations.dataValidation if d.type == "list" and rng in str(d.sqref).split()]


def test_xlsx_validation_lists(tmp_path):
    wb = build(tmp_path)
    reg, rec = wb["Learnings Register"], wb["Recommendations"]
    assert dv_for(reg, col_of(reg, "Priority"))[0].formula1 == '"1,2,3"'
    assert dv_for(reg, col_of(reg, "Counterpart in comparator report"))[0].formula1 == '"yes,partial,none,not compared"'
    assert dv_for(rec, col_of(rec, "Hierarchy-of-controls tier"))[0].formula1 == '"' + ",".join(HOC_TIERS) + '"'


def test_xlsx_print_setup(tmp_path):
    wb = build(tmp_path)
    for ws in wb.worksheets:
        assert ws.page_setup.orientation == "landscape", ws.title
        assert ws.sheet_properties.pageSetUpPr.fitToPage and ws.page_setup.fitToWidth == 1, ws.title
        assert ws.page_setup.fitToHeight == 0, ws.title
        assert ws.oddHeader.center.text.startswith(STATUS) and ws.oddFooter.center.text, ws.title
    for t in ("Learnings Register", "Recommendations", "Learning Plan"):
        assert wb[t].print_title_rows.replace("$", "") == "1:1"  # openpyxl reloads it as $1:$1


def test_xlsx_lr_hyperlinks(tmp_path):
    wb = build(tmp_path)
    rows = {c.value: c.row for c in wb["Learnings Register"]["A"][1:]}
    rec = wb["Recommendations"]
    cells = rec["A"][1:]
    assert len(cells) == 4
    for c in cells:
        assert c.hyperlink.location == f"'Learnings Register'!A{rows[c.value]}"


def test_xlsx_share_resets_docprops(tmp_path, shws):  # noqa: F811
    ws, orgs = shws
    rp = write_doc(tmp_path, copy.deepcopy(MINI_REG), "reg.json")
    sign(ws, rp, "--classification", "external", "--legal-review", "cleared")
    r = _cli("register-xlsx", str(rp), "--ws", str(ws), "--share", "external", "--orgs", str(orgs))
    assert r.returncode == 0, r.stderr
    p = openpyxl.load_workbook(tmp_path / "learnings-register.external.xlsx").properties
    assert not p.creator and not p.lastModifiedBy
    assert p.created.year == 2000 and p.modified.year == 2000
    # an internal copy keeps the ordinary properties
    internal = build(tmp_path).properties
    assert internal.created.year > 2000


def test_xlsx_run_folder_fallback(tmp_path):
    run_dir = tmp_path / "pmlc-runs" / "2024-03-20-plant-pump-fire"
    (run_dir / "workspace").mkdir(parents=True)
    (run_dir / "workspace" / "findings.json").write_text("[]", encoding="utf-8")
    side = run_dir / "elsewhere"
    side.mkdir()
    rp = write_doc(side, copy.deepcopy(MINI_REG), "reg.json")
    run("register-xlsx", str(rp))
    assert (side / "learnings-register.xlsx").exists()
    assert not (run_dir / "learnings-register.xlsx").exists()
    assert not (run_dir / "outputs").exists()
    wb = openpyxl.load_workbook(side / "learnings-register.xlsx")
    about = {r[0]: r[1] for r in wb["About"].iter_rows(values_only=True) if len(r) > 1 and r[1] is not None}
    assert "no workspace to check" in about["Review status"]


def test_xlsx_new_columns_only_when_present(tmp_path):
    wb = build(tmp_path, name="plain")
    heads = [c.value for c in wb["Learnings Register"][1]]
    assert "Urgency" not in heads and len(heads) == 28
    reg = two_entry_reg()
    reg["entries"][1].update(urgency="alert", recommendation_class="strategic", regulatory_touchpoint="Factories Act s.7")
    wb = build(tmp_path, reg, "rich")
    ws = wb["Learnings Register"]
    heads = [c.value for c in ws[1]]
    assert heads[28:] == ["Urgency", "Recommendation class", "Regulatory touchpoint"]
    assert [c.value for c in ws[3]][28:] == ["alert", "strategic", "Factories Act s.7"]
    assert list(ws.tables) == ["LearningsRegister"] and ws.tables["LearningsRegister"].ref == ws.dimensions
