"""Task 12 (spec §7.3-§7.7; P5-13, P5-18): the 9M session engine (nine-m-suggest, nine-m-answer, nine-m-replay,
nine-m-actions-template), select --m, the profile rules nine-m-overview, nine-m-cites and m1-balance, the conditional
nine-m-demo banner, the nine-m-map.json side output, the register 9M column and the dashboard 9M strip.
Synthetic exports only (fictional event, no real data), plus one Hazira dry run for §7.7 item 11."""
import copy
import csv
import io
import json
import re
import subprocess
import sys

import openpyxl
import pytest

from test_ninem_seed import REAL_CODES
from test_pack import item
from test_pmlc import make_export, PMLC, run
from test_privacy import share_ws, sign, write_doc
import pmlc  # noqa: E402
import pmlc_pack as pp  # noqa: E402
from pmlc_lib import ninem, ninem_session as ns  # noqa: E402
from pmlc_lib.pack import profiles  # noqa: E402
from pmlc_lib.register.check import check_register  # noqa: E402

FND_RX = re.compile(r"^FND-[0-9A-F]+$")
_REAL_LOAD = profiles.load
NINE_M_PROFILE = {
    "sections": [{"id": "nine-m-overview", "kinds": ["table"], "min": 1, "max": 1},
                 {"id": "m*", "kinds": ["learning", "barrier", "recommendation", "question"], "min": 0, "max": 12},
                 {"id": "not-relevant", "kinds": ["table"], "min": 0, "max": 1},
                 {"id": "organisation-context", "kinds": ["para"], "min": 0, "max": 5}],
    "bottom_line": {"min": 1, "max": 7},
    "rules": ["nine-m-overview", "nine-m-cites", "m1-balance", "no-owner-due", "banner:nine-m-demo"],
}


def cli(*args):
    return subprocess.run([sys.executable, str(PMLC), *args], capture_output=True, text=True)


def nine_m_export(d, name="nine-m.xlsx"):
    """The synthetic export with causal rows, its invented Nanocodes swapped for codes whose prefixes the 9M map
    knows."""
    x = make_export(d, name=name, causal_rows=True)
    wb = openpyxl.load_workbook(x)
    for sheet in wb.sheetnames:
        s = wb[sheet]
        for row in s.iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    for old, new in REAL_CODES.items():
                        if old in c.value:
                            c.value = c.value.replace(old, new)
    wb.save(x)
    return x


@pytest.fixture(scope="module")
def nws(tmp_path_factory):
    """A seeded synthetic workspace: ingest, classify, causal and nine-m-seed (no session yet)."""
    d = tmp_path_factory.mktemp("ninem-session")
    x = nine_m_export(d)
    run("ingest", str(x), "--slug", "test-pump", "--date", "2024-03-20")
    ws = d / "pmlc-runs" / "2024-03-20-test-pump" / "workspace"
    for c in ("classify", "causal", "nine-m-seed"):
        run(c, str(ws))
    return ws


@pytest.fixture()
def fresh(nws, tmp_path):
    """A private copy of the seeded workspace, so session writes never leak between tests."""
    import shutil
    run_dir = tmp_path / "pmlc-runs" / nws.parent.name
    shutil.copytree(nws.parent, run_dir)
    return run_dir / "workspace"


def session_of(ws):
    return json.loads((ws / "nine_m_session.json").read_text(encoding="utf-8"))


def ctx(ws):
    w = ns.load_session_ws(ws)
    return w, w["nine_m_seed"]


# ----------------------------------------------------------------------------- suggest

@pytest.mark.parametrize("step", ns.NINE_M_STEPS)
def test_suggest_deterministic_order_and_cites(nws, step):
    """§7.7 items 4 and 10: the same call gives the same list; 2-4 suggestions plus None; every fact-bearing
    suggestion cites a finding of this workspace."""
    w, seed = ctx(nws)
    ids = {f["id"] for f in w["findings"]}
    for m in ninem.NINE_M_CODES:
        a = ns.suggest(w, seed, m, step)
        assert a == ns.suggest(w, seed, m, step)
        assert 2 <= len(a) - 1 <= 4, (m, step, a)
        for s in a[:-1]:
            fnds = [x for x in s["cites"] if x.startswith("FND-")]
            assert all(x in ids for x in fnds)
            if not ns.is_procedural(s["text"]):
                assert fnds, (m, step, s)


def test_suggest_rating_order(nws):
    """Candidates come in rating order: the relevance suggestion for M3 cites its top-rated finding first."""
    w, seed = ctx(nws)
    first = ns.suggest(w, seed, "M3", "relevance")[0]
    R = {f["id"]: f["rating"] for f in w["findings"]}
    assert first["text"].startswith("Relevant")
    assert [R[x] for x in first["cites"]] == sorted((R[x] for x in first["cites"]), reverse=True)


def test_suggest_always_offers_none(nws):
    w, seed = ctx(nws)
    for m in ninem.NINE_M_CODES:
        for step in ns.NINE_M_STEPS:
            last = ns.suggest(w, seed, m, step, top=2)[-1]
            assert last == {"text": ninem.NINE_M_NONE, "cites": []}
    for step in ns.NINE_M_GLOBAL_STEPS:
        assert ns.suggest(w, seed, None, step, session={"records": []})[-1]["text"] == ninem.NINE_M_NONE


def test_suggest_never_offers_discipline(nws):
    """§7.5: discipline, accountability and consequences are never offered, even when the export's own words carry
    them; training comes only with a stronger tier and says it is weak on its own."""
    w, seed = ctx(nws)
    w = copy.deepcopy(w)
    for f in w["findings"]:
        f["short"] = "Disciplinary action and consequences for the pump technician; " + f["short"]
        f["label"] = "Accountability gap: " + f["label"]
    for m in ninem.NINE_M_CODES:
        for step in ns.NINE_M_STEPS:
            for s in ns.suggest(w, seed, m, step, top=4):
                text = s["text"]
                if step == "subcategory":   # framework names (D9, verbatim) may hold the word; nothing else may
                    text = re.sub(r"M\d\.\d\d [^(]*", "", text)
                assert not re.search(r"disciplin|accountab|consequence|punish|sanction", text, re.I), (m, step, s)
    change = [s["text"] for s in ns.suggest(w, seed, "M3", "change", top=4)]
    assert change[0].startswith("Elimination")
    trained = [t for t in change if "training" in t.lower()]
    assert trained and all("(weak on its own)" in t for t in trained)
    assert all(not t.lower().startswith("training") for t in change)


def test_suggest_subcategory_ids(nws):
    w, seed = ctx(nws)
    subs = {s["id"] for c in ninem.load_nine_m()["chapters"] for s in c["subcategories"]}
    for m in ninem.NINE_M_CODES:
        for s in ns.suggest(w, seed, m, "subcategory"):
            got = re.match(r"(M\d\.\d\d) ", s["text"])
            if got:
                assert got.group(1) in subs and got.group(1).startswith(m + ".")


def test_suggest_refuses_bad_input(nws):
    w, seed = ctx(nws)
    with pytest.raises(ValueError):
        ns.suggest(w, seed, "M10", "relevance")
    with pytest.raises(ValueError):
        ns.suggest(w, seed, "M3", "blame")


# ----------------------------------------------------------------------------- session: record and answer

def test_session_record_and_answer(fresh):
    """§7.7 item 4: a recorded question has 2-4 suggestions plus None; an answer by number takes the suggestion's
    text (source export) and a free-text answer is stored verbatim (source user)."""
    r = cli("nine-m-suggest", str(fresh), "--m", "M3", "--step", "relevance", "--record")
    assert r.returncode == 0, r.stderr
    assert "Q01" in r.stdout and ninem.NINE_M_NONE in r.stdout
    s = session_of(fresh)
    assert s["version"] == 1 and s["mode"] == "interview"
    rec = s["records"][0]
    assert (rec["q"], rec["m"], rec["step"], rec["answer"]) == ("Q01", "M3", "relevance", None)
    assert 2 <= len(rec["suggestions"]) - 1 <= 4 and rec["suggestions"][-1]["text"] == ninem.NINE_M_NONE
    assert cli("nine-m-answer", str(fresh), "--q", "Q01", "--choice", "1").returncode == 0
    a = session_of(fresh)["records"][0]["answer"]
    assert a == {"choice": 1, "text": rec["suggestions"][0]["text"], "source": "export"}
    run("nine-m-suggest", str(fresh), "--m", "M3", "--step", "subcategory", "--record")
    assert cli("nine-m-answer", str(fresh), "--q", "Q02", "--text", "M3.02, the written start-up steps").returncode == 0
    a = session_of(fresh)["records"][1]["answer"]
    assert a == {"choice": None, "text": "M3.02, the written start-up steps", "source": "user"}
    assert ns.m_status(session_of(fresh), "M3") == "relevant"
    # without --record nothing is written
    run("nine-m-suggest", str(fresh), "--m", "M8", "--step", "relevance")
    assert len(session_of(fresh)["records"]) == 2


def test_answer_errors(fresh):
    r = cli("nine-m-answer", str(fresh), "--q", "Q01", "--choice", "1")
    assert r.returncode == 2 and "nine-m-suggest" in r.stderr and "Traceback" not in r.stderr
    run("nine-m-suggest", str(fresh), "--m", "M3", "--step", "relevance", "--record")
    for args, msg in ((("--q", "Q09", "--choice", "1"), "Q09"), (("--q", "Q01", "--choice", "9"), "choice"),
                      (("--q", "Q01", "--choice", "0"), "choice")):
        r = cli("nine-m-answer", str(fresh), *args)
        assert r.returncode == 2 and msg in r.stderr and "Traceback" not in r.stderr
    r = cli("nine-m-answer", str(fresh), "--q", "Q01")
    assert r.returncode == 2


def test_suggest_needs_the_seed(tmp_path):
    x = nine_m_export(tmp_path)
    run("ingest", str(x), "--legacy-out", str(tmp_path / "ws"))
    for args in (("nine-m-suggest", str(tmp_path / "ws"), "--m", "M3", "--step", "relevance"),
                 ("nine-m-replay", str(tmp_path / "ws"), "--accept-recommended")):
        r = cli(*args)
        assert r.returncode == 2 and "run nine-m-seed first" in r.stderr and "Traceback" not in r.stderr


def test_suggest_cli_checks(fresh):
    r = cli("nine-m-suggest", str(fresh), "--step", "relevance")
    assert r.returncode == 2 and "--m" in r.stderr
    r = cli("nine-m-suggest", str(fresh), "--m", "M3", "--step", "subcategory", "--subcategory", "M9.01")
    assert r.returncode == 2 and "M9.01" in r.stderr
    r = cli("nine-m-suggest", str(fresh), "--m", "M3", "--step", "condition", "--subcategory", "M3.99")
    assert r.returncode == 2 and "M3.99" in r.stderr
    r = cli("nine-m-suggest", str(fresh), "--step", "scope")
    assert r.returncode == 0 and "Every M with seed weight" in r.stdout
    r = cli("nine-m-suggest", str(fresh), "--m", "M3", "--step", "relevance", "--json")
    assert json.loads(r.stdout)[-1]["text"] == ninem.NINE_M_NONE


def test_question_budget(fresh):
    """6 questions per M plus 4 global: the seventh question on one M is refused."""
    for _ in range(ns.NINE_M_BUDGET_PER_M):
        run("nine-m-suggest", str(fresh), "--m", "M7", "--step", "condition", "--record")
    r = cli("nine-m-suggest", str(fresh), "--m", "M7", "--step", "change", "--record")
    assert r.returncode == 2 and "budget" in r.stderr
    for _ in range(ns.NINE_M_BUDGET_GLOBAL):
        run("nine-m-suggest", str(fresh), "--step", "gaps", "--record")
    r = cli("nine-m-suggest", str(fresh), "--step", "gaps", "--record")
    assert r.returncode == 2 and "budget" in r.stderr


def test_answer_with_date_is_flagged(fresh, capsys):
    """Review Focus 2: an answer naming an owner and a date is stored verbatim with a D7 note; a pack recommendation
    that carries the owner is a problem even without --strict, and the actions template stays blank."""
    run("nine-m-suggest", str(fresh), "--m", "M9", "--step", "change", "--record")
    r = cli("nine-m-answer", str(fresh), "--q", "Q01", "--text", "the shift manager by 30 June")
    assert r.returncode == 0
    assert "owners and due dates are set by the organisation" in r.stderr
    assert session_of(fresh)["records"][0]["answer"]["text"] == "the shift manager by 30 June"
    w = pmlc.load_ws(fresh)
    pack = nine_m_pack()
    pack["sections"][1]["items"][2]["owner"] = "the shift manager"
    pack["sections"][1]["items"][2]["by_when"] = "30 June"
    res = pp.check_pack(pack, w)
    assert [x for x in res.problems if "owner and due dates are set by the organisation" in x]


# ----------------------------------------------------------------------------- replay

def test_replay_accept_recommended(fresh):
    """OD5-1, §7.7 item 5: the replay takes suggestion 1 at every question, marks the session as a replay, keeps
    the budget, and every M in scope ends relevant or not relevant."""
    r = cli("nine-m-replay", str(fresh), "--accept-recommended")
    assert r.returncode == 0, r.stderr
    s = session_of(fresh)
    assert s["mode"] == "replay"
    assert s == json.loads(json.dumps(s, sort_keys=True))
    assert all(rec["answer"]["choice"] == 1 and rec["answer"]["source"] == "export"
               for rec in s["records"] if rec["step"] != "scope")
    scope = ns.session_scope(s)
    assert scope
    for m in scope:
        assert ns.m_status(s, m) in ("relevant", "not-relevant")
        assert sum(1 for rec in s["records"] if rec["m"] == m) <= ns.NINE_M_BUDGET_PER_M
    assert sum(1 for rec in s["records"] if rec["m"] is None) <= ns.NINE_M_BUDGET_GLOBAL
    for rec in s["records"]:
        assert 2 <= len(rec["suggestions"]) - 1 <= 4
    relevant = [m for m in scope if ns.m_status(s, m) == "relevant"]
    for m in relevant:
        steps = [rec["step"] for rec in s["records"] if rec["m"] == m]
        assert steps == list(ns.NINE_M_STEPS)
    first = (fresh / "nine_m_session.json").read_bytes()
    run("nine-m-replay", str(fresh), "--accept-recommended")
    assert (fresh / "nine_m_session.json").read_bytes() == first          # byte-identical re-run


def test_replay_scope_and_refusals(fresh):
    r = cli("nine-m-replay", str(fresh))
    assert r.returncode == 2 and "--accept-recommended" in r.stderr
    r = cli("nine-m-replay", str(fresh), "--accept-recommended", "--ms", "M3,M12")
    assert r.returncode == 2 and "M12" in r.stderr
    run("nine-m-replay", str(fresh), "--accept-recommended", "--ms", "M8,M3")
    s = session_of(fresh)
    assert ns.session_scope(s) == ["M3", "M8"]
    assert {rec["m"] for rec in s["records"]} <= {None, "M3", "M8"}
    scope = next(rec for rec in s["records"] if rec["step"] == "scope")
    assert scope["answer"]["source"] == "user"
    # an interview is never overwritten by a replay
    (fresh / "nine_m_session.json").unlink()
    run("nine-m-suggest", str(fresh), "--m", "M3", "--step", "relevance", "--record")
    r = cli("nine-m-replay", str(fresh), "--accept-recommended")
    assert r.returncode == 2 and "interview" in r.stderr
    r = cli("nine-m-suggest", str(fresh), "--m", "M3", "--step", "relevance", "--record")
    assert r.returncode == 0


def test_record_refused_on_a_replay(fresh):
    run("nine-m-replay", str(fresh), "--accept-recommended")
    r = cli("nine-m-suggest", str(fresh), "--m", "M3", "--step", "relevance", "--record")
    assert r.returncode == 2 and "replay" in r.stderr


def test_replay_thin_export(tmp_path):
    """Review Focus 3: a thin export (no evidence links, no causal sheet) still replays, with every M decided."""
    from test_pmlc import make_thin_export
    x = make_thin_export(tmp_path)
    run("ingest", str(x), "--legacy-out", str(tmp_path / "ws"))
    run("nine-m-seed", str(tmp_path / "ws"))
    run("nine-m-replay", str(tmp_path / "ws"), "--accept-recommended")
    s = session_of(tmp_path / "ws")
    for m in ns.session_scope(s):
        assert ns.m_status(s, m) in ("relevant", "not-relevant")


# ----------------------------------------------------------------------------- select --m

def test_select_m(nws, tmp_path):
    seed = json.loads((nws / "nine_m_seed.json").read_text(encoding="utf-8"))
    prim = sorted(f for f, v in seed["findings"].items() if v["m_primary"] == "M9")
    sec = sorted(f for f, v in seed["findings"].items() if v["m_primary"] == "M9" or "M9" in v["m_secondary"])
    got = sorted(x["id"] for x in json.loads(run("select", str(nws), "--m", "M9", "--format", "json")))
    assert got == prim and prim
    got = sorted(x["id"] for x in json.loads(run("select", str(nws), "--m", "M9", "--m-secondary", "--format", "json")))
    assert got == sec and set(sec) > set(prim)
    md = run("select", str(nws), "--m", "M9")
    assert "m=['M9']" in md
    assert "m=" not in run("select", str(nws), "--level", "4")
    import shutil
    bare = tmp_path / "bare"
    shutil.copytree(nws, bare)
    (bare / "nine_m_seed.json").unlink()
    r = cli("select", str(bare), "--m", "M3")
    assert r.returncode == 2 and "run nine-m-seed first" in r.stderr
    r = cli("select", str(nws), "--m-secondary")
    assert r.returncode == 2 and "--m" in r.stderr
    r = cli("select", str(nws), "--m", "M10")
    assert r.returncode == 2


# ----------------------------------------------------------------------------- profile rules

@pytest.fixture(autouse=True)
def _nm_profile(monkeypatch):
    """Give every test the nine-m-analysis profile (Task 22 adds it to pack-profiles.json)."""
    monkeypatch.setattr(profiles, "load", lambda path=None: (dict(_REAL_LOAD(), **{"nine-m-analysis": NINE_M_PROFILE})
                                                             if path is None else _REAL_LOAD(path)))


def overview(ms=None):
    rows = [[m, "relevant" if m in (ms or ("M3", "M9")) else "not relevant", "1", "-", "1"] for m in ninem.NINE_M_CODES]
    return {"kind": "table", "columns": ["M", "Relevance", "Conditions", "Top subcategories",
                                                                 "Findings"],
            "rows": rows, "row_cites": [["FND-0002"] for _ in rows], "cites": ["FND-0002"]}


def not_relevant(ms):
    rows = [[m, "No finding points to it, and the review found no condition"] for m in ms]
    return {"kind": "table", "id": "not-relevant-ms", "columns": ["M", "Why not relevant"], "rows": rows,
            "row_cites": [[] for _ in rows]}


def nine_m_pack(demo=False, m1_recs=0):
    """A synthetic nine-m-analysis pack: M3 and M9 relevant, the other seven not relevant."""
    others = [m for m in ninem.NINE_M_CODES if m not in ("M3", "M9")]
    m3 = [item("learning", lesson="Return to service had no checklist", cites=["FND-0002"], group="c-m3-1",
               learning="LR-01"),
          item("barrier", barrier="Return-to-service check", failure_mode="absent", hoc_tier="administrative",
               cites=["FND-0002"], group="c-m3-1"),
          item("recommendation", change="Interlock the restart on a signed test record", indicator="Restarts with a test",
               hoc_tier="engineering", cites=["FND-0002"], group="c-m3-1")]
    m9 = [item("learning", lesson="Seals were bought without a specification", cites=["FND-0003"], group="c-m9-1"),
          item("barrier", barrier="Spares specification", failure_mode="absent", hoc_tier="administrative",
               cites=["FND-0003"], group="c-m9-1"),
          item("recommendation", change="Hold seals without a specification at goods receipt", indicator="Seals held",
               hoc_tier="engineering", cites=["FND-0003"], group="c-m9-1")]
    secs = [{"id": "nine-m-overview", "heading": "The nine Ms", "items": [overview()]},
            {"id": "m3", "heading": "M3 Method", "items": m3},
            {"id": "m9", "heading": "M9 Maintenance", "items": m9},
            {"id": "not-relevant", "heading": "Not relevant", "items": [not_relevant(others)]}]
    if m1_recs:
        secs[0]["items"] = [overview(("M1", "M3", "M9"))]
        secs[3]["items"] = [not_relevant([m for m in others if m != "M1"])]
        m1 = [item("learning", lesson="Start-up was done from memory", cites=["FND-0001"], group="c-m1-1")]
        m1 += [item("recommendation", change=f"Change {i}", indicator="i", hoc_tier="engineering", cites=["FND-0001"],
                    group="c-m1-1") for i in range(m1_recs)]
        secs.insert(1, {"id": "m1", "heading": "M1 Manpower/Mindpower", "items": m1})
    if demo:
        secs.append({"id": "organisation-context", "heading": "Context", "items": [
            {"kind": "para", "text": profiles.BANNERS["nine-m-demo"]}]})
    return {"pack": "nine-m-analysis", "schema": 1, "title": "9M analysis: synthetic pump fire", "audience": "Investigators",
            "status": "Learning purposes: not a statutory document",
            "bottom_line": [{"text": "Return to service had no checklist.", "cites": ["FND-0002"]}],
            "sections": secs}


def profile_msgs(pack, w):
    return profiles.check(pack, NINE_M_PROFILE, None, w=w)


def test_overview_rule(nws):
    w = pmlc.load_ws(nws)
    good = nine_m_pack()
    assert not [x for x in profile_msgs(good, w) if "nine-m-overview" in x or x.startswith("M")], profile_msgs(good, w)
    bad = nine_m_pack()
    bad["sections"][0]["items"][0]["rows"].pop()
    assert [x for x in profile_msgs(bad, w) if "nine-m-overview: 8 rows" in x]
    bad = nine_m_pack()
    rows = bad["sections"][0]["items"][0]["rows"]
    rows[0], rows[1] = rows[1], rows[0]
    assert [x for x in profile_msgs(bad, w) if "nine-m-overview: row 1 is M2, expected M1" in x]
    bad = nine_m_pack()
    bad["sections"][3]["items"][0]["rows"] = [r for r in bad["sections"][3]["items"][0]["rows"] if r[0] != "M5"]
    assert [x for x in profile_msgs(bad, w) if x.startswith("M5: ")]
    bad = nine_m_pack()
    bad["sections"][3]["items"][0]["rows"][0][1] = " "
    assert [x for x in profile_msgs(bad, w) if x.startswith("M1: ")]
    bad = nine_m_pack()
    bad["sections"][1]["items"] = bad["sections"][1]["items"][1:]          # m3 without a learning
    assert [x for x in profile_msgs(bad, w) if x.startswith("M3: ")]
    bad = nine_m_pack()
    bad["sections"][3]["items"][0]["rows"].append(["M3", "also not relevant"])
    assert [x for x in profile_msgs(bad, w) if "M3: both" in x]
    bad = nine_m_pack()
    bad["sections"] = bad["sections"][1:]
    assert [x for x in profile_msgs(bad, w) if "nine-m-overview: no" in x]


def test_nine_m_cites_rule(nws):
    w = pmlc.load_ws(nws)
    seed = w["nine_m_seed"]
    assert not [x for x in profile_msgs(nine_m_pack(), w) if "nine-m-cites" in x]
    outside = next(f for f, v in seed["findings"].items() if "M2" not in v["scores"])
    pack = nine_m_pack()
    pack["sections"][1]["items"][0]["cites"] = ["FND-0002", outside]
    pack["sections"][1]["id"] = "m2"
    pack["sections"][0]["items"][0] = overview(("M2", "M9"))
    pack["sections"][3]["items"][0] = not_relevant([m for m in ninem.NINE_M_CODES if m not in ("M2", "M9")])
    msgs = [x for x in profile_msgs(pack, w) if "nine-m-cites" in x]
    assert msgs and outside in msgs[0] and "M2" in msgs[0]
    bare = dict(w)
    bare.pop("nine_m_seed")
    assert [x for x in profile_msgs(nine_m_pack(), bare) if "run nine-m-seed" in x]


def test_m1_balance_rule(nws):
    w = pmlc.load_ws(nws)
    assert not [x for x in profile_msgs(nine_m_pack(m1_recs=1), w) if "m1-balance" in x]
    msgs = [x for x in profile_msgs(nine_m_pack(m1_recs=3), w) if "m1-balance" in x]
    assert msgs and "M1 holds the most recommendations (3)" in msgs[0]
    # with no other M ahead on level 3-4 findings the rule is quiet
    w2 = copy.deepcopy(w)
    for m, p in w2["nine_m_seed"]["profile"].items():
        p["by_level"] = {"1": 0, "2": 0, "3": 9 if m == "M1" else 0, "4": 0}
    assert not [x for x in profile_msgs(nine_m_pack(m1_recs=3), w2) if "m1-balance" in x]


def test_demo_banner_only_on_replay(fresh):
    """OD5-1: banner:nine-m-demo is required only when the session is a replay."""
    w = pmlc.load_ws(fresh)
    assert not [x for x in profile_msgs(nine_m_pack(), w) if "banner nine-m-demo" in x]       # no session
    run("nine-m-suggest", str(fresh), "--m", "M3", "--step", "relevance", "--record")
    w = pmlc.load_ws(fresh)
    assert not [x for x in profile_msgs(nine_m_pack(), w) if "banner nine-m-demo" in x]       # an interview
    (fresh / "nine_m_session.json").unlink()
    run("nine-m-replay", str(fresh), "--accept-recommended")
    w = pmlc.load_ws(fresh)
    assert [x for x in profile_msgs(nine_m_pack(), w) if "banner nine-m-demo" in x]           # a replay
    assert not [x for x in profile_msgs(nine_m_pack(demo=True), w) if "banner nine-m-demo" in x]
    # the other banners keep their unconditional rule
    prof = dict(NINE_M_PROFILE, rules=["banner:crosswalk"])
    assert [x for x in profiles.check(nine_m_pack(), prof, None, w=w) if "banner crosswalk" in x]


def test_nine_m_pack_checks_strict(fresh):
    run("nine-m-replay", str(fresh), "--accept-recommended")
    w = pmlc.load_ws(fresh)
    res = pp.check_pack(nine_m_pack(demo=True), w, strict=True)
    assert not [x for x in res.problems if "nine-m" in x or "m1-balance" in x or "banner" in x], res.problems


# ----------------------------------------------------------------------------- side output

def test_nine_m_map_side_output(fresh):
    run("nine-m-replay", str(fresh), "--accept-recommended")
    out = fresh.parent / "outputs"
    out.mkdir(exist_ok=True)
    p = out / "nine-m-analysis.json"
    p.write_text(json.dumps(nine_m_pack(demo=True)), encoding="utf-8")
    r = cli("pack", str(fresh), str(p))
    assert r.returncode == 0, r.stderr
    mp = json.loads((out / "nine-m-map.json").read_text(encoding="utf-8"))
    assert set(mp) == {"findings", "conditions"}
    assert mp["conditions"]["c-m3-1"]["m"] == "M3" and mp["conditions"]["c-m9-1"]["m"] == "M9"
    assert mp["conditions"]["c-m3-1"]["learning_id"] == "LR-01"
    assert mp["conditions"]["c-m9-1"]["learning_id"]
    assert mp["conditions"]["c-m3-1"]["recommendation_ids"] == ["m3-3"]
    sub = mp["conditions"]["c-m3-1"]["subcategory"]
    assert sub is None or re.fullmatch(r"M3\.\d\d", sub)
    s = session_of(fresh)
    want = next(re.match(r"(M3\.\d\d)", rec["answer"]["text"]) for rec in s["records"]
                if rec["m"] == "M3" and rec["step"] == "subcategory")
    assert sub == (want.group(1) if want else None)
    f2 = mp["findings"]["FND-0002"]
    assert (f2["m_primary"], f2["condition_id"], f2["subcategory"]) == ("M3", "c-m3-1", sub)
    assert f2["m_secondary"] == []
    assert mp["findings"]["FND-0003"]["m_primary"] == "M9"
    first = (out / "nine-m-map.json").read_bytes()
    assert cli("pack", str(fresh), str(p)).returncode == 0
    assert (out / "nine-m-map.json").read_bytes() == first
    # another pack writes no map
    (out / "nine-m-map.json").unlink()
    other = dict(nine_m_pack(demo=True), pack="test-pack")
    q = out / "test-pack.json"
    q.write_text(json.dumps(other), encoding="utf-8")
    assert cli("pack", str(fresh), str(q)).returncode == 0
    assert not (out / "nine-m-map.json").exists()


def test_nine_m_map_feeds_rec_review(fresh):
    """The side output is what rec-review --nine-m reads (Task 6): the spread counts the M3 and M9 recommendations."""
    from pmlc_lib.rec_review import review
    pack = pp.with_item_ids(nine_m_pack())
    mp = ns.nine_m_map(pack, {"records": []})
    reg = {"entries": [{"lr_id": "LR-01", "recommendations": [{"change": "x", "hoc_tier": "engineering"}]}]}
    spread = review(reg, mp)["spread"]
    assert spread["M3"] == 1 and spread["M9"] == 1 and spread["M1"] == 0


# ----------------------------------------------------------------------------- actions template

def write_pack(ws, pack, name="nine-m-analysis.json"):
    out = ws.parent / "outputs"
    out.mkdir(exist_ok=True)
    p = out / name
    p.write_text(json.dumps(pack), encoding="utf-8")
    return p


def test_actions_template_owner_due_blank(fresh):
    p = write_pack(fresh, nine_m_pack())
    r = cli("nine-m-actions-template", str(fresh), str(p), "--format", "csv")
    assert r.returncode == 0, r.stderr
    f = fresh.parent / "outputs" / "nine-m-analysis.actions-template.csv"
    rows = list(csv.reader(io.StringIO(f.read_bytes().decode("utf-8"))))
    assert rows[0] == ["Recommendation ID", "M", "Change", "HoC tier", "Corrective/Preventive", "Indicator", "Owner",
                       "Due date", "Status", "Notes"]
    assert [r[:2] for r in rows[1:]] == [["m3-3", "M3"], ["m9-3", "M9"]]
    for row in rows[1:]:
        assert row[6] == row[7] == row[8] == ""
    assert b"\r\n" in f.read_bytes()
    r = cli("nine-m-actions-template", str(fresh), str(p))
    assert r.returncode == 0, r.stderr
    wb = openpyxl.load_workbook(fresh.parent / "outputs" / "nine-m-analysis.actions-template.xlsx")
    sh = wb.active
    head = [c.value for c in sh[1]]
    assert head[6:9] == ["Owner", "Due date", "Status"]
    for row in sh.iter_rows(min_row=2, values_only=True):
        assert row[0] and row[6] is None and row[7] is None and row[8] is None


def test_actions_template_no_overwrite(fresh):
    p = write_pack(fresh, nine_m_pack())
    run("nine-m-actions-template", str(fresh), str(p), "--format", "csv")
    f = fresh.parent / "outputs" / "nine-m-analysis.actions-template.csv"
    f.write_text("organisation's own entries", encoding="utf-8")
    r = cli("nine-m-actions-template", str(fresh), str(p), "--format", "csv")
    assert r.returncode == 1 and "already exists" in r.stderr
    assert f.read_text(encoding="utf-8") == "organisation's own entries"
    # owners in the pack are refused before anything is written (D7)
    bad = nine_m_pack()
    bad["sections"][1]["items"][2]["owner"] = "shift manager"
    q = write_pack(fresh, bad, "bad.json")
    r = cli("nine-m-actions-template", str(fresh), str(q))
    assert r.returncode != 0 and not (fresh.parent / "outputs" / "bad.actions-template.xlsx").exists()
    none = write_pack(fresh, dict(nine_m_pack(), sections=nine_m_pack()["sections"][:1]), "none.json")
    r = cli("nine-m-actions-template", str(fresh), str(none))
    assert r.returncode == 1 and "no recommendation" in r.stderr


def test_actions_template_shared_copy(tmp_path):
    """Review Focus 1: on a shared pack the template is masked, named .external. and checked by the name guard."""
    ws, _ = share_ws(tmp_path)
    run("classify", str(ws))
    run("nine-m-seed", str(ws))
    from test_privacy import SINGLE
    names = tmp_path / "names.txt"
    names.write_text("Quillon Vantress\n", encoding="utf-8")
    pack = nine_m_pack()
    pack["pack"] = "test-pack"
    rec = pack["sections"][1]["items"][2]
    rec["change"] = "Interlock the restart checked at 14:35 on 2024-03-20 by the " + SINGLE
    rec["lint_ok"] = [{"rule": "JC7", "reason": "test of masking a singleton role"}]
    p = write_doc(tmp_path, pack, "nm.json")
    sign(ws, p, "--classification", "external", "--legal-review", "cleared")
    r = cli("nine-m-actions-template", str(ws), str(p), "--share", "external", "--names", str(names))
    assert r.returncode == 0, r.stderr
    out = ws.parent / "outputs"
    shared = out / "nm.actions-template.external.xlsx"
    assert shared.is_file() and not (out / "nm.actions-template.xlsx").exists()
    cells = " ".join(str(c.value) for row in openpyxl.load_workbook(shared).active.iter_rows() for c in row)
    assert "14:35" not in cells and "2024-03-20" not in cells and SINGLE not in cells
    bad = copy.deepcopy(pack)
    bad["sections"][1]["items"][2]["change"] = "Ask Quillon Vantress to interlock the restart"
    bad["sections"][1]["items"][2]["lint_ok"] = [{"rule": "JC1", "reason": "test of the name guard"}]
    q = write_doc(tmp_path, bad, "nm-bad.json")
    r = cli("nine-m-actions-template", str(ws), str(q), "--names", str(names), "--format", "csv")
    assert r.returncode == 1 and "refused: a listed name appears in the output" in r.stderr
    assert not (out / "nm-bad.actions-template.csv").exists()


# ----------------------------------------------------------------------------- register and dashboard

def test_register_nine_m_column(nws, tmp_path):
    """§7.7 item 8: check-register accepts M1-M9 and the nine-m.json subcategory ids; register-xlsx shows the 9M
    columns and the 9M x HFACS level group; register-md shows them."""
    w = pmlc.load_ws(nws)
    base = {"lr_id": "LR-01", "title": "Return to service without a test", "finding_ids": ["FND-0002"],
            "evidence_ids": ["EVD-0002"], "priority": 1, "owner_entity": "F02", "hfacs_level": "L3 Supervision",
            "tier_messages": {"T1": "Do not pressurise without the test record."},
            "recommendations": [{"what": "Interlock", "hoc_tier": "engineering", "leading_indicator": "i"}]}
    ok = dict(base, nine_m="M3", nine_m_subcategory="M3.02")
    assert not [x for x in check_register({"entries": [ok]}, w).problems if "nine_m" in x]
    for e, msg in ((dict(base, nine_m="M3", nine_m_subcategory="M3.42"), "not a subcategory id"),
                   (dict(base, nine_m="M3", nine_m_subcategory="M9.02"), "belongs to M9"),
                   (dict(base, nine_m="M10"), "nine_m 'M10'")):
        assert [x for x in check_register({"entries": [e]}, w).problems if msg in x], msg
    free = check_register({"entries": [dict(base, nine_m="M7", nine_m_subcategory="Machine: alarm design")]}, w)
    assert not [x for x in free.problems if "nine_m" in x]
    assert [x for x in free.warnings if "nine_m_subcategory" in x]
    reg = tmp_path / "r.json"
    reg.write_text(json.dumps({"entries": [ok, dict(base, lr_id="LR-02", nine_m="M9", hfacs_level="L4 Resource")]}))
    out = tmp_path / "r.xlsx"
    run("register-xlsx", str(reg), "--out", str(out))
    wb = openpyxl.load_workbook(out)
    head = [c.value for c in wb["Learnings Register"][1]]
    assert "9M" in head and "9M subcategory" in head
    col = head.index("9M")
    assert [r[col] for r in wb["Learnings Register"].iter_rows(min_row=2, values_only=True)] == ["M3", "M9"]
    summary = [tuple(r) for r in wb["Summary"].iter_rows(values_only=True)]
    assert ("Learnings by 9M and HFACS level", "M3 × L3 Supervision", 1) in summary
    assert ("Learnings by 9M and HFACS level", "M9 × L4 Organisational Influences", 1) in summary
    md = run("register-md", str(nws), str(reg), "--out", str(tmp_path / "r.md"))
    text = (tmp_path / "r.md").read_text(encoding="utf-8")
    assert md and "9M: M3" in text and "9M subcategory: M3.02" in text and "Learnings by 9M and HFACS level" in text
    plain = tmp_path / "p.json"
    plain.write_text(json.dumps({"entries": [base]}))
    run("register-xlsx", str(plain), "--out", str(tmp_path / "p.xlsx"))
    wb = openpyxl.load_workbook(tmp_path / "p.xlsx")
    assert "9M" not in [c.value for c in wb["Learnings Register"][1]]
    assert not [r for r in wb["Summary"].iter_rows(values_only=True) if r[0] == "Learnings by 9M and HFACS level"]


def test_dashboard_nine_m_strip(nws, tmp_path):
    """P5-18: with any nine_m value the dashboard shows nine cells, M1-M9, each counting learnings and linking to
    them; without one there is no strip."""
    base = {"title": "t", "finding_ids": ["FND-0001"], "priority": 1, "owner_entity": "F02",
            "hfacs_level": "L3 Supervision", "tier_messages": {"T1": "m"},
            "recommendations": [{"what": "x", "hoc_tier": "engineering"}]}
    reg = tmp_path / "r.json"
    reg.write_text(json.dumps({"entries": [dict(base, lr_id="LR-01", nine_m="M3"), dict(base, lr_id="LR-02", nine_m="M3"),
                                           dict(base, lr_id="LR-03", nine_m="M8"), dict(base, lr_id="LR-04")]}))
    html = tmp_path / "d.html"
    run("dashboard", str(nws), "--out", str(html), "--register", str(reg))
    t = html.read_text(encoding="utf-8")
    strip = t[t.index('id="nine-m-strip"'):]
    strip = strip[:strip.index("</table>")]
    cells = re.findall(r"<td[^>]*>(.*?)</td>", strip)
    assert len(cells) == 9
    assert re.search(r"<b>2</b>", cells[2]) and 'href="#detail/LR-01"' in cells[2] and 'href="#detail/LR-02"' in cells[2]
    assert "<b>1</b>" in cells[7] and 'href="#detail/LR-03"' in cells[7]
    assert "LR-04" not in strip
    for i, m in enumerate(ninem.NINE_M_CODES):
        assert m in re.findall(r"<th[^>]*>(.*?)</th>", strip)[i]
    plain = tmp_path / "p.json"
    plain.write_text(json.dumps({"entries": [dict(base, lr_id="LR-01")]}))
    run("dashboard", str(nws), "--out", str(tmp_path / "p.html"), "--register", str(plain))
    assert "nine-m-strip" not in (tmp_path / "p.html").read_text(encoding="utf-8")


# ----------------------------------------------------------------------------- Hazira dry run (§7.7 item 11)

def test_m8_m9_reachable_by_relevance(tmp_path_factory):
    """On a Hazira replay, Material and Maintenance reach a relevant condition through question 3a, although their
    codes seed only 9 and 27 findings."""
    from test_calibration_p5 import EXPORT
    root = tmp_path_factory.mktemp("hazira-9m")
    run("setup", str(EXPORT), "--slug", "hazira-gas-release-fire", "--date", "2020-09-24", "--root", str(root),
        "--overlay", "india-process-site")
    ws = root / "pmlc-runs" / "2020-09-24-hazira-gas-release-fire" / "workspace"
    run("nine-m-seed", str(ws))
    run("nine-m-replay", str(ws), "--accept-recommended")
    s = session_of(ws)
    w = pmlc.load_ws(ws)
    ids = {f["id"] for f in w["findings"]}
    for m in ("M8", "M9"):
        assert ns.m_status(s, m) == "relevant"
        cond = next(rec for rec in s["records"] if rec["m"] == m and rec["step"] == "condition")
        assert cond["answer"]["choice"] == 1 and cond["suggestions"][0]["cites"]
        assert set(cond["suggestions"][0]["cites"]) <= ids
    for rec in s["records"]:
        for sg in rec["suggestions"]:
            assert all(x in ids for x in sg["cites"] if x.startswith("FND-"))
            if not ns.is_procedural(sg["text"]):
                assert [x for x in sg["cites"] if FND_RX.match(x)], (rec["q"], sg)
    assert sum(1 for rec in s["records"]) <= 9 * ns.NINE_M_BUDGET_PER_M + ns.NINE_M_BUDGET_GLOBAL
