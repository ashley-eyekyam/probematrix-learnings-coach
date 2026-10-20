"""setup, check-all, render-all, index and dossier-index on the synthetic export (fictional event)."""
import json
import shutil
from pathlib import Path

import pytest

from test_pmlc import make_export, run, run_fail

FIX = Path(__file__).resolve().parent / "fixtures" / "packs"


@pytest.fixture
def setup_ws(tmp_path):
    x = make_export(tmp_path)
    out = run("setup", str(x), "--slug", "pump-fire", "--root", str(tmp_path))
    ws = tmp_path / "pmlc-runs" / "2024-03-20-pump-fire" / "workspace"
    return ws, out


def put(ws, *names):
    outd = ws.parent / "outputs"
    outd.mkdir(exist_ok=True)
    for n in names:
        shutil.copy(FIX / f"{n}.json", outd / f"{n}.json")
    return outd


def test_setup_runs_all_steps(setup_ws):
    ws, out = setup_ws
    for n in ("findings.json", "clusters.json", "causal_graph.json", "timeline_candidates.json", "interfaces.json",
              "profile.md"):
        assert (ws / n).is_file(), n
    assert "Next: PMLC plan" in out and str(ws) in out
    assert "4 findings" in out and "Mr Fictional Person" in out


def test_setup_stops_on_failure(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.title = "Plan Summary"
    wb.create_sheet("Action Items")
    wb.create_sheet("Risk Matrix")
    wb.save(tmp_path / "capa.xlsx")
    r = run_fail("setup", str(tmp_path / "capa.xlsx"), "--slug", "x-y", "--date", "2024-03-20")
    assert r.returncode == 1 and "setup stopped at ingest" in r.stderr


def test_check_all_exit_codes(setup_ws):
    ws, _ = setup_ws
    outd = put(ws, "board-brief", "supervisor-coach")
    (outd / "gaps.json").write_text("{}")
    (outd / "notes.jsonl").write_text("x\n")
    out = run("check-all", str(ws))
    assert "board-brief.json: ok" in out and "supervisor-coach.json: ok" in out and "gaps" not in out
    put(ws, "bad-l3-01")
    r = run_fail("check-all", str(ws))
    assert r.returncode == 1 and "bad-l3-01.json: " in r.stdout and "problem(s)" in r.stdout
    j = run_fail("check-all", str(ws), "--json")
    assert json.loads(j.stdout)["ok"] is False


def test_render_all_skips_unchecked(setup_ws):
    ws, _ = setup_ws
    outd = put(ws, "board-brief", "supervisor-coach")
    run("check-pack", str(ws), str(outd / "board-brief.json"))
    out = run("render-all", str(ws))
    assert (outd / "board-brief.html").is_file()
    assert not (outd / "supervisor-coach.html").exists()
    assert "skipped supervisor-coach.json: run check-all first" in out


def test_index_lists_packs_with_status_and_footer(setup_ws):
    ws, _ = setup_ws
    outd = put(ws, "board-brief", "supervisor-coach")
    run("check-all", str(ws))
    run("render-all", str(ws))
    run("index", str(ws))
    html = (outd / "index.html").read_text(encoding="utf-8")
    for t in ("Board Brief: Synthetic pump fire, 2024", "Supervisor Coach: Synthetic pump fire, 2024"):
        assert t in html
    assert "DRAFT" in html and "Learning purposes: not a statutory document" in html
    assert "ProbeMatrix Learnings Coach v" in html and "<script" not in html
    assert 'href="board-brief.html"' in html


def test_dossier_index_is_valid_pack(setup_ws):
    ws, _ = setup_ws
    outd = put(ws, "dossier-F02")
    out = run("dossier-index", str(ws))
    p = outd / "dossier-index.json"
    assert p.is_file() and "dossier-index.json" in out
    d = json.loads(p.read_text())
    assert len(d["sections"]) == 1 and d["sections"][0]["items"][0]["kind"] == "table"
    run("check-pack", str(ws), str(p))
    run("check-all", str(ws))


def test_dossier_index_none(setup_ws):
    ws, _ = setup_ws
    put(ws, "board-brief")
    assert "no dossier" in run_fail("dossier-index", str(ws)).stderr
