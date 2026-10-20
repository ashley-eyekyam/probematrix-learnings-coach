"""The privacy core (pmlc_privacy): names, variants, singleton roles, on the synthetic export."""
import json
import re

import openpyxl
import pytest

from test_pack import item, pack_with
from test_pmlc import make_export, run
import pmlc  # noqa: E402
import pmlc_lib.privacy  # noqa: E402
import pmlc_lint as pl  # noqa: E402
import pmlc_privacy as pv  # noqa: E402

ROLES = "Security Guard - Gate 3, Shift A Engineer (CR 2), Operator"


@pytest.fixture(scope="module")
def sws(tmp_path_factory):
    d = tmp_path_factory.mktemp("priv")
    x = make_export(d)
    wb = openpyxl.load_workbook(x)
    ws = wb["Identified Findings"]
    col = [c.value for c in ws[1]].index("Individuals Involved") + 1
    ws.cell(row=2, column=col).value = ROLES
    wb.save(x)
    run("ingest", str(x), "--legacy-out", str(d / "ws"))
    run("classify", str(d / "ws"))
    return d / "ws"


def test_singletons_written_by_classify(sws):
    p = sws / "singleton_roles.json"
    doc = json.loads(p.read_text())
    assert doc["version"] == 1
    by = {e["role"]: e for e in doc["roles"]}
    assert by["Security Guard - Gate 3"]["reason"] == "qualifier"
    assert by["Shift A Engineer (CR 2)"]["reason"] == "qualifier"
    assert "Operator" not in by
    for e in doc["roles"]:
        assert set(e) == {"role", "reason", "findings", "generalised"}
        assert e["findings"] >= 1 and re.fullmatch(r".+ \(.+\)", e["generalised"])
    assert pv.load_singletons(pmlc.load_ws(sws))["Shift A Engineer (CR 2)"] == by["Shift A Engineer (CR 2)"]["generalised"]
    p.unlink()
    assert pmlc.load_ws(sws)["singletons"] == []
    run("classify", str(sws))
    assert p.exists()


def _qual(role):
    tax = {"entities": [{"id": "F01", "name": "Ops"}]}
    out = pv.build_singletons({role: {"F": ["F01"], "T": "T1", "R": "actor"}}, [{"roles": [role]}], tax, {"T1": "Frontline"})
    return [e["role"] for e in out["roles"]]


@pytest.mark.parametrize("role", ["Control Room Operator", "Control Room Operators", "Shift Supervisor",
                                  "Maintenance Planner (owner)", "Field Operator - Operations",
                                  "Gas Terminal Control Room Operators (Gas Terminal CR 2)"])
def test_crew_and_bare_qualifiers_are_not_singletons(role):
    assert _qual(role) == []


@pytest.mark.parametrize("role", ["Security Guard - Gate 3", "Panel Operator CR-2", "Shift A Supervisor",
                                  "Clerk #4", "Welder (Bay 7)", "Desk B Clerk", "Control Room 2 Operator"])
def test_instance_ids_are_singletons(role):
    assert _qual(role) == [role]


def test_single_holder_role():
    tax = {"entities": [{"id": "F13", "name": "Plant & Asset Leadership"}]}
    rm = {"Plant Head": {"F": ["F13"], "T": "T4", "R": "owner"},
          "Plant Manager": {"F": ["F13"], "T": "T3", "R": "owner"},
          "Site Director": {"F": ["F13"], "T": "T3", "R": "owner"}}
    fs = [{"roles": ["Plant Head", "Plant Manager"]}, {"roles": ["Plant Head", "Site Director"]}]
    out = pv.build_singletons(rm, fs, tax, {"T3": "Function managers", "T4": "Site leadership"})
    assert out["roles"] == [{"role": "Plant Head", "reason": "single-holder", "findings": 2,
                             "generalised": "Plant & Asset Leadership (Site leadership)"}]


def test_jc7_uses_singletons(sws):
    w = pmlc.load_ws(sws)
    p = pack_with(item("para", text="The Security Guard - Gate 3 closed the valve.", cites=["FND-0001"]))
    internal = [h for h in pl.lint(p, w, kind="pack", share="internal") if h.rule == "JC7"]
    external = [h for h in pl.lint(p, w, kind="pack", share="external") if h.rule == "JC7"]
    assert [h.severity for h in internal] == ["warning"]
    assert [h.severity for h in external] == ["error"]
    ordinary = pack_with(item("para", text="The Control Room Operator closed the valve.", cites=["FND-0001"]))
    assert not [h for h in pl.lint(ordinary, w, kind="pack", share="external") if h.rule == "JC7"]
    clean = pack_with(item("para", text="The Operator closed the valve.", cites=["FND-0001"]))
    assert not [h for h in pl.lint(clean, w, kind="pack", share="external") if h.rule == "JC7"]


def test_name_variants():
    v = pv.name_variants("Shri R. K. Sharma")
    assert v[0] == "Shri R. K. Sharma"
    for x in ("Sharma", "R. Sharma", "R. K. Sharma", "Mr Sharma", "Smt. Sharma", "Dr Sharma"):
        assert x in v
    v = pv.name_variants("Mr J Smith")
    for x in ("Smith", "J. Smith", "J Smith", "Mrs. Smith", "Shri Smith"):
        assert x in v
    assert pv.name_variants("John Smith")[2] == "J. Smith"
    assert pv.name_variants("ID #4471") == ["ID #4471"]
    assert "Dr. Anand" in pv.name_variants("Anand")


def test_name_variants_skip_role_words(tmp_path):
    """Final review M1: a surname that is a role noun or a vocabulary word is no variant on its own, so the role
    "Stores/Procurement Officer" survives a names file listing "Fictional Officer"."""
    assert "Officer" in pv.name_variants("Fictional Officer")              # --redact keeps it, and warns (X17)
    v = pv.name_variants("Fictional Officer", skip_role_words=True)
    assert "Officer" not in v and "Fictional Officer" in v and "F. Officer" in v and "Mr Officer" in v
    ctx = pv.ShareCtx("external", ["Fictional Officer"], [], {})
    assert pv.share_text("The Stores/Procurement Officer told Fictional Officer.", ctx) == \
        "The Stores/Procurement Officer told [name]."
    names = tmp_path / "names.txt"
    names.write_text("Fictional Officer\n", encoding="utf-8")
    doc = tmp_path / "doc.txt"
    doc.write_text("The Stores/Procurement Officer signed.", encoding="utf-8")
    r = _cli("scan-names", str(names), str(doc), "--variants")
    assert r.returncode == 0, r.stdout


def test_name_pattern_boundaries():
    assert pv.name_pattern([]) is None
    rx = pv.name_pattern(["Shri R. K. Sharma", "Mr J Smith"], variants=True)
    for text in ("Shri R. K. Sharma left", "report by R. Sharma.", "Mr  Smith said", "smt. sharma", "(J Smith)"):
        assert rx.search(text), text
    for text in ("Sharmaji left", "the Smithson seal", "unSharma", "Smith2"):
        assert not rx.search(text), text
    exact = pv.name_pattern(["R. K. Sharma"])
    assert exact.search("by R. K. Sharma") and not exact.search("by Sharma")


def test_names_file_comments_and_ids(tmp_path, monkeypatch):
    f = tmp_path / "names.txt"
    f.write_text("  # comment\n\nR. K. Sharma\nID #4471\n", encoding="utf-8")
    assert pv.load_names(str(f)) == ["R. K. Sharma", "ID #4471"]
    rx = pv.name_pattern(pv.load_names(str(f)))
    assert rx.search("Prepared by R. K. Sharma") and rx.search("badge ID #4471 scanned")
    monkeypatch.setenv("PMLC_NAMES", str(f))
    assert len(pv.load_names()) == 2
    monkeypatch.setenv("PMLC_ORGS", str(f))
    assert len(pv.load_orgs()) == 2
    monkeypatch.delenv("PMLC_NAMES")
    assert pv.load_names() == []


def test_patterns_and_suffix():
    assert re.search(pv.PATTERNS["email"], "a.b@x.org") and re.search(pv.PATTERNS["phone"], "+91 98765 43210")
    assert re.search(pv.PATTERNS["staff_no"], "Staff No. 4471") and pv.ORG_SUFFIX_RX.search("Acme Pumps Ltd")


# ----------------------------------------------------------------------------- ruling X21: tag and phone shapes

@pytest.mark.parametrize("tag", ["P-101", "AO-101", "PSV-2301", "PSV12A", "F123"])
def test_tag_rx_matches_equipment_tags(tag):
    assert pv.TAG_RX.fullmatch(tag)


@pytest.mark.parametrize("ident", ["F02", "F14", "F99", "F09b", "T4", "T9", "FND-0001", "EVD-0001", "LR-01", "ML-02"])
def test_tag_rx_skips_entity_tier_and_record_ids(ident):
    assert not pv.TAG_RX.search(ident)


def test_entity_ids_neither_masked_nor_flagged(tmp_path):
    ctx = pv.ShareCtx("external", [], [], {})
    assert pv.share_text("F02 and F14 own it; T4 and F99; P-101, AO-101 and PSV-2301 leaked.", ctx) == \
        "F02 and F14 own it; T4 and F99; [equipment tag], [equipment tag] and [equipment tag] leaked."
    doc = {"owner_entity": "F02 Mechanical Maintenance AO-101", "entities": ["F14"], "text": "AO-101"}
    assert pv.share_doc(doc, ctx) == {"owner_entity": "F02 Mechanical Maintenance AO-101", "entities": ["F14"],
                                      "text": "[equipment tag]"}
    reg = {"entries": [{"lr_id": "LR-01", "owner_entity": "F02 Mechanical Maintenance",
                        "lesson": "F02 owns the change; the P-101 seal leaked."}]}
    hits = [h for h in pl.lint(reg, {"findings": []}, kind="register", share="external") if h.rule == "JC10"]
    assert [h.match for h in hits] == ["P-101"]


@pytest.mark.parametrize("ident", ["FND-AE629313", "FND-3AB996FC", "EVD-AB1234", "LR-AB12", "ML-QX123"])
def test_tag_rx_skips_hex_record_ids(ident):
    """Final review I3: the ID guard holds after the hyphen too, so a hex ID's tail is not a tag."""
    assert not pv.TAG_RX.search(ident), ident


def test_hex_ids_survive_share_text_and_lint():
    ctx = pv.ShareCtx("external", [], [], {})
    text = "FND-AE629313 and EVD-AB1234 show pump P-101 leaked."
    assert pv.share_text(text, ctx) == "FND-AE629313 and EVD-AB1234 show pump [equipment tag] leaked."
    reg = {"entries": [{"lr_id": "LR-01", "lesson": text}]}
    hits = [h for h in pl.lint(reg, {"findings": []}, kind="register", share="external") if h.rule == "JC10"]
    assert [h.match for h in hits] == ["P-101"]


@pytest.mark.parametrize("phone", ["+91 98765 43210", "022-2345-6789", "98765 43210", "+44 20 7946 0958"])
def test_phone_pattern_matches_phones(phone):
    assert re.search(pv.PATTERNS["phone"], phone).group(0) == phone


@pytest.mark.parametrize("text", ["2020-09-24", "on 2020-09-24 and 2020-10-01", "24-09-2020", "24/09/2020",
                                  "2020-09-24T14:05:00", "from 2020-09-24 2020-09-25"])
def test_phone_pattern_skips_dates(text):
    assert not re.search(pv.PATTERNS["phone"], text), text


# ----------------------------------------------------------------------------- Task 8: --share and the render guard

import copy  # noqa: E402

from test_pmlc import MINI_REG, PMLC  # noqa: E402

SLUG = "2024-03-20-acme-pump-fire"
ORGS = ("Bharat Seals", "ACME")
SINGLE = "Security Guard - Gate 3"
SEED = {"Short Description": "Pump P-101 restarted at 14:05 hrs on 2024-03-12 after the Acme Pumps Ltd seal change.",
        "AI Justification": "Work order of 12 March 2024 shows no test at 14:05:30 by Bharat Seals [EVD-0001].",
        "Causal Explanation": "Bharat Seals supplied the seal; enabled by SV 1.4.",
        "Individuals Involved": f"{SINGLE}, Pump Technician (seal change crew)"}
SPACK = {"pack": "test-pack", "schema": 1, "title": "Test pack: synthetic pump fire", "audience": "Testers",
         "status": "Learning purposes: not a statutory document",
         "bottom_line": [{"text": "Test before service.", "cites": ["FND-0002"]}],
         "sections": [{"id": "s1", "heading": "Section one", "items": [
             {"kind": "para", "text": "Test before service.", "cites": ["FND-0001"]},
             {"kind": "recommendation", "change": "Interlock the restart on a test record", "indicator": "Restarts tested",
              "hoc_tier": "engineering", "owner": "Plant Manager", "cites": ["FND-0001"]}]}]}
SHARED = "Shared {}: identifying details are masked"
TIME_RX = re.compile(r"\b\d{1,2}:\d{2}\b")


def _cli(*args, env=None):
    import os
    import subprocess
    import sys
    return subprocess.run([sys.executable, str(PMLC), *args], capture_output=True, text=True,
                          env=None if env is None else {**os.environ, **env})


def share_ws(tmp_path, status="completed", edited=None):
    """An approvable synthetic run (completed, one finding edited) whose FND-0001 carries an org, a tag, a time, full
    dates and a singleton role; and an orgs file. Returns (ws, orgs file)."""
    d = tmp_path / "exp"
    d.mkdir()
    x = make_export(d, status=status, edited={"FND-0001": "Yes"} if edited is None else edited,
                    summary=[("Location", "Bharat Seals yard, unit AO-101")])
    wb = openpyxl.load_workbook(x)
    ws = wb["Identified Findings"]
    head = [c.value for c in ws[1]]
    for k, v in SEED.items():
        ws.cell(row=2, column=head.index(k) + 1).value = v
    wb.save(x)
    run("ingest", str(x), "--slug", "acme-pump-fire", "--root", str(tmp_path))
    w = tmp_path / "pmlc-runs" / SLUG / "workspace"
    run("classify", str(w))
    orgs = tmp_path / "orgs.txt"
    orgs.write_text("\n".join(ORGS) + "\n", encoding="utf-8")
    return w, orgs


def write_doc(tmp_path, doc, name="pack.json"):
    p = tmp_path / name
    p.write_text(json.dumps(doc), encoding="utf-8")
    return p


def sign(ws, p, *extra):
    r = _cli("sign-off", str(ws), str(p), "--status", "approved", "--role", "HSE manager", "--date", "2026-10-19", *extra)
    assert r.returncode == 0, r.stderr
    return p


def files(d):
    return sorted(x.name for x in d.iterdir()) if d.exists() else []


@pytest.fixture()
def shws(tmp_path):
    return share_ws(tmp_path)


def test_share_text_masks_in_order():
    ctx = pv.ShareCtx("external", ["R. K. Sharma"], ["Bharat Seals"], {SINGLE: "Security (Frontline)"})
    t = pv.share_text("R. K. Sharma and Mr Fictional Person of Bharat Seals and Acme Pumps Ltd checked P-101 "
                      f"and PSV12A; the {SINGLE} saw it at 14:05 hrs, 14:05:30, 9:15 and 1405 hrs. FND-0001 EVD-0001 F02 "
                      "LR-01 T4.", ctx)
    assert t == ("[name] and [name] of [organisation] and [organisation] checked [equipment tag] and [equipment tag]; "
                 "the Security (Frontline) saw it,, and. FND-0001 EVD-0001 F02 LR-01 T4.")
    for raw, want in [("a 1:10 slope and 3:2 odds", "a 1:10 slope and 3:2 odds"), ("at 9:15 am the alarm", "the alarm"),
                      ("at 9:15 the alarm", "the alarm"), ("alarm 9:15:30 rang", "alarm rang"),
                      ("the 7:45 shift", "the shift"), ("ratio 1:10 mix", "ratio 1:10 mix"), ("a 1:2 slope", "a 1:2 slope"),
                      ("a 1:10 dilution and odds of 3:15", "a 1:10 dilution and odds of 3:15"), ("a ratio of 1:10", "a ratio of 1:10"), ("at 03:12", ""),
                      ("1405 hrs", ""),
                      ("1405 hours", ""), ("1405 tonnes", "1405 tonnes"), ("Mr R. K. Sharma", "[name]"),
                      ("Shri Sharma left", "[name] left")]:
        assert pv.share_text(raw, ctx) == want, raw
    for raw, want in [("on 2024-03-12.", "on March 2024."), ("at 2024-03-12T14:05:00Z", "at March 2024"),
                      ("12 March 2024", "March 2024"), ("3rd Sept 2019", "September 2019"),
                      ("March 12, 2024", "March 2024"), ("12/03/2024", "March 2024"), ("2024-03", "2024-03"),
                      ("March 2024", "March 2024")]:
        assert pv.share_text(raw, ctx) == want, raw
    assert pv.share_text("Restarted at about 14:05 hrs on 2 May 2024 by ACME Pumps Ltd and ACME.",
                         ctx._replace(orgs=["ACME"])) == "Restarted on May 2024 by [organisation] and [organisation]."
    assert pv.share_text("R. K. Sharma at 14:05 on 2024-03-12", ctx._replace(mode="internal")) == \
        "R. K. Sharma at 14:05 on 2024-03-12"


@pytest.mark.parametrize("raw,want", [
    ("At 0216. He informed the operator.", "He informed the operator."),
    ("At 0216 he informed the operator.", "he informed the operator."),
    ("The leak started at about 0300 hrs near the pump.", "The leak started near the pump."),
    ("The leak started at about 0300 near the pump.", "The leak started near the pump."),
    ("Gas was seen around 0220 by the crew.", "Gas was seen by the crew."),
    ("approx. 0235 the alarm sounded", "the alarm sounded"),
    ("The line was open from 0200 until 0415 that night.", "The line was open that night."),
    ("Isolated by 0330, as logged.", "Isolated, as logged."),
    ("Readings at 9.15 am showed a drop.", "Readings showed a drop."),
    ("At 21.30 hrs the valve closed.", "the valve closed."),
    ("The pump stopped 9.15 p.m. on the day.", "The pump stopped on the day."),
    ("The alarm rang. At 0216. He informed the shift.", "The alarm rang. He informed the shift."),
    ("Hot work ran from 0200-0400 near the drain.", "Hot work ran near the drain."),
])
def test_share_text_strips_hhmm_after_a_preposition(raw, want):
    """Final review I5: 'At 0216' and dotted times with am, pm or hrs are clock times too."""
    assert pv.share_text(raw, pv.ShareCtx("external", [], [], {})) == want


@pytest.mark.parametrize("raw", [
    "in 2019 the gasket changed", "since 2018, leaving the gap", "from 2017 to 2020", "by 2020.", "from 2019",
    "until 2021 the plan held", "about 1999 the plant opened", "ran at 1500 rpm", "about 1200 m of pipe",
    "at 2500 kPa", "lasted 2.50 hours", "rated 4.20 bar", "at 2460", "At 0216x", "from 0200.5 m",
    "from 2017-2019 the audits lapsed", "at 2024/25 rates", "Wait... at 1200 rpm.",
])
def test_share_text_keeps_years_and_quantities(raw):
    assert pv.share_text(raw, pv.ShareCtx("external", [], [], {})) == raw


def test_share_doc_drops_keys_and_generalises_owners():
    ctx = pv.ShareCtx("external", [], ["Bharat Seals"], {})
    doc = {"text": "Bharat Seals at 14:05", "suggested_interviews": ["x"], "Key Questions": "q", "generated": "d",
           "causal_hypothesis": "h", "Evidence to Collect": "e",
           "run": SLUG, "votes": [{"vote": "downgrade", "reason": "the Plant Head said so", "voted_at": "2024-03-21"}],
           "owner": "Plant Manager", "items": [{"owner": "Shift Supervisor"}, {"owner": ["Plant Manager"]}], "n": 3}
    before = copy.deepcopy(doc)
    out = pv.share_doc(doc, ctx, owners=lambda o: "Plant & Asset Leadership" if "Plant" in o else None)
    assert doc == before
    assert out == {"text": "[organisation]", "votes": [{"vote": "downgrade", "voted_at": "March 2024"}],
                   "owner": "Plant & Asset Leadership",
                   "items": [{"owner": "Shift Supervisor"}, {"owner": ["Plant & Asset Leadership"]}], "n": 3}
    assert pv.share_doc(doc, ctx._replace(mode="internal")) == doc


def test_share_doc_edges_keep_node_ids_and_role_map_keys_are_masked():
    ctx = pv.ShareCtx("external", [], ["Bharat Seals"], {})
    g = {"kind": "graph", "nodes": [{"id": "N10", "label": "P-101 leak"}, {"id": "N11", "label": "Fire"}],
         "edges": [{"from": "N10", "to": "N11", "label": "at 14:05"}],
         "items": [{"kind": "message", "to": "Bharat Seals crew"}]}
    out = pv.share_doc(g, ctx)
    assert out["nodes"] == [{"id": "N10", "label": "[equipment tag] leak"}, {"id": "N11", "label": "Fire"}]
    assert out["edges"] == [{"from": "N10", "to": "N11", "label": ""}]
    assert out["items"][0]["to"] == "[organisation] crew"
    cl = {"clusters": {"role_map": {"Bharat Seals fitter": {"F": ["F02"], "T": "T1"}},
                       "unresolved_roles": ["Bharat Seals clerk"], "findings": {"FND-0001": {"F": ["F02"]}}}}
    out = pv.share_doc(cl, ctx)["clusters"]
    assert out == {"role_map": {"[organisation] fitter": {"F": ["F02"], "T": "T1"}},
                   "unresolved_roles": ["[organisation] clerk"], "findings": {"FND-0001": {"F": ["F02"]}}}


def test_share_graph_item_renders(tmp_path, shws):
    ws, orgs = shws
    doc = copy.deepcopy(SPACK)
    doc["sections"][0]["items"].append({"kind": "graph", "nodes": [{"id": "N10", "label": "Seal fitted"},
                                                                   {"id": "N11", "label": "Leak"}],
                                        "edges": [{"from": "N10", "to": "N11"}]})
    p = sign(ws, write_doc(tmp_path, doc), "--classification", "external", "--legal-review", "cleared")
    r = _cli("pack", str(ws), str(p), "--share", "external", "--orgs", str(orgs), "--out", str(tmp_path / "out"))
    assert r.returncode == 0, r.stderr
    t = (tmp_path / "out" / "test-pack.external.html").read_text(encoding="utf-8")
    assert "Seal fitted → Leak" in t


def test_shared_copies_drop_investigative_follow_up(tmp_path):
    ws, orgs = share_ws(tmp_path)
    fp = ws / "findings.json"
    F = json.loads(fp.read_text())
    F[0].update(causal_hypothesis="HYPO-MARK", evidence_to_collect="COLLECT-MARK")
    fp.write_text(json.dumps(F), encoding="utf-8")
    for cls in ("external", "public"):
        p = sign(ws, write_doc(tmp_path, SPACK, f"{cls}.json"), "--classification", cls, "--legal-review", "cleared")
        out = tmp_path / cls
        r = _cli("pack", str(ws), str(p), "--share", cls, "--orgs", str(orgs), "--out", str(out))
        assert r.returncode == 0, r.stderr
        t = (out / f"test-pack.{cls}.html").read_text(encoding="utf-8")
        for x in ("HYPO-MARK", "COLLECT-MARK", "causal_hypothesis", "evidence_to_collect"):
            assert x not in t, (cls, x)


def test_owner_generaliser_maps_t4_t5_only(shws):
    g = pmlc.owner_generaliser(pmlc.load_ws(shws[0]))
    assert g("Plant Manager") and g("Plant Manager") != "Plant Manager"
    assert g("Corporate HSE Director")
    assert g("Maintenance Supervisor") is None and g("Pump Technician") is None


def test_share_refused_without_legal_clearance(tmp_path, shws):
    ws, orgs = shws
    p = sign(ws, write_doc(tmp_path, SPACK), "--classification", "external")
    out = tmp_path / "out"
    r = _cli("pack", str(ws), str(p), "--share", "external", "--orgs", str(orgs), "--out", str(out))
    assert r.returncode == 1
    assert r.stderr.strip().splitlines()[-1] == 'cannot share externally: legal_review is not "cleared"'
    assert files(out) == []


def test_share_refusal_lists_every_unmet_condition(tmp_path, shws):
    ws, _ = shws
    p = write_doc(tmp_path, SPACK)
    out = tmp_path / "out"
    r = _cli("pack", str(ws), str(p), "--share", "public", "--embed", "full", "--out", str(out))
    assert r.returncode == 1
    assert r.stderr.strip().splitlines()[-1] == (
        'cannot share publicly: review is draft; classification is "internal" (needs "public"); '
        'legal_review is not "cleared"; --embed full is refused (use --embed summary)')
    assert files(out) == []
    sign(ws, p, "--classification", "external", "--legal-review", "cleared")
    r = _cli("pack", str(ws), str(p), "--share", "public", "--out", str(out))
    assert r.returncode == 1 and r.stderr.strip().endswith('classification is "external" (needs "public")')
    r = _cli("sign-off", str(ws), str(p), "--role", "HSE manager", "--classification", "public")
    assert r.returncode == 0
    r = _cli("pack", str(ws), str(p), "--share", "public", "--out", str(out))
    assert r.returncode == 1                      # sign-off keeps the earlier legal_review, so only the status is unmet
    assert r.stderr.strip().splitlines()[-1] == "cannot share publicly: review is reviewed (needs approved)"
    assert files(out) == []


def test_share_refused_on_forced_draft(tmp_path):
    ws, orgs = share_ws(tmp_path, status="in_progress", edited={})
    doc = dict(copy.deepcopy(SPACK), handling={"status": "approved", "reviewer_role": "HSE manager",
                                               "reviewed_on": "2026-10-19", "classification": "external",
                                               "legal_review": "cleared"})
    doc["handling"]["content_hash"] = pmlc.content_hash(doc)
    p = write_doc(tmp_path, doc)
    out = tmp_path / "out"
    r = _cli("pack", str(ws), str(p), "--share", "external", "--out", str(out))
    assert r.returncode == 1
    assert r.stderr.strip().splitlines()[-1] == (
        "cannot share externally: review is draft (the investigation status is in progress, "
        "no finding is marked as edited by a person, not signed off)")
    assert files(out) == []


def test_share_refused_after_edit(tmp_path, shws):
    """Review Focus 2: one character changed after sign-off, and the external copy is refused."""
    ws, orgs = shws
    p = sign(ws, write_doc(tmp_path, SPACK), "--classification", "external", "--legal-review", "cleared")
    ok = _cli("pack", str(ws), str(p), "--share", "external", "--orgs", str(orgs), "--out", str(tmp_path / "ok"))
    assert ok.returncode == 0, ok.stderr
    doc = json.loads(p.read_text())
    doc["sections"][0]["items"][0]["text"] = "Test before service!"
    p.write_text(json.dumps(doc), encoding="utf-8")
    out = tmp_path / "out"
    r = _cli("pack", str(ws), str(p), "--share", "external", "--orgs", str(orgs), "--out", str(out))
    assert r.returncode == 1
    assert r.stderr.strip().splitlines()[-1] == "cannot share externally: review is draft (edited since sign-off)"
    assert files(out) == []


def _escaped(s):
    return [s, json.dumps(s, ensure_ascii=False)[1:-1].replace("&", "\\u0026")]


def test_share_external_masks(tmp_path, shws):
    ws, orgs = shws
    p = sign(ws, write_doc(tmp_path, SPACK), "--classification", "external", "--legal-review", "cleared")
    raw = p.read_bytes()
    out = tmp_path / "out"
    r = _cli("pack", str(ws), str(p), "--share", "external", "--orgs", str(orgs), "--out", str(out))
    assert r.returncode == 0, r.stderr
    assert files(out) == ["test-pack.external.html", "test-pack.external.md"]
    assert p.read_bytes() == raw
    t = (out / "test-pack.external.html").read_text(encoding="utf-8")
    md = (out / "test-pack.external.md").read_text(encoding="utf-8")
    for leak in ("Bharat Seals", "Acme Pumps", "P-101", SINGLE, "14:05", "2024-03-12", "12 March 2024", "2024-03-20",
                 "2026-10-19", SLUG, "Plant Manager", "suggested_interviews", "key_questions"):
        assert leak not in t and leak not in md, leak
    assert not TIME_RX.search(t)
    for mark in ("[organisation]", "[equipment tag]", "March 2024"):
        assert mark in t, mark
    label = pv.load_singletons(pmlc.load_ws(ws))[SINGLE]
    assert any(x in t for x in _escaped(label))
    assert SHARED.format("external") in t and SHARED.format("external") in md
    assert "Review: approved by HSE manager on October 2026" in t
    assert "Work order of" in t                       # external keeps the justification, masked


def test_share_public_leaks_nothing(tmp_path, shws):
    """Review Focus 5: the whole HTML string (payload, print trail, <title>) carries no identifying detail."""
    ws, orgs = shws
    p = sign(ws, write_doc(tmp_path, SPACK), "--classification", "public", "--legal-review", "cleared")
    out = tmp_path / "out"
    r = _cli("pack", str(ws), str(p), "--share", "public", "--orgs", str(orgs), "--out", str(out))
    assert r.returncode == 0, r.stderr
    t = (out / "test-pack.public.html").read_text(encoding="utf-8")
    assert SLUG not in t and "acme-pump-fire" not in t
    assert not TIME_RX.search(t), TIME_RX.search(t)
    for leak in ("P-101", "AO-101", "Bharat Seals", "Acme Pumps", SINGLE):
        assert leak not in t, leak
    assert not re.search(r"\b\d{4}-\d{2}-\d{2}\b", t), re.search(r"\b\d{4}-\d{2}-\d{2}\b", t)
    assert not re.search(r"\b\d{1,2}(?:st|nd|rd|th)? (?:January|February|March|April|May|June|July|August|September|"
                         r"October|November|December) \d{4}\b", t)
    assert "Work order of" not in t and "supplied the seal" not in t      # AI justification and causal text dropped
    data = json.loads(re.search(r'id="pmlc-pack-data">(.*?)</script>', t, re.S).group(1))
    f1 = data["findings"]["FND-0001"]
    assert f1["label"] == "Pump restarted without leak test" and "justification" not in f1 and "causal" not in f1
    assert "run" not in data and SHARED.format("public") in t
    assert (out / "test-pack.public.md").exists() and not (out / "test-pack.html").exists()


def test_share_embed_full_refused(tmp_path, shws):
    ws, orgs = shws
    p = sign(ws, write_doc(tmp_path, SPACK), "--classification", "external", "--legal-review", "cleared")
    for mode in ("full", "cited"):
        out = tmp_path / mode
        r = _cli("pack", str(ws), str(p), "--share", "external", "--embed", mode, "--out", str(out))
        assert r.returncode == 1
        assert r.stderr.strip().splitlines()[-1] == (f"cannot share externally: --embed {mode} is refused "
                                                     "(use --embed summary)")
        assert files(out) == []
    r = _cli("pack", str(ws), str(p), "--share", "external", "--embed", "summary", "--orgs", str(orgs),
             "--out", str(tmp_path / "summary"))
    assert r.returncode == 0, r.stderr


def test_share_dashboard_external(tmp_path, shws):
    ws, orgs = shws
    reg = copy.deepcopy(MINI_REG)
    reg["entries"][0]["context"] = "At 14:05 on 2024-03-12 the ACME crew restarted P-101."
    rp = write_doc(tmp_path, reg, "reg.json")
    r = _cli("dashboard", str(ws), "--register", str(rp), "--share", "external")
    assert r.returncode == 1 and "cannot share externally: review is draft" in r.stderr
    sign(ws, rp, "--classification", "external", "--legal-review", "cleared")
    r = _cli("dashboard", str(ws), "--register", str(rp), "--share", "external", "--orgs", str(orgs))
    assert r.returncode == 0, r.stderr
    outs = ws.parent / "outputs"
    assert files(outs) == ["organisation-learning-map.external.html"]
    t = (outs / "organisation-learning-map.external.html").read_text(encoding="utf-8")
    for leak in ("ACME", "Bharat Seals", "P-101", "AO-101", SINGLE, "2024-03-12", "2024-03-20", "Acme Pumps"):
        assert leak not in t, leak
    assert not TIME_RX.search(t) and SHARED.format("external") in t and "[organisation]" in t
    assert "export.xlsx" not in t and "learnings-register.xlsx" not in t and "learnings-register.json" not in t


def test_register_xlsx_share_external(tmp_path, shws):
    ws, orgs = shws
    reg = copy.deepcopy(MINI_REG)
    reg["entries"][0]["context"] = "At 14:05 on 2024-03-12 the ACME crew restarted P-101."
    rp = write_doc(tmp_path, reg, "reg.json")
    r = _cli("register-xlsx", str(rp), "--ws", str(ws), "--share", "external")
    assert r.returncode == 1 and "cannot share externally: review is draft" in r.stderr
    assert files(tmp_path / "reg-out") == [] and not (tmp_path / "learnings-register.external.xlsx").exists()
    sign(ws, rp, "--classification", "external", "--legal-review", "cleared")
    r = _cli("register-xlsx", str(rp), "--ws", str(ws), "--share", "external", "--orgs", str(orgs))
    assert r.returncode == 0, r.stderr
    x = tmp_path / "learnings-register.external.xlsx"
    assert x.exists() and not (tmp_path / "learnings-register.xlsx").exists()
    wb = openpyxl.load_workbook(x)
    cells = " ".join(str(c) for s in wb.worksheets for row in s.iter_rows(values_only=True) for c in row if c)
    for leak in ("ACME", "14:05", "2024-03-12", "P-101", "2026-10-19"):
        assert leak not in cells, leak
    assert "[organisation]" in cells and SHARED.format("external") in cells and "export.xlsx" not in cells
    rp2 = write_doc(tmp_path, reg, "reg2.json")            # no workspace can be found: refused
    r = _cli("register-xlsx", str(rp2), "--share", "external")
    assert r.returncode == 1 and "no workspace (pass --ws)" in r.stderr


def test_render_name_guard(tmp_path, shws):
    ws, _ = shws
    names = tmp_path / "names.txt"
    names.write_text("# staff\nFictional Person\n", encoding="utf-8")
    env = {"PMLC_NAMES": str(names)}
    msg = re.compile(r"refused: a listed name appears in the output \((\d+) times\); fix the pack text or use "
                     r"ingest --redact")
    # A name in the pack text, accepted by lint_ok, is still caught in the rendered page (internal mode).
    doc = copy.deepcopy(SPACK)
    doc["sections"][0]["items"][0].update(text="Fictional Person tested the seal.",
                                          lint_ok={"rule": "JC1", "reason": "test of the render guard"})
    p = write_doc(tmp_path, doc)
    out = tmp_path / "out"
    r = _cli("pack", str(ws), str(p), "--out", str(out), env=env)
    assert r.returncode == 1 and msg.search(r.stderr) and int(msg.search(r.stderr).group(1)) >= 2
    assert files(out) == []
    # A name the pack text does not carry: the cited record's roles hold it.
    doc = copy.deepcopy(SPACK)
    doc["sections"][0]["items"][0]["cites"] = ["FND-0004"]
    p = write_doc(tmp_path, doc)
    r = _cli("pack", str(ws), str(p), "--out", str(out), "--names", str(names))
    assert r.returncode == 1 and msg.search(r.stderr) and files(out) == []
    # Without the names file the same pack renders.
    assert _cli("pack", str(ws), str(p), "--out", str(out)).returncode == 0
    # A name that HTML escaping changes (an apostrophe) is still caught, with no Markdown written.
    names.write_text("Fictional O'Brien\n", encoding="utf-8")
    doc = copy.deepcopy(SPACK)
    doc["sections"][0]["items"][0].update(text="Fictional O'Brien tested the seal.",
                                          lint_ok={"rule": "JC1", "reason": "test of the render guard"})
    p = write_doc(tmp_path, doc)
    out2 = tmp_path / "out2"
    r = _cli("pack", str(ws), str(p), "--out", str(out2), "--no-md", "--names", str(names))
    assert r.returncode == 1 and msg.search(r.stderr) and files(out2) == []
    names.write_text("# staff\nFictional Person\n", encoding="utf-8")
    # The dashboard has the same guard.
    reg = copy.deepcopy(MINI_REG)
    reg["entries"][0]["lesson"] += " Fictional Person saw it."
    rp = write_doc(tmp_path, reg, "reg.json")
    dash = tmp_path / "map.html"
    r = _cli("dashboard", str(ws), "--register", str(rp), "--out", str(dash), env=env)
    assert r.returncode == 1 and msg.search(r.stderr) and not dash.exists()


def test_scan_names_uses_privacy_core(tmp_path):
    n = tmp_path / "names.txt"
    n.write_text("# list\n\nR. K. Sharma\n", encoding="utf-8")
    t = tmp_path / "t.md"
    t.write_text("by R. K. Sharma\n", encoding="utf-8")
    from test_pmlc import PMLC
    import subprocess
    import sys
    r = subprocess.run([sys.executable, str(PMLC), "scan-names", str(n), str(t)], capture_output=True, text=True)
    assert r.returncode == 1 and "1 hit" in r.stdout


# ---- Task 10: scan-names hardening and private-output ignore rules ----------------------------------------------
def _scan(tmp_path, names, *args, env=None):
    import subprocess
    import sys
    from test_pmlc import PMLC
    n = tmp_path / "names10.txt"
    n.write_text(names, encoding="utf-8")
    return subprocess.run([sys.executable, str(PMLC), "scan-names", str(n), *map(str, args)],
                          capture_output=True, text=True, env=env)


def _zip(path, members):
    import zipfile
    with zipfile.ZipFile(path, "w") as z:
        for k, v in members.items():
            z.writestr(k, v)
    return path


def test_scan_names_variants(tmp_path):
    t = tmp_path / "t.md"
    t.write_text("Mr Sharma signed. Checked by R. Sharma.\n", encoding="utf-8")
    assert _scan(tmp_path, "Ravi Sharma\n", t).returncode == 0
    r = _scan(tmp_path, "Ravi Sharma\n", t, "--variants")
    assert r.returncode == 1 and "2 hit" in r.stdout and "Mr Sharma" in r.stdout and "R. Sharma" in r.stdout


def test_scan_names_docprops_and_comments(tmp_path):
    d = _zip(tmp_path / "a.docx", {
        "word/document.xml": "<w:t>clean</w:t>",
        "docProps/core.xml": "<dc:creator>Ravi Sharma</dc:creator>",
        "word/comments.xml": "<w:t>ask Ravi Sharma</w:t>",
        "word/header1.xml": "<w:t>Ravi Sharma</w:t>"})
    r = _scan(tmp_path, "Ravi Sharma\n", d)
    assert r.returncode == 1 and "3 hit" in r.stdout and "docProps/core.xml" in r.stdout
    assert "word/comments.xml" in r.stdout and "word/header1.xml" in r.stdout
    x = tmp_path / "b.xlsx"
    wb = openpyxl.Workbook()
    wb.properties.creator = "Ravi Sharma"
    wb.save(x)
    r = _scan(tmp_path, "Ravi Sharma\n", x)
    assert r.returncode == 1 and "docProps/core.xml" in r.stdout
    p = _zip(tmp_path / "c.pptx", {"docProps/core.xml": "<dc:creator>Ravi Sharma</dc:creator>"})
    assert _scan(tmp_path, "Ravi Sharma\n", p).returncode == 1


def test_scan_names_patterns(tmp_path):
    t = tmp_path / "t.md"
    t.write_text("mail a.b@x.org or call +91 98765 43210; Staff No. 4471\n", encoding="utf-8")
    assert _scan(tmp_path, "Nobody Here\n", t).returncode == 0
    r = _scan(tmp_path, "Nobody Here\n", t, "--patterns")
    assert r.returncode == 1
    for lab in ("pattern:email", "pattern:phone", "pattern:staff_no"):
        assert lab in r.stdout


def test_scan_names_pdf_metadata(tmp_path):
    import shutil
    if not shutil.which("pdfinfo"):
        pytest.skip("pdfinfo not installed")
    pdf = tmp_path / "m.pdf"
    pdf.write_bytes(
        b"%PDF-1.4\n1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj\n"
        b"3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 10 10]>>endobj\n"
        b"4 0 obj<</Author(Ravi Sharma)/Creator(Writer)>>endobj\n"
        b"trailer<</Root 1 0 R/Info 4 0 R/Size 5>>\n%%EOF\n")
    r = _scan(tmp_path, "Ravi Sharma\n", pdf)
    assert "pdf:Author" in r.stdout and "1 hit" in r.stdout and r.returncode == 1


def test_scan_names_pdf_without_pdfinfo(tmp_path):
    import os
    pdf = tmp_path / "m.pdf"
    pdf.write_bytes(b"%PDF-1.4\n")
    env = dict(os.environ, PATH=str(tmp_path / "empty"))
    r = _scan(tmp_path, "Ravi Sharma\n", pdf, env=env)
    assert "note: pdfinfo not found: PDF metadata not checked" in r.stdout


def test_gitignore_private_outputs():
    import shutil
    import subprocess
    if not shutil.which("git"):
        pytest.skip("git not installed")
    root = pmlc.HERE.parent.parent.parent

    def ignored(p):
        return subprocess.run(["git", "check-ignore", "-q", p], cwd=root).returncode == 0
    for p in ("pmlc-runs/x/outputs/a.pdf", "docs/x.docx", "x.pptx"):
        assert ignored(p), p
    for p in ("samples/x/a.pdf", "docs/briefs/a.pdf", "samples/x/outputs/a.docx"):
        assert not ignored(p), p


def test_scan_names_unreadable_office_files_are_unread(tmp_path):
    for nm in ("bad.docx", "bad.xlsx", "bad.pptx"):
        f = tmp_path / nm
        f.write_text("not a zip", encoding="utf-8")
        r = _scan(tmp_path, "Ravi Sharma\n", f)
        assert r.returncode == 2 and nm in r.stdout and "not scanned" in r.stdout, nm


def test_scan_names_split_runs_and_body(tmp_path):
    d = _zip(tmp_path / "s.docx", {
        "word/document.xml": "<w:p><w:r><w:t>Ra</w:t></w:r><w:r><w:rPr><w:b/></w:rPr><w:t>vi Sharma</w:t></w:r></w:p>"})
    r = _scan(tmp_path, "Ravi Sharma\n", d)
    assert r.returncode == 1 and "1 hit" in r.stdout and "word/document.xml" in r.stdout
    e = _zip(tmp_path / "e.pptx", {"ppt/slides/slide1.xml": "<a:p><a:r><a:t>R&amp;D: Ravi</a:t></a:r></a:p>"
                                   "<a:p><a:r><a:t>Sharma</a:t></a:r></a:p>"})
    assert _scan(tmp_path, "Ravi Sharma\n", e).returncode == 1  # a paragraph break counts as whitespace (conservative)
    assert _scan(tmp_path, "R&D\n", e).returncode == 1  # entity unescaped


def test_scan_names_names_only_in_attributes(tmp_path):
    cases = {
        "a.docx": ("word/comments.xml", '<w:comment w:id="0" w:author="Ravi Sharma" w:initials="RS"><w:p/></w:comment>'),
        "b.docx": ("word/people.xml", '<w15:person w15:author="Ravi Sharma"><w15:presenceInfo/></w15:person>'),
        "c.pptx": ("ppt/commentAuthors.xml", '<p:cmAuthor id="0" name="Ravi Sharma" initials="RS"/>'),
        "f.pptx": ("ppt/authors.xml", '<p188:author id="{1}" name="Ravi Sharma" userId="x"/>'),
        "g.xlsx": ("xl/persons/person.xml", '<person displayName="Ravi Sharma" id="{1}" userId="x"/>'),
        "d.pptx": ("ppt/comments/modernComment_1_2.xml", "<p188:cm author='Ravi Sharma'/>"),
    }
    for nm, (part, xml) in cases.items():
        r = _scan(tmp_path, "Ravi Sharma\n", _zip(tmp_path / nm, {part: xml}))
        assert r.returncode == 1 and part in r.stdout, nm
    n = _zip(tmp_path / "h.docx", {"word/comments.xml": '<w:comment username="Ravi Sharma" filename="Ravi Sharma"/>'})
    assert _scan(tmp_path, "Ravi Sharma\n", n).returncode == 0  # username= / filename= are not name attributes
    z = _zip(tmp_path / "e.docx", {"word/comments.xml": '<w:comment w:author="A &amp; Ravi Sharma"/>'})
    assert _scan(tmp_path, "A & Ravi Sharma\n", z).returncode == 1


def test_scan_names_oversize_member(tmp_path, monkeypatch):
    monkeypatch.setattr(pmlc_lib.privacy, "_MAX_MEMBER", 10)
    d = _zip(tmp_path / "big.docx", {"docProps/core.xml": "<x>" + "a" * 50 + "</x>"})
    out = list(pmlc._iter_texts(d, bad := []))
    assert out == [] and bad == [str(d)]


def test_pdf_meta_failure_is_not_checked(tmp_path):
    import shutil
    if not shutil.which("pdfinfo"):
        pytest.skip("pdfinfo not installed")
    f = tmp_path / "junk.pdf"
    f.write_bytes(b"junk")
    assert pmlc._pdf_meta(f) is False


# ----------------------------------------------------------------------------- public-sample tooling (Task 9)

def _sample(tmp_path, *extra, hidden_name=False):
    import subprocess
    import sys
    from openpyxl.comments import Comment
    from test_pmlc import ROOT
    memo = "To: Qux Zorbin\nFrom: Qux Zorbin\nCc: boss\n" + "The pump seal failed and leaked. " * 8
    x = make_export(tmp_path, summary=[("Investigating Officer", "Qux Zorbin"), ("Supervisor at Time", "Qux Zorbin")],
                    narrative=memo)
    wb = openpyxl.load_workbook(x)
    wb.properties.creator = "Qux Zorbin"
    wb.properties.lastModifiedBy = "Qux Zorbin"
    ws = wb["Identified Findings"]
    ws["A2"].comment = Comment("note by Qux Zorbin", "Qux Zorbin")
    ws.oddHeader.center.text = "Qux Zorbin"
    ws.oddFooter.left.text = "Qux Zorbin"
    if hidden_name:
        ws.append(["FND-9999", "", "Hidden Qux Zorbin"])
        ws.row_dimensions[ws.max_row].hidden = True
    wb.create_sheet("Personnel").append(["Qux Zorbin"])
    wb["Personnel"].sheet_state = "hidden"
    wb.save(x)
    out = tmp_path / "out.xlsx"
    r = subprocess.run([sys.executable, str(ROOT / "tools" / "make-public-sample.py"), str(x), str(out), *extra],
                       capture_output=True, text=True)
    return out, r


def test_make_public_sample_strips_metadata(tmp_path):
    import subprocess
    import sys
    import zipfile
    from test_pmlc import PMLC
    out, r = _sample(tmp_path)
    assert r.returncode == 0, r.stderr + r.stdout
    with zipfile.ZipFile(out) as z:
        assert not [n for n in z.namelist() if "comment" in n.lower()]
        blob = " ".join(z.read(n).decode("utf-8", "ignore") for n in z.namelist())
    assert "Qux" not in blob
    names = tmp_path / "n.txt"
    names.write_text("Qux Zorbin\n")
    r = subprocess.run([sys.executable, str(PMLC), "scan-names", str(names), str(out)], capture_output=True, text=True)
    assert r.returncode == 0 and "0 hit" in r.stdout
    # a hidden row in a kept sheet is unhidden, so the --names scan sees it and refuses to write
    h = tmp_path / "h"
    h.mkdir()
    out2, r2 = _sample(h, "--names", str(names), hidden_name=True)
    assert r2.returncode != 0 and "Hidden Qux" in r2.stdout and not out2.exists()


def test_make_public_sample_blanks_fields_and_memo_lines(tmp_path):
    out, r = _sample(tmp_path)
    assert r.returncode == 0, r.stderr
    s = openpyxl.load_workbook(out)["Incident Summary"]
    vals = {r[0].value: r[1].value for r in s.iter_rows() if r[0].value}
    assert vals["Investigating Officer"] is None and vals["Supervisor at Time"] is None
    text = " ".join(str(c.value) for r in s.iter_rows() for c in r if c.value)
    assert "To:" not in text and "From:" not in text and "Cc:" not in text and "pump seal failed" in text
    k = tmp_path / "k"
    k.mkdir()
    out2, r2 = _sample(k, "--keep-field", "Investigating Officer")
    assert r2.returncode == 0, r2.stderr
    s2 = openpyxl.load_workbook(out2)["Incident Summary"]
    vals2 = {r[0].value: r[1].value for r in s2.iter_rows() if r[0].value}
    assert vals2["Investigating Officer"] == "Qux Zorbin" and vals2["Supervisor at Time"] is None


def test_initials_without_space_match(tmp_path):
    rx = pv.name_pattern(["R. K. Sharma"], variants=True)
    for t in ("R.K. Sharma", "R. K. Sharma", "R K Sharma", "Shri Sharma", "r.k. sharma"):
        assert rx.fullmatch(t), t
    assert not rx.search("RKSharmaji")


def test_keep_field_unknown_warns_and_label_tolerant(tmp_path):
    import subprocess
    import sys
    from test_pmlc import ROOT
    x = make_export(tmp_path, summary=[(" investigating officer ", "Qux Zorbin")])
    out = tmp_path / "o.xlsx"
    r = subprocess.run([sys.executable, str(ROOT / "tools" / "make-public-sample.py"), str(x), str(out),
                        "--keep-field", "Nonsense"], capture_output=True, text=True)
    assert r.returncode == 0 and "--keep-field 'nonsense' matches none" in r.stderr
    s = openpyxl.load_workbook(out)["Incident Summary"]
    assert "Qux" not in " ".join(str(c.value) for row in s.iter_rows() for c in row if c.value)


# ----------------------------------------------------------------------------- final review minors

def test_lint_ok_rule_ids_survive_share_doc():
    """M2: a lint_ok rule ID such as JC10 is an ID, not an equipment tag."""
    ctx = pv.ShareCtx("external", [], [], {})
    doc = {"kind": "para", "text": "P-101", "lint_ok": {"rule": "JC10", "reason": "a drawing number"}}
    assert pv.share_doc(doc, ctx) == {"kind": "para", "text": "[equipment tag]",
                                      "lint_ok": {"rule": "JC10", "reason": "a drawing number"}}


def test_sign_off_lints_at_the_target_classification(shws, tmp_path):
    shws = shws[0]
    """M3: sign-off --classification external runs the external rules, so an organisation name (JC9, an error
    outside) is refused at sign-off, not later at pack --share external."""
    doc = copy.deepcopy(SPACK)
    doc["sections"][0]["items"][0]["text"] += " Acme Pumps Ltd supplied the seal."
    p = write_doc(tmp_path, doc)
    r = _cli("sign-off", str(shws), str(p), "--role", "HSE manager", "--status", "approved",
             "--classification", "external", "--legal-review", "cleared")
    assert r.returncode == 1 and "JC9" in r.stderr, r.stderr
    assert "handling" not in json.loads(p.read_text())
    assert _cli("sign-off", str(shws), str(p), "--role", "HSE manager").returncode == 0   # internal: a warning only


def test_register_xlsx_writes_formula_like_text_as_text(tmp_path):
    """M6: a cell string starting with = + - or @ is text in the workbook, never a formula."""
    reg = copy.deepcopy(MINI_REG)
    reg["title"] = "=HYPERLINK(\"http://x\",\"t\")"
    reg["entries"][0]["title"] = "=1+1"
    reg["entries"][0]["lesson"] = "+cmd|' /C calc'!A0"
    reg["entries"][0]["recommendations"][0]["what"] = "@SUM(A1)"
    p = write_doc(tmp_path, reg, "reg.json")
    out = tmp_path / "reg.xlsx"
    run("register-xlsx", str(p), "--out", str(out))
    wb = openpyxl.load_workbook(out)
    cells = [c for s in wb.worksheets for row in s.iter_rows() for c in row if isinstance(c.value, str)]
    risky = [c for c in cells if c.value[:1] in "=+-@"]
    assert {c.value for c in risky} >= {"=1+1", "+cmd|' /C calc'!A0", "@SUM(A1)", reg["title"]}
    assert all(c.data_type == "s" for c in risky), [(c.coordinate, c.data_type) for c in risky]


def test_pack_given_as_register_is_refused(shws, tmp_path):
    """M9: a pack JSON passed where a register goes gives one clear line, not a traceback."""
    shws = shws[0]
    p = write_doc(tmp_path, copy.deepcopy(SPACK))
    for args in (("dashboard", str(shws), "--register", str(p), "--out", str(tmp_path / "m.html")),
                 ("register-xlsx", str(p), "--out", str(tmp_path / "r.xlsx"))):
        r = _cli(*args)
        assert r.returncode == 1 and "Traceback" not in r.stderr, r.stderr
        assert r.stderr.strip() == f"{p} is a pack, not a register"


# ----------------------------------------------------------------------------- Task 19: deferred hygiene, part B

@pytest.mark.parametrize("raw,want", [
    ("Gas was seen on the scale at 9:15.", "Gas was seen on the scale."),                                # A44
    ("Check the seals, e.g., the flange seal.", "Check the seals, e.g., the flange seal."),            # A45
    ("Check the seals, e.g., the flange seal, at 9:15 pm.", "Check the seals, e.g., the flange seal."),
    ("The leak started at 9:15 pm. Then the alarm sounded.", "The leak started. Then the alarm sounded."),  # A46
    ("The leak started at 9.15 p.m. Then the alarm sounded.", "The leak started. Then the alarm sounded."),
    ("The tank, at 14:05. Then", "The tank. Then"),
    ("The valve was open since 0300.", "The valve was open."),                                           # A47
    ("The site has about 1200 workers.", "The site has about 1200 workers."),
    ("The audit in 2019 found gaps.", "The audit in 2019 found gaps."),
    ("The plant has run since 2018, leaving the gap.", "The plant has run since 2018, leaving the gap."),
    ("The ratio: 1:10.", "The ratio: 1:10."),
    ("Mix (ratio 1:10) was used.", "Mix (ratio 1:10) was used."),
    ("A 1:10 dilution. The mix at 9:15 rang.", "A 1:10 dilution. The mix rang."),
])
def test_strip_times_cases(raw, want):
    assert pv.strip_times(raw) == want


@pytest.mark.parametrize("role", ["Maintenance Manager (Dec-2019)", "Operations Head - 24 May 2020",
                                  "Site Engineer (2019)", "Planner (12/03/2020)", "Planner - 2020-03-12"])
def test_a_date_only_qualifier_is_not_a_singleton(role):
    assert _qual(role) == []


def test_a_date_with_a_real_qualifier_is_still_a_singleton():
    assert _qual("Shift A Engineer (CR 2) (Dec-2019)") == ["Shift A Engineer (CR 2) (Dec-2019)"]


def test_name_pattern_accepts_initials_without_spaces():
    for names in (["R.K. Vandermere"], ["R. K. Vandermere"], ["RK Vandermere"]):
        rx = pv.name_pattern(names, variants=True)
        for text in ("by R.K. Vandermere", "by R. K. Vandermere", "by R K Vandermere", "by Vandermere"):
            assert rx.search(text), (names, text)
    exact = pv.name_pattern(["R.K. Vandermere"])
    assert exact.search("by R. K. Vandermere") and exact.search("by R.K. Vandermere")
    assert "R. K. Vandermere" in pv.name_variants("R.K. Vandermere")


def test_singleton_roles_json_goes_stale_after_a_changed_export(tmp_path):
    from test_pmlc import _changed_export, _status_row
    x = make_export(tmp_path)
    w = tmp_path / "ws"
    run("ingest", str(x), "--legacy-out", str(w))
    run("classify", str(w))
    assert "current" in _status_row(run("status", str(w)), "singleton_roles.json")
    run("ingest", str(_changed_export(tmp_path)), "--legacy-out", str(w), "--force")
    assert "stale" in _status_row(run("status", str(w)), "singleton_roles.json")
    run("classify", str(w))
    assert "current" in _status_row(run("status", str(w)), "singleton_roles.json")


def test_share_doc_merges_role_map_values_when_masked_keys_collide():
    """A43: two roles that mask to the same key keep both roles' entities, tiers and relations."""
    ctx = pv.ShareCtx("external", [], ["Quillon Seals", "ACME"], {})
    cl = {"role_map": {"Quillon Seals fitter": {"F": ["F02"], "T": "T1", "R": ["actor"]},
                       "ACME fitter": {"F": ["F03", "F02"], "T": "T2", "R": ["actor", "owner"], "source": "overlay"}}}
    out = pv.share_doc(cl, ctx)["role_map"]
    assert list(out) == ["[organisation] fitter"]
    assert out["[organisation] fitter"] == {"F": ["F02", "F03"], "T": "T1", "R": ["actor", "owner"], "source": "overlay"}


def test_slug_date_is_coarsened_to_the_year():
    """A51: an ISO date at the start of a slug keeps the year only; a lone date still becomes 'Month YYYY'."""
    ctx = pv.ShareCtx("public", [], [], {})
    assert pv.share_text("run 2020-09-24-hazira-gas-release-fire done", ctx) == "run 2020-hazira-gas-release-fire done"
    assert pv.coarsen_dates("2020-09-24-hazira") == "2020-hazira"
    assert pv.share_text("on 2020-09-24 only", ctx) == "on September 2020 only"


def test_scan_names_patterns_ignore_the_brand_address(tmp_path):
    names = tmp_path / "names.txt"
    names.write_text("Fictional Person\n", encoding="utf-8")
    doc = tmp_path / "doc.md"
    doc.write_text(pmlc.BRAND + "\nMail jo.example@acme.test about it.\n", encoding="utf-8")
    r = _cli("scan-names", str(names), str(doc), "--patterns")
    assert "jo.example@acme.test" in r.stdout and "ashley@eyekyam.com" not in r.stdout
    assert "1 hit(s)." in r.stdout


def test_office_member_cap_reads_with_a_byte_limit(tmp_path, monkeypatch, capsys):
    """A35: the member cap counts bytes actually read, not the size the zip header declares."""
    import io
    import zipfile

    class FakeInfo:
        filename = "docProps/core.xml"
        file_size = 1                                   # a lying header

    class FakeZip:
        def __init__(self, *a, **k):
            pass

        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

        def infolist(self):
            return [FakeInfo()]

        def open(self, name, *a, **k):
            return io.BytesIO(b"<x>" + b"a" * 500 + b"</x>")

        def read(self, name):
            raise AssertionError("unbounded read")

    monkeypatch.setattr(zipfile, "ZipFile", FakeZip)
    monkeypatch.setattr(pmlc_lib.privacy, "_MAX_MEMBER", 100)
    unread = []
    got = list(pmlc_lib.privacy._office_xml_texts(tmp_path / "f.docx", unread))
    assert got == [] and unread == [str(tmp_path / "f.docx")]
    assert "over 50 MB" in capsys.readouterr().out


def test_redact_map_marks_an_ambiguous_variant(tmp_path):
    """_canon collision guard: one surname shared by two people with different roles is redacted, not guessed."""
    from pmlc_lib.ingest.xlsx import build_redact_map
    repl, multi = build_redact_map({"A. Vandermere": "Fitter", "B. Vandermere": "Officer", "C. Quillon": "Clerk"})
    assert repl["vandermere"] == "[redacted]" and repl["quillon"] == "Clerk"
    assert repl["a.vandermere".replace(".", "")] == "Fitter" and "A. Vandermere" in multi


def _pack_record(ws, p):
    from pmlc_lib.workspace import load_manifest, pack_record
    return pack_record(load_manifest(ws), p)


def test_a_refused_share_records_nothing_in_the_manifest(tmp_path, shws):
    """The share gate runs before the passing check is recorded; an internal render still records it."""
    ws, _ = shws
    p = write_doc(tmp_path, SPACK, "test-pack.json")
    r = _cli("pack", str(ws), str(p), "--share", "external", "--out", str(tmp_path / "out"))
    assert r.returncode == 1 and "cannot share externally" in r.stderr
    assert _pack_record(ws, p) is None
    r = _cli("pack", str(ws), str(p), "--out", str(tmp_path / "out2"))
    assert r.returncode == 0, r.stderr
    assert _pack_record(ws, p)["ok"] is True


def test_a_failing_pack_is_still_recorded_when_a_share_is_asked_for(tmp_path, shws):
    ws, _ = shws
    bad = copy.deepcopy(SPACK)
    bad["status"] = "Draft"
    p = write_doc(tmp_path, bad, "test-pack.json")
    r = _cli("pack", str(ws), str(p), "--share", "external", "--out", str(tmp_path / "out"))
    assert r.returncode == 1
    assert _pack_record(ws, p)["ok"] is False


def test_share_out_may_not_be_the_internal_copy(tmp_path, shws):
    ws, orgs = shws
    reg = copy.deepcopy(MINI_REG)
    rp = write_doc(tmp_path, reg, "reg.json")
    sign(ws, rp, "--classification", "external", "--legal-review", "cleared")
    internal = ws.parent / "outputs" / "organisation-learning-map.html"
    r = _cli("dashboard", str(ws), "--register", str(rp), "--share", "external", "--orgs", str(orgs), "--out", str(internal))
    assert r.returncode == 1 and "internal copy" in r.stderr and not internal.exists()
    x = tmp_path / "learnings-register.xlsx"
    r = _cli("register-xlsx", str(rp), "--ws", str(ws), "--share", "external", "--orgs", str(orgs), "--out", str(x))
    assert r.returncode == 1 and "internal copy" in r.stderr and not x.exists()
    other = tmp_path / "elsewhere.xlsx"
    r = _cli("register-xlsx", str(rp), "--ws", str(ws), "--share", "external", "--orgs", str(orgs), "--out", str(other))
    assert r.returncode == 0, r.stderr
    assert other.exists()


def test_summary_mode_builds_the_embedded_records_once(monkeypatch, tmp_path, shws):
    import pmlc_lib.pack.html as ph
    ws, _ = shws
    p = write_doc(tmp_path, SPACK)
    calls = []
    real = ph.embedded_records
    monkeypatch.setattr(ph, "embedded_records", lambda *a, **k: calls.append(1) or real(*a, **k))
    import argparse
    ns = argparse.Namespace(ws=str(ws), pack=str(p), cmd="pack", share=None, embed="summary", out=str(tmp_path / "o"),
                            no_md=True, names=None, orgs=None)
    ph.cmd_pack(ns)
    assert len(calls) == 1
