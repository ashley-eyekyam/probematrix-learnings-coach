"""Engine tests on a small synthetic export (fictional event, no real data)."""
import argparse
import json
import re
import shutil
import subprocess
import sys
from pathlib import Path

import pytest

openpyxl = pytest.importorskip("openpyxl")
ROOT = Path(__file__).resolve().parents[1]
PMLC = ROOT / "skills" / "learnings-coach" / "scripts" / "pmlc.py"
sys.path.insert(0, str(PMLC.parent))
import pmlc  # noqa: E402  (engine module; standard-library imports only)

FIND_COLS = ["Finding ID", "Nanocode", "Name", "Level", "Level Name", "Category", "Subcategory", "Classification",
             "Rating", "Instance ID", "Instance Label", "Short Description", "AI Justification", "Causal Explanation",
             "Individuals Involved", "Linked Evidence", "Evidence to Collect", "Key Questions", "Suggested Interviews",
             "Causal Hypothesis", "User Edited", "Created At", "Updated At"]

ROWS = [
    ["FND-0001", "AB 1.1", "Step Omitted", 1, "Unsafe Acts", "Unsafe Acts", "Errors", "substantiated", 3, 1,
     "Pump restarted without leak test", "Pump restarted without leak test after seal change on 12/03/2024.",
     "Work order shows no test [EVD-0001].", "Enabled by SV 1.4 and OR 2.2.",
     "Pump Technician (seal change crew), Maintenance Supervisor (return-to-service approval)", "EVD-0001",
     None, None, None, None, "No", None, None],
    ["FND-0002", "SV 1.4", "No Guidance Provided", 3, "Supervision", "Supervision", "Supervision Inadequate",
     "substantiated", 4, 1, "No return-to-service checklist", "No checklist for return to service.",
     "Procedure index has no entry [EVD-0002].", "Led to AB 1.1.",
     "Maintenance Supervisor, Maintenance Manager (procedure owner)", "EVD-0002", None, None, None, None, "No",
     None, None],
    ["FND-0003", "OR 2.2", "Spares Programme Failure", 4, "Organisational Influences", "Organisational Influences",
     "Resource Management", "substantiated", 4, 1, "Seals bought locally without specification",
     "Local seals without specification since March 2022.", "Purchase records [EVD-0003].",
     "This condition enabled SV 1.4 and AB 1.1.",
     "Stores/Procurement Officer, Plant Manager (budget owner), Reliability Engineer (absent)", "EVD-0003, EVD-0001",
     None, None, None, None, "No", None, None],
    ["FND-0004", "PE 2.1", "Alarm Not Configured", 2, "Preconditions", "Preconditions", "Technological Environment",
     "substantiated", 2, 1, "Gas alarm display-only", "Gas detector alarm had no shutdown.", "DCS list [EVD-0002].",
     "Delayed response; see OR 2.2.", "Control Room Operators, Emergency Response planners, Mr Fictional Person",
     "EVD-0002", None, None, None, None, "No", None, None],
]

MINI_REG = {"title": "Synthetic register", "source": "export.xlsx", "entries": [
    {"lr_id": "LR-01", "title": "Return to service without a test", "finding_ids": ["FND-0001", "FND-0002"],
     "evidence_ids": ["EVD-0001"], "priority": 1, "owner_entity": "F02 Mechanical Maintenance",
     "hfacs_level": "L3 Supervision Inadequate", "tiers": ["T1", "T2"],
     "lesson": "If a joint is disturbed, test it before service, because untested joints leak (EVD-0001).",
     "precursors": "Earlier seal leaks", "custom_field": "kept as given",
     "tier_messages": {"T1": "Do not pressurise without the test record.", "T2": "Hold point before acceptance."},
     "recommendations": [
         {"what": "Interlock pressurisation on a signed test record", "hoc_tier": "Engineering control: interlock",
          "leading_indicator": "Share of restarts with a test record", "owner_note": "extra key"},
         {"what": "y </script><script>alert(1)</script> & <b>", "hoc_tier": "Administrative (procedure)",
          "leading_indicator": "i"}]}]}


def make_export(d: Path) -> Path:
    wb = openpyxl.Workbook()
    s = wb.active
    s.title = "Incident Summary"
    for r in [("INCIDENT SUMMARY", None), (None, None), ("Field", "Value"), ("Incident Title", "Synthetic pump fire"),
              ("Date", "2024-03-20"), ("Investigation Status", "completed")]:
        s.append(r)
    f = wb.create_sheet("Identified Findings")
    f.append(FIND_COLS)
    for r in ROWS:
        f.append(r)
    e = wb.create_sheet("Evidence Links")
    e.append(["Finding Code", "Finding Name", "Evidence ID", "Evidence Type", "Relevance", "Specific Reference",
              "Evidence Description"])
    e.append(["AB 1.1", "Step Omitted", "EVD-0001", "document", "Shows no test", "no leak test recorded", "Work order"])
    e.append(["SV 1.4", "No Guidance Provided", "EVD-0002", "document", "No checklist", "index has no entry", "Index"])
    e.append(["OR 2.2", "Spares Programme Failure", "EVD-0003", "log_record", "Local buys", "bought locally", "PO list"])
    c = wb.create_sheet("Causal Chain")
    c.append(["Link ID", "From Code", "From Level", "To Code"])
    c.append(["No causal links found", None, None, None])
    x = d / "export.xlsx"
    wb.save(x)
    return x


@pytest.fixture(scope="module")
def ws(tmp_path_factory):
    d = tmp_path_factory.mktemp("pmlc")
    x = make_export(d)
    run("ingest", str(x), "--legacy-out", str(d / "ws"))
    run("classify", str(d / "ws"))
    return d / "ws"


def run(*args):
    r = subprocess.run([sys.executable, str(PMLC), *args], capture_output=True, text=True)
    assert r.returncode == 0, r.stderr + r.stdout
    return r.stdout


def run_fail(*args):
    r = subprocess.run([sys.executable, str(PMLC), *args], capture_output=True, text=True)
    assert r.returncode != 0, r.stdout
    return r


def test_ingest(ws):
    F = json.loads((ws / "findings.json").read_text())
    assert len(F) == 4
    assert F[0]["roles"] == ["Pump Technician (seal change crew)", "Maintenance Supervisor (return-to-service approval)"]
    E = json.loads((ws / "evidence.json").read_text())
    assert set(E) == {"EVD-0001", "EVD-0002", "EVD-0003"}
    inc = json.loads((ws / "incident.json").read_text())
    assert inc["fields"]["Incident Title"] == "Synthetic pump fire"
    assert inc["causal_chain_sheet"] == []


def test_classify(ws):
    cl = json.loads((ws / "clusters.json").read_text())
    rm = cl["role_map"]
    assert "F02" in rm["Pump Technician (seal change crew)"]["F"]
    assert rm["Maintenance Supervisor (return-to-service approval)"]["T"] == "T2"
    assert rm["Maintenance Supervisor (return-to-service approval)"]["R"] == "approver"
    assert rm["Reliability Engineer (absent)"]["R"] == "absent"
    assert "F06" in rm["Stores/Procurement Officer"]["F"]
    assert "F09b" in rm["Emergency Response planners"]["F"]
    assert "Mr Fictional Person" in cl["unresolved_roles"]


def test_overrides(ws, tmp_path):
    o = tmp_path / "o.json"
    o.write_text(json.dumps({"Mr Fictional Person": {"F": ["F01"], "T": "T1", "R": "actor"}}))
    run("classify", str(ws), "--overrides", str(o))
    cl = json.loads((ws / "clusters.json").read_text())
    assert "Mr Fictional Person" not in cl["unresolved_roles"]


def test_overrides_survive_plain_reclassify(tmp_path):
    x = make_export(tmp_path)
    w = tmp_path / "ws"
    run("ingest", str(x), "--legacy-out", str(w))
    run("classify", str(w))
    (w / "role_overrides.json").write_text(
        json.dumps({"Mr Fictional Person": {"F": ["F01"], "T": "T1", "R": "actor"},
                    "Stores/Procurement Officer": {"T": "T1"}}))
    out = run("classify", str(w))  # plain re-run, as training-designer used to do
    assert "Applied 2 override(s) from" in out
    cl = json.loads((w / "clusters.json").read_text())
    assert cl["role_map"]["Mr Fictional Person"]["source"] == "override"
    assert cl["overrides"]["applied"] == 2 and cl["overrides"]["stale"] == []
    assert cl["overrides"]["file"].endswith("role_overrides.json")
    assert not (w / "unresolved_roles.json").exists()
    run("classify", str(w), "--no-overrides")
    cl = json.loads((w / "clusters.json").read_text())
    assert cl["role_map"]["Mr Fictional Person"]["source"] == "rule"
    assert cl["overrides"] == {"file": None, "applied": 0, "stale": []}
    assert (w / "unresolved_roles.json").exists()


def test_overrides_bad_json_and_stale_keys(tmp_path):
    x = make_export(tmp_path)
    w = tmp_path / "ws"
    run("ingest", str(x), "--legacy-out", str(w))
    (w / "role_overrides.json").write_text("{not json")
    r = run_fail("classify", str(w))
    assert "role_overrides.json" in r.stderr and "--no-overrides" in r.stderr
    run("classify", str(w), "--no-overrides")
    (w / "role_overrides.json").write_text(json.dumps({"Nobody Here": {"F": ["F01"]}}))
    out = run("classify", str(w))
    assert "warning: 1 override(s) match no role: Nobody Here" in out
    cl = json.loads((w / "clusters.json").read_text())
    assert cl["overrides"]["stale"] == ["Nobody Here"] and cl["overrides"]["applied"] == 0


def test_explicit_overrides_file_recorded(tmp_path):
    x = make_export(tmp_path)
    w = tmp_path / "ws"
    run("ingest", str(x), "--legacy-out", str(w))
    o = tmp_path / "o.json"
    o.write_text(json.dumps({"Mr Fictional Person": {"F": ["F01"], "T": "T1", "R": "actor"}}))
    run("classify", str(w), "--overrides", str(o))
    cl = json.loads((w / "clusters.json").read_text())
    assert cl["overrides"]["file"] == str(o) and cl["overrides"]["applied"] == 1


@pytest.fixture
def fresh_ws(tmp_path):
    x = make_export(tmp_path)
    w = tmp_path / "ws"
    run("ingest", str(x), "--legacy-out", str(w))
    return w


def test_override_f_as_string_is_one_entity(fresh_ws):                              # final review 1
    (fresh_ws / "role_overrides.json").write_text(json.dumps({"Mr Fictional Person": {"F": "F05", "T": "T1"}}))
    run("classify", str(fresh_ws))
    cl = json.loads((fresh_ws / "clusters.json").read_text())
    assert cl["role_map"]["Mr Fictional Person"]["F"] == ["F05"]
    assert not {"F", "0", "5"} & set(cl["entities"])


@pytest.mark.parametrize("bad, words", [
    ({"Mr Fictional Person": "F05"}, ["role 'Mr Fictional Person'", "must be an object"]),
    (["F05"], ["must be a JSON object of roles"]),
    ({"Mr Fictional Person": {"F": [5]}}, ["F must be a list"]),
    ({"Mr Fictional Person": {"T": ["T1"]}}, ["T must be one ID"]),
    ({"Mr Fictional Person": {"F": ["F77"]}}, ["unknown F ID(s) F77"]),
    ({"Mr Fictional Person": {"T": "T7"}}, ["unknown T ID(s) T7"]),
    ({"Mr Fictional Person": {"R": "boss"}}, ["unknown R ID(s) boss"]),
    ({"Mr Fictional Person": {"Tier": "T1"}}, ["unknown key(s) Tier"]),
])
def test_malformed_overrides_exit_with_a_sentence(fresh_ws, bad, words):           # final review 1
    (fresh_ws / "role_overrides.json").write_text(json.dumps(bad))
    r = run_fail("classify", str(fresh_ws))
    assert "Traceback" not in r.stderr and "role_overrides.json" in r.stderr
    for w in words:
        assert w in r.stderr, r.stderr


def test_override_fallback_ids_are_valid(fresh_ws):                                 # final review 1
    (fresh_ws / "role_overrides.json").write_text(json.dumps({"Mr Fictional Person": {"F": [], "T": "T9"}}))
    run("classify", str(fresh_ws))
    o = fresh_ws / "o.json"
    o.write_text(json.dumps({"Mr Fictional Person": {"F": ["F99"], "R": "actor"}}))
    run("classify", str(fresh_ws), "--overrides", str(o))


def test_missing_explicit_overrides_file(fresh_ws, tmp_path):                       # final review 1
    r = run_fail("classify", str(fresh_ws), "--overrides", str(tmp_path / "missing.json"))
    assert "Traceback" not in r.stderr
    assert "Overrides file" in r.stderr and "missing.json not found" in r.stderr and "--no-overrides" in r.stderr


def test_training_designer_never_classifies():
    assert "PMLC classify" not in (ROOT / "skills/training-designer/SKILL.md").read_text()


def test_causal(ws):
    out = run("causal", str(ws))
    g = json.loads((ws / "causal_graph.json").read_text())
    pairs = {(e["from"], e["to"]) for e in g["edges"]}
    assert ("SV 1.4", "AB 1.1") in pairs
    assert ("OR 2.2", "SV 1.4") in pairs
    assert (ws / "causal_graph.mmd").read_text().startswith("flowchart")
    assert "distinct links" in out


def test_select_and_timeline(ws):
    out = run("select", str(ws), "--entity", "F02", "--brief")
    assert "FND-0001" in out
    out = run("select", str(ws), "--level", "4", "--format", "json")
    assert json.loads(out)[0]["id"] == "FND-0003"
    run("timeline", str(ws), "--before", "2024-03-20")
    t = json.loads((ws / "timeline_candidates.json").read_text())
    assert any(x["date"] == "2024-03-12" for x in t)


def test_interfaces_dashboard_register(ws, tmp_path):
    run("interfaces", str(ws))
    assert json.loads((ws / "interfaces.json").read_text())["pairs"]
    html = tmp_path / "d.html"
    run("dashboard", str(ws), "--out", str(html))
    assert "ProbeMatrix Learnings Coach" in html.read_text()
    reg = tmp_path / "r.json"
    reg.write_text(json.dumps({"entries": [
        {"lr_id": "LR-01", "title": "Return to service without a test", "finding_ids": ["FND-0001"], "priority": 1,
         "owner_entity": "F02 Mechanical Maintenance", "hfacs_level": "L3 Supervision Inadequate",
         "tier_messages": {"T1": "Do not pressurise without the test record.", "T2": "Hold point before acceptance."},
         "recommendations": [{"what": "x", "hoc_tier": "Engineering control: interlock"},
                             {"what": "y", "hoc_tier": "Administrative (procedure)"}]}]}))
    out = tmp_path / "r.xlsx"
    run("register-xlsx", str(reg), "--out", str(out))
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Learnings Register", "Recommendations", "Learning Plan", "Summary", "About"]
    summary = [tuple(r) for r in wb["Summary"].iter_rows(values_only=True)]
    assert ("Recommendations by control tier", "engineering", 1) in summary
    assert ("Learnings by HFACS level", "L3 Supervision", 1) in summary
    assert ("Learnings by cluster", "F02", 1) in summary
    assert wb["Learning Plan"]["D2"].value == "Do not pressurise without the test record."
    html2 = tmp_path / "d2.html"
    run("dashboard", str(ws), "--out", str(html2), "--register", str(reg))
    t = html2.read_text()
    assert "Learning plan by level" in t and "Hold point before acceptance." in t and "50%" in t


def test_check_register_and_overrides(ws, tmp_path):
    good = {"entries": [{"lr_id": "LR-01", "finding_ids": ["FND-0001", "FND-0002"], "evidence_ids": ["EVD-0001"],
                         "hfacs_level": "L3 x", "priority": 1, "owner_entity": "F02 Mechanical Maintenance",
                         "tier_messages": {"T2": "m"},
                         "recommendations": [{"what": "w", "hoc_tier": "engineering", "leading_indicator": "i"}]}]}
    p = tmp_path / "g.json"
    p.write_text(json.dumps(good))
    out = run("check-register", str(ws), str(p))
    assert "0 problem" in out and "2 of 4 findings covered" in out
    bad = {"entries": [{"lr_id": "LR-02", "finding_ids": ["FND-9999"], "recommendations": [{"hoc_tier": "maybe"}]}]}
    p.write_text(json.dumps(bad))
    r = subprocess.run([sys.executable, str(PMLC), "check-register", str(ws), str(p)], capture_output=True, text=True)
    assert r.returncode == 1 and "unknown finding FND-9999" in r.stdout and "no leading_indicator" in r.stdout
    h = tmp_path / "d.html"
    run("dashboard", str(ws), "--out", str(h), "--date", "2024-03-21", "--title", "Pump fire")
    assert "2024-03-21" in h.read_text() and "Pump fire" in h.read_text()


def test_scan_names(ws, tmp_path):
    n = tmp_path / "names.txt"
    n.write_text("Fictional Person\n")
    r = subprocess.run([sys.executable, str(PMLC), "scan-names", str(n), str(ws / "findings.json")],
                       capture_output=True, text=True)
    assert r.returncode == 1 and "1 hit" in r.stdout


def test_scan_names_reports_files_it_cannot_read(tmp_path):
    n = tmp_path / "names.txt"
    n.write_text("Fictional Person\n")
    d = tmp_path / "out"
    d.mkdir()
    (d / "notes.md").write_text("roles only\n")
    (d / "report.pdf").write_bytes(b"%PDF-1.4 Fictional Person")
    r = subprocess.run([sys.executable, str(PMLC), "scan-names", str(n), str(d)], capture_output=True, text=True)
    assert r.returncode == 2 and "0 hit(s)" in r.stdout and "report.pdf" in r.stdout and "not scanned" in r.stdout
    assert "pdftotext" in r.stdout                                    # tells you how to scan it
    (d / "report.pdf").unlink()
    r = subprocess.run([sys.executable, str(PMLC), "scan-names", str(n), str(d)], capture_output=True, text=True)
    assert r.returncode == 0 and "not scanned" not in r.stdout


def test_slug_rules():
    assert pmlc.run_slug("2020-09-24", "hazira-gas-release-fire") == "2020-09-24-hazira-gas-release-fire"
    assert pmlc.run_slug("2020-09-24 00:00:00", "a1") == "2020-09-24-a1"
    for bad in ["Hazira", "gas_fire", "-fire", "fire-", "gas--fire", "", None]:
        with pytest.raises(ValueError):
            pmlc.run_slug("2020-09-24", bad)
    for bad_date in ["24/09/2020", "2020-13-01", "2024-03-211", "", None]:
        with pytest.raises(ValueError):
            pmlc.run_slug(bad_date, "fire")
    assert pmlc.valid_run_slug("2020-09-24-hazira-fire")
    assert not pmlc.valid_run_slug("2020-02-30-x")
    assert not pmlc.valid_run_slug("hazira-2020")
    assert not pmlc.valid_run_slug("2020-09-24")


def test_outputs_dir(tmp_path):
    std = tmp_path / "pmlc-runs" / "2020-09-24-hazira-fire" / "workspace"
    assert pmlc.is_standard_ws(std)
    assert pmlc.outputs_dir(std) == (tmp_path / "pmlc-runs" / "2020-09-24-hazira-fire" / "outputs").resolve()
    legacy = tmp_path / "pmlc-hazira2020"
    assert not pmlc.is_standard_ws(legacy)
    assert pmlc.outputs_dir(legacy) == legacy.resolve() / "outputs"
    assert not pmlc.is_standard_ws(tmp_path / "runs" / "2020-09-24-x" / "workspace")
    assert not pmlc.is_standard_ws(tmp_path / "pmlc-runs" / "Bad" / "workspace")


def test_ingest_standard_layout(tmp_path):
    x = make_export(tmp_path)
    out = run("ingest", str(x), "--slug", "pump-fire")
    run_dir = tmp_path / "pmlc-runs" / "2024-03-20-pump-fire"
    assert (run_dir / "workspace" / "findings.json").exists()
    assert (run_dir / "outputs").is_dir()
    assert "outputs:" in out
    r = run_fail("ingest", str(x), "--slug", "pump-fire")
    assert "already holds a workspace" in r.stderr
    run("ingest", str(x), "--slug", "pump-fire", "--force")


def test_ingest_date_override_and_root(tmp_path):
    x = make_export(tmp_path)
    root = tmp_path / "elsewhere"
    run("ingest", str(x), "--slug", "pump-fire", "--date", "2024-03-21", "--root", str(root))
    assert (root / "pmlc-runs" / "2024-03-21-pump-fire" / "workspace" / "findings.json").exists()


def test_ingest_explicit_standard_out(tmp_path):
    x = make_export(tmp_path)
    ws = tmp_path / "pmlc-runs" / "2024-03-20-pump-fire" / "workspace"
    run("ingest", str(x), "--out", str(ws))
    assert (ws / "findings.json").exists() and (ws.parent / "outputs").is_dir()


@pytest.mark.parametrize("args, msg", [
    ([], "--slug"),
    (["--slug", "Pump Fire"], "lowercase"),
    (["--slug", "pump-fire", "--date", "20/03/2024"], "--date"),
    (["--out", "somewhere/ws"], "pmlc-runs"),
])
def test_ingest_refusals(tmp_path, args, msg):
    x = make_export(tmp_path)
    r = run_fail("ingest", str(x), *args)
    assert msg in r.stderr
    assert not (tmp_path / "pmlc-runs").exists()


def test_default_output_paths(tmp_path):
    x = make_export(tmp_path)
    run("ingest", str(x), "--slug", "pump-fire")
    run_dir = tmp_path / "pmlc-runs" / "2024-03-20-pump-fire"
    ws = run_dir / "workspace"
    run("classify", str(ws))
    run("dashboard", str(ws))
    assert (run_dir / "outputs" / "organisation-learning-map.html").exists()
    reg = run_dir / "outputs" / "learnings-register.json"
    reg.write_text(json.dumps(MINI_REG))
    run("register-xlsx", str(reg))
    assert (run_dir / "outputs" / "learnings-register.xlsx").exists()
    out = run("run-path", str(ws))
    assert "slug:      2024-03-20-pump-fire" in out
    assert f"outputs:   {(run_dir / 'outputs').resolve()}" in out
    assert "legacy layout" in run("run-path", str(tmp_path))


def test_downgraded_findings():
    w = {"incident": {"review_votes": [{"finding": "FND-1", "vote": "Downgrade"},
                                       {"finding": "FND-2", "vote": "upgrade"},
                                       {"finding": "FND-3", "vote": "remove"}]}}
    assert pmlc.downgraded_findings(w) == {"FND-1", "FND-3"}
    assert pmlc.downgraded_findings({}) == set()


def test_detail_data(ws):
    w = pmlc.load_ws(ws)
    d = pmlc.build_detail_data(w, MINI_REG)
    assert d["register"] == {"title": "Synthetic register", "source": "export.xlsx"}
    assert d["entries"][0]["custom_field"] == "kept as given"
    assert d["entries"][0]["recommendations"][0]["owner_note"] == "extra key"
    assert set(d["findings"]) == {"FND-0001", "FND-0002", "FND-0003", "FND-0004"}
    assert d["findings"]["FND-0001"]["clusters"]["F"]
    assert d["findings"]["FND-0001"]["votes"] == []
    assert set(d["evidence"]) == {"EVD-0001", "EVD-0002", "EVD-0003"}
    assert {u["id"]: u["reason"] for u in d["uncovered"]} == {
        "FND-0003": "not covered by any learning", "FND-0004": "not covered by any learning"}
    assert d["derived"]["LR-01"] == {"level": "L3 Supervision", "entity": "F02",
                                     "hoc": ["engineering", "administrative"], "strongest": "engineering"}
    assert d["totals"] == {"findings": 4, "evidence": 3}
    assert d["tier_labels"]["T5"] == "Corporate / board"
    assert d["labels"]["worst_case"] == "Credible worst case"
    w2 = dict(w, incident=dict(w["incident"], review_votes=[{"finding": "FND-0004", "vote": "downgrade", "reason": "r"}]))
    d2 = pmlc.build_detail_data(w2, MINI_REG)
    assert {u["id"]: u["reason"] for u in d2["uncovered"]}["FND-0004"] == "reviewer downgrade vote"
    assert d2["findings"]["FND-0004"]["votes"][0]["vote"] == "downgrade"


def test_detail_data_minimal_workspace(ws):
    w = {"findings": pmlc.load_ws(ws)["findings"]}
    d = pmlc.build_detail_data(w, [{"lr_id": "LR-09", "finding_ids": ["FND-0001", "FND-9999"]}])
    assert d["register"] == {} and d["evidence"] == {}
    assert "FND-9999" not in d["findings"] and d["findings"]["FND-0001"]["clusters"] == {}
    assert d["derived"]["LR-09"]["strongest"] == "unclassified"


def test_embed_json_is_script_safe():
    s = pmlc.embed_json({"t": "</script><b>&"})
    assert "<" not in s and ">" not in s and "&" not in s
    assert json.loads(s) == {"t": "</script><b>&"}


def test_sample_folders(tmp_path):
    if shutil.which("git") is None:
        pytest.skip("git is not available")
    inside = subprocess.run(["git", "rev-parse", "--is-inside-work-tree"], cwd=ROOT, capture_output=True, text=True)
    if inside.returncode != 0:
        pytest.skip("repository root is not a git work tree")
    tracked = subprocess.run(["git", "ls-files", "samples"], cwd=ROOT, capture_output=True, text=True).stdout.split()
    tracked_dirs = {Path(x).parts[1] for x in tracked if len(Path(x).parts) > 2}
    assert [p for p in pmlc.sample_folder_problems(ROOT / "samples") if p.split(":")[0] in tracked_dirs] == []
    (tmp_path / "Hazira_2020").mkdir()
    good = tmp_path / "2020-09-24-hazira-gas-release-fire"
    good.mkdir()
    probs = pmlc.sample_folder_problems(tmp_path)
    assert any("Hazira_2020" in p and "folder name" in p for p in probs)
    assert any(good.name in p and "README.md" in p for p in probs)
    (good / "README.md").write_text("x")
    (good / f"{good.name}_probematrix-findings.xlsx").write_bytes(b"x")
    assert [p for p in pmlc.sample_folder_problems(tmp_path) if good.name in p] == []
    r = subprocess.run([sys.executable, str(PMLC), "check-samples", str(tmp_path)], capture_output=True, text=True)
    assert r.returncode == 1 and "Hazira_2020" in r.stdout


def test_make_public_sample_slug(tmp_path):
    x = make_export(tmp_path)
    names = tmp_path / "names.txt"
    names.write_text("Nobody Here\n")
    tool = ROOT / "tools" / "make-public-sample.py"
    r = subprocess.run([sys.executable, str(tool), str(x), "--slug", "Bad_Slug", "--names", str(names)],
                       capture_output=True, text=True)
    assert r.returncode != 0 and "--slug" in r.stderr
    root = tmp_path / "samples"
    r = subprocess.run([sys.executable, str(tool), str(x), "--slug", "2024-03-20-pump-fire", "--names", str(names),
                        "--samples-root", str(root)], capture_output=True, text=True)
    assert r.returncode == 0, r.stderr + r.stdout
    assert (root / "2024-03-20-pump-fire" / "2024-03-20-pump-fire_probematrix-findings.xlsx").exists()


def test_ingest_review_fixes(tmp_path):
    x = make_export(tmp_path)
    ws = tmp_path / "pmlc-runs" / "2024-03-20-pump-fire" / "workspace"
    run("ingest", str(x), "--out", str(ws))
    assert "already holds a workspace" in run_fail("ingest", str(x), "--out", str(ws)).stderr
    assert "not allowed with" in run_fail("ingest", str(x), "--slug", "a", "--legacy-out", str(tmp_path / "lw")).stderr
    r = subprocess.run([sys.executable, str(PMLC), "ingest", str(x), "--legacy-out", str(ws), "--force"],
                       capture_output=True, text=True)
    assert r.returncode == 0 and "Warning" not in r.stderr
    assert "No such folder" in run_fail("check-samples", str(tmp_path / "nope")).stderr


def test_ingest_excel_datetime_date(tmp_path):
    import datetime as _dt
    x = make_export(tmp_path)
    wb = openpyxl.load_workbook(x)
    for row in wb["Incident Summary"].iter_rows():
        if row[0].value == "Date":
            row[1].value = _dt.datetime(2024, 3, 20, 0, 0)
    wb.save(x)
    run("ingest", str(x), "--slug", "pump-fire")
    assert (tmp_path / "pmlc-runs" / "2024-03-20-pump-fire" / "workspace" / "findings.json").exists()


def test_check_register_flags_downgraded_citation(ws, tmp_path):
    import shutil as _sh
    w2 = tmp_path / "w2"
    _sh.copytree(ws, w2)
    inc = json.loads((w2 / "incident.json").read_text())
    inc["review_votes"] = [{"finding": "FND-0001", "vote": "downgrade", "reason": "r"}, {"finding": "FND-0002"}]
    (w2 / "incident.json").write_text(json.dumps(inc))
    reg = tmp_path / "r.json"
    reg.write_text(json.dumps(MINI_REG))
    r = subprocess.run([sys.executable, str(PMLC), "check-register", str(w2), str(reg)], capture_output=True, text=True)
    assert "cites FND-0001, which has a reviewer downgrade vote" in r.stdout
    assert "FND-0001 (r)" in run("profile", str(w2))


def test_make_public_sample_dst_and_missing(tmp_path):
    x = make_export(tmp_path)
    names = tmp_path / "names.txt"
    names.write_text("Nobody Here\n")
    tool = ROOT / "tools" / "make-public-sample.py"
    out = tmp_path / "plain.xlsx"
    r = subprocess.run([sys.executable, str(tool), str(x), str(out), "--names", str(names)], capture_output=True, text=True)
    assert r.returncode == 0 and out.exists(), r.stderr
    r = subprocess.run([sys.executable, str(tool), str(x), "--names", str(names)], capture_output=True, text=True)
    assert r.returncode != 0 and "--slug" in r.stderr


def test_check_samples_needs_no_scan_record(tmp_path):
    good = tmp_path / "2020-09-24-hazira-gas-release-fire"
    good.mkdir()
    (good / "README.md").write_text("x")
    (good / f"{good.name}_probematrix-findings.xlsx").write_bytes(b"x")
    r = subprocess.run([sys.executable, str(PMLC), "check-samples", str(tmp_path)], capture_output=True, text=True)
    assert r.returncode == 0, r.stdout + r.stderr


def test_make_public_sample_without_names(tmp_path):
    x = make_export(tmp_path)
    tool = ROOT / "tools" / "make-public-sample.py"
    root = tmp_path / "samples"
    r = subprocess.run([sys.executable, str(tool), str(x), "--slug", "2024-03-20-pump-fire",
                        "--samples-root", str(root)], capture_output=True, text=True)
    assert r.returncode == 0, r.stderr + r.stdout
    assert "name scan skipped" in r.stdout
    assert (root / "2024-03-20-pump-fire" / "2024-03-20-pump-fire_probematrix-findings.xlsx").exists()


def test_docs_no_zero_hit_rule():
    for rel in ("AGENTS.md", "samples/README.md", "skills/learnings-coach/SKILL.md"):
        assert "zero-hit" not in (ROOT / rel).read_text(encoding="utf-8").lower(), rel


def _data_block(t):
    m = re.search(r'<script type="application/json" id="pmlc-data">(.*?)</script>', t, re.S)
    assert m, "no pmlc-data block"
    return json.loads(m.group(1))


def test_dashboard_two_tabs(ws, tmp_path):
    reg = tmp_path / "r.json"
    reg.write_text(json.dumps(MINI_REG))
    out = tmp_path / "d.html"
    run("dashboard", str(ws), "--register", str(reg), "--out", str(out))
    t = out.read_text(encoding="utf-8")
    assert 'role="tablist"' in t and 'id="tab-overview"' in t and 'id="tab-detail"' in t
    assert 'href="#detail/LR-01"' in t
    assert "Learning plan by level" in t and "Who learns, at which level" in t
    d = _data_block(t)
    assert d["entries"][0]["recommendations"][1]["what"].startswith("y </script>")
    assert set(d["findings"]) == {"FND-0001", "FND-0002", "FND-0003", "FND-0004"}
    assert t.count("</script>") == 2
    assert not re.search(r'<(script|link)[^>]+(src|href)="https?:', t)
    assert "function init()" in t and ".tabbar{" in t
    assert "ProbeMatrix Learnings Coach" in t


def test_dashboard_without_register_has_no_tabs(ws, tmp_path):
    out = tmp_path / "d.html"
    run("dashboard", str(ws), "--out", str(out))
    t = out.read_text(encoding="utf-8")
    assert 'role="tablist"' not in t and "pmlc-data" not in t and "<script" not in t
    assert "--register" in t and "Who learns, at which level" in t


def test_citation_records(ws):
    w = pmlc.load_ws(ws)
    w = dict(w, incident=dict(w["incident"], review_votes=[{"finding": "FND-0003", "vote": "downgrade", "reason": "r"}]))
    f, e = pmlc.citation_records(w, ["FND-0003", "FND-9999", "FND-0003"], ["EVD-0002", "EVD-0404"])
    assert list(f) == ["FND-0003"]
    assert f["FND-0003"]["votes"][0]["vote"] == "downgrade" and f["FND-0003"]["clusters"]["F"]
    assert list(e) == ["EVD-0002", "EVD-0003", "EVD-0001"]
    assert pmlc.citation_records({"findings": []}, ["FND-0001"]) == ({}, {})


def test_map_uses_shared_tokens_and_wide_rule(ws, tmp_path):
    out = tmp_path / "d.html"
    run("dashboard", str(ws), "--out", str(out))
    t = out.read_text(encoding="utf-8")
    assert pmlc.PAGE_TOKENS_CSS in t
    assert pmlc.WIDE_CSS in t and "@media (min-width:1400px)" in t
    assert t.index("main{max-width:1180px") < t.index(pmlc.WIDE_CSS)


def test_version_single_source():
    assert re.fullmatch(r"\d+\.\d+\.\d+((a|rc)\d+)?", pmlc.__version__)
    assert f"v{pmlc.__version__}" in pmlc.BRAND
    assert run("--version").strip() == f"pmlc {pmlc.__version__}"


def test_build_parser_has_every_subcommand():
    sub = next(a for a in pmlc.build_parser()._actions if isinstance(a, argparse._SubParsersAction))
    assert {"ingest", "classify", "select", "check-register", "check-pack", "pack"} <= set(sub.choices)


def test_release_script_bumps_all_three(tmp_path):
    files = ["skills/learnings-coach/scripts/pmlc.py", ".claude-plugin/plugin.json", ".claude-plugin/marketplace.json"]
    for f in files:
        (tmp_path / f).parent.mkdir(parents=True, exist_ok=True)
        shutil.copy(ROOT / f, tmp_path / f)
    for pep, sem in (("2.0.0a1", "2.0.0-alpha.1"), ("2.0.0rc1", "2.0.0-rc.1"), ("2.0.0", "2.0.0")):  # ruling 25
        r = subprocess.run([sys.executable, str(ROOT / "tools" / "release.py"), pep, "--root", str(tmp_path)],
                           capture_output=True, text=True)
        assert r.returncode == 0, r.stderr + r.stdout
        assert f'__version__ = "{pep}"' in (tmp_path / files[0]).read_text(encoding="utf-8")
        for f in files[1:]:
            assert f'"version": "{sem}"' in (tmp_path / f).read_text(encoding="utf-8"), f
        assert f"git tag -a v{sem}" in r.stdout
    before = {f: (tmp_path / f).read_bytes() for f in files}
    bad = subprocess.run([sys.executable, str(ROOT / "tools" / "release.py"), "banana", "--root", str(tmp_path)],
                         capture_output=True, text=True)
    assert bad.returncode != 0
    assert {f: (tmp_path / f).read_bytes() for f in files} == before


# ----------------------------------------------------------------------------- taxonomy v2.0: tags, F14 split, overlays

@pytest.mark.parametrize("role, tier", [("Plant General Manager", "T4"), ("General Manager (Operations)", "T4"),
                                        ("Deputy General Manager", "T3"), ("Assistant General Manager", "T3")])
def test_general_manager_is_plant_leadership(role, tier):                            # final review 10
    assert pmlc.map_role(role, pmlc._load_taxonomy())["T"] == tier


def test_sign_no_longer_matches_design():
    tax = pmlc._load_taxonomy()
    assert pmlc.map_role("F&G System Designers", tax)["R"] != "approver"
    assert pmlc.map_role("Permit signatory", tax)["R"] == "approver"
    assert pmlc.map_role("Supervisors who signed the permit", tax)["R"] == "approver"
    assert pmlc.map_role("Supervisor (calibration sign-off)", tax)["R"] == "approver"
    assert pmlc.map_role("Line Supervisors who assigned them", tax)["R"] != "approver"


def test_rank_beats_policy():
    tax = pmlc._load_taxonomy()
    assert pmlc.map_role("Plant Management (policy owner)", tax)["T"] == "T4"
    assert pmlc.map_role("Corporate policy owner", tax)["T"] == "T5"
    assert pmlc.map_role("HR/Training Supervisor", tax)["T"] == "T2"
    assert pmlc.map_role("Policy owner", tax)["T"] == "T3"          # 'owner' is a T3 match
    assert pmlc.map_role("Policy", tax)["T"] == "T5"                # weak: only when no tier matches
    assert pmlc.map_role("Plant management responsible for MOC governance", tax)["T"] == "T4"
    assert pmlc.map_role("Governance", tax)["T"] == "T5"            # governance is weak T5 (ruling R6)


def test_f07_document_token_is_narrow():
    tax = pmlc._load_taxonomy()
    assert "F07" not in pmlc.map_role("Supervisors who signed incomplete calibration documents", tax)["F"]
    assert "F07" not in pmlc.map_role("Maintenance documentation supervision", tax)["F"]
    for r in ("Document control", "Engineering Document Control", "Engineering documentation control",
              "Document control supervisors", "Engineering/Document Control Team"):
        assert "F07" in pmlc.map_role(r, tax)["F"], r


def test_f14_split():
    tax = pmlc._load_taxonomy()
    assert pmlc.map_role("Corporate Director", tax)["F"] == ["F14"]
    assert pmlc.map_role("Statutory inspectorate", tax)["F"] == ["F15"]
    assert pmlc.map_role("OEM / seal manufacturer", tax)["F"] == ["F16"]
    assert pmlc.map_role("Contractor crew", tax)["F"][-1] == "F17"
    tags = {e["id"]: e["tags"] for e in tax["entities"]}
    assert tags["F17"] == ["contractor", "supply"] and tags["F16"] == ["supply"] and tags["F15"] == ["external"]
    assert all(tags[e] for e in tags)


def test_site_tokens_need_the_overlay():
    assert "F10" not in pmlc.map_role("CISF constable", pmlc._load_taxonomy())["F"]
    assert "F10" in pmlc.map_role("CISF constable", pmlc._load_taxonomy(overlays=["india-process-site"]))["F"]
    ov = pmlc._load_taxonomy(overlays=["india-process-site"])
    assert pmlc.map_role("CISF constable", ov)["T"] == "T1"
    assert pmlc.map_role("Resident Engineer", ov)["T"] == "T4"
    assert pmlc.map_role("SIC", ov)["F"] == ["F09b"]
    assert pmlc.map_role("SIC", pmlc._load_taxonomy())["F"] == []


def test_overlay_errors(tmp_path):
    with pytest.raises(SystemExit, match="not found"):
        pmlc._load_taxonomy(overlays=["no-such-overlay"])
    bad = tmp_path / "bad.json"
    bad.write_text(json.dumps({"name": "bad", "extend": {"entities": {"F77": "x"}}}), encoding="utf-8")
    with pytest.raises(SystemExit, match="F77"):
        pmlc._load_taxonomy(overlays=[str(bad)])


def test_classify_records_taxonomy_and_warns_on_change(tmp_path):
    x = make_export(tmp_path)
    w = tmp_path / "ws"
    run("ingest", str(x), "--legacy-out", str(w))
    out = run("classify", str(w))
    assert "warning: previous classify" not in out
    cl = json.loads((w / "clusters.json").read_text())
    assert cl["taxonomy"]["file"] == "taxonomy.json" and cl["taxonomy"]["overlays"] == []
    assert cl["taxonomy"]["tags"]["F09b"] == ["emergency"]
    out = run("classify", str(w), "--overlay", "india-process-site")
    assert "warning: previous classify used taxonomy.json; this run uses taxonomy.json + india-process-site" in out
    cl = json.loads((w / "clusters.json").read_text())
    assert cl["taxonomy"]["overlays"] == ["india-process-site"]
    assert "warning: previous classify" not in run("classify", str(w), "--overlay", "india-process-site")


def _sources(w):
    return json.loads((w / "manifest.json").read_text())["files"]["clusters.json"]["sources"]["taxonomy"]


def test_classify_remembers_overlay_and_taxonomy(fresh_ws, tmp_path):              # ruling 24
    w = fresh_ws
    run("classify", str(w), "--overlay", "india-process-site")
    out = run("classify", str(w))
    assert ("Reusing taxonomy taxonomy.json and overlays india-process-site from the previous classify (pass "
            "--taxonomy/--overlay to change, --no-overlay for none).") in out
    assert "warning: previous classify" not in out
    cl = json.loads((w / "clusters.json").read_text())
    assert cl["taxonomy"]["overlays"] == ["india-process-site"]
    assert len(_sources(w)) == 2 and _sources(w)[1].endswith("india-process-site.json")
    assert "current" in _status_row(run("status", str(w)), "clusters.json")
    out = run("classify", str(w), "--no-overlay")
    assert "warning: previous classify used taxonomy.json + india-process-site; this run uses taxonomy.json" in out
    assert json.loads((w / "clusters.json").read_text())["taxonomy"]["overlays"] == []
    assert len(_sources(w)) == 1
    assert "Reusing" not in run("classify", str(w))
    assert json.loads((w / "clusters.json").read_text())["taxonomy"]["overlays"] == []
    assert "cannot be combined" in run_fail("classify", str(w), "--no-overlay", "--overlay", "x").stderr
    tax = tmp_path / "tax.json"
    tax.write_bytes((pmlc.REFS / "taxonomy.json").read_bytes())
    run("classify", str(w), "--taxonomy", str(tax), "--overlay", "india-process-site")
    out = run("classify", str(w))
    assert f"Reusing taxonomy {tax} and overlays india-process-site" in out and "warning:" not in out
    assert _sources(w)[0] == str(tax.resolve()) and len(_sources(w)) == 2
    out = run("classify", str(w), "--no-overlay")
    assert f"Reusing taxonomy {tax} and overlays none" in out
    cl = json.loads((w / "clusters.json").read_text())
    assert cl["taxonomy"]["file"] == str(tax) and cl["taxonomy"]["overlays"] == []
    tax.unlink()
    assert "no longer there" in run_fail("classify", str(w)).stderr
    run("classify", str(w), "--taxonomy", str(pmlc.REFS / "taxonomy.json"))


def test_select_by_tag_and_unknown_tag(ws):
    ids = {d["id"] for d in json.loads(run("select", str(ws), "--tag", "emergency", "--format", "json"))}
    assert "FND-0004" in ids
    both = {d["id"] for d in json.loads(run("select", str(ws), "--tag", "emergency", "--entity", "F02",
                                            "--format", "json"))}
    f02 = {d["id"] for d in json.loads(run("select", str(ws), "--entity", "F02", "--format", "json"))}
    assert both == ids | f02
    r = run_fail("select", str(ws), "--tag", "nonsense")
    assert "nonsense" in r.stderr and "Tags in use" in r.stderr


def test_custom_taxonomy_renamed_ids_still_select_by_tag(tmp_path, ws):
    tax = json.loads((pmlc.REFS / "taxonomy.json").read_text(encoding="utf-8"))
    for i, e in enumerate(tax["entities"]):
        e["id"] = f"X{i}"
    p = tmp_path / "renamed.json"
    p.write_text(json.dumps(tax), encoding="utf-8")
    w2 = tmp_path / "ws2"
    shutil.copytree(ws, w2)
    run("classify", str(w2), "--taxonomy", str(p))
    cl = json.loads((w2 / "clusters.json").read_text())
    assert "F09b" not in cl["entities"] and cl["taxonomy"]["file"] == str(p)
    sel = lambda d: {x["id"] for x in json.loads(run("select", str(d), "--tag", "emergency", "--format", "json"))}  # noqa: E731
    assert sel(w2) == sel(ws)


def test_old_clusters_without_tags(ws, tmp_path):
    w2 = tmp_path / "old"
    shutil.copytree(ws, w2)
    cl = json.loads((w2 / "clusters.json").read_text())
    del cl["taxonomy"]["tags"]
    (w2 / "clusters.json").write_text(json.dumps(cl), encoding="utf-8")
    r = run_fail("select", str(w2), "--tag", "emergency")
    assert "defines no tags: add `tags` to its entities (see references/taxonomy.json)" in r.stderr
    assert "re-run classify if clusters.json predates v2.0" in r.stderr


def test_entities_for_tags():
    cl = {"taxonomy": {"file": "t.json", "tags": {"A": ["x"], "B": ["y", "x"], "C": ["z"]}}}
    assert pmlc.entities_for_tags(cl, ["x"]) == ["A", "B"]
    assert pmlc.entities_for_tags(cl, ["z", "y"]) == ["B", "C"]
    with pytest.raises(SystemExit, match=r"No entity has tag\(s\) \['q'\] in taxonomy t.json"):
        pmlc.entities_for_tags(cl, ["x", "q"])
    with pytest.raises(SystemExit, match=r"^Taxonomy t.json defines no tags"):
        pmlc.entities_for_tags({"taxonomy": {"file": "t.json", "tags": {"A": []}}}, ["x"])


def _ids(*args):
    return {f["id"] for f in json.loads(run(*args, "--format", "json"))}


def test_top_rated_and_run_path(ws):
    assert _ids("select", str(ws), "--top-rated") == {"FND-0002", "FND-0003"}
    out = run("run-path", str(ws))
    assert "top_rating: 4" in out and "incident_date: 2024-03-20" in out


def test_ingest_records_incident_date(tmp_path):
    x = make_export(tmp_path)
    run("ingest", str(x), "--legacy-out", str(tmp_path / "w1"), "--date", "2024-03-21")
    assert json.loads((tmp_path / "w1" / "incident.json").read_text())["date"] == "2024-03-21"
    run("ingest", str(x), "--legacy-out", str(tmp_path / "w2"))
    assert json.loads((tmp_path / "w2" / "incident.json").read_text())["date"] == "2024-03-20"


def test_subcategory_prefix_exclude_downgraded(tmp_path):
    x = make_export(tmp_path)
    wb = openpyxl.load_workbook(x)
    c = wb.create_sheet("Collaboration Data")
    for r in [("VOTES",), ("Finding ID", "Vote", "Reason"), ("FND-0004", "downgrade", "weak evidence")]:
        c.append(r)
    wb.save(x)
    w = tmp_path / "w"
    run("ingest", str(x), "--legacy-out", str(w))
    run("classify", str(w))
    assert _ids("select", str(w)) == {"FND-0001", "FND-0002", "FND-0003", "FND-0004"}
    assert _ids("select", str(w), "--exclude-downgraded") == {"FND-0001", "FND-0002", "FND-0003"}
    md = run("select", str(w), "--exclude-downgraded", "--brief")
    assert md.rstrip().splitlines()[-1] == "Excluded (reviewer downgrade): FND-0004"
    r = subprocess.run([sys.executable, str(PMLC), "select", str(w), "--exclude-downgraded", "--format", "json"],
                       capture_output=True, text=True)
    assert "Excluded (reviewer downgrade): FND-0004" in r.stderr
    assert json.loads(r.stdout)
    assert _ids("select", str(w), "--subcategory", "errors") == {"FND-0001"}
    assert _ids("select", str(w), "--prefix", "SV", "OR") == {"FND-0002", "FND-0003"}


def test_entity_mode_all_and_tags_synthetic(ws):
    both = _ids("select", str(ws), "--entity", "F02", "F06", "--entity-mode", "all")
    anyof = _ids("select", str(ws), "--entity", "F02", "F06")
    assert both <= anyof


def test_timeline_before_incident(ws):
    run("timeline", str(ws), "--before", "incident")
    a = (ws / "timeline_candidates.json").read_text()
    run("timeline", str(ws), "--before", "2024-03-20")
    assert (ws / "timeline_candidates.json").read_text() == a


def test_timeline_before_incident_without_date(tmp_path):
    x = make_export(tmp_path)
    run("ingest", str(x), "--legacy-out", str(tmp_path / "w"))
    p = tmp_path / "w" / "incident.json"
    d = json.loads(p.read_text())
    d.pop("date")
    d["fields"].pop("Date")
    p.write_text(json.dumps(d))
    r = run_fail("timeline", str(tmp_path / "w"), "--before", "incident")
    assert "No incident date in this workspace; pass --date to ingest." in r.stderr


def test_select_args_rejects_unknown_filter():
    assert pmlc.select_args(tag=["x"]).tag == ["x"] and pmlc.select_args().entity_mode == "any"
    with pytest.raises(TypeError):
        pmlc.select_args(bogus=1)


def test_no_placeholders_left():
    for p in (ROOT / "skills").glob("*/SKILL.md"):
        t = p.read_text(encoding="utf-8")
        assert "<top" not in t and "<incident date" not in t, p


def _select_proc(*args):
    r = subprocess.run([sys.executable, str(PMLC), "select", *map(str, args)], capture_output=True, text=True)
    assert r.returncode == 0, r.stderr
    return r


def test_budget_too_small_keeps_none(ws):
    r = _select_proc(ws, "--budget-chars", "10")
    assert "# Selected findings (0)" in r.stdout and "### FND" not in r.stdout and "## FND" not in r.stdout
    assert "Dropped for budget: 4 findings (lowest kept rating none): FND-0003, FND-0002, FND-0001, FND-0004" in r.stdout
    m = re.search(r"below the minimum output \((\d+) chars\); no findings kept", r.stderr)
    assert m and "select: budget 10 is below" in r.stderr and len(r.stdout) == int(m.group(1))
    assert "Dropped for budget: 4 findings" in r.stderr
    j = _select_proc(ws, "--budget-chars", "10", "--format", "json")
    assert json.loads(j.stdout) == [] and "Dropped for budget: 4 findings" in j.stderr


def test_budget_keeps_rating_order_and_counts_footers(ws):
    full = _select_proc(ws, "--brief").stdout
    r = _select_proc(ws, "--brief", "--budget-chars", str(len(full) - 1))
    assert len(r.stdout) <= len(full) - 1
    assert "Dropped for budget: " in r.stdout and "Dropped for budget: " in r.stderr
    assert "below the minimum" not in r.stderr
    assert re.search(r"select: \d+ findings, %d chars" % len(r.stdout), r.stderr)
    ok = _select_proc(ws, "--brief", "--budget-chars", str(len(full)))
    assert ok.stdout == full and "Dropped" not in ok.stdout


def test_fields_group_count(ws):
    items = json.loads(run("select", str(ws), "--fields", "id,short", "--format", "json"))
    assert items and all(set(d) == {"id", "short"} for d in items)
    r = run_fail("select", str(ws), "--fields", "id,bogus")
    assert "label, short, justification" in r.stderr
    md = run("select", str(ws), "--fields", "short")
    assert "**Short:**" in md and "**Justification:**" not in md and "## FND-0002" in md
    cl = json.loads((ws / "clusters.json").read_text())["findings"]
    want = {}
    for c in cl.values():
        for e in c["F"]:
            want[e] = want.get(e, 0) + 1
    assert json.loads(run("select", str(ws), "--count-by", "entity", "--format", "json")) == want
    assert "| F02 |" in run("select", str(ws), "--count-by", "entity")
    md = run("select", str(ws), "--group-by", "level")
    assert md.index("\n## L3") < md.index("### FND-0002")
    g = json.loads(run("select", str(ws), "--group-by", "level", "--format", "json"))
    assert {d["id"] for d in g["L3"]} == {"FND-0002"}
    r = _select_proc(ws, "--count-by", "tier", "--budget-chars", "5")
    assert "ignores --budget-chars" in r.stderr


def test_size_report_on_stderr(ws):
    r = _select_proc(ws)
    m = re.search(r"select: (\d+) findings, (\d+) chars", r.stderr)
    assert m and int(m.group(2)) == len(r.stdout)


def _edit_export(src, dst, fn):
    wb = openpyxl.load_workbook(src)
    fn(wb)
    wb.save(dst)
    return dst


def test_ingest_refuses_capa_planner(tmp_path):
    def mk(wb):
        wb.create_sheet("Corrective Actions")
        wb.create_sheet("Preventive Actions")
    x = _edit_export(make_export(tmp_path), tmp_path / "capa.xlsx", mk)
    r = run_fail("ingest", str(x), "--slug", "pump-fire", "--root", str(tmp_path / "r"))
    assert 'This is a CAPA Planner export, not a findings export. Export "Identified Findings"' in r.stderr
    assert not (tmp_path / "r").exists()


def test_renamed_header_reported_and_strict(tmp_path):
    def mk(wb):
        s = wb["Identified Findings"]
        for c in s[1]:
            if c.value == "Rating":
                c.value = "Score"
    x = _edit_export(make_export(tmp_path), tmp_path / "renamed.xlsx", mk)
    out = run("ingest", str(x), "--slug", "pump-fire", "--root", str(tmp_path / "r"))
    assert "missing: Identified Findings: Rating" in out
    assert "unknown: Identified Findings: Score" in out
    assert re.search(r"schema: fingerprint [0-9a-f]{12};", out)
    r = run_fail("ingest", str(x), "--slug", "pump-fire", "--root", str(tmp_path / "r2"), "--strict")
    assert "Rating" in r.stderr
    inc = json.loads(next((tmp_path / "r").rglob("incident.json")).read_text())
    assert inc["schema"]["missing"]["Identified Findings"] == ["Rating"]
    assert inc["schema"]["unknown"] == {"Identified Findings": ["Score"]}


def test_strict_tolerates_optional_missing(tmp_path):
    def mk(wb):
        s = wb["Identified Findings"]
        for c in s[1]:
            if c.value == "Created At":
                c.value = "Created"
    x = _edit_export(make_export(tmp_path), tmp_path / "opt.xlsx", mk)
    run("ingest", str(x), "--slug", "pump-fire", "--root", str(tmp_path / "r"), "--strict")


def test_no_stats_sheet_note(tmp_path):
    out = run("ingest", str(make_export(tmp_path)), "--slug", "pump-fire", "--root", str(tmp_path / "r"))
    assert "cross-check skipped" in out
    assert "schema: fingerprint " in out
    inc = json.loads(next((tmp_path / "r").rglob("incident.json")).read_text())
    assert inc["analysis_stats"] is None


def test_stats_mismatch_warns(tmp_path):
    def mk(wb):
        s = wb.create_sheet("Analysis Statistics")
        for r in [("ANALYSIS STATISTICS", None), ("Level", "Level Name", "Total Findings"), (1, "Unsafe Acts", 5),
                  ("Total", None, 9)]:
            s.append(r)
    x = _edit_export(make_export(tmp_path), tmp_path / "st.xlsx", mk)
    out = run("ingest", str(x), "--slug", "pump-fire", "--root", str(tmp_path / "r"))
    assert "warning: Analysis Statistics says 9 findings, ingested 4" in out


# ----------------------------------------------------------------------------- manifest and status

PACKS = ROOT / "tests" / "fixtures" / "packs"

def _status_row(out, name):
    return next(ln for ln in out.splitlines() if ln.startswith(f"| {name} "))


def _changed_export(d: Path) -> Path:
    import openpyxl as ox
    x = make_export(d)
    wb = ox.load_workbook(x)
    wb["Identified Findings"].cell(row=2, column=12).value = "A different short description of the event."
    y = d / "changed.xlsx"
    wb.save(y)
    return y


def test_stale_after_force_ingest(tmp_path):
    x = make_export(tmp_path)
    w = tmp_path / "ws"
    run("ingest", str(x), "--legacy-out", str(w))
    run("classify", str(w))
    run("causal", str(w))
    assert "current" in _status_row(run("status", str(w)), "causal_graph.json")
    run("ingest", str(_changed_export(tmp_path)), "--legacy-out", str(w), "--force")
    out = run("status", str(w))
    assert "stale" in _status_row(out, "causal_graph.json")
    assert "stale" in _status_row(out, "clusters.json")
    assert "stale" in _status_row(out, "causal_graph.mmd")
    run("causal", str(w))
    assert "current" in _status_row(run("status", str(w)), "causal_graph.json")


def test_same_export_not_stale(tmp_path):
    x = make_export(tmp_path)
    w = tmp_path / "ws"
    run("ingest", str(x), "--legacy-out", str(w))
    run("classify", str(w))
    run("causal", str(w))
    run("interfaces", str(w))
    run("timeline", str(w))
    run("ingest", str(x), "--legacy-out", str(w), "--force")
    out = run("status", str(w))
    for name in ("findings.json", "clusters.json", "causal_graph.json", "interfaces.json", "timeline_candidates.json"):
        assert "current" in _status_row(out, name), out


def test_stale_propagates_and_warns(tmp_path):
    x = make_export(tmp_path)
    w = tmp_path / "ws"
    run("ingest", str(x), "--legacy-out", str(w))
    run("classify", str(w))
    run("interfaces", str(w))
    run("ingest", str(_changed_export(tmp_path)), "--legacy-out", str(w), "--force")
    assert "stale" in _status_row(run("status", str(w)), "interfaces.json")
    r = subprocess.run([sys.executable, str(PMLC), "dashboard", str(w)], capture_output=True, text=True)
    assert r.returncode == 0, r.stderr
    assert "warning: stale input clusters.json" in r.stderr and "re-run pmlc classify" in r.stderr
    pack = PACKS / "supervisor-coach.json"
    r = subprocess.run([sys.executable, str(PMLC), "pack", str(w), str(pack)], capture_output=True, text=True)
    assert "warning: stale input" in r.stderr


def test_pack_check_recorded(tmp_path):
    x = make_export(tmp_path)
    w = tmp_path / "ws"
    run("ingest", str(x), "--legacy-out", str(w))
    p = tmp_path / "mine.json"
    p.write_bytes((PACKS / "supervisor-coach.json").read_bytes())
    r = subprocess.run([sys.executable, str(PMLC), "check-pack", str(w), str(p)], capture_output=True, text=True)
    out = run("status", str(w))
    row = next(ln for ln in out.splitlines() if ln.startswith("| supervisor-coach "))
    assert row.rstrip().endswith("| no |"), (row, r.stdout, r.stderr)
    p.write_text(p.read_text(encoding="utf-8") + "\n", encoding="utf-8")
    row = next(ln for ln in run("status", str(w)).splitlines() if ln.startswith("| supervisor-coach "))
    assert row.rstrip().endswith("| yes |")
    p.unlink()
    row = next(ln for ln in run("status", str(w)).splitlines() if ln.startswith("| supervisor-coach "))
    assert row.rstrip().endswith("| missing |")


def test_status_on_pre_manifest_workspace(tmp_path):
    x = make_export(tmp_path)
    w = tmp_path / "ws"
    run("ingest", str(x), "--legacy-out", str(w))
    run("classify", str(w))
    (w / "manifest.json").unlink()
    out = run("status", str(w))
    assert "unknown (built before manifest)" in _status_row(out, "findings.json")
    assert "unknown (built before manifest)" in _status_row(out, "clusters.json")


def test_status_unrecorded_pack_and_changed_taxonomy(tmp_path):
    x = make_export(tmp_path)
    w = tmp_path / "ws"
    run("ingest", str(x), "--legacy-out", str(w))
    tax = tmp_path / "tax.json"
    tax.write_bytes((pmlc.REFS / "taxonomy.json").read_bytes())
    ov = tmp_path / "ov.json"
    ov.write_text("{}", encoding="utf-8")
    run("classify", str(w), "--taxonomy", str(tax), "--overrides", str(ov))
    assert "current" in _status_row(run("status", str(w)), "clusters.json")
    ov.write_text('{"x": {}}', encoding="utf-8")
    out = run("status", str(w))
    assert "stale" in _status_row(out, "clusters.json") and "overrides" in _status_row(out, "clusters.json")
    ov.unlink()
    assert "missing" in _status_row(run("status", str(w)), "clusters.json")
    (w / "outputs").mkdir(exist_ok=True)
    (w / "outputs" / "loose.json").write_text('{"pack": "loose"}', encoding="utf-8")
    assert "| loose | never | n/a |" in run("status", str(w))


def test_manifest_helpers(tmp_path):
    assert pmlc.stale_files(tmp_path) == []
    (tmp_path / "findings.json").write_text("[]", encoding="utf-8")
    (tmp_path / "x.json").write_text("1", encoding="utf-8")
    pmlc.record_derived(tmp_path, "x.json", {"findings.json": pmlc.file_sha256(tmp_path / "findings.json")})
    assert pmlc.stale_files(tmp_path) == []
    (tmp_path / "findings.json").write_text("[1]", encoding="utf-8")
    assert pmlc.stale_files(tmp_path)[0][0] == "x.json"


def test_new_overrides_file_marks_clusters_stale(tmp_path):
    x = make_export(tmp_path)
    w = tmp_path / "ws"
    run("ingest", str(x), "--legacy-out", str(w))
    run("classify", str(w))
    assert "current" in _status_row(run("status", str(w)), "clusters.json")
    (w / "role_overrides.json").write_text("{}", encoding="utf-8")
    assert "stale" in _status_row(run("status", str(w)), "clusters.json")
    run("classify", str(w))
    assert "current" in _status_row(run("status", str(w)), "clusters.json")
    (w / "role_overrides.json").unlink()
    run("classify", str(w), "--no-overrides")
    (w / "role_overrides.json").write_text("{}", encoding="utf-8")
    assert "current" in _status_row(run("status", str(w)), "clusters.json")


def test_months_before(ws):
    run("timeline", str(ws))
    t = json.loads((ws / "timeline_candidates.json").read_text())
    seal = [x for x in t if x["date"] == "2022-03-01"]
    assert seal and all(x["months_before"] == 24 for x in seal)
    assert all(isinstance(x["months_before"], int) for x in t if x["precision"] != "year")
    assert all(x["months_before"] is None and isinstance(x["years_before"], int) for x in t if x["precision"] == "year")


def test_months_between():
    d = pmlc.dt.date
    assert pmlc.months_between(d(2022, 3, 1), d(2024, 3, 20)) == 24
    assert pmlc.months_between(d(2022, 3, 21), d(2024, 3, 20)) == 23
    assert pmlc.months_between(d(2024, 5, 1), d(2024, 3, 20)) == -2


def test_timeline_entity_filter(ws):
    run("timeline", str(ws), "--entity", "F06")
    t = json.loads((ws / "timeline_candidates.json").read_text())
    f3 = {"FND-0003"} | set(next(f for f in pmlc.load_ws(ws)["findings"] if f["id"] == "FND-0003")["evidence"])
    assert t and all(set(x["sources"]) <= f3 for x in t)


def test_timeline_min_precision(ws):
    run("timeline", str(ws), "--min-precision", "month")
    t = json.loads((ws / "timeline_candidates.json").read_text())
    assert t and all(x["precision"] != "year" for x in t)
    run("timeline", str(ws), "--min-precision", "day")
    assert all(x["precision"] == "day" for x in json.loads((ws / "timeline_candidates.json").read_text()))


def test_clock_times_and_event_day(tmp_path):
    x = make_export(tmp_path)
    run("ingest", str(x), "--legacy-out", str(tmp_path / "w"))
    w = tmp_path / "w"
    fp = w / "findings.json"
    fs = json.loads(fp.read_text())
    fs[0]["short"] += " On 20 March 2024 14:05 hrs the seal was changed. At ~02:35 the alarm sounded. Pressure 2.5 bar."
    fs[1]["short"] += " Earlier on 5 January 2023 at 09:10 a check was missed. Then at 03:01:20 we left."
    fp.write_text(json.dumps(fs))
    out = run("timeline", str(w), "--on", "incident")
    t = json.loads((w / "timeline_candidates.json").read_text())
    assert any(x["date"] == "2024-03-20" and x.get("time") == "14:05" for x in t)
    seq = json.loads((w / "timeline_event_day.json").read_text())
    times = [x["time"] for x in seq]
    assert times == sorted(times) and "02:35" in times and "03:01" in times and "14:05" in times
    assert "09:10" not in times and "02:30" not in times and "2:50" not in times
    assert "| 02:35 |" in out
    assert "timeline_event_day.json" in json.loads((w / "manifest.json").read_text())["files"]
    run("timeline", str(w), "--on", "2023-01-05")
    other = [x["time"] for x in json.loads((w / "timeline_event_day.json").read_text())]
    assert other == ["09:10"]                    # undated 02:35 / 03:01 belong to the incident day only (ruling 26)


def test_on_bad_day(ws):
    assert "YYYY-MM-DD" in run_fail("timeline", str(ws), "--on", "yesterday").stderr


def test_on_incident_without_date(tmp_path):
    x = make_export(tmp_path)
    run("ingest", str(x), "--legacy-out", str(tmp_path / "w"))
    p = tmp_path / "w" / "incident.json"
    d = json.loads(p.read_text())
    d.pop("date")
    d["fields"].pop("Date")
    p.write_text(json.dumps(d))
    r = run_fail("timeline", str(tmp_path / "w"), "--on", "incident")
    assert "No incident date in this workspace; pass --date to ingest." in r.stderr


def test_entity_needs_clusters(tmp_path):
    x = make_export(tmp_path)
    run("ingest", str(x), "--legacy-out", str(tmp_path / "w"))
    assert "classify" in run_fail("timeline", str(tmp_path / "w"), "--entity", "F06").stderr


def test_vocab_drift_is_a_warning(ws, tmp_path):
    reg = json.loads(json.dumps(MINI_REG))
    reg["entries"][0]["confidence"] = "very high"
    p = tmp_path / "reg.json"
    p.write_text(json.dumps(reg))
    r = subprocess.run([sys.executable, str(PMLC), "check-register", str(ws), str(p)], capture_output=True, text=True)
    assert "warning: LR-01: confidence 'very high' is not one of" in r.stdout
    base = tmp_path / "base.json"
    base.write_text(json.dumps(MINI_REG))
    r0 = subprocess.run([sys.executable, str(PMLC), "check-register", str(ws), str(base)],
                        capture_output=True, text=True)
    assert r.returncode == r0.returncode
    assert "; 0 problem(s)." in r.stdout or r0.returncode != 0
    assert "warning: LR-01: confidence" not in r0.stdout


def test_register_warnings_rules():
    w = {"evidence": {"EVD-0001": {}}, "findings": [{"subcategory": "Organizational Process"}]}
    e = {"lr_id": "LR-09", "evidence_strength": "Direct", "causal_role": ["initiating", "bogus"],
         "transferability": "", "performance_mode": "rule-based; violation", "comparison": "maybe",
         "tier_messages": {"T1": "x", "T9": "y"},
         "hfacs_level": "L4 Organisational Process: procedures (x)",
         "recommendations": [{"evidence_basis": "EVD-0001, EVD-9999"}]}
    out = pmlc.register_warnings([e], w)
    text = "\n".join(out)
    assert "causal_role 'bogus'" in text and "comparison 'maybe'" in text
    assert "T9" in text and "EVD-9999" in text
    assert "evidence_strength" not in text and "transferability" not in text and "performance_mode" not in text
    assert "hfacs_level" not in text and "EVD-0001" not in text
    e2 = {"lr_id": "LR-10", "performance_mode": "guessing", "hfacs_level": "L4 Something Else (x)"}
    t2 = "\n".join(pmlc.register_warnings([e2], w))
    assert "performance_mode 'guessing'" in t2 and "hfacs_level" in t2
    assert pmlc.register_warnings([{"lr_id": "LR-11"}], w) == []


SKILL_FILES = sorted((ROOT / "skills").glob("*/SKILL.md"))


def test_sub_skills_state_dependency():
    subs = [p for p in SKILL_FILES if p.parent.name != "learnings-coach"]
    assert len(subs) == 22
    for p in subs:
        engine = [ln for ln in p.read_text(encoding="utf-8").splitlines() if ln.startswith("Engine:")]
        assert engine, p
        assert "installed alongside" in engine[0], p


def test_no_repo_only_paths_in_skills():
    pat = re.compile(r"\b(samples|tools|docs|tests|evals)/")
    for p in SKILL_FILES:
        for ln in p.read_text(encoding="utf-8").splitlines():
            if pat.search(ln):
                assert "(repository only)" in ln, "%s: %s" % (p, ln)


def test_installer_checks_python(tmp_path):
    bash = shutil.which("bash")
    if not bash:
        pytest.skip("bash not available")
    fake = tmp_path / "bin"
    fake.mkdir()
    py = fake / "python3"
    py.write_text("#!/bin/sh\nexit 1\n")
    py.chmod(0o755)
    dest = tmp_path / "dest"
    env = {"PATH": "%s:/usr/bin:/bin" % fake, "CODEX_SKILLS_DIR": str(dest), "HOME": str(tmp_path)}
    r = subprocess.run([bash, str(ROOT / "tools" / "install-codex.sh")], env=env,
                       capture_output=True, text=True)
    assert r.returncode == 1
    assert "3.9" in r.stdout + r.stderr
    assert not dest.exists()


def _variant(tmp_path, edit):
    """Build the synthetic export, let edit(workbook) change it, ingest it, and return the workspace."""
    x = make_export(tmp_path)
    wb = openpyxl.load_workbook(x)
    edit(wb)
    wb.save(x)
    w = tmp_path / "w"
    run("ingest", str(x), "--legacy-out", str(w))
    return w


def _incident(w):
    return json.loads((w / "incident.json").read_text())


LONG_NARR = "A seal failed and gas escaped at the pump house on a night shift. " * 5  # over 200 characters


def test_narrative_capture_both_layouts(tmp_path):
    assert len(LONG_NARR) > 200

    def same_row(wb):
        wb["Incident Summary"].append(("INCIDENT NARRATIVE", LONG_NARR))

    def next_row(wb):
        s = wb["Incident Summary"]
        s.append(("INCIDENT NARRATIVE",))
        s.append((LONG_NARR,))

    def short_ignored(wb):
        s = wb["Incident Summary"]
        s.append(("INCIDENT NARRATIVE",))
        s.append(("too short to be the narrative",))

    for n, edit in (("a", same_row), ("b", next_row)):
        d = tmp_path / n
        d.mkdir()
        assert _incident(_variant(d, edit))["narrative"] == LONG_NARR.strip()
    d = tmp_path / "c"
    d.mkdir()
    assert _incident(_variant(d, short_ignored))["narrative"] == ""


def test_votes_ingest(tmp_path):
    def edit(wb):
        c = wb.create_sheet("Collaboration Data")
        for r in [("COMMENTS",), ("Finding ID", "Author", "Comment"), ("FND-0001", "Reviewer A", "a comment row"),
                  ("VOTES",), ("Finding ID", "Vote", "Reason"), ("FND-0002", "upgrade", "strong evidence"),
                  ("FND-0004", "downgrade", "weak evidence"), ("COMMENTS",), ("FND-0003", "Reviewer B", "late")]:
            c.append(r)

    votes = _incident(_variant(tmp_path, edit))["review_votes"]
    assert votes == [{"finding": "FND-0002", "vote": "upgrade", "reason": "strong evidence"},
                     {"finding": "FND-0004", "vote": "downgrade", "reason": "weak evidence"}]


def test_profile_flags(tmp_path):
    def edit(wb):
        s = wb["Incident Summary"]
        s.append(("Incident Type", "Explosion"))
        s["B6"] = "in_progress"
        s.append(("INCIDENT NARRATIVE", LONG_NARR + "It happened on 05/04/2023."))
        c = wb.create_sheet("Collaboration Data")
        for r in [("VOTES",), ("Finding ID", "Vote", "Reason"), ("FND-0004", "downgrade", "weak evidence")]:
            c.append(r)

    w = _variant(tmp_path, edit)
    out = run("profile", str(w))
    flags = out.split("## Data-quality flags")[1]
    assert "Summary-sheet Date 2024-03-20 does not appear in the narrative, which mentions 2023-04-05" in flags
    assert "Incident Type 'Explosion' is not mentioned in the narrative" in flags
    assert "Investigation Status is 'in_progress'; label outputs as draft learnings." in flags
    assert "Reviewer downgrade votes on: FND-0004 (weak evidence)." in flags


def test_profile_no_status_flag_when_completed(ws):
    flags = run("profile", str(ws)).split("## Data-quality flags")[1]
    assert "in_progress" not in flags and "downgrade votes" not in flags


def test_causal_sheet_merge(tmp_path):
    def edit(wb):
        c = wb["Causal Chain"]
        c.delete_rows(2)
        c.append(["CL-1", "PE 2.1", 2, "AB 1.1"])
        c.append(["CL-2", "PE 2.1", 2, None])  # missing To Code: not kept at ingest

    w = _variant(tmp_path, edit)
    assert _incident(w)["causal_chain_sheet"] == [
        {"Link ID": "CL-1", "From Code": "PE 2.1", "From Level": "2", "To Code": "AB 1.1"}]
    run("causal", str(w))
    g = json.loads((w / "causal_graph.json").read_text())
    sheet = [e for e in g["edges"] if "causal-chain-sheet" in e["findings"]]
    assert [(e["from"], e["to"], e["from_level"], e["to_level"]) for e in sheet] == [("PE 2.1", "AB 1.1", 2, 1)]


def test_select_full_md_out(ws, tmp_path):
    out = tmp_path / "sel.md"
    run("select", str(ws), "--entity", "F02", "--out", str(out))
    md = out.read_text()
    assert "**Justification:** Work order shows no test [EVD-0001]." in md
    assert "**Causal explanation:** Enabled by SV 1.4 and OR 2.2." in md
    assert '- `EVD-0001` (document): Shows no test Quote: "no leak test recorded"' in md
    brief = tmp_path / "brief.md"
    run("select", str(ws), "--entity", "F02", "--brief", "--out", str(brief))
    assert "Quote:" not in brief.read_text() and "**Justification:**" not in brief.read_text()


# ----------------------------------------------------------------------------- select presets

def _skill_calls():
    out = {}
    for md in sorted((ROOT / "skills").glob("*/SKILL.md")):
        calls = pmlc.skill_select_lines(md.read_text(encoding="utf-8"))
        if calls:
            out[md.parent.name] = calls
    return out


def test_presets_match_skill_lines():
    presets = json.loads((ROOT / "skills/learnings-coach/references/select-presets.json").read_text(encoding="utf-8"))
    assert {k: [p["args"] for p in v] for k, v in presets.items()} == _skill_calls()
    for skill, entries in presets.items():
        names = [p["name"] for p in entries]
        assert len(names) == len(set(names)) and all(re.fullmatch(r"[a-z0-9-]+", n) for n in names), skill


def test_skill_select_lines_parsing():
    t = ('`PMLC select <ws> --code "<code>"`\n'
         'PMLC select <ws> --text "a|b c" --brief   # why\n'
         'PMLC select <ws> --top-rated --format json --out <ws>/top.json\n'
         'see (`PMLC select <ws> --level 4`) now\n'
         'PMLC select <ws> --entity <F-id> --brief\n')
    assert pmlc.skill_select_lines(t) == [["--text", "a|b c", "--brief"], ["--top-rated", "--format", "json"],
                                          ["--level", "4"]]


def test_preset_runs(ws):
    out = run("select", str(ws), "--preset", "supervisor-coach")
    n = len(json.loads((ROOT / "skills/learnings-coach/references/select-presets.json").read_text())["supervisor-coach"])
    assert out.count("# Preset supervisor-coach:") == n and n >= 2
    one = run("select", str(ws), "--preset", "supervisor-coach:" + json.loads(
        (ROOT / "skills/learnings-coach/references/select-presets.json").read_text())["supervisor-coach"][0]["name"],
        "--brief")
    assert one.count("# Preset") == 1 and "# Selected findings" in one


def test_preset_out_and_unknown(ws, tmp_path):
    f = tmp_path / "p.md"
    run("select", str(ws), "--preset", "frontline-learning-cards", "--out", str(f), "--format", "json")
    assert f.read_text().count("# Preset frontline-learning-cards:") == 3
    assert "Unknown preset skill" in run_fail("select", str(ws), "--preset", "nope").stderr
    assert "Unknown preset" in run_fail("select", str(ws), "--preset", "manager-coach:nope").stderr


@pytest.mark.parametrize("flag", [["--tier", "T1"], ["--top-rated"], ["--entity-mode", "any"], ["--max", "3"],
                                  ["--exclude-downgraded"], ["--text", "x"], ["--level", "1"]])
def test_preset_with_filter_refused(ws, flag):
    r = run_fail("select", str(ws), "--preset", "supervisor-coach", *flag)
    assert "--preset cannot be combined with filters; edit the preset or drop --preset." in r.stderr


def _preset_bundles(text):
    parts = re.split(r"^# Preset [^\n]*\n", text, flags=re.M)[1:]
    return [json.loads(p) for p in parts]


def test_preset_keeps_own_format(ws):
    assert len(_preset_bundles(run("select", str(ws), "--preset", "investigator-coach:top-json"))) == 1
    md = run("select", str(ws), "--preset", "investigator-coach:top-json", "--format", "md")
    assert "# Selected findings" in md


def test_preset_json_bundles_split_on_headings(ws):
    out = run("select", str(ws), "--preset", "supervisor-coach", "--format", "json")
    assert len(_preset_bundles(out)) == 3


# ----------------------------------------------------------------------------- docs drift guard

def _subcommands() -> list[str]:
    parser = pmlc.build_parser()
    sub = next(a for a in parser._actions if isinstance(a, argparse._SubParsersAction))
    return list(sub.choices)


def test_every_subcommand_documented():
    doc = pmlc.__doc__.split("Subcommands", 1)[1].split("Only dependency", 1)[0]
    listed = re.findall(r"^  ([a-z][a-z-]*)\s", doc, flags=re.M)
    docs = "\n".join((ROOT / n).read_text(encoding="utf-8") for n in ("AGENTS.md", "README.md"))
    names = _subcommands()
    assert sorted(listed) == sorted(names), "docstring Subcommands block differs from the parser"
    missing = [n for n in names if not re.search(r"(?<![\w-])" + re.escape(n) + r"(?![\w-])", docs)]
    assert not missing, f"not in AGENTS.md or README.md: {missing}"


_NUMERIC_FLAGS = {"--level", "--top", "--max", "--min-rating", "--per-date", "--show", "--budget-chars"}


def _skill_argv(line: str) -> list[str]:
    import shlex
    line = re.sub(r"\s+#.*$", "", line.strip())
    line = re.sub(r"\[(--[^\]]*)\]", r"\1", line)  # [--opt value] marks an optional part
    flags = "|".join(re.escape(f) for f in sorted(_NUMERIC_FLAGS))
    line = re.sub(rf"({flags})(\s+)<[^>]*>", r"\1\g<2>1", line)
    line = re.sub(r"(--format\s+)<[^>]*>", r"\1md", line)
    return shlex.split(re.sub(r"<[^>]*>", "X", line))[1:]


def _skill_command_lines():
    """(file, line number, command) for each `PMLC ...` line in a fenced block; `\\` continuations are joined."""
    for md in sorted((ROOT / "skills").glob("*/SKILL.md")):
        fenced, cur = False, None
        for n, line in enumerate(md.read_text(encoding="utf-8").splitlines(), 1):
            if line.lstrip().startswith("```"):
                fenced = not fenced
            elif cur:
                cur = (cur[0], cur[1] + " " + line.strip())
            elif fenced and line.lstrip().startswith("PMLC "):
                cur = (n, line.strip())
            if cur and not cur[1].endswith("\\"):
                yield md, cur[0], cur[1]
                cur = None
            elif cur:
                cur = (cur[0], cur[1][:-1].rstrip())


def test_skill_commands_parse(capsys):
    seen = 0
    for md, n, line in _skill_command_lines():
        seen += 1
        try:
            pmlc.build_parser().parse_args(_skill_argv(line))
        except ValueError as e:
            pytest.fail(f"{md.relative_to(ROOT)}:{n}: {line.strip()}\n  {e}")
        except SystemExit:
            err = capsys.readouterr().err.strip().splitlines()
            pytest.fail(f"{md.relative_to(ROOT)}:{n}: {line.strip()}\n  {err[-1] if err else 'parse error'}")
    assert seen > 30


def test_skill_argv_placeholders():
    assert _skill_argv("PMLC select <ws> --level <n> --format <f>  # c") == ["select", "X", "--level", "1", "--format", "md"]


# ----------------------------------------------------------------------------- UAT inputs (tools/make-uat-inputs.py)

def _sheets_and_header(x):
    wb = openpyxl.load_workbook(x, read_only=True)
    head = [c.value for c in next(wb["Identified Findings"].iter_rows(max_row=1))]
    return wb.sheetnames, head


@pytest.mark.parametrize("where", ["fresh", "committed"])
def test_make_uat_inputs(tmp_path, where):
    if where == "fresh":
        r = subprocess.run([sys.executable, str(ROOT / "tools" / "make-uat-inputs.py"), "--out", str(tmp_path)],
                           capture_output=True, text=True, cwd=ROOT)
        assert r.returncode == 0, r.stderr
        d = tmp_path
    else:
        d = ROOT / "docs" / "uat" / "inputs"
    base = ["Incident Summary", "Identified Findings", "Evidence Links", "Causal Chain"]
    sheets, head = _sheets_and_header(d / "uat-votes.xlsx")
    assert sheets == base + ["Collaboration Data"] and head == FIND_COLS
    votes = [tuple(r) for r in openpyxl.load_workbook(d / "uat-votes.xlsx")["Collaboration Data"].iter_rows(values_only=True)]
    assert ("FND-0004", "downgrade", "weak evidence") in votes
    sheets, head = _sheets_and_header(d / "uat-renamed-header.xlsx")
    assert sheets == base and "Rating" not in head and head[FIND_COLS.index("Rating")] == "Score"
    sheets, head = _sheets_and_header(d / "uat-capa-planner.xlsx")
    assert sheets == base + ["Corrective Actions", "Preventive Actions"] and head == FIND_COLS
    for n in ("uat-votes-pack.json", "uat-unlinked-pack.json"):
        assert json.loads((d / n).read_text(encoding="utf-8"))["schema"] == 1


def test_dashboard_with_no_ratings(tmp_path):
    """No finding has a rating (Rating column renamed): dashboard must not crash on max() of nothing."""
    import openpyxl
    x = make_export(tmp_path)
    wb = openpyxl.load_workbook(x)
    sheet = wb["Identified Findings"]
    for cell in sheet[1]:
        if cell.value == "Rating":
            cell.value = "Score"
    wb.save(x)
    run("ingest", str(x), "--slug", "pump-renamed")
    ws_dir = tmp_path / "pmlc-runs"
    wsp = next(ws_dir.glob("*/workspace"))
    run("classify", str(wsp))
    run("dashboard", str(wsp))
    assert (wsp.parent / "outputs" / "organisation-learning-map.html").is_file()
