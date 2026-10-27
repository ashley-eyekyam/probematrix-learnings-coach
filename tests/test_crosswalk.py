"""Phase 5, Task 9: crosswalk data (RBPS 20 elements, ISO 45001 clauses) and the `crosswalk` command (L7-01).
Synthetic data only."""
import copy
import json
import subprocess
import sys

import openpyxl
import pytest

from test_items import register, wrap
from test_pmlc import make_export, PMLC, run
from test_privacy import share_ws, sign, write_doc
from test_references import _crosswalk, _crosswalk_requirements
import pmlc_lib.crosswalk as CW

#: ids that carried verified true before Task 9 and have no source key (checked in Task 3)
LEGACY_VERIFIED = {
    "ISO45001-5.4", "ISO45001-7.2", "ISO45001-7.3", "ISO45001-7.4", "ISO45001-9.1", "ISO45001-9.2", "ISO45001-9.3",
    "ISO45001-10.2", "PSM-(m)(3)", "PSM-(m)(4)", "PSM-(m)(5)", "PSM-(m)(6)", "PSM-(m)(7)", "PSM-(o)(4)", "PSM-(o)(5)",
    "RMP-68.81", "SEVESO-AnnexIII(vi)", "SEVESO-AnnexIII(vii)", "COMAH-SMS", "RBPS-incident-investigation",
    "RBPS-measurement-and-metrics", "RBPS-auditing", "RBPS-management-review", "ICMM-CCM-8", "ICMM-CCM-9",
    "IN-FactoriesAct-s.41G", "IN-FactoriesAct-s.88", "IN-OSHCode2020-s.10", "IN-OSHCode2020-s.11",
    "IN-OSHCode2020-s.22", "IN-MSIHC-r.5", "IN-ERDMP", "IN-OISD",
}


def cli(*args):
    return subprocess.run([sys.executable, str(PMLC), *args], capture_output=True, text=True)


@pytest.fixture(scope="module")
def cws(tmp_path_factory):
    d = tmp_path_factory.mktemp("cw")
    x = make_export(d, causal_rows=True)
    run("ingest", str(x), "--slug", "test-pump", "--date", "2024-03-20")
    ws = d / "pmlc-runs" / "2024-03-20-test-pump" / "workspace"
    out = ws.parent / "outputs"
    out.mkdir(exist_ok=True)
    reg = register()
    reg["entries"][0]["clauses"] = ["ISO45001-10.2", "PSM-(m)(5)"]
    reg["entries"][0]["lesson"] += " Brief the contractor crews."
    (out / "learnings-register.json").write_text(json.dumps(reg), encoding="utf-8")
    return ws


def go(ws, *extra):
    r = cli("crosswalk", str(ws), "--json", *extra)
    assert r.returncode == 0, r.stderr + r.stdout
    return json.loads(r.stdout)


def rows(d):
    return {r["id"]: r for r in d["rows"]}


# ----------------------------------------------------------------------------- data

def test_rbps_has_20_elements():
    reqs = _crosswalk()["frameworks"]["rbps"]["requirements"]
    assert len(reqs) == 20
    assert {"RBPS-incident-investigation", "RBPS-measurement-and-metrics", "RBPS-auditing",
            "RBPS-management-review"} <= {r["id"] for r in reqs}
    assert "RBPS-management-of-change" in {r["id"] for r in reqs}


def test_iso_clauses_present():
    ids = {r["id"] for r in _crosswalk()["frameworks"]["iso45001"]["requirements"]}
    want = {f"ISO45001-{c}" for c in ("4.1 4.2 4.3 4.4 5.1 5.2 5.3 5.4 6.1 6.2 7.1 7.2 7.3 7.4 7.5 8.1 8.2 9.1 9.2 9.3 "
                                      "10.1 10.2 10.3").split()}
    assert want <= ids, sorted(want - ids)


def test_verified_needs_source():
    bad = []
    for r in _crosswalk_requirements():
        if r["verified"] and not str(r.get("source", "")).startswith("https://") and r["id"] not in LEGACY_VERIFIED:
            bad.append(r["id"])
        if "topics" in r:
            import re
            re.compile(r["topics"])
    assert not bad, bad


# ----------------------------------------------------------------------------- command

def test_status_rules(cws):
    d = go(cws, "--register", str(cws.parent / "outputs" / "learnings-register.json"), "--framework", "iso45001")
    r = rows(d)
    assert d["framework"] == "iso45001" and d["title"]
    assert r["ISO45001-10.2"]["status"] == "evidenced here"
    assert r["ISO45001-10.2"]["evidence_here"] == ["LR-01"]
    assert r["ISO45001-10.2"]["cites"] == ["FND-0001", "FND-0002"]
    assert r["ISO45001-7.4"]["status"] == "partly"          # the "brief" topic hits a lesson text
    assert r["ISO45001-8.2"]["status"] == "needs organisation evidence"
    assert r["ISO45001-8.2"]["evidence_here"] == [] and r["ISO45001-8.2"]["organisation_must_show"]
    assert all(x["status"] != "not applicable" for x in d["rows"])
    assert go(cws, "--register", str(cws.parent / "outputs" / "learnings-register.json"),
              "--framework", "iso45001") == d                 # byte-identical on re-run


def test_default_register_and_psm(cws):
    d = go(cws, "--framework", "psm")
    assert rows(d)["PSM-(m)(5)"]["status"] == "evidenced here"


def test_org_files_move_rows_up(cws, tmp_path):
    from pmlc_lib.templates import ADOPTION_COLUMNS
    reg = json.loads((cws.parent / "outputs" / "learnings-register.json").read_text(encoding="utf-8"))
    base = rows(CW.crosswalk(CW.load_ws(cws), reg, "iso45001"))
    adoption = [dict(zip(ADOPTION_COLUMNS, ["LR-01", "Test", "adopt", "", "", ""]))]
    up = rows(CW.crosswalk(CW.load_ws(cws), reg, "iso45001", adoption=adoption))
    assert base["ISO45001-7.3"]["status"] == "needs organisation evidence"
    assert up["ISO45001-7.3"]["status"] == "partly"
    assert "adoption.csv" in up["ISO45001-7.3"]["organisation_must_show"]
    assert up["ISO45001-8.2"]["status"] == base["ISO45001-8.2"]["status"]
    dist = rows(CW.crosswalk(CW.load_ws(cws), reg, "psm", distribution=[{"pack": "x"}]))
    assert "distribution.csv" in dist["PSM-(m)(6)"]["organisation_must_show"]
    # no file, no wording about it
    assert "adoption.csv" not in base["ISO45001-7.3"]["organisation_must_show"]


def test_system_csv(cws, tmp_path):
    f = tmp_path / "sys.csv"
    f.write_text("id,title,topics\nORG-1,Our restart rule,restart|test\nORG-2,Our permit rule\n", encoding="utf-8")
    d = go(cws, "--system-csv", str(f))
    r = rows(d)
    assert list(r) == ["ORG-1", "ORG-2"]
    assert r["ORG-1"]["status"] == "partly"
    reg = json.loads((cws.parent / "outputs" / "learnings-register.json").read_text(encoding="utf-8"))
    reg["entries"][0]["clauses"] = ["ORG-1"]
    rp = tmp_path / "reg2.json"
    rp.write_text(json.dumps(reg), encoding="utf-8")
    assert rows(go(cws, "--system-csv", str(f), "--register", str(rp)))["ORG-1"]["status"] == "evidenced here"
    assert r["ORG-2"]["status"] == "needs organisation evidence"
    bad = tmp_path / "bad.csv"
    bad.write_text("name\nx\n", encoding="utf-8")
    p = cli("crosswalk", str(cws), "--system-csv", str(bad))
    assert p.returncode != 0 and "id,title" in p.stderr and "Traceback" not in p.stderr
    assert cli("crosswalk", str(cws)).returncode != 0          # neither framework nor system csv


def test_crosswalk_refuses_pack(cws, tmp_path):
    p = write_doc(tmp_path, wrap({"kind": "para", "text": "x", "cites": ["FND-0001"]}), "pk.json")
    r = cli("crosswalk", str(cws), "--register", str(p), "--framework", "psm")
    assert r.returncode != 0 and "Learnings Register" in r.stderr and "Traceback" not in r.stderr
    nolr = tmp_path / "e.json"
    nolr.write_text(json.dumps({"title": "x"}), encoding="utf-8")
    r = cli("crosswalk", str(cws), "--register", str(nolr), "--framework", "psm")
    assert r.returncode != 0 and "Learnings Register" in r.stderr and "Traceback" not in r.stderr
    for bad in ({"entries": []}, [1, 2]):             # final review I2: an empty table would read as a result
        nolr.write_text(json.dumps(bad), encoding="utf-8")
        r = cli("crosswalk", str(cws), "--register", str(nolr), "--framework", "iso45001")
        assert r.returncode != 0 and "Traceback" not in r.stderr and r.stdout == ""
        assert "crosswalk needs a Learnings Register with at least one learning" in r.stderr


def test_crosswalk_thin_export(tmp_path):
    x = make_export(tmp_path)
    run("ingest", str(x), "--slug", "thin", "--date", "2024-03-20")
    ws = tmp_path / "pmlc-runs" / "2024-03-20-thin" / "workspace"
    reg = tmp_path / "r.json"
    reg.write_text(json.dumps({"entries": [{"lr_id": "LR-01", "title": "Zzz"}]}), encoding="utf-8")
    d = go(ws, "--register", str(reg), "--framework", "rbps")
    assert len(d["rows"]) == 20 and all(r["status"] == "needs organisation evidence" for r in d["rows"])


@pytest.mark.parametrize("fw", ["iso45001", "psm", "rmp", "seveso", "rbps", "icmm", "india"])
def test_table_item_pastes_into_pack(cws, tmp_path, fw):
    r = cli("crosswalk", str(cws), "--framework", fw, "--table-item")
    assert r.returncode == 0, r.stderr
    p = write_doc(tmp_path, wrap(json.loads(r.stdout)), "cw.json")
    c = cli("check-pack", str(cws), str(p), "--strict")
    assert c.returncode == 0, c.stdout + c.stderr


def test_table_item(cws):
    r = cli("crosswalk", str(cws), "--framework", "iso45001", "--table-item")
    assert r.returncode == 0, r.stderr
    it = json.loads(r.stdout)
    assert it["kind"] == "table" and it["columns"] == ["Requirement", "Status", "Evidence here",
                                                        "Organisation must show"]
    assert len(it["rows"]) == len(it["row_cites"]) and it["rows"][0][0].startswith("ISO45001-")
    # Task 19: the compliance-crosswalk pack names its section crosswalk-<framework>, and pack ids are unique across
    # sections and items, so the table carries its own id, as the evidence-audit tables do
    assert it["id"] == "crosswalk-iso45001-table"
    md = cli("crosswalk", str(cws), "--framework", "iso45001")
    assert md.returncode == 0 and "| Requirement |" in md.stdout


# ----------------------------------------------------------------------------- pack-xlsx --kind table

def table_pack(text="Who checked the restart?", req="ISO45001-10.2 Corrective action"):
    d = copy.deepcopy(wrap({"kind": "table", "id": "t1", "columns": ["Requirement", "Status"],
                            "rows": [[req, text]], "row_cites": [["FND-0001"]],
                            "cites": ["FND-0001"],
                            "lint_ok": [{"rule": "JC1", "reason": "test of the name guard"},
                                        {"rule": "JC7", "reason": "test of masking a singleton role"}]}))
    return d


def test_pack_xlsx_table_shared_copy(tmp_path):
    ws, _ = share_ws(tmp_path)
    names = tmp_path / "names.txt"
    names.write_text("Quillon Vantress\n", encoding="utf-8")
    bad = write_doc(tmp_path, table_pack("Did Quillon Vantress sign?"), "bad.json")
    r = cli("pack-xlsx", str(ws), str(bad), "--kind", "table", "--names", str(names), "--out", str(tmp_path / "b.xlsx"))
    assert r.returncode == 1 and "refused: a listed name appears in the output" in r.stderr
    from test_privacy import SINGLE
    p = write_doc(tmp_path, table_pack("Checked at 14:35 on 2024-03-20 by the " + SINGLE), "good.json")
    sign(ws, p, "--classification", "external", "--legal-review", "cleared")
    internal = tmp_path / "i.xlsx"
    assert cli("pack-xlsx", str(ws), str(p), "--kind", "table", "--out", str(internal)).returncode == 0
    wb = openpyxl.load_workbook(internal)
    sh = wb["t1"]
    assert [c.value for c in sh[1]] == ["Requirement", "Status", "cites"]
    assert sh["C2"].value == "FND-0001"
    r = cli("pack-xlsx", str(ws), str(p), "--kind", "table", "--share", "external", "--names", str(names))
    assert r.returncode == 0, r.stderr
    shared = ws.parent / "outputs" / "test-pack-tables.external.xlsx"
    assert shared.is_file() and not (ws.parent / "outputs" / "test-pack-tables.xlsx").exists()
    cells = " ".join(str(c.value) for row in openpyxl.load_workbook(shared)["t1"].iter_rows() for c in row)
    assert "ISO45001-10.2" in cells and "[equipment tag]" not in cells
    assert "14:35" not in cells and "2024-03-20" not in cells and SINGLE not in cells
    assert "Shared external" in openpyxl.load_workbook(shared)["t1"].oddHeader.center.text
